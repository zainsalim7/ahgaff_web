# Phase 5 (Pilot, SHADOW MODE) regression tests for the advanced access center:
# run/reset (admin only), summary/results/filters, export xlsx, view_pilot sub-permission gating,
# and no mutation of the 4 real users' legacy fields (role/custom_permissions/faculty_id/department_id).
import os
import io
import pytest
import requests
from bson import ObjectId
from pymongo import MongoClient
from dotenv import load_dotenv

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')
if not BASE_URL:
    with open('/app/frontend/.env') as f:
        for line in f:
            if line.startswith('REACT_APP_BACKEND_URL='):
                BASE_URL = line.split('=', 1)[1].strip().rstrip('/')
API = f"{BASE_URL}/api/advanced-access"

load_dotenv('/app/backend/.env')
DB = MongoClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]

REAL_USERS = {
    "Salim": "698f71ba5c5620a4449eeb7e",
    "Saeed": "698f79d75c5620a4449eeb89",
    "teacher180156": "698f671431cfba92c49a550e",
    "president": "6ab13e3709466a2b7b88502d",
}
LEGACY_FIELDS = ["role", "custom_permissions", "permissions", "faculty_id", "department_id", "role_id", "is_active", "hr_scope"]


def _login(u, p):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": u, "password": p})
    if r.status_code != 200:
        pytest.skip(f"login failed for {u}: {r.status_code}")
    d = r.json()
    return {"Authorization": f"Bearer {d.get('access_token') or d.get('token')}"}


def _snapshot():
    out = {}
    for name, uid in REAL_USERS.items():
        u = DB.users.find_one({"_id": ObjectId(uid)}) or {}
        out[name] = {k: u.get(k) for k in LEGACY_FIELDS}
    return out


@pytest.fixture(scope="module")
def admin():
    return _login("admin", "admin123")


@pytest.fixture(scope="module")
def baseline():
    return _snapshot()


# --- run / summary ---
class TestPilotRun:
    def test_run_pilot_admin(self, admin, baseline):
        r = requests.post(f"{API}/pilot/run", headers=admin, timeout=180)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["total_cases"] > 50
        assert d["fail_count"] == 0
        assert d["security_issues_count"] == 0
        assert d["success_rate"] == 100

    def test_summary_has_4_users(self, admin):
        d = requests.get(f"{API}/pilot/summary", headers=admin).json()
        assert d["has_run"] is True
        assert "_id" not in d
        assert set(REAL_USERS.keys()) <= set(d["by_user"].keys()), d["by_user"].keys()
        assert d["by_user"]["president"]["baseline"] == d["by_user"]["president"]["total"]
        assert d["total_cases"] == sum(v["total"] for v in d["by_user"].values())

    def test_president_no_assignment(self, admin):
        n = DB.advanced_user_assignments.count_documents({"user_id": REAL_USERS["president"], "context_label_ar": {"$regex": "^PILOT_"}})
        assert n == 0

    def test_saeed_three_active_pilot_assignments(self):
        act = list(DB.advanced_user_assignments.find({"user_id": REAL_USERS["Saeed"], "context_label_ar": {"$regex": "^PILOT_"}, "status": "active"}))
        assert len(act) == 3, [a["context_label_ar"] for a in act]
        assert sorted(a["assignment_type"] for a in act) == ["academic", "administrative", "administrative"]

    def test_run_non_admin_forbidden(self):
        h = _login("Saeed", "test1234")
        assert requests.post(f"{API}/pilot/run", headers=h).status_code == 403
        assert requests.post(f"{API}/pilot/reset", headers=h).status_code == 403


# --- results filters ---
class TestPilotFilters:
    def test_filters(self, admin):
        allr = requests.get(f"{API}/pilot/results", headers=admin).json()
        assert len(allr) > 0 and "_id" not in allr[0]
        for params, key, val in [({"user_key": "saeed"}, "user_key", "saeed"), ({"severity": "info"}, "severity", "info"),
                                 ({"status": "baseline"}, "status", "baseline"), ({"status": "diff"}, "status", "diff")]:
            r = requests.get(f"{API}/pilot/results", headers=admin, params=params).json()
            assert all(x[key] == val for x in r)
            assert len(r) == sum(1 for x in allr if x[key] == val)

    def test_permission_key_filter(self, admin):
        allr = requests.get(f"{API}/pilot/results", headers=admin).json()
        pk = next(x["permission_key"] for x in allr if x["permission_key"] and x["permission_key"][0].isalpha())
        r = requests.get(f"{API}/pilot/results", headers=admin, params={"permission_key": pk}).json()
        assert len(r) >= 1 and all(pk in x["permission_key"] for x in r)

    def test_permission_key_regex_special_chars(self, admin):
        # unescaped $regex — "(" should not crash server
        r = requests.get(f"{API}/pilot/results", headers=admin, params={"permission_key": "("})
        assert r.status_code == 200, f"status {r.status_code}"


# --- export ---
class TestPilotExport:
    def test_export_xlsx(self, admin):
        from openpyxl import load_workbook
        r = requests.get(f"{API}/pilot/export.xlsx", headers=admin)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        total = requests.get(f"{API}/pilot/summary", headers=admin).json()["total_cases"]
        assert ws.max_row - 1 == total


# --- view_pilot gating ---
class TestViewPilotGate:
    def test_view_only_tester_denied_then_allowed(self, admin):
        tid = REAL_USERS["teacher180156"]
        try:
            assert requests.post(f"{API}/testers", headers=admin, json={"user_id": tid, "sub_permissions": ["view"]}).status_code == 200
            th = _login("teacher180156", "teacher123")
            assert requests.get(f"{API}/pilot/summary", headers=th).status_code == 403
            assert requests.get(f"{API}/pilot/results", headers=th).status_code == 403
            assert requests.get(f"{API}/pilot/export.xlsx", headers=th).status_code == 403
            assert requests.post(f"{API}/testers", headers=admin, json={"user_id": tid, "sub_permissions": ["view", "view_pilot"]}).status_code == 200
            assert requests.get(f"{API}/pilot/summary", headers=th).status_code == 200
            assert requests.get(f"{API}/pilot/results", headers=th).status_code == 200
            assert requests.get(f"{API}/pilot/export.xlsx", headers=th).status_code == 200
            assert requests.post(f"{API}/pilot/run", headers=th).status_code == 403
        finally:
            requests.delete(f"{API}/testers/{tid}", headers=admin)
        th = _login("teacher180156", "teacher123")
        assert requests.get(f"{API}/pilot/summary", headers=th).status_code == 403

    def test_saeed_no_view_pilot(self):
        h = _login("Saeed", "test1234")
        assert requests.get(f"{API}/pilot/summary", headers=h).status_code == 403

    def test_salim_denied(self):
        h = _login("Salim", "test1234")
        assert requests.get(f"{API}/pilot/summary", headers=h).status_code == 403


# --- reset then re-run; legacy data untouched ---
class TestPilotResetAndIntegrity:
    def test_reset_cleans_everything(self, admin):
        r = requests.post(f"{API}/pilot/reset", headers=admin)
        assert r.status_code == 200
        assert requests.get(f"{API}/pilot/summary", headers=admin).json() == {"has_run": False}
        assert requests.get(f"{API}/pilot/results", headers=admin).json() == []
        roles = requests.get(f"{API}/roles", headers=admin).json()
        assert not [x for x in roles if (x.get("name_ar") or "").startswith("PILOT_")]
        asg = requests.get(f"{API}/assignments", headers=admin).json()
        assert not [a for a in asg if (a.get("context_label_ar") or "").startswith("PILOT_") and a.get("status") == "active"]

    def test_rerun_and_legacy_unchanged(self, admin, baseline):
        r = requests.post(f"{API}/pilot/run", headers=admin, timeout=180)
        assert r.status_code == 200
        assert _snapshot() == baseline
