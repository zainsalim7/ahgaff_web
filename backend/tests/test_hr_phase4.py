"""HR Phase 4 tests: Employee Documents, Appraisal PDF, Annual Report"""
import io
import os
import pytest
import requests
from datetime import datetime

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL missing"

EMP100_ID = "6ab41cd4bee9a1cf08bef181"
EMP200_ID = "6ab41cffbee9a1cf08bef190"
APPR_ID = "6ab42193ce78fea2f1416d28"
EXISTING_DOC_ID = "6ab42ee7efb8b139d9e771c8"


def _login(u, p):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": u, "password": p}, timeout=30)
    assert r.status_code == 200, f"login {u} → {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_h():
    return {"Authorization": f"Bearer {_login('admin', 'admin123')}"}


@pytest.fixture(scope="module")
def emp100_h():
    return {"Authorization": f"Bearer {_login('EMP-100', 'EMP-100')}"}


@pytest.fixture(scope="module")
def emp200_h():
    return {"Authorization": f"Bearer {_login('EMP-200', 'EMP-200')}"}


# ─────── Documents ───────
def _tiny_pdf():
    return (b"%PDF-1.4\n1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n"
            b"2 0 obj<< /Type /Pages /Count 1 /Kids [3 0 R] >>endobj\n"
            b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] >>endobj\n"
            b"xref\n0 4\n0000000000 65535 f\ntrailer<< /Size 4 /Root 1 0 R >>\nstartxref\n0\n%%EOF")


class TestDocumentsMeta:
    def test_meta(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/meta", headers=admin_h)
        assert r.status_code == 200
        d = r.json()
        assert "types" in d and "contract" in d["types"]
        assert d["max_mb"] == 10
        assert "application/pdf" in d["allowed"]


class TestDocumentUpload:
    upload_id = None

    def test_upload_wrong_content_type(self, admin_h):
        r = requests.post(
            f"{BASE_URL}/api/hr/documents/{EMP100_ID}",
            headers=admin_h,
            files={"file": ("x.txt", b"hello", "text/plain")},
            data={"type": "contract", "title": "TEST_bad"},
        )
        assert r.status_code == 400

    def test_upload_bad_type(self, admin_h):
        r = requests.post(
            f"{BASE_URL}/api/hr/documents/{EMP100_ID}",
            headers=admin_h,
            files={"file": ("x.pdf", _tiny_pdf(), "application/pdf")},
            data={"type": "bogus", "title": "TEST_bad"},
        )
        assert r.status_code == 400

    def test_upload_as_emp100_forbidden(self, emp100_h):
        r = requests.post(
            f"{BASE_URL}/api/hr/documents/{EMP100_ID}",
            headers=emp100_h,
            files={"file": ("x.pdf", _tiny_pdf(), "application/pdf")},
            data={"type": "contract"},
        )
        assert r.status_code == 403

    def test_upload_success_expiry_soon(self, admin_h):
        # expiry within 60 days
        from datetime import date, timedelta
        exp = (date.today() + timedelta(days=30)).strftime("%Y-%m-%d")
        r = requests.post(
            f"{BASE_URL}/api/hr/documents/{EMP100_ID}",
            headers=admin_h,
            files={"file": ("test.pdf", _tiny_pdf(), "application/pdf")},
            data={"type": "certificate", "title": "TEST_cert_phase4", "expiry_date": exp, "notes": "TEST"},
        )
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["document"]["expiry_state"] == "soon"
        TestDocumentUpload.upload_id = j["id"]


class TestDocumentList:
    def test_list_admin(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/{EMP100_ID}", headers=admin_h)
        assert r.status_code == 200
        assert "items" in r.json()

    def test_list_own_emp100(self, emp100_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/{EMP100_ID}", headers=emp100_h)
        assert r.status_code == 200

    def test_list_other_emp100(self, emp100_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/{EMP200_ID}", headers=emp100_h)
        assert r.status_code == 403

    def test_my(self, emp100_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/my", headers=emp100_h)
        assert r.status_code == 200
        assert isinstance(r.json()["items"], list)


class TestDocumentFile:
    def test_download(self, admin_h):
        did = TestDocumentUpload.upload_id
        assert did
        r = requests.get(f"{BASE_URL}/api/hr/documents/file/{did}", headers=admin_h)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert len(r.content) > 10

    def test_update_meta(self, admin_h):
        did = TestDocumentUpload.upload_id
        r = requests.put(f"{BASE_URL}/api/hr/documents/{did}", headers=admin_h,
                         json={"type": "certificate", "title": "TEST_updated", "notes": "new"})
        assert r.status_code == 200

    def test_expiring(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/expiring?days=60", headers=admin_h)
        assert r.status_code == 200
        items = r.json()["items"]
        assert isinstance(items, list) and len(items) >= 1
        # existing doc with 2026-11-01 must be present (server date 2026-09-23 → 39 days)
        assert any(i["id"] == EXISTING_DOC_ID for i in items) or any(
            i["id"] == TestDocumentUpload.upload_id for i in items
        )

    def test_alerts_daily(self, admin_h):
        r = requests.post(f"{BASE_URL}/api/hr/alerts/run-now?kind=daily", headers=admin_h)
        assert r.status_code == 200
        assert "documents_expiring" in r.json()

    def test_alerts_weekly(self, admin_h):
        r = requests.post(f"{BASE_URL}/api/hr/alerts/run-now?kind=weekly", headers=admin_h)
        assert r.status_code == 200
        assert "expiring_documents" in r.json()

    def test_delete(self, admin_h):
        did = TestDocumentUpload.upload_id
        r = requests.delete(f"{BASE_URL}/api/hr/documents/{did}", headers=admin_h)
        assert r.status_code == 200
        r2 = requests.get(f"{BASE_URL}/api/hr/documents/{EMP100_ID}", headers=admin_h)
        assert did not in [i["id"] for i in r2.json()["items"]]


# ─────── Appraisal PDF ───────
class TestAppraisalPDF:
    verify_token = None

    def test_admin_pdf(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/appraisals/{APPR_ID}/pdf", headers=admin_h)
        assert r.status_code == 200, r.text[:300]
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert len(r.content) > 10 * 1024

    def test_emp100_pdf(self, emp100_h):
        r = requests.get(f"{BASE_URL}/api/hr/appraisals/{APPR_ID}/pdf", headers=emp100_h)
        assert r.status_code == 200

    def test_emp200_pdf(self, emp200_h):
        r = requests.get(f"{BASE_URL}/api/hr/appraisals/{APPR_ID}/pdf", headers=emp200_h)
        assert r.status_code == 200

    def test_draft_pdf_400(self, admin_h):
        # create draft for EMP-200
        r = requests.post(f"{BASE_URL}/api/hr/appraisals", headers=admin_h,
                          json={"employee_id": EMP200_ID, "year": 2025})
        assert r.status_code in (200, 201, 400, 409), r.text
        # find its id
        if r.status_code in (200, 201):
            draft_id = r.json().get("id") or r.json().get("appraisal", {}).get("id")
        else:
            # already exists, find it
            lst = requests.get(f"{BASE_URL}/api/hr/appraisals?employee_id={EMP200_ID}&year=2025", headers=admin_h)
            drafts = [a for a in lst.json().get("items", []) if a.get("status") == "draft"]
            draft_id = drafts[0]["id"] if drafts else None
        if not draft_id:
            pytest.skip("no draft available")
        p = requests.get(f"{BASE_URL}/api/hr/appraisals/{draft_id}/pdf", headers=admin_h)
        assert p.status_code == 400
        d = requests.delete(f"{BASE_URL}/api/hr/appraisals/{draft_id}", headers=admin_h)
        assert d.status_code in (200, 204)


class TestPublicVerify:
    def test_invalid_token_no_auth(self):
        r = requests.get(f"{BASE_URL}/api/hr/verify/appraisal/thisisnotarealtokenxxx")
        assert r.status_code == 200
        assert r.json().get("valid") is False

    def test_real_token_no_auth(self, admin_h):
        # Generate PDF first to ensure verify_token exists
        requests.get(f"{BASE_URL}/api/hr/appraisals/{APPR_ID}/pdf", headers=admin_h)
        # Query verify token via mongo
        import subprocess
        cmd = ("python3 -c \"import asyncio,os;from motor.motor_asyncio import AsyncIOMotorClient;from bson import ObjectId;"
               "c=AsyncIOMotorClient(os.environ['MONGO_URL']);"
               f"d=c[os.environ['DB_NAME']];"
               f"print(asyncio.run(d.hr_appraisals.find_one({{'_id':ObjectId('{APPR_ID}')}})).get('verify_token'))\"")
        # simpler: use pymongo
        from pymongo import MongoClient
        from bson import ObjectId
        # Read env directly
        with open("/app/backend/.env") as f:
            env = dict(l.strip().split("=", 1) for l in f if "=" in l and not l.startswith("#"))
        mongo_url = env["MONGO_URL"].strip('"')
        db_name = env["DB_NAME"].strip('"')
        mc = MongoClient(mongo_url)
        a = mc[db_name]["hr_appraisals"].find_one({"_id": ObjectId(APPR_ID)})
        token = a.get("verify_token")
        assert token, "verify_token missing on appraisal"
        r = requests.get(f"{BASE_URL}/api/hr/verify/appraisal/{token}")
        assert r.status_code == 200
        j = r.json()
        assert j.get("valid") is True
        assert "employee_name" in j and "grade" in j and "total_score" in j


# ─────── Annual Report ───────
class TestAnnualReport:
    def test_annual_admin(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/reports/annual?year=2026", headers=admin_h)
        assert r.status_code == 200, r.text[:300]
        d = r.json()
        for k in ("workforce", "attendance", "leaves", "appraisals", "tasks"):
            assert k in d
        wf = d["workforce"]
        for k in ("total", "active", "hires", "exits", "by_category", "by_status", "by_contract", "by_unit"):
            assert k in wf
        assert len(d["attendance"]["months"]) == 9  # Jan..Sep 2026
        assert "avg_rate" in d["attendance"]
        assert "distribution" in d["appraisals"] and len(d["appraisals"]["distribution"]) == 5
        assert "per_unit" in d["appraisals"]

    def test_annual_unit_filter(self, admin_h):
        units = requests.get(f"{BASE_URL}/api/hr/org-units", headers=admin_h).json()
        items = units if isinstance(units, list) else (units.get("items") or units.get("units") or [])
        assert items
        uid = items[0].get("id") or items[0].get("_id")
        r = requests.get(f"{BASE_URL}/api/hr/reports/annual?year=2026&org_unit_id={uid}", headers=admin_h)
        assert r.status_code == 200

    def test_annual_emp100_forbidden(self, emp100_h):
        r = requests.get(f"{BASE_URL}/api/hr/reports/annual?year=2026", headers=emp100_h)
        assert r.status_code == 403

    def test_export_pdf(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/reports/annual/export?fmt=pdf&year=2026", headers=admin_h)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert len(r.content) > 5000

    def test_export_xlsx(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/reports/annual/export?fmt=xlsx&year=2026", headers=admin_h)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        # verify multiple sheets
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert len(wb.sheetnames) > 3


# ─────── Re-upload one doc for FE test if needed ───────
class TestReuploadForFE:
    def test_ensure_existing_doc(self, admin_h):
        r = requests.get(f"{BASE_URL}/api/hr/documents/{EMP100_ID}", headers=admin_h)
        items = r.json().get("items", [])
        if not any(i["id"] == EXISTING_DOC_ID for i in items):
            # Re-upload a contract for EMP-100 so FE tests find at least one
            from datetime import date, timedelta
            exp = (date(2026, 11, 1)).strftime("%Y-%m-%d")
            up = requests.post(
                f"{BASE_URL}/api/hr/documents/{EMP100_ID}",
                headers=admin_h,
                files={"file": ("contract.pdf", _tiny_pdf(), "application/pdf")},
                data={"type": "contract", "title": "TEST_contract_reup", "expiry_date": exp},
            )
            assert up.status_code == 200
