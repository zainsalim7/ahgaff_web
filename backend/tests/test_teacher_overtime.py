"""Backend tests for Teacher Overtime Sheet (كشف الساعات الإضافية للأساتذة)"""
import os
import subprocess
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://schedule-hub-272.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"


def _mongo_eval(expr: str) -> str:
    r = subprocess.run(
        ["mongosh", "test_database", "--quiet", "--eval", expr],
        capture_output=True, text=True, timeout=15,
    )
    return (r.stdout or "").strip()


@pytest.fixture(scope="session")
def faculty_id() -> str:
    fid = _mongo_eval("print(db.teachers.findOne({seed_ot:true}).faculty_id)")
    assert fid and len(fid) == 24, f"invalid faculty id from seed: {fid!r}"
    return fid


@pytest.fixture(scope="session")
def admin_token() -> str:
    r = requests.post(f"{API}/auth/login", json={"username": "admin", "password": "admin123"}, timeout=15)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="session")
def teacher_token() -> str:
    r = requests.post(f"{API}/auth/login", json={"username": "teacher180156", "password": "teacher123"}, timeout=15)
    if r.status_code != 200:
        pytest.skip("teacher login failed")
    return r.json()["access_token"]


PARAMS_BASE = {"start_date": "2026-08-15", "end_date": "2026-09-24"}


# ---------- Policies ----------
class TestPolicies:
    def test_list_policies(self, admin_headers):
        r = requests.get(f"{API}/reports/teacher-workload/policies", headers=admin_headers, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "items" in data and "titles" in data
        assert data["threshold"] == 75
        default = [x for x in data["items"] if x.get("is_default")]
        assert default and default[0]["faculty_id"] == ""
        assert any(x.get("inherits_default") in (True, False) for x in data["items"] if not x.get("is_default"))

    def test_policy_divisor_validation(self, admin_headers, faculty_id):
        r = requests.put(
            f"{API}/reports/teacher-workload/policy",
            headers=admin_headers,
            json={"faculty_id": faculty_id, "bonus_divisor": 0},
            timeout=20,
        )
        assert r.status_code == 400, r.text

    def test_teacher_forbidden_policies(self, teacher_token):
        r = requests.get(
            f"{API}/reports/teacher-workload/policies",
            headers={"Authorization": f"Bearer {teacher_token}"},
            timeout=15,
        )
        assert r.status_code == 403


# ---------- Report ----------
class TestReport:
    def _get(self, headers, faculty_id, **extra):
        params = {**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "true", **extra}
        return requests.get(f"{API}/reports/teacher-workload", headers=headers, params=params, timeout=60)

    def test_period_weeks(self, admin_headers, faculty_id):
        r = self._get(admin_headers, faculty_id)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data.get("period", {}).get("total_weeks") == 6

    def test_seeded_teacher_baaza(self, admin_headers, faculty_id):
        r = self._get(admin_headers, faculty_id)
        data = r.json()
        teachers = data.get("teachers", [])
        baaz = next((t for t in teachers if "باوزير" in (t.get("teacher_name") or "")), None)
        assert baaz, f"teacher باوزير not found. names={[t.get('teacher_name') for t in teachers]}"
        s = baaz["summary"]
        assert s.get("weekly_hours") == 8
        assert s.get("required_hours") == 48
        assert abs(s.get("total_actual_hours", 0) - 58.5) < 0.5
        assert s.get("bonus_enabled") is True
        assert abs(s.get("bonus_hours", 0) - 9.75) < 0.5
        assert abs(s.get("hours_after_bonus", 0) - 68.25) < 0.5
        assert abs(s.get("overtime_hours", 0) - 20.25) < 0.5
        assert s.get("hourly_rate") == 5000
        assert abs(s.get("overtime_amount", 0) - 101250) < 1

        # Course-level checks
        courses = {c.get("course_name"): c for c in baaz.get("courses", [])}
        fiqh = next((c for name, c in courses.items() if "فقه العبادات" in name), None)
        assert fiqh, f"فقه العبادات missing: {list(courses.keys())}"
        assert "بدأ من الأسبوع 3" in (fiqh.get("note") or "")
        assert abs((fiqh.get("expected_hours") or 0) - 12) < 0.5
        assert abs((fiqh.get("actual_hours") or 0) - 12) < 0.5

        maqasid = next((c for name, c in courses.items() if "مقاصد" in name), None)
        assert maqasid, "مقاصد الشريعة missing"
        assert "غياب" in (maqasid.get("note") or "") or "عذر" in (maqasid.get("note") or "")

    def test_seeded_teacher_bin_sumait(self, admin_headers, faculty_id):
        r = self._get(admin_headers, faculty_id)
        data = r.json()
        teachers = data.get("teachers", [])
        bs = next((t for t in teachers if "سميط" in (t.get("teacher_name") or "")), None)
        assert bs, "teacher بن سميط not found"
        s = bs["summary"]
        assert s.get("weekly_hours") == 9
        assert s.get("required_hours") == 54
        assert abs(s.get("total_actual_hours", 0) - 105) < 1
        assert abs(s.get("bonus_hours", 0) - 17.5) < 1
        assert abs(s.get("overtime_hours", 0) - 68.5) < 1
        assert s.get("hourly_rate") == 6000
        assert abs(s.get("overtime_amount", 0) - 411000) < 5

    def test_summary_totals(self, admin_headers, faculty_id):
        r = self._get(admin_headers, faculty_id)
        s = r.json().get("summary", {})
        assert "total_overtime_hours" in s
        assert "total_overtime_amount" in s
        assert s.get("threshold") == 75 or r.json().get("threshold") == 75

    def test_hide_empty_toggle(self, admin_headers, faculty_id):
        params = {**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "false"}
        r = requests.get(f"{API}/reports/teacher-workload", headers=admin_headers, params=params, timeout=60)
        assert r.status_code == 200
        teachers = r.json().get("teachers", [])
        empty = [t for t in teachers if "لا محاضرات" in (t.get("note") or "")]
        # may or may not have empties; just ensure call works and includes field
        assert isinstance(teachers, list)
        _ = empty

    def test_legacy_fields_present(self, admin_headers, faculty_id):
        r = self._get(admin_headers, faculty_id)
        teachers = r.json().get("teachers", [])
        assert teachers
        t = teachers[0]
        s = t.get("summary", {})
        for k in ("required_hours", "total_scheduled_hours", "total_actual_hours", "difference_hours", "completion_rate"):
            assert k in s, f"legacy field {k} missing in summary"
        if t.get("courses"):
            c = t["courses"][0]
            for k in ("scheduled_lectures", "executed_lectures", "scheduled_hours", "actual_hours"):
                assert k in c, f"legacy course field {k} missing"


# ---------- Policy toggle flow (disable → re-enable) ----------
class TestPolicyToggleFlow:
    def test_disable_bonus_then_restore(self, admin_headers, faculty_id):
        # Disable
        r = requests.put(
            f"{API}/reports/teacher-workload/policy",
            headers=admin_headers,
            json={"faculty_id": faculty_id, "bonus_enabled": False, "bonus_divisor": 6,
                  "rates": {"أستاذ مساعد": 5000, "أستاذ مشارك": 6000}, "currency": "ر.ي"},
            timeout=20,
        )
        assert r.status_code == 200, r.text
        rep = requests.get(f"{API}/reports/teacher-workload", headers=admin_headers,
                           params={**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "true"}, timeout=60)
        teachers = rep.json().get("teachers", [])
        baaz = next((t for t in teachers if "باوزير" in (t.get("teacher_name") or "")), None)
        assert baaz
        s = baaz["summary"]
        assert abs(s.get("bonus_hours", 0)) < 0.01
        # overtime = actual - required = 58.5 - 48 = 10.5
        assert abs(s.get("overtime_hours", 0) - 10.5) < 0.5

        # Restore (re-enable)
        r2 = requests.put(
            f"{API}/reports/teacher-workload/policy",
            headers=admin_headers,
            json={"faculty_id": faculty_id, "bonus_enabled": True, "bonus_divisor": 6,
                  "rates": {"أستاذ مساعد": 5000, "أستاذ مشارك": 6000}, "currency": "ر.ي"},
            timeout=20,
        )
        assert r2.status_code == 200
        rep2 = requests.get(f"{API}/reports/teacher-workload", headers=admin_headers,
                            params={**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "true"}, timeout=60)
        baaz2 = next((t for t in rep2.json().get("teachers", []) if "باوزير" in (t.get("teacher_name") or "")), None)
        assert baaz2 and abs(baaz2["summary"].get("bonus_hours", 0) - 9.75) < 0.5


# ---------- Exports ----------
class TestExports:
    def test_sheet_pdf(self, admin_headers, faculty_id):
        r = requests.get(
            f"{API}/export/report/teacher-workload/sheet-pdf",
            headers=admin_headers,
            params={**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "true"},
            timeout=60,
        )
        assert r.status_code == 200, r.text[:300]
        assert "application/pdf" in r.headers.get("content-type", "")
        assert len(r.content) > 1000

    def test_sheet_excel(self, admin_headers, faculty_id):
        r = requests.get(
            f"{API}/export/report/teacher-workload/sheet-excel",
            headers=admin_headers,
            params={**PARAMS_BASE, "faculty_id": faculty_id, "hide_empty": "true"},
            timeout=60,
        )
        assert r.status_code == 200
        assert "sheet" in r.headers.get("content-type", "").lower() or "openxml" in r.headers.get("content-type", "").lower()
        # Verify sheets
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert "كشف الساعات الإضافية" in wb.sheetnames, f"sheets={wb.sheetnames}"
        assert "المستحقات المالية" in wb.sheetnames, f"sheets={wb.sheetnames}"

    def test_legacy_pdf(self, admin_headers, faculty_id):
        r = requests.get(
            f"{API}/export/report/teacher-workload/pdf",
            headers=admin_headers,
            params={**PARAMS_BASE, "faculty_id": faculty_id},
            timeout=60,
        )
        assert r.status_code == 200

    def test_legacy_excel(self, admin_headers, faculty_id):
        r = requests.get(
            f"{API}/export/report/teacher-workload/excel",
            headers=admin_headers,
            params={**PARAMS_BASE, "faculty_id": faculty_id},
            timeout=60,
        )
        assert r.status_code == 200

    def test_legacy_excel_monthly(self, admin_headers, faculty_id):
        r = requests.get(
            f"{API}/export/report/teacher-workload/excel",
            headers=admin_headers,
            params={**PARAMS_BASE, "faculty_id": faculty_id, "monthly": "true"},
            timeout=60,
        )
        assert r.status_code == 200


# ---------- Teacher scope ----------
class TestTeacherScope:
    def test_teacher_sees_only_own(self, teacher_token):
        r = requests.get(
            f"{API}/reports/teacher-workload",
            headers={"Authorization": f"Bearer {teacher_token}"},
            params=PARAMS_BASE,
            timeout=30,
        )
        if r.status_code in (403, 404):
            pytest.skip(f"teacher cannot access report ({r.status_code})")
        assert r.status_code == 200
        teachers = r.json().get("teachers", [])
        assert len(teachers) <= 1
