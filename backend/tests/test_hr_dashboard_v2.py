"""HR dedicated dashboard endpoint tests (iteration 79)."""
import os
import io
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")


def _login(username: str, password: str) -> str:
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": username, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed for {username}: {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_headers():
    return {"Authorization": f"Bearer {_login('admin', 'admin123')}"}


@pytest.fixture(scope="module")
def dean_headers():
    return {"Authorization": f"Bearer {_login('Salim', 'test1234')}"}


@pytest.fixture(scope="module")
def head_headers():
    return {"Authorization": f"Bearer {_login('Saeed', 'test1234')}"}


# ── /api/dashboard/management/hr shape ────────────────────────────
@pytest.mark.parametrize("period", ["week", "month", "day"])
def test_hr_endpoint_structure(admin_headers, period):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": period}, headers=admin_headers, timeout=90)
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    for k in ("view", "period_label", "date_from", "date_to", "sections", "units", "hr"):
        assert k in body, f"missing top-level {k}"
    assert body["view"] == "hr"
    assert isinstance(body["units"], list) and len(body["units"]) > 0
    for u in body["units"]:
        for k in ("id", "name", "type", "type_label"):
            assert k in u, f"unit missing {k}: {u}"
        assert "parent_id" in u  # can be None

    hr = body["hr"]
    assert hr is not None
    # scope
    assert hr["scope"]["org_unit_id"] is None
    assert hr["scope"]["label"] == "كل الوحدات التنظيمية"
    # headcount sums
    hc = hr["headcount"]
    assert sum(c["count"] for c in hc["by_category"]) == hc["total_active"], "by_category sum mismatch"
    for row in hc["by_unit"]:
        assert row["academic"] + row["administrative"] + row["other"] == row["total"], f"row sum mismatch: {row}"
    # alerts
    alerts = hr["alerts"]
    assert isinstance(alerts, list)
    keys = {a.get("key") for a in alerts}
    for req_key in ("hr_absent", "hr_pending_leaves", "hr_overdue_tasks", "hr_contracts", "hr_appraisals"):
        assert req_key in keys, f"missing alert key {req_key} in {keys}"
    # today fields
    for k in ("present_today", "absent_today", "on_leave_today", "pending_leaves",
              "expiring_contracts", "overdue_tasks_list", "pending_appraisals_list"):
        assert k in hr, f"missing hr.{k}"
    # period
    assert "period" in hr and "chart" in hr["period"]


# ── unit scoping ──────────────────────────────────────────────────
def test_hr_endpoint_faculty_scope(admin_headers):
    r0 = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                      params={"period": "month"}, headers=admin_headers, timeout=60)
    assert r0.status_code == 200
    b0 = r0.json()
    unscoped_total = b0["hr"]["headcount"]["total_active"]
    fac = next((u for u in b0["units"] if u["type"] == "faculty"), None)
    assert fac is not None, "no faculty unit found"

    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": "month", "org_unit_id": fac["id"]},
                     headers=admin_headers, timeout=60)
    assert r.status_code == 200
    body = r.json()
    hr = body["hr"]
    assert hr["scope"]["org_unit_id"] == fac["id"]
    assert hr["scope"]["label"] == fac["name"]
    assert hr["employees_active"] == hr["headcount"]["total_active"]
    assert hr["headcount"]["total_active"] <= unscoped_total


def test_hr_endpoint_department_scope(admin_headers):
    r0 = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                      params={"period": "month"}, headers=admin_headers, timeout=60)
    b0 = r0.json()
    dept = next((u for u in b0["units"] if u["type"] == "department"), None)
    assert dept is not None, "no department unit found"
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": "month", "org_unit_id": dept["id"]},
                     headers=admin_headers, timeout=60)
    assert r.status_code == 200
    hr = r.json()["hr"]
    assert hr["scope"]["org_unit_id"] == dept["id"]
    assert hr["scope"]["label"] == dept["name"]


def test_hr_endpoint_invalid_org_unit_id_falls_back(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": "week", "org_unit_id": "does-not-exist-xyz"},
                     headers=admin_headers, timeout=60)
    assert r.status_code == 200
    hr = r.json()["hr"]
    assert hr["scope"]["org_unit_id"] is None


# ── Academic endpoint no longer contains hr block ────────────────
def test_academic_endpoint_no_hr_block(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management",
                     params={"period": "week"}, headers=admin_headers, timeout=60)
    assert r.status_code == 200
    body = r.json()
    assert "hr" not in body or body.get("hr") in (None, {}), f"academic response still contains hr: {list(body.keys())}"
    assert body["sections"]["hr"] is True  # capability flag still true for admin
    # regression: academic parts present
    for k in ("numbers", "alerts", "chart", "teachers", "students", "rooms", "finance"):
        assert k in body, f"missing academic key {k}"


# ── RBAC on new endpoint ─────────────────────────────────────────
def test_dean_hr_endpoint_forbidden(dean_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": "week"}, headers=dean_headers, timeout=30)
    assert r.status_code == 403, f"expected 403 for dean, got {r.status_code}"
    # academic dashboard: sections.hr false
    r2 = requests.get(f"{BASE_URL}/api/dashboard/management",
                      params={"period": "week"}, headers=dean_headers, timeout=30)
    assert r2.status_code == 200
    assert r2.json()["sections"]["hr"] is False


def test_head_hr_endpoint_forbidden(head_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                     params={"period": "week"}, headers=head_headers, timeout=30)
    assert r.status_code == 403
    r2 = requests.get(f"{BASE_URL}/api/dashboard/management",
                      params={"period": "week"}, headers=head_headers, timeout=30)
    assert r2.json()["sections"]["hr"] is False


# ── HR export ────────────────────────────────────────────────────
def test_hr_export_excel(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr/export",
                     params={"fmt": "excel", "period": "month"}, headers=admin_headers, timeout=120)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    assert any("الملخص" in s for s in names), f"missing summary sheet: {names}"
    assert any("الموظفون حسب الفئة" in s for s in names), f"missing category sheet: {names}"
    assert any("موظفو الوحدات" in s for s in names), f"missing units sheet: {names}"
    assert any(s.startswith("الدوام الإداري") for s in names), f"missing dawaam sheet: {names}"


def test_hr_export_excel_with_org_unit(admin_headers):
    r0 = requests.get(f"{BASE_URL}/api/dashboard/management/hr",
                      params={"period": "month"}, headers=admin_headers, timeout=60)
    fac = next((u for u in r0.json()["units"] if u["type"] == "faculty"), None)
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr/export",
                     params={"fmt": "excel", "period": "month", "org_unit_id": fac["id"]},
                     headers=admin_headers, timeout=120)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "")


def test_hr_export_pdf(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/hr/export",
                     params={"fmt": "pdf", "period": "month"}, headers=admin_headers, timeout=120)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    assert len(r.content) > 500


def test_academic_export_no_hr_sheets(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/export",
                     params={"fmt": "excel", "period": "month"}, headers=admin_headers, timeout=120)
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    assert not any("موظفو الوحدات" in s for s in names), f"academic export must NOT contain HR sheets: {names}"
