"""HR dashboard section tests (iteration 78)."""
import os
import io
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")


def _login(username: str, password: str) -> str:
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": username, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed for {username}: {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_headers():
    return {"Authorization": f"Bearer {_login('admin', 'admin123')}"}


@pytest.fixture(scope="module")
def dean_headers():
    return {"Authorization": f"Bearer {_login('Salim', 'test1234')}"}


@pytest.fixture(scope="module")
def head_headers():
    return {"Authorization": f"Bearer {_login('Saeed', 'test1234')}"}


# ── Permissions catalog ───────────────────────────────────────────
def test_permissions_include_dashboard_hr(admin_headers):
    r = requests.get(f"{BASE_URL}/api/permissions/all", headers=admin_headers, timeout=30)
    assert r.status_code == 200
    data = r.json()
    perms = data if isinstance(data, list) else data.get("permissions", [])
    match = [p for p in perms if p.get("key") == "dashboard_hr"]
    assert match, "dashboard_hr not found in /api/permissions/all"
    assert match[0].get("category") == "لوحة القيادة"


# ── HR section structure + sums for each period ───────────────────
@pytest.mark.parametrize("period", ["week", "month", "day"])
def test_admin_dashboard_hr_structure(admin_headers, period):
    r = requests.get(f"{BASE_URL}/api/dashboard/management", params={"period": period}, headers=admin_headers, timeout=90)
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["sections"]["hr"] is True, "sections.hr must be true for admin"
    hr = body["hr"]
    assert hr is not None, "hr block missing for admin"

    # today fields
    for k in ("date", "employees_active", "present_today"):
        assert k in hr, f"missing hr.{k}"

    # headcount
    hc = hr["headcount"]
    assert hc is not None
    for k in ("total_active", "total_all", "by_category", "by_contract", "by_status",
              "by_unit", "by_unit_type", "units_count", "units_with_staff"):
        assert k in hc, f"headcount missing {k}"
    assert len(hc["by_category"]) == 4
    cat_keys = {c["key"] for c in hc["by_category"]}
    assert cat_keys == {"academic", "administrative", "technical", "service"}
    # sum(by_category.count) == total_active
    assert sum(c["count"] for c in hc["by_category"]) == hc["total_active"], "by_category sum != total_active"
    # by_unit sums
    for row in hc["by_unit"]:
        for k in ("name", "type_label", "total", "academic", "administrative", "other", "sub_units"):
            assert k in row, f"by_unit row missing {k}"
        assert row["academic"] + row["administrative"] + row["other"] == row["total"], f"unit row sum mismatch: {row}"
    assert sum(u["total"] for u in hc["by_unit"]) == hc["total_active"], "by_unit total sum != total_active"

    # period block
    p = hr["period"]
    assert p is not None
    for k in ("work_days", "expected", "attended", "present", "late", "absent", "excused", "unmarked",
              "late_minutes", "late_hours", "commitment_rate", "leaves_approved", "leave_days",
              "leave_types", "tasks_done", "tasks_created", "chart", "top_employees",
              "bottom_employees", "units_attendance"):
        assert k in p, f"period missing {k}"
    expected_group = "unit" if period == "day" else "date"
    assert p["chart"]["group_by"] == expected_group, f"chart.group_by expected {expected_group} for period={period}"
    assert isinstance(p["chart"]["points"], list)


# ── RBAC ──────────────────────────────────────────────────────────
def test_dean_no_hr_section(dean_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management", params={"period": "week"}, headers=dean_headers, timeout=60)
    assert r.status_code == 200
    body = r.json()
    # Dean not admin/president; dashboard_hr not in default DEAN perms → hr should be False
    assert body["sections"]["hr"] is False, "Dean should not have HR section by default"
    assert body["hr"] is None


def test_department_head_no_hr_section(head_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management", params={"period": "week"}, headers=head_headers, timeout=60)
    assert r.status_code == 200
    body = r.json()
    assert body["sections"]["hr"] is False, "Dept head should not have HR section by default"
    assert body["hr"] is None


# ── Export ────────────────────────────────────────────────────────
def test_export_excel_contains_hr_sheets(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/export",
                     params={"fmt": "excel", "period": "month"}, headers=admin_headers, timeout=120)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    sheet_names = wb.sheetnames
    assert any("الموظفون حسب الفئة" in s for s in sheet_names), f"missing category sheet: {sheet_names}"
    assert any("موظفو الوحدات" in s for s in sheet_names), f"missing units sheet: {sheet_names}"
    assert any(s.startswith("الدوام الإداري") for s in sheet_names), f"missing dawaam sheet: {sheet_names}"


def test_export_pdf(admin_headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/management/export",
                     params={"fmt": "pdf", "period": "month"}, headers=admin_headers, timeout=120)
    assert r.status_code == 200
    assert r.headers.get("content-type", "").startswith("application/pdf")
    assert len(r.content) > 1000
    assert r.content[:4] == b"%PDF"
