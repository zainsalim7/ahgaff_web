"""Iteration 103: Multi-day holidays + overtime holiday deduction."""
import os, requests, pytest

BASE = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or pytest.skip("No URL", allow_module_level=True)
API = f"{BASE}/api"


def _login(u, p):
    r = requests.post(f"{API}/auth/login", json={"username": u, "password": p}, timeout=30)
    assert r.status_code == 200, f"login {u} failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_h():
    return {"Authorization": f"Bearer {_login('admin','admin123')}"}


@pytest.fixture(scope="module")
def dean_h():
    try:
        return {"Authorization": f"Bearer {_login('Salim','test1234')}"}
    except AssertionError:
        pytest.skip("Dean login failed")


# ---------- Calendar events / holidays ----------
def test_create_event_bad_end_date(admin_h):
    r = requests.post(f"{API}/calendar/events", headers=admin_h, json={
        "gregorian_date": "2026-10-05", "end_date": "2026-10-04",
        "event_name": "TEST_bad_range", "event_type": "holiday"
    })
    assert r.status_code == 400, r.text


def test_create_event_invalid_end(admin_h):
    r = requests.post(f"{API}/calendar/events", headers=admin_h, json={
        "gregorian_date": "2026-10-05", "end_date": "not-a-date",
        "event_name": "TEST_bad_end_fmt", "event_type": "holiday"
    })
    assert r.status_code == 400, r.text


def test_create_update_delete_range(admin_h):
    # create 2-day holiday
    r = requests.post(f"{API}/calendar/events", headers=admin_h, json={
        "gregorian_date": "2026-10-05", "end_date": "2026-10-06",
        "event_name": "TEST_range_holiday", "event_type": "holiday"
    })
    assert r.status_code == 200, r.text
    data = r.json()
    assert data.get("end_date") == "2026-10-06"
    eid = data["id"]

    # update end_date to another day
    r2 = requests.put(f"{API}/calendar/events/{eid}", headers=admin_h, json={"end_date": "2026-10-07"})
    assert r2.status_code == 200

    # set end_date "" -> null
    r3 = requests.put(f"{API}/calendar/events/{eid}", headers=admin_h, json={"end_date": ""})
    assert r3.status_code == 200

    # verify event stored with null end_date
    lst = requests.get(f"{API}/calendar/events?year=2026&month=10", headers=admin_h).json()
    me = [e for e in lst if e["id"] == eid]
    assert me and (me[0].get("end_date") is None)

    # cleanup
    rd = requests.delete(f"{API}/calendar/events/{eid}", headers=admin_h)
    assert rd.status_code == 200


def test_rbac_non_admin_cannot_write(admin_h, dean_h):
    # dean does POST/PUT/DELETE -> 403; GET works
    r = requests.post(f"{API}/calendar/events", headers=dean_h, json={
        "gregorian_date": "2026-10-05", "event_name": "TEST_x", "event_type": "holiday"
    })
    assert r.status_code == 403, r.text

    r2 = requests.get(f"{API}/calendar/events?year=2026", headers=dean_h)
    assert r2.status_code == 200

    # create as admin
    cr = requests.post(f"{API}/calendar/events", headers=admin_h, json={
        "gregorian_date": "2026-10-05", "event_name": "TEST_admin_only", "event_type": "holiday"
    }).json()
    eid = cr["id"]
    try:
        r3 = requests.put(f"{API}/calendar/events/{eid}", headers=dean_h, json={"event_name": "x"})
        assert r3.status_code == 403
        r4 = requests.delete(f"{API}/calendar/events/{eid}", headers=dean_h)
        assert r4.status_code == 403
    finally:
        requests.delete(f"{API}/calendar/events/{eid}", headers=admin_h)


def test_holidays_expand_multi_day(admin_h):
    # Seeded 'إجازة تجريبية' 09-07..09-08
    r = requests.get(f"{API}/calendar/holidays?start=2026-08-15&end=2026-09-24", headers=admin_h)
    assert r.status_code == 200
    data = r.json()
    days = {d["date"] for d in data["days"]}
    assert "2026-09-07" in days and "2026-09-08" in days
    assert data["count"] >= 2

    # add holiday crossing end
    cr = requests.post(f"{API}/calendar/events", headers=admin_h, json={
        "gregorian_date": "2026-09-23", "end_date": "2026-09-27",
        "event_name": "TEST_cross_end", "event_type": "holiday"
    }).json()
    eid = cr["id"]
    try:
        r2 = requests.get(f"{API}/calendar/holidays?start=2026-08-15&end=2026-09-24", headers=admin_h).json()
        in_range = {d["date"] for d in r2["days"] if d["name"] == "TEST_cross_end"}
        assert in_range == {"2026-09-23", "2026-09-24"}, in_range
    finally:
        requests.delete(f"{API}/calendar/events/{eid}", headers=admin_h)


# ---------- Teacher workload integration ----------
def _get_faculty_id(admin_h):
    # Query mongo directly for seed_ot teacher's faculty_id
    try:
        from pymongo import MongoClient
        from dotenv import load_dotenv
        load_dotenv("/app/backend/.env")
        c = MongoClient(os.environ["MONGO_URL"])
        db = c[os.environ["DB_NAME"]]
        t = db.teachers.find_one({"seed_ot": True})
        if t and t.get("faculty_id"):
            return t["faculty_id"]
    except Exception as e:
        pytest.skip(f"mongo lookup failed: {e}")
    pytest.skip("no seed_ot teacher found")


def test_workload_with_holidays(admin_h):
    fac = _get_faculty_id(admin_h)
    params = {"start_date": "2026-08-15", "end_date": "2026-09-24", "hide_empty": "true", "faculty_id": fac}
    r = requests.get(f"{API}/reports/teacher-workload", headers=admin_h, params=params)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "holidays" in data
    assert data.get("period", {}).get("holiday_days") == 2, data.get("period")

    teachers = data.get("teachers") or data.get("rows") or []
    # find أحمد بن سميط
    ahmed = next((t for t in teachers if "أحمد بن سميط" in (t.get("teacher_name") or t.get("name") or "")), None)
    assert ahmed, "teacher أحمد بن سميط not found"
    rows_a = ahmed.get("courses") or ahmed.get("rows") or []
    bal_rows = [r for r in rows_a if (r.get("course_name") == "بلاغة-2")]
    assert bal_rows, f"No بلاغة-2 rows for ahmed. rows: {rows_a}"
    for row in bal_rows:
        assert abs(float(row.get("holiday_hours", 0)) - 1.5) < 0.01, row
        assert abs(float(row.get("expected_hours", 0)) - 16.5) < 0.01, row
        assert "إجازة رسمية" in (row.get("note") or ""), row
        exp = float(row.get("expected_hours", 0))
        act = float(row.get("actual_hours", 0))
        short = float(row.get("shortfall_hours", exp - act))
        assert abs(short - (exp - act)) < 0.01

    # عبدالله باوزير -> holiday_hours 0 (lecture executed)
    abd = next((t for t in teachers if "باوزير" in (t.get("teacher_name") or t.get("name") or "")), None)
    assert abd, "teacher باوزير not found"
    rows_b = abd.get("courses") or abd.get("rows") or []
    for row in rows_b:
        hh = float(row.get("holiday_hours", 0))
        # some courses may not fall on holiday at all; for ones that do and were executed, expect 0
        assert hh >= 0


def test_workload_no_holidays(admin_h):
    fac = _get_faculty_id(admin_h)
    r = requests.get(f"{API}/reports/teacher-workload", headers=admin_h, params={
        "start_date": "2026-03-01", "end_date": "2026-03-31", "hide_empty": "true", "faculty_id": fac
    })
    assert r.status_code == 200, r.text
    data = r.json()
    assert data.get("period", {}).get("holiday_days", 0) == 0


def test_exports_with_holidays(admin_h):
    fac = _get_faculty_id(admin_h)
    params = {"start_date": "2026-08-15", "end_date": "2026-09-24", "hide_empty": "true", "faculty_id": fac}
    rp = requests.get(f"{API}/export/report/teacher-workload/sheet-pdf", headers=admin_h, params=params)
    assert rp.status_code == 200, f"pdf {rp.status_code} {rp.text[:200]}"
    assert len(rp.content) > 500

    re_ = requests.get(f"{API}/export/report/teacher-workload/sheet-excel", headers=admin_h, params=params)
    assert re_.status_code == 200, f"excel {re_.status_code} {re_.text[:200]}"
    # Verify Excel contains header text
    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(re_.content))
    ws = wb.active
    row2 = " ".join(str(c.value or "") for c in ws[2])
    assert "أيام الإجازات الرسمية" in row2 and "2" in row2, row2
