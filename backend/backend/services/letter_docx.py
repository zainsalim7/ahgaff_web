"""📝 تصدير الخطاب إلى Word (DOCX) — نفس أقسام PDF بتنسيق قابل للتحرير"""
import base64
import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .letter_pdf import DEFAULT_LAYOUT, filter_table
from .rich_text_pdf import FONTS, is_html, parse_rich

_ALIGN = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def _rtl(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); pPr.append(b)
    return p


def _run(p, text, size=13, bold=False, underline=False, color=None, font="amiri", latin=False):
    r = p.add_run(text)
    fam = "Arial" if latin else FONTS.get(font, FONTS["amiri"])[3]
    r.font.name = fam; r.font.size = Pt(size); r.font.bold = bold; r.font.underline = underline; r.font.rtl = not latin
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), fam)
    if color:
        c = str(color).strip()
        try:
            if c.startswith("#") and len(c) in (4, 7):
                c = c[1:]; c = "".join(ch * 2 for ch in c) if len(c) == 3 else c
                r.font.color.rgb = RGBColor.from_string(c.upper())
            elif c.startswith("rgb"):
                nums = [int(float(x)) for x in c[c.index("(") + 1:c.index(")")].split(",")[:3]]
                r.font.color.rgb = RGBColor(*nums)
        except Exception:
            pass
    return r


def _para(doc, align="right", space_after=4, leading=1.5):
    p = _rtl(doc.add_paragraph()); p.alignment = _ALIGN.get(align, WD_ALIGN_PARAGRAPH.RIGHT)
    pf = p.paragraph_format; pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = leading
    return p


def _rich(doc, html, default_size=13, default_align="right", leading=1.5, first_indent_mm=0):
    for i, para in enumerate(parse_rich(html, default_size=default_size, default_font="amiri", default_align=default_align)):
        p = _para(doc, para.get("align") or default_align, leading=leading)
        if first_indent_mm and (para.get("align") or default_align) != "center":
            p.paragraph_format.first_line_indent = Mm(first_indent_mm)
        for r in para["runs"]:
            _run(p, r["text"], size=r.get("size") or default_size, bold=r.get("bold", False), underline=r.get("underline", False), color=r.get("color"), font=r.get("font") or "amiri")


def _plain(doc, text, size=13, align="right", bold=False, underline=False, leading=1.5):
    p = _para(doc, align, leading=leading); _run(p, text, size=size, bold=bold, underline=underline); return p


def _hr(p, size=12, bottom=True):
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom" if bottom else "w:top")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:space"), "1"); el.set(qn("w:color"), "000000")
    bdr.append(el); pPr.append(bdr)


def _img_stream(b64: str):
    if not b64:
        return None
    try:
        raw = b64.split(",", 1)[1] if "," in b64 else b64
        return io.BytesIO(base64.b64decode(raw))
    except Exception:
        return None


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_color); tcPr.append(sh)


def _table_rtl(tbl):
    tblPr = tbl._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); tblPr.append(bv)


def build_letter_docx(s: dict, settings: dict, letterhead: bool = True) -> bytes:
    L = {**DEFAULT_LAYOUT, **(settings.get("layout") or {}), **(s.get("layout") or {})}
    f = lambda k: float(L.get(k) if L.get(k) not in (None, "") else DEFAULT_LAYOUT.get(k, 0))
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.left_margin = sec.right_margin = Mm(f("margin_side")); sec.top_margin = Mm(12); sec.bottom_margin = Mm(18)
    st = doc.styles["Normal"]; st.font.name = "Amiri"; st.font.size = Pt(f("body_font"))

    if letterhead:
        logo = _img_stream(settings.get("logo_base64") or "")
        if logo is None:
            d = Path(__file__).parent.parent / "assets" / "university_logo.jpeg"
            logo = str(d) if d.exists() else None
        if logo:
            p = _para(doc, "center", space_after=2); p.add_run().add_picture(logo, height=Mm(22))
        ht = doc.add_table(rows=1, cols=2); ht.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = ht.rows[0].cells
        pr = right.paragraphs[0]; _rtl(pr); pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(pr, settings.get("org_name") or "جامعة الأحقاف", size=f("header_font_ar1"), bold=True)
        pr2 = right.add_paragraph(); _rtl(pr2); pr2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _run(pr2, settings.get("office_name") or "", size=f("header_font_ar2"))
        pl = left.paragraphs[0]; pl.alignment = WD_ALIGN_PARAGRAPH.LEFT; _run(pl, settings.get("org_name_en") or "AL-AHGAFF UNIVERSITY", size=f("header_font_en1"), bold=True, latin=True)
        pl2 = left.add_paragraph(); pl2.alignment = WD_ALIGN_PARAGRAPH.LEFT; _run(pl2, settings.get("office_name_en") or "", size=f("header_font_en2"), latin=True)
        hr = _para(doc, "right", space_after=6); _hr(hr, 14)

    issued = (s.get("issued_at") or "")[:10]
    hijri = ""
    try:
        from hijridate import Gregorian
        yy, m, d = [int(x) for x in issued.split("-")]; hj = Gregorian(yy, m, d).to_hijri(); hijri = f"{hj.year}/{hj.month:02d}/{hj.day:02d}هـ"
    except Exception:
        pass
    rt = doc.add_table(rows=1, cols=2); rt.alignment = WD_TABLE_ALIGNMENT.CENTER
    cl, cr = rt.rows[0].cells
    pn = cl.paragraphs[0]; pn.alignment = WD_ALIGN_PARAGRAPH.LEFT; _rtl(pn); _run(pn, f"الرقم : {s.get('number_display', '')}", size=f("ref_font"), bold=True, color="#008022")
    pd = cr.paragraphs[0]; pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl(pd); _run(pd, f"التاريخ: {hijri}", size=f("ref_font") - 1.5, bold=True, color="#008022")
    pd2 = cr.add_paragraph(); pd2.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl(pd2); _run(pd2, f"الموافق: {issued.replace('-', '/')}م", size=f("ref_font") - 1.5, bold=True, color="#008022")
    _para(doc, space_after=6)

    secs = {k: v for k, v in (s.get("sections") or {}).items() if v and is_html(v) and any(r["text"].strip() for p in parse_rich(v) for r in p["runs"])}
    rec = s.get("recipient") or {}
    r_align = str(L.get("recipient_align") or "right")
    if "recipient" in secs:
        _rich(doc, secs["recipient"], default_size=f("recipient_font"), default_align=r_align)
    else:
        _plain(doc, f"إلى: {rec.get('name') or rec.get('title') or ''}   {rec.get('suffix') or ''}".rstrip(), size=f("recipient_font"), align=r_align, bold=bool(L.get("recipient_bold", True)))
        if rec.get("name") and rec.get("title"):
            _plain(doc, rec["title"], size=f("recipient_sub_font"), align=r_align)
        if rec.get("organization") and rec.get("organization") != rec.get("title"):
            _plain(doc, rec["organization"], size=f("recipient_sub_font"), align=r_align)
    if "greeting" in secs:
        _rich(doc, secs["greeting"], default_size=f("greeting_font"), default_align=str(L.get("greeting_align") or "right"))
    elif L.get("show_greeting"):
        _plain(doc, "السلام عليكم ورحمة الله وبركاته،", size=f("greeting_font"), align=str(L.get("greeting_align") or "right"))
    if "subject" in secs:
        _rich(doc, secs["subject"], default_size=14, default_align="center")
    else:
        _plain(doc, f"الموضوع: {s.get('subject', '')}", size=14, align="center", bold=True, underline=True)

    def table(tbl):
        tbl = filter_table(tbl, L.get("table_columns"))
        headers, rows = tbl.get("headers") or [], tbl.get("rows") or []
        if not headers or not rows:
            return
        t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; _table_rtl(t)
        tf = f("table_font"); hb = str(L.get("table_header_bg") or "#EDF2FA").lstrip("#")
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; _shade(c, hb if len(hb) == 6 else "EDF2FA")
            p = c.paragraphs[0]; _rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _run(p, str(h), size=tf, bold=True, color=L.get("table_header_color"))
        for r in rows:
            cells = t.add_row().cells
            for i in range(len(headers)):
                p = cells[i].paragraphs[0]; _rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _run(p, str(r[i]) if i < len(r) else "", size=tf, color=L.get("table_text_color"))
        _para(doc, space_after=4)

    body = s.get("body") or ""
    parts = body.split("{جدول_الأسماء}")
    has_table = bool((s.get("table") or {}).get("rows"))
    for pi, part in enumerate(parts):
        if is_html(part):
            _rich(doc, part, default_size=f("body_font"), default_align="right", leading=f("body_leading"), first_indent_mm=f("body_indent"))
        else:
            for line in part.strip("\n").split("\n"):
                if line.strip():
                    _plain(doc, line, size=f("body_font"), leading=f("body_leading"))
        if pi < len(parts) - 1 and has_table:
            table(s["table"])
    if "{جدول_الأسماء}" not in body and has_table and s.get("auto_table", True):
        table(s["table"])

    if "closing" in secs:
        _rich(doc, secs["closing"], default_size=f("body_font"))
    elif L.get("show_closing"):
        _plain(doc, settings.get("closing") or "وتفضلوا بقبول فائق الاحترام والتقدير،", size=f("body_font"))
    _para(doc, space_after=10)

    s_align = str(L.get("signature_align") or "left")
    if "signature" in secs:
        _rich(doc, secs["signature"], default_size=f("signature_title_font"), default_align=s_align if s_align != "left" else "left")
    else:
        _plain(doc, (s.get("signatory_title") or "").strip(), size=f("signature_title_font"), align=s_align, bold=bool(L.get("signature_title_bold", True)))
        _plain(doc, (s.get("signatory_name") or "").strip(), size=f("signature_name_font"), align=s_align, bold=bool(L.get("signature_name_bold", False)))
    sig = _img_stream(settings.get("signature_base64") or "")
    if sig:
        p = _para(doc, s_align); p.add_run().add_picture(sig, width=Mm(38))

    if s.get("verify_url") and s.get("verify_url") != "DRAFT-PREVIEW":
        p = _para(doc, "left", space_after=0); _run(p, f"للتحقق: {s['verify_url']}", size=8, latin=True, color="#64748b")
    if letterhead:
        fp = _para(doc, "center", space_after=0); _hr(fp, 6, bottom=False)
        parts_ = [settings.get("address")] + [f"{l}: {settings[k]}" for k, l in (("phones", "تلفون"), ("fax", "فاكس")) if settings.get(k)] + ([f"ص.ب ({settings['po_box']})"] if settings.get("po_box") else []) + [settings.get("website")]
        _run(fp, " — ".join(x for x in parts_ if x), size=9)
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()
