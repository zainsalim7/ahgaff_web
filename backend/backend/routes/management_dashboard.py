"""📊 لوحة القيادة للإدارة العليا (رئيس الجامعة / العميد / رئيس القسم) — تجميع بنداء واحد مع احترام النطاق"""
import io
import os
import asyncio
import logging
from datetime import datetime, timedelta, timezone
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from bson import ObjectId

from .deps import get_db, get_current_user, get_scope_filter, has_permission, export_filename, export_headers
from models.permissions import UserRole, READ_ONLY_ROLES, Permission

router = APIRouter()

YEMEN_TZ = timezone(timedelta(hours=3))
LATE_TEACHER_MINUTES = 15
MISSED_LECTURE_GRACE_MINUTES = 30
LOW_ATTENDANCE_THRESHOLD = 75
LOW_ATTENDANCE_MIN_LECTURES = 3
PERIOD_LABELS = {"day": "اليوم", "week": "آخر 7 أيام", "month": "آخر 30 يوماً"}
STATUS_AR = {"completed": "منفَّذة", "scheduled": "مجدولة", "cancelled": "ملغاة", "absent": "غياب الأستاذ"}


def _now_yemen() -> datetime:
    return datetime.now(YEMEN_TZ).replace(tzinfo=None)


def _period_range(period: str):
    today = _now_yemen().date()
    if period == "day":
        return today, today
    if period == "week":
        return today - timedelta(days=6), today
    return today - timedelta(days=29), today


def _is_management(user: dict) -> bool:
    return user.get("role") not in (UserRole.TEACHER, "teacher", "student")


SECTION_PERMS = {"alerts": Permission.DASHBOARD_ALERTS, "attendance": Permission.DASHBOARD_ATTENDANCE, "teachers": Permission.DASHBOARD_TEACHERS,
                 "students": Permission.DASHBOARD_STUDENTS, "rooms": Permission.DASHBOARD_ROOMS, "finance": Permission.DASHBOARD_FINANCE,
                 "hr": Permission.DASHBOARD_HR, "export": Permission.DASHBOARD_EXPORT}


def dashboard_sections(user: dict) -> dict:
    """الأجزاء المسموحة: الأدمن ورئيس الجامعة كل شيء؛ غيرهم حسب صلاحيات فئة «لوحة القيادة» (الأرقام العامة للجميع)"""
    full = user.get("role") in (UserRole.ADMIN, UserRole.UNIVERSITY_PRESIDENT)
    return {k: (full or has_permission(user, p)) for k, p in SECTION_PERMS.items()}


async def _resolve_scope(db, user: dict, faculty_id: Optional[str], department_id: Optional[str]) -> dict:
    """النطاق المسموح (كل الأقسام للأدمن) ثم تضييقه بالفلتر اليدوي ضمن المسموح فقط"""
    is_admin = user.get("role") in (UserRole.ADMIN, UserRole.UNIVERSITY_PRESIDENT)  # نطاق الجامعة كلها
    dq = {} if is_admin else await get_scope_filter(user, "departments")
    depts = await db.departments.find(dq, {"name": 1, "faculty_id": 1}).to_list(500)
    fac_ids = sorted({d.get("faculty_id") for d in depts if d.get("faculty_id")})
    faculties = []
    if fac_ids:
        oids = [ObjectId(f) for f in fac_ids if ObjectId.is_valid(f)]
        faculties = [{"id": str(f["_id"]), "name": f.get("name", "")}
                     for f in await db.faculties.find({"_id": {"$in": oids}}, {"name": 1}).to_list(100)]
    allowed = [{"id": str(d["_id"]), "name": d.get("name", ""), "faculty_id": d.get("faculty_id")} for d in depts]
    selected = allowed
    if is_admin:
        label = "كل الجامعة"
    elif len(allowed) == 1:
        label = allowed[0]["name"].strip()
    else:
        label = "كل الكليات" if len(faculties) > 1 else (faculties[0]["name"].strip() if faculties else "نطاقي")
    if faculty_id:
        if faculty_id not in fac_ids:
            raise HTTPException(status_code=403, detail="هذه الكلية خارج نطاق صلاحيتك")
        selected = [d for d in allowed if d["faculty_id"] == faculty_id]
        label = next((f["name"].strip() for f in faculties if f["id"] == faculty_id), label)
    if department_id:
        sel = [d for d in selected if d["id"] == department_id]
        if not sel:
            raise HTTPException(status_code=403, detail="هذا القسم خارج نطاق صلاحيتك")
        selected = sel
        label = sel[0]["name"].strip()
    dept_ids = None if (is_admin and not faculty_id and not department_id) else [d["id"] for d in selected]
    return {"is_admin": is_admin, "faculties": faculties, "departments": allowed, "selected": selected,
            "dept_ids": dept_ids, "faculty_ids": fac_ids, "label": label,
            "can_filter": is_admin or len(allowed) > 1}


def _dept_q(dept_ids):
    return {} if dept_ids is None else {"department_id": {"$in": dept_ids}}


async def _scope_courses(db, dept_ids):
    q = {} if dept_ids is None else {"$or": [{"department_id": {"$in": dept_ids}}, {"shared_links.department_id": {"$in": dept_ids}}]}
    return await db.courses.find(q, {"name": 1, "code": 1, "teacher_id": 1, "department_id": 1, "semester_id": 1, "is_active": 1}).to_list(20000)


def _rate(present, late, absent):
    tot = present + late + absent
    return round((present + late) * 100 / tot, 1) if tot else None


async def _teacher_names(db, ids):
    ids = [i for i in set(ids) if i and ObjectId.is_valid(i)]
    if not ids:
        return {}
    oids = [ObjectId(i) for i in ids]
    out = {str(t["_id"]): t.get("full_name", "") for t in await db.teachers.find({"_id": {"$in": oids}}, {"full_name": 1}).to_list(5000)}
    missing = [ObjectId(i) for i in ids if i not in out]
    if missing:
        for u in await db.users.find({"_id": {"$in": missing}}, {"full_name": 1}).to_list(5000):
            out[str(u["_id"])] = u.get("full_name", "")
    return out


async def build_dashboard(db, user: dict, period: str, faculty_id: Optional[str], department_id: Optional[str]) -> dict:
    period = period if period in PERIOD_LABELS else "week"
    sections = dashboard_sections(user)
    scope = await _resolve_scope(db, user, faculty_id, department_id)
    dept_ids = scope["dept_ids"]
    d_from, d_to = _period_range(period)
    s_from, s_to = d_from.isoformat(), d_to.isoformat()
    now = _now_yemen()
    today = now.date().isoformat()

    active_sem = await db.semesters.find_one({"$or": [{"status": "active"}, {"is_active": True}]})
    active_sem_id = str(active_sem["_id"]) if active_sem else None

    courses = await _scope_courses(db, dept_ids)
    course_map = {str(c["_id"]): c for c in courses}
    course_ids = list(course_map.keys())
    active_courses = [c for c in courses if c.get("is_active", True) and (not active_sem_id or c.get("semester_id") == active_sem_id)]

    stu_q = {**_dept_q(dept_ids), "is_active": True, "is_alumni": {"$ne": True}, "status": {"$nin": ["graduated"]}}
    tch_q = {**_dept_q(dept_ids), "is_active": {"$ne": False}}
    lec_q = {"date": {"$gte": s_from, "$lte": s_to}}
    today_q = {"date": today}
    if dept_ids is not None:
        lec_q["course_id"] = {"$in": course_ids}
        today_q["course_id"] = {"$in": course_ids}
    lec_proj = {"course_id": 1, "date": 1, "start_time": 1, "end_time": 1, "status": 1, "attendance_started_at": 1, "teacher_id": 1, "room": 1}

    students_count, teachers_count, lectures, today_lectures = await asyncio.gather(
        db.students.count_documents(stu_q),
        db.teachers.count_documents(tch_q),
        db.lectures.find(lec_q, lec_proj).to_list(50000),
        db.lectures.find(today_q, lec_proj).to_list(5000),
    )

    # ── الحضور في الفترة
    completed_ids = [str(l["_id"]) for l in lectures if l.get("status") == "completed"]
    att_by_status, att_by_student, att_by_lecture = {}, {}, {}
    if completed_ids:
        pipe = [{"$match": {"lecture_id": {"$in": completed_ids}}},
                {"$group": {"_id": {"s": "$status", "st": "$student_id", "l": "$lecture_id"}, "n": {"$sum": 1}}}]
        async for r in db.attendance.aggregate(pipe):
            st, sid, lid, n = r["_id"].get("s"), r["_id"].get("st"), r["_id"].get("l"), r["n"]
            att_by_status[st] = att_by_status.get(st, 0) + n
            att_by_student.setdefault(sid, {}).setdefault(st, 0)
            att_by_student[sid][st] += n
            att_by_lecture.setdefault(lid, {}).setdefault(st, 0)
            att_by_lecture[lid][st] += n
    present, late, absent = att_by_status.get("present", 0), att_by_status.get("late", 0), att_by_status.get("absent", 0)
    excused = att_by_status.get("excused", 0)

    lec_status = {k: 0 for k in STATUS_AR}
    for l in lectures:
        lec_status[l.get("status", "scheduled")] = lec_status.get(l.get("status", "scheduled"), 0) + 1

    # ── المخطط
    chart_points = []
    if period == "day":
        group_by = "department" if len(scope["selected"]) > 1 or dept_ids is None else "course"
        buckets = {}
        for l in lectures:
            c = course_map.get(l["course_id"])
            if not c:
                continue
            key = c.get("department_id") if group_by == "department" else str(c["_id"])
            b = buckets.setdefault(key, {"lectures": 0, "completed": 0, "present": 0, "late": 0, "absent": 0})
            b["lectures"] += 1
            if l.get("status") == "completed":
                b["completed"] += 1
                a = att_by_lecture.get(str(l["_id"]), {})
                b["present"] += a.get("present", 0); b["late"] += a.get("late", 0); b["absent"] += a.get("absent", 0)
        if group_by == "department":
            names = {d["id"]: d["name"] for d in scope["departments"]}
            if dept_ids is None:
                names = {str(d["_id"]): d.get("name", "") for d in await db.departments.find({}, {"name": 1}).to_list(500)}
        else:
            names = {k: (v.get("name") or v.get("code") or "") for k, v in course_map.items()}
        for k, b in buckets.items():
            chart_points.append({"label": names.get(k, "غير محدد"), **b, "rate": _rate(b["present"], b["late"], b["absent"])})
        chart_points.sort(key=lambda p: -(p["lectures"]))
        chart_points = chart_points[:12]
    else:
        group_by = "date"
        by_date = {}
        cur = d_from
        while cur <= d_to:
            by_date[cur.isoformat()] = {"lectures": 0, "completed": 0, "present": 0, "late": 0, "absent": 0}
            cur += timedelta(days=1)
        for l in lectures:
            b = by_date.get(l.get("date"))
            if b is None:
                continue
            b["lectures"] += 1
            if l.get("status") == "completed":
                b["completed"] += 1
                a = att_by_lecture.get(str(l["_id"]), {})
                b["present"] += a.get("present", 0); b["late"] += a.get("late", 0); b["absent"] += a.get("absent", 0)
        days_ar = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
        for k, b in by_date.items():
            dt = datetime.strptime(k, "%Y-%m-%d")
            lbl = days_ar[dt.weekday()] if period == "week" else f"{dt.day}/{dt.month}"
            chart_points.append({"label": lbl, "date": k, **b, "rate": _rate(b["present"], b["late"], b["absent"])})

    # ── التنبيهات
    low_students = []
    for sid, m in att_by_student.items():
        p, la, ab = m.get("present", 0), m.get("late", 0), m.get("absent", 0)
        if p + la + ab >= LOW_ATTENDANCE_MIN_LECTURES:
            r = _rate(p, la, ab)
            if r is not None and r < LOW_ATTENDANCE_THRESHOLD:
                low_students.append((sid, r, p + la + ab, ab))
    low_students.sort(key=lambda x: x[1])
    low_list = []
    if low_students:
        oids = [ObjectId(s[0]) for s in low_students[:10] if ObjectId.is_valid(s[0])]
        smap = {str(s["_id"]): s for s in await db.students.find({"_id": {"$in": oids}}, {"full_name": 1, "student_id": 1, "department_id": 1, "level": 1}).to_list(50)}
        dnames = {d["id"]: d["name"] for d in scope["departments"]}
        for sid, r, tot, ab in low_students[:10]:
            s = smap.get(sid, {})
            low_list.append({"id": sid, "name": s.get("full_name", "—"), "student_id": s.get("student_id", ""),
                             "department": dnames.get(s.get("department_id"), ""), "level": s.get("level", ""),
                             "rate": r, "lectures": tot, "absent": ab})

    late_teachers = {}
    for l in lectures:
        if l.get("status") != "completed" or not l.get("attendance_started_at"):
            continue
        try:
            start = datetime.strptime(f"{l['date']}T{l['start_time']}", "%Y-%m-%dT%H:%M")
            st = l["attendance_started_at"]
            st = datetime.fromisoformat(st) if isinstance(st, str) else st
            if st.tzinfo is not None:
                st = st.astimezone(YEMEN_TZ).replace(tzinfo=None)
            if st.date() != start.date():
                continue
            delay = int((st - start).total_seconds() // 60)
        except Exception:
            continue
        if delay >= LATE_TEACHER_MINUTES:
            c = course_map.get(l["course_id"], {})
            tid = l.get("teacher_id") or c.get("teacher_id")
            e = late_teachers.setdefault(tid, {"teacher_id": tid, "count": 0, "max_delay": 0, "total_delay": 0})
            e["count"] += 1; e["total_delay"] += delay; e["max_delay"] = max(e["max_delay"], delay)
    tnames = await _teacher_names(db, list(late_teachers.keys()))
    late_list = sorted([{**v, "name": tnames.get(k, "غير معروف")} for k, v in late_teachers.items()], key=lambda x: -x["count"])[:10]

    now_hm = now.strftime("%H:%M")
    missed_today = []
    for l in today_lectures:
        if l.get("status") != "scheduled":
            continue
        try:
            st = datetime.strptime(f"{today}T{l['start_time']}", "%Y-%m-%dT%H:%M")
        except Exception:
            continue
        if now >= st + timedelta(minutes=MISSED_LECTURE_GRACE_MINUTES):
            c = course_map.get(l["course_id"]) or {}
            missed_today.append({"lecture_id": str(l["_id"]), "course": c.get("name", "—"), "time": f"{l.get('start_time', '')}-{l.get('end_time', '')}",
                                 "room": l.get("room", ""), "teacher_id": l.get("teacher_id") or c.get("teacher_id")})
    if missed_today:
        mn = await _teacher_names(db, [m["teacher_id"] for m in missed_today])
        for m in missed_today:
            m["teacher"] = mn.get(m.pop("teacher_id"), "—")

    today_done = sum(1 for l in today_lectures if l.get("status") == "completed")
    today_cancelled = sum(1 for l in today_lectures if l.get("status") in ("cancelled", "absent"))
    today_upcoming = sum(1 for l in today_lectures if l.get("status") == "scheduled" and l.get("start_time", "") > now_hm)

    # ── المالية
    finance = None
    if sections["finance"] and (scope["is_admin"] or has_permission(user, "manage_fee_receipts")):
        from .fee_receipts import _allowed_type_ids
        finance = await _finance(db, dept_ids, active_sem, students_count, await _allowed_type_ids(db, user))

    # ── سجل النشاط
    act_q = {"action": {"$nin": ["view_page", "view_report"]}}
    if dept_ids is not None:
        ors = [{"department_id": {"$in": dept_ids}}]
        fids = {d["faculty_id"] for d in scope["selected"] if d.get("faculty_id")}
        if fids and not department_id:
            ors.append({"faculty_id": {"$in": list(fids)}, "department_id": {"$in": [None, ""]}})
        act_q["$or"] = ors
    activity = []
    async for a in db.activity_logs.find(act_q, {"username": 1, "action_ar": 1, "action": 1, "entity_name": 1, "entity_type": 1, "timestamp": 1, "user_role": 1}).sort("timestamp", -1).limit(15):
        ts = a.get("timestamp")
        if isinstance(ts, datetime):
            ts = (ts.replace(tzinfo=timezone.utc) if ts.tzinfo is None else ts).astimezone(YEMEN_TZ).strftime("%Y-%m-%d %H:%M")
        activity.append({"id": str(a["_id"]), "username": a.get("username", ""), "role": a.get("user_role", ""),
                         "action": a.get("action_ar") or a.get("action", ""), "entity": a.get("entity_name") or "", "time": ts or ""})

    cancelled_n = lec_status.get("cancelled", 0) + lec_status.get("absent", 0)
    cancel_pct = round(cancelled_n * 100 / len(lectures), 1) if lectures else 0
    alerts = [
        {"key": "low_attendance", "level": "danger" if low_students else "ok", "count": len(low_students),
         "title": f"طلاب حضورهم أقل من {LOW_ATTENDANCE_THRESHOLD}%", "hint": f"خلال {PERIOD_LABELS[period]} (≥{LOW_ATTENDANCE_MIN_LECTURES} محاضرات)", "items": low_list, "route": "/report-warnings"},
        {"key": "missed_today", "level": "danger" if missed_today else "ok", "count": len(missed_today),
         "title": "محاضرات اليوم لم يُسجَّل لها حضور", "hint": f"تجاوزت بدايتها {MISSED_LECTURE_GRACE_MINUTES} دقيقة", "items": missed_today, "route": "/schedule"},
        {"key": "late_teachers", "level": "warning" if late_list else "ok", "count": len(late_list),
         "title": "أساتذة تأخروا في بدء التحضير", "hint": f"تأخير ≥{LATE_TEACHER_MINUTES} دقيقة خلال {PERIOD_LABELS[period]}", "items": late_list, "route": "/report-teacher-delays"},
        {"key": "cancelled", "level": "warning" if cancel_pct >= 20 and cancelled_n else ("info" if cancelled_n else "ok"), "count": cancelled_n,
         "title": "محاضرات ملغاة / غياب أستاذ", "hint": f"{cancel_pct}% من محاضرات {PERIOD_LABELS[period]}", "items": [], "route": "/report-lesson-completion"},
    ]
    if finance is not None:
        pend = sum(t["pending"] for t in finance["types"])
        alerts.append({"key": "pending_fees", "level": "warning" if pend else "ok", "count": pend,
                       "title": "سندات مالية بانتظار التعميد", "hint": finance["academic_year"] or "", "items": [], "route": "/fee-receipts"})

    phase2 = await _phase2(db, scope, dept_ids, lectures, today_lectures, course_map, active_courses, att_by_lecture, att_by_student, stu_q, period)
    for k in ("teachers", "students", "rooms"):
        if not sections[k]:
            phase2[k] = None

    return {
        "generated_at": now.strftime("%Y-%m-%d %H:%M"),
        "sections": sections,
        **phase2,
        "period": period, "period_label": PERIOD_LABELS[period], "date_from": s_from, "date_to": s_to,
        "scope": {"label": scope["label"], "is_admin": scope["is_admin"], "can_filter": scope["can_filter"],
                  "read_only": user.get("role") in READ_ONLY_ROLES,
                  "faculties": scope["faculties"], "departments": scope["departments"],
                  "faculty_id": faculty_id, "department_id": department_id},
        "semester": {"name": active_sem.get("name", "") if active_sem else "", "academic_year": active_sem.get("academic_year", "") if active_sem else ""},
        "numbers": {
            "students": students_count, "teachers": teachers_count, "courses": len(active_courses),
            "departments": len(scope["selected"]) if dept_ids is not None else await db.departments.count_documents({}),
            "faculties": len({d["faculty_id"] for d in scope["selected"] if d.get("faculty_id")}) if dept_ids is not None else await db.faculties.count_documents({}),
            "lectures_today": len(today_lectures), "today_done": today_done, "today_cancelled": today_cancelled, "today_upcoming": today_upcoming,
            "lectures_period": len(lectures), "lectures_status": lec_status,
            "attendance_rate": _rate(present, late, absent), "present": present, "late": late, "absent": absent, "excused": excused,
        },
        "chart": {"group_by": group_by, "points": chart_points} if sections["attendance"] else None,
        "alerts": alerts if sections["alerts"] else [],
        "finance": finance,
        "activity": activity,
    }


STATUS_LABELS = {"active": "نشط", "suspended": "مجمّد", "frozen": "مجمّد", "repeat": "إعادة", "repeater": "إعادة",
                 "dismissed": "مفصول", "expelled": "مفصول", "graduated": "خريج", "withdrawn": "منسحب", "transferred": "منقول"}


def _delay_minutes(l: dict):
    try:
        start = datetime.strptime(f"{l['date']}T{l['start_time']}", "%Y-%m-%dT%H:%M")
        st = l.get("attendance_started_at")
        st = datetime.fromisoformat(st) if isinstance(st, str) else st
        if st is None:
            return None
        if st.tzinfo is not None:
            st = st.astimezone(YEMEN_TZ).replace(tzinfo=None)
        if st.date() != start.date():
            return None
        return max(0, int((st - start).total_seconds() // 60))
    except Exception:
        return None


def _lec_minutes(l: dict) -> int:
    try:
        a = datetime.strptime(l["start_time"], "%H:%M"); b = datetime.strptime(l["end_time"], "%H:%M")
        m = int((b - a).total_seconds() // 60)
        return m if 0 < m < 600 else 90
    except Exception:
        return 90


async def _phase2(db, scope, dept_ids, lectures, today_lectures, course_map, active_courses, att_by_lecture, att_by_student, stu_q, period) -> dict:
    """المرحلة 2: إحصائيات الأساتذة + الطلاب + القاعات/الجدول"""
    dnames = {d["id"]: d["name"].strip() for d in scope["departments"]}
    if dept_ids is None:
        dnames = {str(d["_id"]): (d.get("name") or "").strip() for d in await db.departments.find({}, {"name": 1}).to_list(500)}

    # ── 5) الأساتذة
    tstats = {}
    for l in lectures:
        c = course_map.get(l["course_id"], {})
        tid = l.get("teacher_id") or c.get("teacher_id")
        if not tid:
            continue
        e = tstats.setdefault(tid, {"teacher_id": tid, "scheduled": 0, "completed": 0, "cancelled": 0, "absent": 0, "upcoming": 0, "delays": [], "minutes": 0, "courses": set()})
        e["courses"].add(l["course_id"])
        st = l.get("status", "scheduled")
        e["scheduled"] += 1
        if st == "completed":
            e["completed"] += 1; e["minutes"] += _lec_minutes(l)
            d = _delay_minutes(l)
            if d is not None:
                e["delays"].append(d)
        elif st == "cancelled":
            e["cancelled"] += 1
        elif st == "absent":
            e["absent"] += 1
        else:
            e["upcoming"] += 1
    tnames = await _teacher_names(db, list(tstats.keys()))
    teachers_list = []
    for tid, e in tstats.items():
        due = e["completed"] + e["cancelled"] + e["absent"]
        teachers_list.append({
            "teacher_id": tid, "name": tnames.get(tid, "غير معروف"), "courses": len(e["courses"]),
            "scheduled": e["scheduled"], "completed": e["completed"], "cancelled": e["cancelled"], "absent": e["absent"], "upcoming": e["upcoming"],
            "commitment": round(e["completed"] * 100 / due, 1) if due else None,
            "avg_delay": round(sum(e["delays"]) / len(e["delays"]), 1) if e["delays"] else 0,
            "late_count": sum(1 for d in e["delays"] if d >= LATE_TEACHER_MINUTES),
            "hours": round(e["minutes"] / 60, 1),
        })
    teachers_list.sort(key=lambda x: (-(x["commitment"] or 0), -x["completed"]))
    with_due = [t for t in teachers_list if t["commitment"] is not None]
    teachers = {
        "count": len(teachers_list), "total_hours": round(sum(t["hours"] for t in teachers_list), 1),
        "avg_commitment": round(sum(t["commitment"] for t in with_due) / len(with_due), 1) if with_due else None,
        "avg_delay": round(sum(t["avg_delay"] * max(t["completed"], 1) for t in teachers_list) / max(sum(max(t["completed"], 1) for t in teachers_list), 1), 1) if teachers_list else 0,
        "top": with_due[:5], "bottom": sorted(with_due, key=lambda x: (x["commitment"], -x["absent"]))[:5],
        "rows": teachers_list[:50],
    }

    # ── 7) الطلاب
    by_dept, by_level, by_status = {}, {}, {}
    base_q = dict(stu_q); base_q.pop("status", None); base_q.pop("is_active", None); base_q.pop("is_alumni", None)
    async for r in db.students.aggregate([{"$match": base_q}, {"$group": {"_id": {"d": "$department_id", "l": "$level", "s": "$status", "a": "$is_active", "al": "$is_alumni"}, "n": {"$sum": 1}}}]):
        k, n = r["_id"], r["n"]
        st = k.get("s") or ("graduated" if k.get("al") else ("active" if k.get("a", True) else "inactive"))
        by_status[st] = by_status.get(st, 0) + n
        if st == "active" and k.get("a", True) is not False:
            by_dept[k.get("d")] = by_dept.get(k.get("d"), 0) + n
            by_level[str(k.get("l") or "?")] = by_level.get(str(k.get("l") or "?"), 0) + n
    dept_att = {}
    for l in lectures:
        if l.get("status") != "completed":
            continue
        dep = course_map.get(l["course_id"], {}).get("department_id")
        a = att_by_lecture.get(str(l["_id"]), {})
        b = dept_att.setdefault(dep, {"present": 0, "late": 0, "absent": 0})
        b["present"] += a.get("present", 0); b["late"] += a.get("late", 0); b["absent"] += a.get("absent", 0)
    dist = []
    for dep, n in sorted(by_dept.items(), key=lambda x: -x[1]):
        a = dept_att.get(dep, {"present": 0, "late": 0, "absent": 0})
        dist.append({"department_id": dep, "name": dnames.get(dep, "غير محدد"), "students": n, "rate": _rate(a["present"], a["late"], a["absent"])})
    warn = depr = 0
    for m in att_by_student.values():
        p, la, ab = m.get("present", 0), m.get("late", 0), m.get("absent", 0)
        tot = p + la + ab
        if tot >= LOW_ATTENDANCE_MIN_LECTURES:
            pct_abs = ab * 100 / tot
            if pct_abs > 40:
                depr += 1
            elif pct_abs > 25:
                warn += 1
    students = {
        "by_department": dist[:15],
        "by_level": [{"level": k, "students": v} for k, v in sorted(by_level.items(), key=lambda x: x[0])],
        "by_status": [{"status": k, "label": STATUS_LABELS.get(k, k), "count": v} for k, v in sorted(by_status.items(), key=lambda x: -x[1])],
        "warned": warn, "deprived": depr, "evaluated": sum(1 for m in att_by_student.values() if sum(m.values()) >= LOW_ATTENDANCE_MIN_LECTURES),
    }

    # ── 8) القاعات والجدول
    settings = await db.schedule_settings.find_one({"_id": "global"}) or {}
    n_days = len(settings.get("working_days") or []) or 6
    n_slots = len(settings.get("time_slots") or []) or 5
    capacity = n_days * n_slots
    slot_q = {} if dept_ids is None else {"department_id": {"$in": dept_ids}}
    slots = await db.weekly_schedule.find(slot_q, {"room_id": 1, "room": 1, "day": 1, "slot_number": 1, "course_id": 1, "merge_group_id": 1}).to_list(50000)
    room_docs = {str(r["_id"]): r.get("name") or r.get("room_number") or "" for r in await db.rooms.find({}, {"name": 1, "room_number": 1}).to_list(2000)}
    room_use, seen = {}, set()
    scheduled_course_ids = set()
    for sl in slots:
        scheduled_course_ids.add(sl.get("course_id"))
        key = (sl.get("room_id") or sl.get("room"), sl.get("day"), str(sl.get("slot_number")), sl.get("merge_group_id") or str(sl["_id"]))
        if key in seen:
            continue
        seen.add(key)
        rid = sl.get("room_id") or sl.get("room")
        if not rid:
            continue
        room_use[rid] = room_use.get(rid, 0) + 1
    rooms_list = sorted([{"room": room_docs.get(rid, str(rid)), "slots": n, "occupancy": round(min(n, capacity) * 100 / capacity, 1)} for rid, n in room_use.items()], key=lambda x: -x["slots"])
    unscheduled = [{"id": str(c["_id"]), "name": c.get("name", ""), "code": c.get("code", ""), "department": dnames.get(c.get("department_id"), "")}
                   for c in active_courses if str(c["_id"]) not in scheduled_course_ids]
    heat = {}
    for l in today_lectures:
        k = l.get("start_time", "—")
        h = heat.setdefault(k, {"time": k, "total": 0, "completed": 0, "cancelled": 0, "scheduled": 0})
        h["total"] += 1
        st = l.get("status", "scheduled")
        h["completed" if st == "completed" else "cancelled" if st in ("cancelled", "absent") else "scheduled"] += 1
    rooms = {
        "capacity_per_room": capacity, "rooms_used": len(rooms_list), "total_slots": len(seen),
        "avg_occupancy": round(sum(r["occupancy"] for r in rooms_list) / len(rooms_list), 1) if rooms_list else 0,
        "most_used": rooms_list[:5], "least_used": sorted(rooms_list, key=lambda x: x["slots"])[:5], "rows": rooms_list[:40],
        "unscheduled_courses": unscheduled[:30], "unscheduled_count": len(unscheduled),
        "today_heatmap": sorted(heat.values(), key=lambda x: x["time"]),
    }
    return {"teachers": teachers, "students": students, "rooms": rooms}


async def _finance(db, dept_ids, active_sem, total_students, allowed_types=None) -> dict:
    year = (active_sem or {}).get("academic_year") or ""
    if not year:
        try:
            from .fee_receipts import _academic_year
            year = await _academic_year(db)
        except Exception:
            year = ""
    sid_filter = None
    if dept_ids is not None:
        sid_filter = [str(s["_id"]) for s in await db.students.find({"department_id": {"$in": dept_ids}}, {"_id": 1}).to_list(50000)]
    types = await db.fee_types.find({"is_active": {"$ne": False}}).to_list(100)
    if allowed_types is not None:
        types = [t for t in types if str(t["_id"]) in allowed_types]
    out = []
    tot_amount = 0.0
    for t in types:
        tid = str(t["_id"])
        base = {"type_id": tid}
        if year:
            base["academic_year"] = year
        if sid_filter is not None:
            base["student_id"] = {"$in": sid_filter}
        approved, pending, rejected = await asyncio.gather(
            db.fee_receipts.count_documents({**base, "status": "approved"}),
            db.fee_receipts.count_documents({**base, "status": "pending"}),
            db.fee_receipts.count_documents({**base, "status": "rejected"}),
        )
        recurring = bool(t.get("recurring"))
        paid_students = len(await db.fee_receipts.distinct("student_id", {**base, "status": "approved"})) if recurring else approved
        amt = 0.0
        async for r in db.fee_receipts.aggregate([{"$match": {**base, "status": "approved", "amount": {"$type": "number"}}}, {"$group": {"_id": None, "s": {"$sum": "$amount"}}}]):
            amt = float(r.get("s") or 0)
        tot_amount += amt
        not_paid = max(total_students - paid_students, 0) if recurring else max(total_students - approved - pending - rejected, 0)
        out.append({"type_id": tid, "name": t.get("name", ""), "recurring": recurring, "approved": approved, "pending": pending,
                    "rejected": rejected, "paid_students": paid_students, "not_paid": not_paid, "amount": amt,
                    "paid_pct": round(paid_students * 100 / total_students, 1) if total_students else 0})
    return {"academic_year": year, "total_students": total_students, "types": out, "total_amount": tot_amount}


@router.get("/dashboard/management")
async def management_dashboard(period: str = "week", faculty_id: Optional[str] = None, department_id: Optional[str] = None,
                               current_user: dict = Depends(get_current_user)):
    if not _is_management(current_user):
        raise HTTPException(status_code=403, detail="لوحة القيادة متاحة للإدارة فقط")
    return await build_dashboard(get_db(), current_user, period, faculty_id or None, department_id or None)


@router.get("/dashboard/management/export")
async def management_dashboard_export(fmt: str = "pdf", period: str = "week", faculty_id: Optional[str] = None,
                                      department_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _is_management(current_user):
        raise HTTPException(status_code=403, detail="لوحة القيادة متاحة للإدارة فقط")
    if not dashboard_sections(current_user)["export"]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية تصدير لوحة القيادة")
    d = await build_dashboard(get_db(), current_user, period, faculty_id or None, department_id or None)
    return _stream(d, fmt, "لوحة القيادة", d["scope"]["label"])


async def build_hr_dashboard(db, user: dict, period: str, org_unit_id: Optional[str]) -> dict:
    """عرض «شؤون الموظفين» الكامل للوحة القيادة (يتطلب صلاحية dashboard_hr)"""
    from .hr_dashboard import hr_dashboard_section, hr_units_list
    period = period if period in PERIOD_LABELS else "week"
    d_from, d_to = _period_range(period)
    hr = await hr_dashboard_section(db, period, d_from, d_to, org_unit_id or None)
    return {"view": "hr", "generated_at": _now_yemen().strftime("%Y-%m-%d %H:%M"), "period": period, "period_label": PERIOD_LABELS[period],
            "date_from": d_from.isoformat(), "date_to": d_to.isoformat(), "sections": dashboard_sections(user),
            "read_only": user.get("role") in READ_ONLY_ROLES, "units": await hr_units_list(db), "hr": hr}


def _hr_guard(user: dict):
    if not _is_management(user):
        raise HTTPException(status_code=403, detail="لوحة القيادة متاحة للإدارة فقط")
    if not dashboard_sections(user)["hr"]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية عرض شؤون الموظفين في لوحة القيادة")


@router.get("/dashboard/management/hr")
async def management_dashboard_hr(period: str = "week", org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _hr_guard(current_user)
    return await build_hr_dashboard(get_db(), current_user, period, org_unit_id)


@router.get("/dashboard/management/hr/export")
async def management_dashboard_hr_export(fmt: str = "pdf", period: str = "week", org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _hr_guard(current_user)
    if not dashboard_sections(current_user)["export"]:
        raise HTTPException(status_code=403, detail="ليس لديك صلاحية تصدير لوحة القيادة")
    d = await build_hr_dashboard(get_db(), current_user, period, org_unit_id)
    return _stream(d, fmt, "لوحة القيادة - شؤون الموظفين", d["hr"]["scope"]["label"])


def _stream(d: dict, fmt: str, base: str, scope_label: str):
    if fmt == "excel":
        buf = _build_excel(d)
        fname = export_filename(base, scope_label, d["period_label"], ext="xlsx")
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        buf = _build_pdf(d)
        fname = export_filename(base, scope_label, d["period_label"], ext="pdf")
        media = "application/pdf"
    return StreamingResponse(buf, media_type=media, headers=export_headers(fname))


def _numbers_rows(d: dict):
    n = d["numbers"]
    rate = f"{n['attendance_rate']}%" if n["attendance_rate"] is not None else "—"
    return [["الطلاب النشطون", n["students"]], ["الأساتذة", n["teachers"]], ["مقررات الفصل النشط", n["courses"]],
            ["الأقسام", n["departments"]], ["الكليات", n["faculties"]],
            ["محاضرات اليوم", n["lectures_today"]], ["منفَّذة اليوم", n["today_done"]], ["ملغاة اليوم", n["today_cancelled"]], ["قادمة اليوم", n["today_upcoming"]],
            [f"محاضرات {d['period_label']}", n["lectures_period"]], ["منفَّذة", n["lectures_status"].get("completed", 0)],
            ["ملغاة", n["lectures_status"].get("cancelled", 0)], ["غياب أستاذ", n["lectures_status"].get("absent", 0)],
            ["حاضر", n["present"]], ["متأخر", n["late"]], ["غائب", n["absent"]], ["نسبة الحضور", rate]]


def _phase2_tables(d: dict):
    """جداول المرحلة 2 (الأساتذة/الطلاب/القاعات) للتصدير — (عنوان، صفوف برأس)"""
    out = []
    t, st, r = d.get("teachers"), d.get("students"), d.get("rooms")
    if t:
        out.append(("الأساتذة", [["الأستاذ", "مقررات", "مجدولة", "منفَّذة", "ملغاة", "غياب", "مرات التأخر", "متوسط التأخير (د)", "ساعات", "الالتزام %"]] +
                    [[x["name"], x["courses"], x["scheduled"], x["completed"], x["cancelled"], x["absent"], x["late_count"], x["avg_delay"], x["hours"], x["commitment"] if x["commitment"] is not None else "—"] for x in t["rows"]]))
    if st:
        out.append(("الطلاب حسب القسم", [["القسم", "الطلاب النشطون", "نسبة الحضور %"]] + [[x["name"], x["students"], x["rate"] if x["rate"] is not None else "—"] for x in st["by_department"]]))
        out.append(("حالات الطلاب", [["الحالة", "العدد"]] + [[x["label"], x["count"]] for x in st["by_status"]] +
                    [["إنذار (غياب >25%)", st["warned"]], ["حرمان (غياب >40%)", st["deprived"]]]))
    if r:
        out.append(("إشغال القاعات", [["القاعة", "الخانات الأسبوعية", "الإشغال %"]] + [[x["room"], x["slots"], x["occupancy"]] for x in r["rows"]]))
        if r["unscheduled_courses"]:
            out.append(("مقررات غير مدرجة", [["المقرر", "الرمز", "القسم"]] + [[x["name"], x["code"], x["department"]] for x in r["unscheduled_courses"]]))
    return out


def _hr_tables(d: dict):
    """جداول شؤون الموظفين للتصدير"""
    h = d.get("hr")
    if not h:
        return []
    out = []
    hc, p = h.get("headcount"), h.get("period")
    if hc:
        out.append(("الموظفون حسب الفئة", [["الفئة", "العدد"]] + [[c["label"], c["count"]] for c in hc["by_category"]] + [["الإجمالي على رأس العمل", hc["total_active"]]]))
        out.append(("موظفو الوحدات", [["الوحدة", "النوع", "الإجمالي", "أكاديمي", "إداري", "غيرهم"]] + [[u["name"], u["type_label"], u["total"], u["academic"], u["administrative"], u["other"]] for u in hc["by_unit"]]))
    if p:
        out.append((f"الدوام الإداري — {d['period_label']}", [["البيان", "القيمة"], ["أيام العمل", p["work_days"]], ["نسبة الالتزام %", p["commitment_rate"] if p["commitment_rate"] is not None else "—"],
                    ["حضور", p["present"]], ["متأخر", p["late"]], ["غياب", p["absent"]], ["بعذر", p["excused"]], ["لم يُسجَّل", p["unmarked"]], ["ساعات التأخير", p["late_hours"]],
                    ["إجازات معتمدة", p["leaves_approved"]], ["أيام الإجازات", p["leave_days"]], ["مهام منجزة", p["tasks_done"]], ["مهام جديدة", p["tasks_created"]]]))
        if p["chart"]["points"]:
            gb = "التاريخ" if p["chart"]["group_by"] == "date" else "الوحدة"
            out.append(("مخطط الدوام الإداري", [[gb, "حاضر", "متأخر", "غائب", "النسبة %"]] + [[x.get("date") or x["label"], x["present"], x["late"], x["absent"], x["rate"] if x["rate"] is not None else "—"] for x in p["chart"]["points"]]))
        if p["units_attendance"]:
            out.append(("الوحدات الأعلى غياباً", [["الوحدة", "حاضر", "متأخر", "غائب", "دقائق التأخير", "النسبة %"]] + [[u["unit"], u["present"], u["late"], u["absent"], u["late_minutes"], u["rate"] if u["rate"] is not None else "—"] for u in p["units_attendance"]]))
        if p["bottom_employees"]:
            out.append(("الأقل التزاماً", [["الموظف", "الوحدة", "حاضر", "متأخر", "غائب", "دقائق التأخير", "النسبة %"]] + [[e["name"], e["unit"], e["present"], e["late"], e["absent"], e["late_minutes"], e["rate"] if e["rate"] is not None else "—"] for e in p["bottom_employees"]]))
    return out


def _hr_summary_rows(h: dict):
    hc, p = h.get("headcount") or {}, h.get("period") or {}
    cats = " · ".join(f"{c['label']} {c['count']}" for c in hc.get("by_category", []))
    rows = [["الموظفون على رأس العمل", hc.get("total_active", h["employees_active"])], ["حسب الفئة", cats], ["الوحدات التنظيمية", f"{hc.get('units_with_staff', 0)} بها موظفون من {hc.get('units_count', 0)}"],
            ["حضر اليوم", h["present_today"]], ["متأخر اليوم", h["late_today"]], ["غائب اليوم", len(h["absent_today"])], ["في إجازة اليوم", len(h["on_leave_today"])], ["لم يُسجَّل اليوم", h["unmarked_today"]],
            ["طلبات إجازة معلّقة", h["pending_leaves_count"]], ["مهام متأخرة", h["overdue_tasks"]], ["تقييمات للاعتماد", h["pending_appraisals"]], ["عقود تنتهي خلال 60 يوماً", len(h["expiring_contracts"])]]
    if p:
        rows += [["نسبة الالتزام بالدوام", f"{p['commitment_rate']}%" if p["commitment_rate"] is not None else "—"], ["أيام العمل في الفترة", p["work_days"]], ["أيام حضور / تأخر / غياب", f"{p['attended']} / {p['late']} / {p['absent']}"],
                 ["ساعات التأخير", p["late_hours"]], ["إجازات معتمدة (أيام)", f"{p['leaves_approved']} ({p['leave_days']})"], ["مهام منجزة / جديدة", f"{p['tasks_done']} / {p['tasks_created']}"]]
    return rows


def _hr_today_tables(h: dict):
    out = []
    if h["absent_today"]:
        out.append(("الغائبون اليوم", [["الموظف", "الوحدة", "ملاحظة"]] + [[i["employee_name"], i["org_unit_name"], i.get("note", "")] for i in h["absent_today"]]))
    if h["on_leave_today"]:
        out.append(("في إجازة اليوم", [["الموظف", "الوحدة", "النوع", "حتى"]] + [[i["employee_name"], i["org_unit_name"], i["type"], i["end_date"]] for i in h["on_leave_today"]]))
    if h["pending_leaves"]:
        out.append(("طلبات إجازة معلّقة", [["الموظف", "الوحدة", "النوع", "من", "الأيام", "لدى"]] + [[i["employee_name"], i["org_unit_name"], i["type"], i["start_date"], i["days"], "المدير" if i["status"] == "pending" else "HR"] for i in h["pending_leaves"]]))
    if h["expiring_contracts"]:
        out.append(("عقود تنتهي", [["الموظف", "الوحدة", "تاريخ الانتهاء"]] + [[i["employee_name"], i["org_unit_name"], i["contract_end_date"]] for i in h["expiring_contracts"]]))
    if h.get("overdue_tasks_list"):
        out.append(("مهام متأخرة", [["المهمة", "المكلَّف", "الاستحقاق"]] + [[i["title"], i["employee_name"], i["due_date"]] for i in h["overdue_tasks_list"]]))
    return out


def _build_excel(d: dict) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1565C0")

    def sheet(title, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:30]
        ws.sheet_view.rightToLeft = True
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill; c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(14, min(50, max(len(str(c.value or "")) for c in col) + 4))
        return ws

    if d.get("view") == "hr":
        h = d["hr"]
        ws = sheet("الملخص", [["البيان", "القيمة"]] + _hr_summary_rows(h), first=True)
        ws.insert_rows(1, 3)
        ws["A1"] = f"لوحة القيادة — شؤون الموظفين — {h['scope']['label']} — {d['period_label']} ({d['date_from']} → {d['date_to']})"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"تاريخ الإصدار: {d['generated_at']}"
        if h["alerts"]:
            sheet("التنبيهات", [["التنبيه", "العدد", "التفاصيل"]] + [[a["title"], a["count"], a["hint"]] for a in h["alerts"]])
        for title, rows in _hr_tables(d):
            sheet(title, rows)
        for title, rows in _hr_today_tables(h):
            sheet(title, rows)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    ws = sheet("الأرقام", [["البيان", "القيمة"]] + _numbers_rows(d), first=True)
    ws.insert_rows(1, 3)
    ws["A1"] = f"لوحة القيادة — {d['scope']['label']} — {d['period_label']} ({d['date_from']} → {d['date_to']})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{d['semester']['name']} {d['semester']['academic_year']}   |   تاريخ الإصدار: {d['generated_at']}"
    if d["alerts"]:
        sheet("التنبيهات", [["التنبيه", "العدد", "التفاصيل"]] + [[a["title"], a["count"], a["hint"]] for a in d["alerts"]])
    if d["alerts"] and d["alerts"][0]["items"]:
        sheet("طلاب حضور منخفض", [["الطالب", "رقم القيد", "القسم", "المستوى", "المحاضرات", "الغياب", "النسبة %"]] +
              [[s["name"], s["student_id"], s["department"], s["level"], s["lectures"], s["absent"], s["rate"]] for s in d["alerts"][0]["items"]])
    lt = next((a for a in d["alerts"] if a["key"] == "late_teachers"), None)
    if lt and lt["items"]:
        sheet("تأخر الأساتذة", [["الأستاذ", "مرات التأخر", "أقصى تأخير (د)", "مجموع التأخير (د)"]] +
              [[t["name"], t["count"], t["max_delay"], t["total_delay"]] for t in lt["items"]])
    gb = {"date": "التاريخ", "department": "القسم", "course": "المقرر"}[d["chart"]["group_by"]] if d["chart"] else ""
    if d["chart"]: sheet("الحضور", [[gb, "المحاضرات", "المنفَّذة", "حاضر", "متأخر", "غائب", "النسبة %"]] +
          [[p.get("date") or p["label"], p["lectures"], p["completed"], p["present"], p["late"], p["absent"], p["rate"] if p["rate"] is not None else "—"] for p in d["chart"]["points"]])
    if d["finance"]:
        sheet("المالية", [["نوع الرسوم", "معتمد", "معلق", "مرفوض", "طلاب دافعون", "غير دافعين", "نسبة الدفع %", "المبالغ المعتمدة"]] +
              [[t["name"] + (" (متكرر)" if t["recurring"] else ""), t["approved"], t["pending"], t["rejected"], t["paid_students"], t["not_paid"], t["paid_pct"], t["amount"]] for t in d["finance"]["types"]])
    for title, rows in _phase2_tables(d):
        sheet(title, rows)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _build_pdf(d: dict) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import arabic_reshaper
    from bidi.algorithm import get_display

    font = "Helvetica"
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    try:
        pdfmetrics.registerFont(TTFont("Amiri", os.path.join(here, "fonts", "Amiri-Regular.ttf"))); font = "Amiri"
    except Exception:
        pass

    def ar(t):
        try:
            return get_display(arabic_reshaper.reshape(str(t if t is not None else "")))
        except Exception:
            return str(t or "")

    title = ParagraphStyle("t", fontName=font, fontSize=16, alignment=TA_CENTER, textColor=colors.HexColor("#1565c0"), spaceAfter=4)
    sub = ParagraphStyle("s", fontName=font, fontSize=10, alignment=TA_CENTER, textColor=colors.HexColor("#607d8b"), spaceAfter=8)
    sec = ParagraphStyle("c", fontName=font, fontSize=12, alignment=TA_RIGHT, textColor=colors.HexColor("#1565c0"), spaceBefore=8, spaceAfter=4)

    def grid(rows, widths, head_bg="#1565C0", fs=8.5):
        rows = [[ar(c) for c in reversed(r)] for r in rows]
        t = Table(rows, colWidths=list(reversed(widths)), repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), fs),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head_bg)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d4e4")), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    if d.get("view") == "hr":
        h = d["hr"]
        el = [Paragraph(ar(f"جامعة الأحقاف — لوحة القيادة: شؤون الموظفين — {h['scope']['label']}"), title),
              Paragraph(ar(f"{d['period_label']} ({d['date_from']} → {d['date_to']})   |   تاريخ الإصدار: {d['generated_at']}"), sub)]
        nr = _hr_summary_rows(h)
        half = (len(nr) + 1) // 2
        rows = [["البيان", "القيمة", "البيان", "القيمة"]]
        for i in range(half):
            a = nr[i]; b = nr[i + half] if i + half < len(nr) else ["", ""]
            rows.append([a[0], a[1], b[0], b[1]])
        el += [Paragraph(ar("الملخص"), sec), grid(rows, [55 * mm, 35 * mm, 55 * mm, 35 * mm], head_bg="#6d28d9", fs=9), Spacer(1, 3 * mm)]
        if h["alerts"]:
            el += [Paragraph(ar("التنبيهات الإدارية"), sec), grid([["التنبيه", "العدد", "التفاصيل"]] + [[a["title"], a["count"], a["hint"]] for a in h["alerts"]], [90 * mm, 25 * mm, 90 * mm], head_bg="#b91c1c")]
        for title_, rows_ in _hr_tables(d) + _hr_today_tables(h):
            if len(rows_) > 1:
                n = len(rows_[0]); w = (270 * mm) / n
                el += [Paragraph(ar(title_), sec), grid(rows_, [w] * n, head_bg="#6d28d9", fs=8)]
        buf = io.BytesIO()
        SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm).build(el)
        buf.seek(0)
        return buf

    el = [Paragraph(ar(f"جامعة الأحقاف — لوحة القيادة: {d['scope']['label']}"), title),
          Paragraph(ar(f"{d['period_label']} ({d['date_from']} → {d['date_to']})   |   {d['semester']['name']} {d['semester']['academic_year']}   |   تاريخ الإصدار: {d['generated_at']}"), sub)]
    nr = _numbers_rows(d)
    half = (len(nr) + 1) // 2
    rows = [["البيان", "القيمة", "البيان", "القيمة"]]
    for i in range(half):
        a = nr[i]; b = nr[i + half] if i + half < len(nr) else ["", ""]
        rows.append([a[0], a[1], b[0], b[1]])
    el += [Paragraph(ar("الأرقام الرئيسية"), sec), grid(rows, [60 * mm, 30 * mm, 60 * mm, 30 * mm], fs=9.5), Spacer(1, 3 * mm)]
    if d["alerts"]:
        el += [Paragraph(ar("التنبيهات"), sec), grid([["التنبيه", "العدد", "التفاصيل"]] + [[a["title"], a["count"], a["hint"]] for a in d["alerts"]], [90 * mm, 25 * mm, 90 * mm])]
    if d["alerts"] and d["alerts"][0]["items"]:
        el += [Paragraph(ar("طلاب حضورهم منخفض (أدنى 10)"), sec),
               grid([["الطالب", "رقم القيد", "القسم", "المستوى", "المحاضرات", "الغياب", "النسبة"]] +
                    [[s["name"], s["student_id"], s["department"], s["level"], s["lectures"], s["absent"], f"{s['rate']}%"] for s in d["alerts"][0]["items"]],
                    [60 * mm, 28 * mm, 55 * mm, 18 * mm, 22 * mm, 18 * mm, 22 * mm])]
    lt = next((a for a in d["alerts"] if a["key"] == "late_teachers"), None)
    if lt and lt["items"]:
        el += [Paragraph(ar("أساتذة تأخروا في بدء التحضير"), sec),
               grid([["الأستاذ", "مرات التأخر", "أقصى تأخير (د)", "مجموع التأخير (د)"]] + [[t["name"], t["count"], t["max_delay"], t["total_delay"]] for t in lt["items"]],
                    [80 * mm, 30 * mm, 35 * mm, 35 * mm])]
    gb = {"date": "التاريخ", "department": "القسم", "course": "المقرر"}[d["chart"]["group_by"]] if d["chart"] else ""
    if d["chart"]: el += [Paragraph(ar(f"الحضور حسب {gb}"), sec),
           grid([[gb, "المحاضرات", "المنفَّذة", "حاضر", "متأخر", "غائب", "النسبة"]] +
                [[p.get("date") or p["label"], p["lectures"], p["completed"], p["present"], p["late"], p["absent"], f"{p['rate']}%" if p["rate"] is not None else "—"] for p in d["chart"]["points"]],
                [50 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm], head_bg="#37474f")]
    if d["finance"]:
        el += [Paragraph(ar(f"المالية — {d['finance']['academic_year']}"), sec),
               grid([["نوع الرسوم", "معتمد", "معلق", "مرفوض", "دافعون", "غير دافعين", "نسبة الدفع", "المبالغ المعتمدة"]] +
                    [[t["name"] + (" (متكرر)" if t["recurring"] else ""), t["approved"], t["pending"], t["rejected"], t["paid_students"], t["not_paid"], f"{t['paid_pct']}%", f"{t['amount']:,.0f}"] for t in d["finance"]["types"]],
                    [60 * mm, 22 * mm, 22 * mm, 22 * mm, 22 * mm, 25 * mm, 25 * mm, 32 * mm], head_bg="#2e7d32")]
    for title, rows in _phase2_tables(d):
        if len(rows) > 1:
            n = len(rows[0]); w = (270 * mm) / n
            el += [Paragraph(ar(title), sec), grid(rows, [w] * n, head_bg="#00695c", fs=8)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm).build(el)
    buf.seek(0)
    return buf
