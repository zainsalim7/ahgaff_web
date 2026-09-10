"""📋 تقرير حضور الطالب المفصّل: كل المحاضرات المدرجة للفصل النشط + تصدير PDF/Excel"""
import io
import os
from datetime import datetime, date as _date
from typing import Optional
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from bson import ObjectId

from .deps import get_db, get_current_user, has_permission, get_scope_filter, export_filename, export_headers
from models.permissions import Permission

router = APIRouter(tags=["تقرير الطالب"])

DAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
STUDENT_STATUS_AR = {"present": "حاضر", "absent": "غائب", "late": "متأخر", "excused": "معذور"}
LECTURE_STATUS_AR = {"completed": "مُنفَّذة", "scheduled": "مجدولة", "cancelled": "ملغاة", "absent": "غياب الأستاذ", "missed": "لم تُسجَّل"}
LEVEL_AR = {1: "الأول", 2: "الثاني", 3: "الثالث", 4: "الرابع", 5: "الخامس", 6: "السادس"}


def _day_ar(d: str) -> str:
    try:
        return DAYS_AR[datetime.strptime(d, "%Y-%m-%d").weekday()]
    except Exception:
        return ""


def _norm_date(d) -> Optional[str]:
    if not d or not isinstance(d, str):
        return None
    p = d.strip().split("-")
    if len(p) == 3 and len(p[0]) == 4:
        return d.strip()
    if len(p) == 3:
        return f"{int(p[2]):04d}-{int(p[1]):02d}-{int(p[0]):02d}"
    return None


async def _assert_scope(db, current_user: dict, student: dict):
    if current_user.get("role") == "admin":
        return
    q = await get_scope_filter(current_user, "students")
    if q and not await db.students.find_one({"$and": [q, {"_id": student["_id"]}]}, {"_id": 1}):
        raise HTTPException(status_code=403, detail="هذا الطالب خارج نطاق صلاحيتك")


async def build_student_report(db, student_id: str) -> dict:
    student = None
    if ObjectId.is_valid(student_id):
        student = await db.students.find_one({"_id": ObjectId(student_id)})
    if not student:
        student = await db.students.find_one({"student_id": student_id})
    if not student:
        raise HTTPException(status_code=404, detail="الطالب غير موجود")
    sid = str(student["_id"])

    sem = await db.semesters.find_one({"$or": [{"status": "active"}, {"is_active": True}]}) or {}
    sem_id = str(sem["_id"]) if sem else None
    sem_start, sem_end = _norm_date(sem.get("start_date")), _norm_date(sem.get("end_date"))

    dept = await db.departments.find_one({"_id": ObjectId(student["department_id"])}) if ObjectId.is_valid(student.get("department_id") or "") else None
    fac_id = student.get("faculty_id") or (dept or {}).get("faculty_id")
    fac = await db.faculties.find_one({"_id": ObjectId(fac_id)}) if ObjectId.is_valid(fac_id or "") else None

    enrollments = await db.enrollments.find({"student_id": sid}).to_list(200)
    today = _date.today().isoformat()
    courses_out, totals = [], {"lectures": 0, "executed": 0, "present": 0, "absent": 0, "late": 0, "upcoming": 0, "cancelled": 0}
    for e in enrollments:
        if not ObjectId.is_valid(e.get("course_id") or ""):
            continue
        course = await db.courses.find_one({"_id": ObjectId(e["course_id"])})
        if not course:
            continue
        c_sem = course.get("semester_id")
        if sem_id and c_sem and c_sem != sem_id:
            continue
        cid = str(course["_id"])
        lq: dict = {"course_id": cid}
        if sem_id:
            ors = [{"semester_id": sem_id}]
            if sem_start and sem_end:
                ors.append({"semester_id": {"$in": [None, ""]}, "date": {"$gte": sem_start, "$lte": sem_end}})
            elif c_sem == sem_id or not c_sem:
                ors.append({"semester_id": {"$in": [None, ""]}})
            lq["$or"] = ors
        lectures = await db.lectures.find(lq).sort([("date", 1), ("start_time", 1)]).to_list(500)
        lids = [str(l["_id"]) for l in lectures]
        att = {a["lecture_id"]: a.get("status") for a in await db.attendance.find(
            {"student_id": sid, "course_id": cid, "lecture_id": {"$in": lids}}, {"lecture_id": 1, "status": 1}).to_list(1000)}

        teacher_name = course.get("teacher_name") or ""
        if not teacher_name and course.get("teacher_id"):
            t = None
            if ObjectId.is_valid(course["teacher_id"]):
                t = await db.teachers.find_one({"_id": ObjectId(course["teacher_id"])}) or await db.users.find_one({"_id": ObjectId(course["teacher_id"])})
            teacher_name = (t or {}).get("full_name", "")

        rows, st = [], {"executed": 0, "present": 0, "absent": 0, "late": 0, "upcoming": 0, "cancelled": 0}
        for i, l in enumerate(lectures, 1):
            lid = str(l["_id"])
            lstatus = l.get("status") or "scheduled"
            if l.get("is_cancelled"):
                lstatus = "cancelled"
            s_status, s_label, counted = None, "—", False
            if lstatus in ("cancelled", "absent"):
                st["cancelled"] += 1
            elif lid in att:
                s_status = att[lid]
                s_label = STUDENT_STATUS_AR.get(s_status, s_status)
                counted = True
                lstatus = "completed"
            elif lstatus == "completed":
                s_status, s_label, counted = "absent", "غائب", True
            elif l.get("date", "") >= today:
                s_label = "لم تُنعقد بعد"
                st["upcoming"] += 1
            else:
                lstatus, s_label = "missed", "لم تُسجَّل"
            if counted:
                st["executed"] += 1
                if s_status == "present" or s_status == "excused":
                    st["present"] += 1
                elif s_status == "late":
                    st["late"] += 1
                else:
                    st["absent"] += 1
            rows.append({
                "n": i, "id": lid, "date": l.get("date", ""), "day": _day_ar(l.get("date", "")),
                "time": f"{l.get('start_time', '')} - {l.get('end_time', '')}".strip(" -"),
                "room": l.get("room") or "", "topic": l.get("lesson_title") or l.get("topic") or "",
                "lecture_status": lstatus, "lecture_status_label": LECTURE_STATUS_AR.get(lstatus, lstatus),
                "student_status": s_status, "student_status_label": s_label,
                "cancellation_reason": l.get("cancellation_reason", "") if lstatus == "cancelled" else "",
            })
        rate = round((st["present"] + st["late"] * 0.5) / st["executed"] * 100, 1) if st["executed"] else None
        for k in st:
            totals[k] += st[k]
        totals["lectures"] += len(lectures)
        courses_out.append({
            "course_id": cid, "course_name": course.get("name", ""), "course_code": course.get("code", ""),
            "teacher_name": teacher_name, "total_lectures": len(lectures), **st,
            "attendance_rate": rate, "warning": rate is not None and rate < 75, "lectures": rows,
        })
    courses_out.sort(key=lambda c: c["course_name"])
    overall = round((totals["present"] + totals["late"] * 0.5) / totals["executed"] * 100, 1) if totals["executed"] else None
    return {
        "student": {
            "id": sid, "student_id": student.get("student_id", ""), "full_name": student.get("full_name", ""),
            "faculty_name": (fac or {}).get("name", ""), "department_name": (dept or {}).get("name", ""),
            "department_id": student.get("department_id"), "level": student.get("level"),
            "level_label": LEVEL_AR.get(student.get("level"), str(student.get("level") or "")),
            "section": student.get("section") or "",
        },
        "semester": {"id": sem_id, "name": sem.get("name", ""), "academic_year": sem.get("academic_year", ""),
                     "start_date": sem_start, "end_date": sem_end},
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "summary": {"total_courses": len(courses_out), "total_lectures": totals["lectures"], "executed": totals["executed"],
                    "present": totals["present"], "absent": totals["absent"], "late": totals["late"],
                    "upcoming": totals["upcoming"], "cancelled": totals["cancelled"], "overall_attendance_rate": overall},
        "courses": courses_out,
    }


def _check(current_user):
    if not has_permission(current_user, Permission.VIEW_REPORTS) and not has_permission(current_user, Permission.VIEW_ATTENDANCE):
        raise HTTPException(status_code=403, detail="غير مصرح لك")


@router.get("/reports/student/{student_id}/detailed")
async def student_detailed_report(student_id: str, current_user: dict = Depends(get_current_user)):
    _check(current_user)
    db = get_db()
    data = await build_student_report(db, student_id)
    await _assert_scope(db, current_user, {"_id": ObjectId(data["student"]["id"])})
    return data


def _fname_parts(d: dict, view: str = "detailed"):
    s = d["student"]
    return ("تقرير حضور الطالب" + (" (مختصر)" if view == "summary" else ""), s["full_name"], s["department_name"],
            f"المستوى {s['level']}" if s.get("level") else "", f"شعبة {s['section']}" if s.get("section") else "")


@router.get("/reports/student/{student_id}/detailed/export")
async def student_detailed_export(student_id: str, fmt: str = "excel", view: str = "detailed",
                                  current_user: dict = Depends(get_current_user)):
    """view=detailed (كل المحاضرات) | summary (جدول واحد بكل المقررات)"""
    _check(current_user)
    db = get_db()
    d = await build_student_report(db, student_id)
    await _assert_scope(db, current_user, {"_id": ObjectId(d["student"]["id"])})
    summary_only = view == "summary"
    if fmt == "pdf":
        buf = _build_pdf(d, summary_only)
        return StreamingResponse(buf, media_type="application/pdf", headers=export_headers(export_filename(*_fname_parts(d, view), ext="pdf")))
    buf = _build_excel(d, summary_only)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=export_headers(export_filename(*_fname_parts(d, view), ext="xlsx")))


def _build_excel(d: dict, summary_only: bool = False) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    s, sm, sem = d["student"], d["summary"], d["semester"]
    hdr_fill = PatternFill("solid", fgColor="1565C0")
    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص"
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:H1")
    ws["A1"] = f"تقرير حضور الطالب — {sem.get('name', '')} {sem.get('academic_year', '')}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    info = [("اسم الطالب", s["full_name"]), ("رقم القيد", s["student_id"]), ("الكلية", s["faculty_name"]),
            ("القسم", s["department_name"]), ("المستوى", s["level_label"]), ("الشعبة", s["section"]),
            ("تاريخ التقرير", d["generated_at"]),
            ("نسبة الحضور العامة", f"{sm['overall_attendance_rate']}%" if sm["overall_attendance_rate"] is not None else "—")]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r0 = len(info) + 5
    heads = ["م", "المقرر", "الرمز", "الأستاذ", "المدرجة", "المنفَّذة", "حاضر", "غائب", "متأخر", "قادمة", "ملغاة", "نسبة الحضور"]
    for ci, h in enumerate(heads, 1):
        c = ws.cell(row=r0, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill; c.alignment = Alignment(horizontal="center")
    for i, c in enumerate(d["courses"], 1):
        vals = [i, c["course_name"], c["course_code"], c["teacher_name"], c["total_lectures"], c["executed"], c["present"],
                c["absent"], c["late"], c["upcoming"], c["cancelled"], f"{c['attendance_rate']}%" if c["attendance_rate"] is not None else "—"]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=r0 + i, column=ci, value=v).alignment = Alignment(horizontal="center")
    tr = r0 + len(d["courses"]) + 1
    tot = ["", "الإجمالي", "", "", sm["total_lectures"], sm["executed"], sm["present"], sm["absent"], sm["late"], sm["upcoming"], sm["cancelled"],
           f"{sm['overall_attendance_rate']}%" if sm["overall_attendance_rate"] is not None else "—"]
    for ci, v in enumerate(tot, 1):
        ws.cell(row=tr, column=ci, value=v).font = Font(bold=True)
    for ci, w in enumerate([5, 30, 10, 24, 10, 10, 8, 8, 8, 8, 8, 12], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    if summary_only:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
    ws2 = wb.create_sheet("المحاضرات")
    ws2.sheet_view.rightToLeft = True
    heads2 = ["م", "المقرر", "الرمز", "التاريخ", "اليوم", "الوقت", "القاعة", "الموضوع", "حالة المحاضرة", "حالة الطالب", "سبب الإلغاء"]
    for ci, h in enumerate(heads2, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill; c.alignment = Alignment(horizontal="center")
    row, n = 2, 1
    for c in d["courses"]:
        for l in c["lectures"]:
            vals = [n, c["course_name"], c["course_code"], l["date"], l["day"], l["time"], l["room"], l["topic"],
                    l["lecture_status_label"], l["student_status_label"], l["cancellation_reason"]]
            for ci, v in enumerate(vals, 1):
                ws2.cell(row=row, column=ci, value=v).alignment = Alignment(horizontal="center")
            row += 1; n += 1
    for ci, w in enumerate([5, 28, 10, 12, 10, 14, 12, 28, 14, 14, 24], 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_pdf(d: dict, summary_only: bool = False) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import arabic_reshaper
    from bidi.algorithm import get_display

    font = "Helvetica"
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    fp = os.path.join(here, "fonts", "Amiri-Regular.ttf")
    try:
        pdfmetrics.registerFont(TTFont("Amiri", fp)); font = "Amiri"
    except Exception:
        pass

    def ar(t):
        try:
            return get_display(arabic_reshaper.reshape(str(t if t is not None else "")))
        except Exception:
            return str(t or "")

    s, sm, sem = d["student"], d["summary"], d["semester"]
    title = ParagraphStyle("t", fontName=font, fontSize=16, alignment=TA_CENTER, textColor=colors.HexColor("#1565c0"), spaceAfter=4)
    sub = ParagraphStyle("s", fontName=font, fontSize=10, alignment=TA_CENTER, textColor=colors.HexColor("#607d8b"), spaceAfter=8)
    name = ParagraphStyle("n", fontName=font, fontSize=15, alignment=TA_RIGHT, textColor=colors.HexColor("#1a2540"), spaceAfter=4)
    norm = ParagraphStyle("p", fontName=font, fontSize=10, alignment=TA_RIGHT, leading=15)
    sec = ParagraphStyle("c", fontName=font, fontSize=12, alignment=TA_RIGHT, textColor=colors.HexColor("#1565c0"), spaceBefore=8, spaceAfter=4)

    def grid(rows, widths, head_bg="#1565C0", font_size=8.5):
        rows = [[ar(c) for c in reversed(r)] for r in rows]
        t = Table(rows, colWidths=list(reversed(widths)), repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head_bg)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d4e4")), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    el = [Paragraph(ar("جامعة الأحقاف — تقرير حضور الطالب"), title),
          Paragraph(ar(f"{sem.get('name', '')} {sem.get('academic_year', '')}" + (f"  ({sem['start_date']} → {sem['end_date']})" if sem.get("start_date") else "") + f"   |   تاريخ التقرير: {d['generated_at']}"), sub),
          Paragraph(ar(s["full_name"]), name),
          Paragraph(ar(f"رقم القيد: {s['student_id']}   |   الكلية: {s['faculty_name']}   |   القسم: {s['department_name']}   |   المستوى: {s['level_label']}   |   الشعبة: {s['section'] or '—'}"), norm),
          Spacer(1, 4 * mm)]
    rate_txt = f"{sm['overall_attendance_rate']}%" if sm["overall_attendance_rate"] is not None else "—"
    el.append(grid([["المقررات", "المحاضرات المدرجة", "المنفَّذة", "حاضر", "غائب", "متأخر", "قادمة", "ملغاة", "نسبة الحضور"],
                    [sm["total_courses"], sm["total_lectures"], sm["executed"], sm["present"], sm["absent"], sm["late"], sm["upcoming"], sm["cancelled"], rate_txt]],
                   [28 * mm] * 9, font_size=9.5))
    el.append(Spacer(1, 5 * mm))
    if summary_only:
        rows = [["#", "المقرر", "الرمز", "الأستاذ", "المدرجة", "المنفَّذة", "حاضر", "غائب", "متأخر", "قادمة", "ملغاة", "نسبة الحضور"]]
        for i, c in enumerate(d["courses"], 1):
            r = f"{c['attendance_rate']}%" if c["attendance_rate"] is not None else "—"
            rows.append([i, c["course_name"], c["course_code"], c["teacher_name"], c["total_lectures"], c["executed"], c["present"],
                         c["absent"], c["late"], c["upcoming"], c["cancelled"], r + (" ⚠" if c["warning"] else "")])
        rows.append(["", "الإجمالي", "", "", sm["total_lectures"], sm["executed"], sm["present"], sm["absent"], sm["late"], sm["upcoming"], sm["cancelled"], rate_txt])
        el.append(Paragraph(ar("ملخص المقررات — الفصل النشط"), sec))
        el.append(grid(rows, [8 * mm, 62 * mm, 20 * mm, 44 * mm, 16 * mm, 16 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm, 22 * mm], font_size=9))
        d["courses"] = []
    for c in d["courses"]:
        r = f"{c['attendance_rate']}%" if c["attendance_rate"] is not None else "—"
        warn = "  ⚠ أقل من 75%" if c["warning"] else ""
        head = [Paragraph(ar(f"{c['course_name']} ({c['course_code']})" + (f" — {c['teacher_name']}" if c["teacher_name"] else "")), sec),
                Paragraph(ar(f"المدرجة: {c['total_lectures']}   المنفَّذة: {c['executed']}   حاضر: {c['present']}   غائب: {c['absent']}   متأخر: {c['late']}   قادمة: {c['upcoming']}   ملغاة: {c['cancelled']}   نسبة الحضور: {r}{warn}"), norm)]
        rows = [["#", "التاريخ", "اليوم", "الوقت", "القاعة", "الموضوع", "حالة المحاضرة", "حالة الطالب"]]
        for l in c["lectures"]:
            rows.append([l["n"], l["date"], l["day"], l["time"], l["room"], l["topic"][:40], l["lecture_status_label"], l["student_status_label"]])
        el.append(KeepTogether(head + [grid(rows, [8 * mm, 22 * mm, 18 * mm, 26 * mm, 22 * mm, 78 * mm, 26 * mm, 26 * mm], head_bg="#37474f")]))
        el.append(Spacer(1, 3 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm).build(el)
    buf.seek(0)
    return buf
