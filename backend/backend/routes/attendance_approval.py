"""
Attendance Change Approval Routes
مسارات اعتماد تعديلات الحضور خارج المهلة
"""
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from typing import List, Optional
import re
from datetime import datetime, timezone, timedelta
from bson import ObjectId

from .deps import get_db, get_current_user, has_permission, log_activity
from models.permissions import Permission, UserRole

router = APIRouter()


# ─────────────────────────────────────────────────────────────
# Models
# ─────────────────────────────────────────────────────────────

class ChangeItem(BaseModel):
    student_id: str
    old_status: Optional[str] = None
    new_status: str
    reason: Optional[str] = None


class BulkChangeRequest(BaseModel):
    lecture_id: str
    changes: List[ChangeItem]
    reason: Optional[str] = None


class RejectPayload(BaseModel):
    review_notes: Optional[str] = None


class BatchActionPayload(BaseModel):
    request_ids: List[str]
    review_notes: Optional[str] = None


# ─────────────────────────────────────────────────────────────
# Internal Helper — يُستدعى من endpoint تسجيل الحضور
# ─────────────────────────────────────────────────────────────

async def create_change_requests_for_diff(
    lecture: dict,
    course: dict,
    new_records: list,
    current_user: dict,
    reason: Optional[str] = None,
) -> dict:
    """
    ينشئ طلبات اعتماد لكل تغيير في حالة طالب في محاضرة.
    يُقارن الحالات القديمة (من db.attendance) بالحالات الجديدة (new_records).

    Args:
        lecture: وثيقة المحاضرة
        course: وثيقة المقرر
        new_records: قائمة سجلات جديدة [{student_id, status}]
        current_user: المستخدم مقدّم الطلب
        reason: سبب اختياري لطلبات هذه الجلسة

    Returns:
        dict {created, skipped, request_ids}
    """
    db = get_db()
    lecture_id = str(lecture["_id"])
    course_id = str(course["_id"])
    faculty_id = course.get("faculty_id") or ""
    department_id = course.get("department_id") or ""

    # جلب الحالات الحالية للطلاب في هذه المحاضرة
    existing = await db.attendance.find({"lecture_id": lecture_id}).to_list(5000)
    current_status_by_student = {r["student_id"]: r.get("status") for r in existing}
    attendance_id_by_student = {r["student_id"]: str(r["_id"]) for r in existing}

    created = 0
    skipped = 0
    request_ids: List[str] = []

    now = datetime.now(timezone.utc)

    for record in new_records:
        sid = record.student_id if hasattr(record, "student_id") else record.get("student_id")
        new_st = record.status if hasattr(record, "status") else record.get("status")
        old_st = current_status_by_student.get(sid)

        # لا فرق فعلي → تخطَّ
        if old_st == new_st:
            skipped += 1
            continue

        # ألغِ أي طلب pending سابق لنفس (lecture, student) — يُستبدل بالجديد
        await db.attendance_change_requests.update_many(
            {
                "lecture_id": lecture_id,
                "student_id": sid,
                "status": "pending",
            },
            {"$set": {
                "status": "cancelled",
                "reviewed_at": now,
                "review_notes": "استُبدل بطلب جديد",
            }}
        )

        # جلب اسم الطالب
        try:
            student = await db.students.find_one({"_id": ObjectId(sid)})
        except Exception:
            student = None
        student_name = (student.get("full_name") if student else "") or ""

        doc = {
            "attendance_id": attendance_id_by_student.get(sid),
            "lecture_id": lecture_id,
            "course_id": course_id,
            "course_name": course.get("name", ""),
            "course_code": course.get("code", ""),
            "lecture_date": lecture.get("date", ""),
            "lecture_start_time": lecture.get("start_time", ""),
            "student_id": sid,
            "student_name": student_name,
            "faculty_id": faculty_id,
            "department_id": department_id,
            "old_status": old_st,
            "new_status": new_st,
            "reason": reason,
            "requested_by": current_user["id"],
            "requested_by_name": current_user.get("full_name", ""),
            "requested_by_role": current_user.get("role", ""),
            "requested_at": now,
            "status": "pending",
            "reviewed_by": None,
            "reviewed_by_name": None,
            "reviewed_at": None,
            "review_notes": None,
        }
        result = await db.attendance_change_requests.insert_one(doc)
        created += 1
        request_ids.append(str(result.inserted_id))

    if created > 0:
        await log_activity(
            current_user, "request_attendance_changes", "attendance_change_request",
            None,
            f"طلب اعتماد {created} تعديل حضور في محاضرة {course.get('name', '')}",
            {"lecture_id": lecture_id, "created": created, "skipped": skipped}
        )

    return {"created": created, "skipped": skipped, "request_ids": request_ids}


# ─────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────

def _serialize_request(doc: dict) -> dict:
    return {
        "id": str(doc["_id"]),
        "attendance_id": doc.get("attendance_id"),
        "lecture_id": doc.get("lecture_id"),
        "course_id": doc.get("course_id"),
        "course_name": doc.get("course_name", ""),
        "course_code": doc.get("course_code", ""),
        "lecture_date": doc.get("lecture_date", ""),
        "lecture_start_time": doc.get("lecture_start_time", ""),
        "student_id": doc.get("student_id"),
        "student_name": doc.get("student_name", ""),
        "faculty_id": doc.get("faculty_id"),
        "department_id": doc.get("department_id"),
        "old_status": doc.get("old_status"),
        "new_status": doc.get("new_status"),
        "reason": doc.get("reason"),
        "requested_by": doc.get("requested_by"),
        "requested_by_name": doc.get("requested_by_name", ""),
        "requested_by_role": doc.get("requested_by_role", ""),
        "requested_at": doc.get("requested_at").isoformat() if doc.get("requested_at") else None,
        "status": doc.get("status"),
        "reviewed_by": doc.get("reviewed_by"),
        "reviewed_by_name": doc.get("reviewed_by_name", ""),
        "reviewed_at": doc.get("reviewed_at").isoformat() if doc.get("reviewed_at") else None,
        "review_notes": doc.get("review_notes"),
    }


def _can_approve(user: dict) -> bool:
    """المدير أو من يملك APPROVE_ATTENDANCE_CHANGES."""
    if user["role"] == UserRole.ADMIN:
        return True
    return has_permission(user, Permission.APPROVE_ATTENDANCE_CHANGES)


async def _apply_change_to_attendance(request_doc: dict, current_user: dict):
    """يطبّق التغيير المعتَمد على جدول attendance."""
    db = get_db()
    lecture_id = request_doc["lecture_id"]
    student_id = request_doc["student_id"]
    new_status = request_doc["new_status"]

    # ابحث عن السجل الأصلي (قد يكون attendance_id قديم أو أُعيد إنشاؤه)
    existing = await db.attendance.find_one({
        "lecture_id": lecture_id,
        "student_id": student_id,
    })
    now = datetime.now(timezone.utc)
    if existing:
        await db.attendance.update_one(
            {"_id": existing["_id"]},
            {"$set": {
                "status": new_status,
                "updated_at": now,
                "updated_by": current_user["id"],
                "last_change_approved_by": current_user["id"],
                "last_change_request_id": str(request_doc["_id"]),
            }}
        )
    else:
        # إن لم يوجد سجل حالي، أنشئه
        await db.attendance.insert_one({
            "lecture_id": lecture_id,
            "course_id": request_doc.get("course_id"),
            "student_id": student_id,
            "status": new_status,
            "date": now,
            "recorded_by": request_doc.get("requested_by"),
            "method": "approval_change",
            "created_at": now,
            "last_change_approved_by": current_user["id"],
            "last_change_request_id": str(request_doc["_id"]),
        })


# ─────────────────────────────────────────────────────────────
# GET Endpoints
# ─────────────────────────────────────────────────────────────

def _scope_query(current_user: dict) -> dict:
    """نطاق العميد → كليته فقط."""
    if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
        return {"faculty_id": current_user["faculty_id"]}
    return {}


def _rx(s: str) -> dict:
    return {"$regex": re.escape(s.strip()), "$options": "i"}


async def _build_filter_query(
    db, current_user: dict, *, lecture_id=None, faculty_id=None, department_id=None,
    level=None, section=None, course_id=None, semester_id=None, requested_by=None,
    new_status=None, q=None, lecture_from=None, lecture_to=None,
    requested_from=None, requested_to=None,
) -> dict:
    """يبني استعلام Mongo من كل الفلاتر (بدون فلتر الحالة)."""
    query: dict = dict(_scope_query(current_user))
    if lecture_id:
        query["lecture_id"] = lecture_id
    if faculty_id and "faculty_id" not in query:
        query["faculty_id"] = faculty_id
    if department_id:
        query["department_id"] = department_id
    if requested_by:
        query["requested_by"] = requested_by
    if new_status:
        query["new_status"] = new_status

    # المستوى / الشعبة / الفصل → عبر المقررات
    course_ids: Optional[set] = None
    if level or section or semester_id:
        cq: dict = {}
        if level:
            cq["level"] = int(level)
        if section:
            cq["section"] = section
        if semester_id:
            cq["semester_id"] = semester_id
        if department_id:
            cq["department_id"] = department_id
        if query.get("faculty_id"):
            cq["faculty_id"] = query["faculty_id"]
        cdocs = await db.courses.find(cq, {"_id": 1}).to_list(5000)
        course_ids = {str(c["_id"]) for c in cdocs}
    if course_id:
        course_ids = {course_id} & course_ids if course_ids is not None else {course_id}
    if course_ids is not None:
        query["course_id"] = {"$in": list(course_ids)} if course_ids else "__none__"

    if lecture_from or lecture_to:
        rng: dict = {}
        if lecture_from:
            rng["$gte"] = lecture_from
        if lecture_to:
            rng["$lte"] = lecture_to
        query["lecture_date"] = rng
    if requested_from or requested_to:
        rng = {}
        if requested_from:
            rng["$gte"] = datetime.fromisoformat(requested_from).replace(tzinfo=timezone.utc)
        if requested_to:
            rng["$lte"] = datetime.fromisoformat(requested_to).replace(hour=23, minute=59, second=59, tzinfo=timezone.utc)
        query["requested_at"] = rng

    if q and q.strip():
        ors = [
            {"student_name": _rx(q)},
            {"course_name": _rx(q)},
            {"course_code": _rx(q)},
            {"requested_by_name": _rx(q)},
        ]
        sdocs = await db.students.find({"student_id": _rx(q)}, {"_id": 1}).to_list(500)
        if sdocs:
            ors.append({"student_id": {"$in": [str(s["_id"]) for s in sdocs]}})
        query["$or"] = ors
    return query


_SORTS = {
    "newest": [("requested_at", -1)],
    "oldest": [("requested_at", 1)],
    "lecture_date": [("lecture_date", -1), ("lecture_start_time", -1)],
    "student_name": [("student_name", 1)],
}


async def _enrich_items(db, items: List[dict]) -> List[dict]:
    """يضيف المستوى/الشعبة/الرقم الجامعي/أسماء الكلية والقسم."""
    if not items:
        return items
    cids = {i["course_id"] for i in items if i.get("course_id")}
    sids = {i["student_id"] for i in items if i.get("student_id")}
    fids = {i["faculty_id"] for i in items if i.get("faculty_id")}
    dids = {i["department_id"] for i in items if i.get("department_id")}

    def _oids(ids):
        out = []
        for x in ids:
            try:
                out.append(ObjectId(x))
            except Exception:
                pass
        return out

    courses = {str(c["_id"]): c for c in await db.courses.find({"_id": {"$in": _oids(cids)}}, {"level": 1, "section": 1}).to_list(5000)}
    students = {str(s["_id"]): s for s in await db.students.find({"_id": {"$in": _oids(sids)}}, {"student_id": 1}).to_list(5000)}
    faculties = {str(f["_id"]): f.get("name", "") for f in await db.faculties.find({"_id": {"$in": _oids(fids)}}, {"name": 1}).to_list(500)}
    departments = {str(d["_id"]): d.get("name", "") for d in await db.departments.find({"_id": {"$in": _oids(dids)}}, {"name": 1}).to_list(500)}
    for i in items:
        c = courses.get(i.get("course_id") or "", {})
        i["level"] = c.get("level")
        i["section"] = c.get("section") or ""
        i["student_number"] = (students.get(i.get("student_id") or "", {}) or {}).get("student_id", "")
        i["faculty_name"] = faculties.get(i.get("faculty_id") or "", "")
        i["department_name"] = departments.get(i.get("department_id") or "", "")
    return items


@router.get("/attendance-changes")
async def list_change_requests(
    status: Optional[str] = "pending",
    lecture_id: Optional[str] = None,
    faculty_id: Optional[str] = None,
    department_id: Optional[str] = None,
    level: Optional[str] = None,
    section: Optional[str] = None,
    course_id: Optional[str] = None,
    semester_id: Optional[str] = None,
    requested_by: Optional[str] = None,
    new_status: Optional[str] = None,
    q: Optional[str] = None,
    lecture_from: Optional[str] = None,
    lecture_to: Optional[str] = None,
    requested_from: Optional[str] = None,
    requested_to: Optional[str] = None,
    sort: str = "newest",
    page: Optional[int] = None,
    page_size: int = 50,
    current_user: dict = Depends(get_current_user),
):
    """
    قائمة طلبات اعتماد التعديل مع فلاتر وترقيم صفحات.
    - المدير: يرى الكل. العميد: كليته فقط. غير ذلك: 403.
    - بدون `page` → يعيد كل النتائج (توافق مع التطبيق).
    """
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك بعرض طلبات الاعتماد")

    db = get_db()
    base = await _build_filter_query(
        db, current_user, lecture_id=lecture_id, faculty_id=faculty_id, department_id=department_id,
        level=level, section=section, course_id=course_id, semester_id=semester_id,
        requested_by=requested_by, new_status=new_status, q=q, lecture_from=lecture_from,
        lecture_to=lecture_to, requested_from=requested_from, requested_to=requested_to,
    )
    query = dict(base)
    if status and status != "all":
        query["status"] = status

    # عدّادات الحالات على نفس الفلاتر (بدون الحالة)
    agg = await db.attendance_change_requests.aggregate([
        {"$match": base}, {"$group": {"_id": "$status", "n": {"$sum": 1}}},
    ]).to_list(20)
    counts = {"pending": 0, "approved": 0, "rejected": 0, "cancelled": 0}
    for a in agg:
        counts[a["_id"] or "unknown"] = a["n"]
    counts["all"] = sum(v for k, v in counts.items() if k != "all")

    total = await db.attendance_change_requests.count_documents(query)
    cursor = db.attendance_change_requests.find(query).sort(_SORTS.get(sort, _SORTS["newest"]))
    if page:
        page_size = max(1, min(page_size, 200))
        cursor = cursor.skip((page - 1) * page_size).limit(page_size)
        docs = await cursor.to_list(page_size)
    else:
        docs = await cursor.to_list(2000)
    items = await _enrich_items(db, [_serialize_request(d) for d in docs])
    return {
        "items": items, "total": total, "counts": counts,
        "page": page or 1, "page_size": page_size if page else total,
        "pages": max(1, -(-total // page_size)) if page else 1,
    }


@router.get("/attendance-changes/filter-options")
async def filter_options(current_user: dict = Depends(get_current_user)):
    """خيارات الفلاتر المستمدة من الطلبات الموجودة فعلاً (كليات/أقسام/مقررات/مقدّمون/فصول)."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    scope = _scope_query(current_user)
    rows = await db.attendance_change_requests.aggregate([
        {"$match": scope},
        {"$group": {
            "_id": "$course_id",
            "faculty_id": {"$first": "$faculty_id"},
            "department_id": {"$first": "$department_id"},
            "course_name": {"$first": "$course_name"},
            "course_code": {"$first": "$course_code"},
            "n": {"$sum": 1},
        }},
    ]).to_list(5000)
    requesters = await db.attendance_change_requests.aggregate([
        {"$match": scope},
        {"$group": {"_id": "$requested_by", "name": {"$first": "$requested_by_name"}, "role": {"$first": "$requested_by_role"}, "n": {"$sum": 1}}},
        {"$sort": {"name": 1}},
    ]).to_list(2000)

    oids = []
    for r in rows:
        try:
            oids.append(ObjectId(r["_id"]))
        except Exception:
            pass
    cmap = {str(c["_id"]): c for c in await db.courses.find({"_id": {"$in": oids}}, {"level": 1, "section": 1, "semester_id": 1}).to_list(5000)}

    fids = {r["faculty_id"] for r in rows if r.get("faculty_id")}
    dids = {r["department_id"] for r in rows if r.get("department_id")}
    sem_ids = {c.get("semester_id") for c in cmap.values() if c.get("semester_id")}

    def _oids(ids):
        out = []
        for x in ids:
            try:
                out.append(ObjectId(x))
            except Exception:
                pass
        return out

    faculties = [{"id": str(f["_id"]), "name": f.get("name", "")} for f in await db.faculties.find({"_id": {"$in": _oids(fids)}}, {"name": 1}).sort("name", 1).to_list(500)]
    departments = [{"id": str(d["_id"]), "name": d.get("name", ""), "faculty_id": d.get("faculty_id")} for d in await db.departments.find({"_id": {"$in": _oids(dids)}}, {"name": 1, "faculty_id": 1}).sort("name", 1).to_list(500)]
    semesters = [{"id": str(s["_id"]), "name": s.get("name", ""), "is_active": bool(s.get("is_active"))} for s in await db.semesters.find({"_id": {"$in": _oids(sem_ids)}}, {"name": 1, "is_active": 1}).to_list(100)]

    courses = []
    for r in rows:
        c = cmap.get(r["_id"] or "", {})
        courses.append({
            "id": r["_id"], "name": r.get("course_name", ""), "code": r.get("course_code", ""),
            "faculty_id": r.get("faculty_id"), "department_id": r.get("department_id"),
            "level": c.get("level"), "section": c.get("section") or "", "semester_id": c.get("semester_id"),
            "count": r["n"],
        })
    courses.sort(key=lambda x: (x["name"] or ""))
    return {
        "faculties": faculties, "departments": departments, "courses": courses, "semesters": semesters,
        "requesters": [{"id": r["_id"], "name": r.get("name") or "", "role": r.get("role") or "", "count": r["n"]} for r in requesters if r["_id"]],
    }


@router.get("/attendance-changes/stats")
async def approval_stats(current_user: dict = Depends(get_current_user)):
    """بطاقات إحصائية: المعلّق لكل كلية، أقدم طلب معلّق، طلبات اليوم، المعالَج هذا الأسبوع."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    scope = _scope_query(current_user)
    now = datetime.now(timezone.utc)
    pending_q = {**scope, "status": "pending"}

    by_fac = await db.attendance_change_requests.aggregate([
        {"$match": pending_q}, {"$group": {"_id": "$faculty_id", "n": {"$sum": 1}}}, {"$sort": {"n": -1}},
    ]).to_list(100)
    fac_names = {}
    oids = []
    for b in by_fac:
        try:
            oids.append(ObjectId(b["_id"]))
        except Exception:
            pass
    for f in await db.faculties.find({"_id": {"$in": oids}}, {"name": 1}).to_list(200):
        fac_names[str(f["_id"])] = f.get("name", "")

    oldest = await db.attendance_change_requests.find_one(pending_q, sort=[("requested_at", 1)])
    oldest_days = None
    if oldest and oldest.get("requested_at"):
        ra = oldest["requested_at"]
        if ra.tzinfo is None:
            ra = ra.replace(tzinfo=timezone.utc)
        oldest_days = (now - ra).days

    start_today = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_ago = now - timedelta(days=7)
    today_n = await db.attendance_change_requests.count_documents({**scope, "requested_at": {"$gte": start_today}})
    reviewed_week = await db.attendance_change_requests.count_documents({**scope, "status": {"$in": ["approved", "rejected"]}, "reviewed_at": {"$gte": week_ago}})
    return {
        "pending_total": sum(b["n"] for b in by_fac),
        "pending_by_faculty": [{"faculty_id": b["_id"], "faculty_name": fac_names.get(b["_id"] or "", "غير محدد"), "count": b["n"]} for b in by_fac],
        "oldest_pending_days": oldest_days,
        "oldest_pending_at": oldest["requested_at"].isoformat() if oldest and oldest.get("requested_at") else None,
        "requested_today": today_n,
        "reviewed_last_7_days": reviewed_week,
    }


@router.get("/attendance-changes/export")
async def export_change_requests(
    status: Optional[str] = "pending",
    faculty_id: Optional[str] = None,
    department_id: Optional[str] = None,
    level: Optional[str] = None,
    section: Optional[str] = None,
    course_id: Optional[str] = None,
    semester_id: Optional[str] = None,
    requested_by: Optional[str] = None,
    new_status: Optional[str] = None,
    q: Optional[str] = None,
    lecture_from: Optional[str] = None,
    lecture_to: Optional[str] = None,
    requested_from: Optional[str] = None,
    requested_to: Optional[str] = None,
    sort: str = "newest",
    current_user: dict = Depends(get_current_user),
):
    """تصدير النتائج المفلترة إلى Excel."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    from io import BytesIO
    from urllib.parse import quote
    from fastapi.responses import StreamingResponse
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from .deps import export_stamp

    db = get_db()
    query = await _build_filter_query(
        db, current_user, faculty_id=faculty_id, department_id=department_id, level=level,
        section=section, course_id=course_id, semester_id=semester_id, requested_by=requested_by,
        new_status=new_status, q=q, lecture_from=lecture_from, lecture_to=lecture_to,
        requested_from=requested_from, requested_to=requested_to,
    )
    if status and status != "all":
        query["status"] = status
    docs = await db.attendance_change_requests.find(query).sort(_SORTS.get(sort, _SORTS["newest"])).to_list(5000)
    items = await _enrich_items(db, [_serialize_request(d) for d in docs])

    ST = {"present": "حاضر", "absent": "غائب", "late": "متأخر", "excused": "مأذون"}
    RQ = {"pending": "قيد الانتظار", "approved": "معتمد", "rejected": "مرفوض", "cancelled": "ملغي"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "طلبات تعديل الحضور"
    ws.sheet_view.rightToLeft = True
    BLUE = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
    GREY = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
    thin = Side(style='thin', color='B0BEC5')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    headers = ["#", "الطالب", "الرقم الجامعي", "المقرر", "الرمز", "الكلية", "القسم", "المستوى", "الشعبة",
               "تاريخ المحاضرة", "الوقت", "من", "إلى", "السبب", "مقدّم الطلب", "تاريخ الطلب", "الحالة", "المراجع", "ملاحظات المراجع"]
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=len(headers))
    t = ws.cell(row=1, column=1, value="طلبات تعديل الحضور")
    t.font = Font(bold=True, size=15, color='1565C0'); t.alignment = center
    ws.merge_cells(start_row=2, end_row=2, start_column=1, end_column=len(headers))
    s = ws.cell(row=2, column=1, value=f"الإجمالي: {len(items)} · الحالة: {RQ.get(status or '', 'الكل')}")
    s.font = Font(bold=True, size=11, color='5B6678'); s.alignment = center
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill = BLUE; c.font = Font(bold=True, color='FFFFFF', size=11); c.alignment = center; c.border = border
    row = 5
    for idx, it in enumerate(items, start=1):
        vals = [idx, it["student_name"], it.get("student_number", ""), it["course_name"], it["course_code"],
                it.get("faculty_name", ""), it.get("department_name", ""), it.get("level") or "", it.get("section") or "",
                it["lecture_date"], it["lecture_start_time"], ST.get(it["old_status"] or "", "—"), ST.get(it["new_status"], it["new_status"]),
                it.get("reason") or "", it["requested_by_name"], (it.get("requested_at") or "")[:16].replace("T", " "),
                RQ.get(it["status"], it["status"]), it.get("reviewed_by_name") or "", it.get("review_notes") or ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.alignment = center; c.border = border
            if idx % 2 == 0:
                c.fill = GREY
        row += 1
    widths = [4, 24, 12, 22, 10, 18, 18, 8, 8, 12, 8, 9, 9, 20, 20, 16, 12, 18, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    out = BytesIO(); wb.save(out); out.seek(0)
    fn = quote(f"طلبات تعديل الحضور - {export_stamp()}.xlsx")
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fn}", "X-Filename": fn})


@router.get("/attendance-changes/pending-count")
async def pending_count(current_user: dict = Depends(get_current_user)):
    """عدد الطلبات المعلّقة — للـ Badge على أيقونة الإشعارات."""
    if not _can_approve(current_user):
        return {"count": 0}
    db = get_db()
    query: dict = {"status": "pending"}
    if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
        query["faculty_id"] = current_user["faculty_id"]
    count = await db.attendance_change_requests.count_documents(query)
    return {"count": count}


@router.get("/attendance-changes/mine")
async def my_change_requests(
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    """تاريخ طلبات المستخدم الحالي."""
    db = get_db()
    query: dict = {"requested_by": current_user["id"]}
    if status and status != "all":
        query["status"] = status
    docs = await db.attendance_change_requests.find(query).sort("requested_at", -1).to_list(2000)
    return {"items": [_serialize_request(d) for d in docs], "total": len(docs)}


@router.get("/attendance-changes/lecture/{lecture_id}/pending")
async def pending_for_lecture(
    lecture_id: str,
    current_user: dict = Depends(get_current_user),
):
    """طلبات معلّقة لمحاضرة محددة — لعرض شارات ⏳ على الواجهة."""
    db = get_db()
    docs = await db.attendance_change_requests.find({
        "lecture_id": lecture_id,
        "status": "pending",
    }).to_list(2000)
    return {"items": [_serialize_request(d) for d in docs], "total": len(docs)}


# ─────────────────────────────────────────────────────────────
# Approve / Reject / Cancel
# ─────────────────────────────────────────────────────────────

@router.post("/attendance-changes/batch/approve")
async def batch_approve(
    payload: BatchActionPayload,
    current_user: dict = Depends(get_current_user),
):
    """اعتماد جماعي لعدة طلبات."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك بالاعتماد")
    db = get_db()
    approved = 0
    errors: List[dict] = []
    for req_id in payload.request_ids:
        try:
            doc = await db.attendance_change_requests.find_one({"_id": ObjectId(req_id)})
            if not doc or doc.get("status") != "pending":
                errors.append({"id": req_id, "reason": "غير موجود أو ليس معلّقاً"})
                continue
            if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
                if doc.get("faculty_id") and doc["faculty_id"] != current_user["faculty_id"]:
                    errors.append({"id": req_id, "reason": "خارج نطاقك"})
                    continue
            await _apply_change_to_attendance(doc, current_user)
            await db.attendance_change_requests.update_one(
                {"_id": doc["_id"]},
                {"$set": {
                    "status": "approved",
                    "reviewed_by": current_user["id"],
                    "reviewed_by_name": current_user.get("full_name", ""),
                    "reviewed_at": datetime.now(timezone.utc),
                }}
            )
            approved += 1
        except Exception as e:
            errors.append({"id": req_id, "reason": str(e)})
    await log_activity(
        current_user, "batch_approve_attendance_changes", "attendance_change_request",
        None, f"اعتماد جماعي: {approved} من {len(payload.request_ids)}"
    )
    return {"success": True, "approved": approved, "errors": errors}


@router.post("/attendance-changes/batch/reject")
async def batch_reject(
    payload: BatchActionPayload,
    current_user: dict = Depends(get_current_user),
):
    """رفض جماعي لعدة طلبات."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك بالرفض")
    db = get_db()
    rejected = 0
    errors: List[dict] = []
    for req_id in payload.request_ids:
        try:
            doc = await db.attendance_change_requests.find_one({"_id": ObjectId(req_id)})
            if not doc or doc.get("status") != "pending":
                errors.append({"id": req_id, "reason": "غير موجود أو ليس معلّقاً"})
                continue
            if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
                if doc.get("faculty_id") and doc["faculty_id"] != current_user["faculty_id"]:
                    errors.append({"id": req_id, "reason": "خارج نطاقك"})
                    continue
            await db.attendance_change_requests.update_one(
                {"_id": doc["_id"]},
                {"$set": {
                    "status": "rejected",
                    "reviewed_by": current_user["id"],
                    "reviewed_by_name": current_user.get("full_name", ""),
                    "reviewed_at": datetime.now(timezone.utc),
                    "review_notes": payload.review_notes or "",
                }}
            )
            rejected += 1
        except Exception as e:
            errors.append({"id": req_id, "reason": str(e)})
    return {"success": True, "rejected": rejected, "errors": errors}


@router.post("/attendance-changes/{req_id}/approve")
async def approve_change(
    req_id: str,
    current_user: dict = Depends(get_current_user),
):
    """اعتماد طلب تعديل حضور."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك بالاعتماد")

    db = get_db()
    doc = await db.attendance_change_requests.find_one({"_id": ObjectId(req_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    if doc.get("status") != "pending":
        raise HTTPException(status_code=400, detail=f"الطلب ليس معلّقاً (الحالة: {doc.get('status')})")

    # نطاق العميد: كليته فقط
    if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
        if doc.get("faculty_id") and doc["faculty_id"] != current_user["faculty_id"]:
            raise HTTPException(status_code=403, detail="خارج نطاق كليتك")

    # طبّق التغيير على attendance
    await _apply_change_to_attendance(doc, current_user)

    # حدّث حالة الطلب
    await db.attendance_change_requests.update_one(
        {"_id": doc["_id"]},
        {"$set": {
            "status": "approved",
            "reviewed_by": current_user["id"],
            "reviewed_by_name": current_user.get("full_name", ""),
            "reviewed_at": datetime.now(timezone.utc),
        }}
    )
    await log_activity(
        current_user, "approve_attendance_change", "attendance_change_request",
        req_id,
        f"اعتماد تعديل حضور الطالب {doc.get('student_name', '')} → {doc.get('new_status')}"
    )
    return {"success": True, "message": "تم اعتماد الطلب وتطبيق التغيير"}


@router.post("/attendance-changes/{req_id}/reject")
async def reject_change(
    req_id: str,
    payload: RejectPayload,
    current_user: dict = Depends(get_current_user),
):
    """رفض طلب تعديل حضور."""
    if not _can_approve(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك بالرفض")

    db = get_db()
    doc = await db.attendance_change_requests.find_one({"_id": ObjectId(req_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    if doc.get("status") != "pending":
        raise HTTPException(status_code=400, detail=f"الطلب ليس معلّقاً (الحالة: {doc.get('status')})")
    if current_user["role"] != UserRole.ADMIN and current_user.get("faculty_id"):
        if doc.get("faculty_id") and doc["faculty_id"] != current_user["faculty_id"]:
            raise HTTPException(status_code=403, detail="خارج نطاق كليتك")

    await db.attendance_change_requests.update_one(
        {"_id": doc["_id"]},
        {"$set": {
            "status": "rejected",
            "reviewed_by": current_user["id"],
            "reviewed_by_name": current_user.get("full_name", ""),
            "reviewed_at": datetime.now(timezone.utc),
            "review_notes": payload.review_notes or "",
        }}
    )
    await log_activity(
        current_user, "reject_attendance_change", "attendance_change_request",
        req_id,
        f"رفض تعديل حضور {doc.get('student_name', '')}: {payload.review_notes or ''}"
    )
    return {"success": True, "message": "تم رفض الطلب"}


@router.delete("/attendance-changes/{req_id}")
async def cancel_request(
    req_id: str,
    current_user: dict = Depends(get_current_user),
):
    """إلغاء ذاتي — مقدّم الطلب فقط، وهو معلّق."""
    db = get_db()
    doc = await db.attendance_change_requests.find_one({"_id": ObjectId(req_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    if doc.get("status") != "pending":
        raise HTTPException(status_code=400, detail="لا يمكن إلغاء طلب غير معلّق")
    if doc.get("requested_by") != current_user["id"]:
        raise HTTPException(status_code=403, detail="لا يمكنك إلغاء طلب غيرك")

    await db.attendance_change_requests.update_one(
        {"_id": doc["_id"]},
        {"$set": {
            "status": "cancelled",
            "reviewed_at": datetime.now(timezone.utc),
            "review_notes": "أُلغِي بواسطة المقدّم",
        }}
    )
    return {"success": True, "message": "تم إلغاء الطلب"}
