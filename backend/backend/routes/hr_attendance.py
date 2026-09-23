"""🕐 شؤون الموظفين — الحضور الإداري: إعدادات الدوام، الكشف اليومي، تسجيل ذاتي، التقرير الشهري"""
import io
from datetime import datetime, timedelta, date
from typing import Optional, List

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, log_activity, export_headers, export_filename
from .hr_common import (P_ATTEND, YEMEN_TZ, AR_DAYS, DEFAULT_SETTINGS, _now, _today, _oid, _ser, _can_view, _guard, parse_date, get_hr_settings,
                        is_work_day, holiday_name, find_my_employee, enrich_employee_refs)

router = APIRouter(prefix="/hr/attendance", tags=["شؤون الموظفين - الحضور الإداري"])

ATT_STATUS = {"present": "حاضر", "late": "متأخر", "half_day": "نصف يوم", "absent": "غائب", "excused": "غياب بعذر", "leave": "إجازة", "mission": "مهمة رسمية", "holiday": "عطلة"}
MANUAL = ("present", "late", "half_day", "absent", "excused", "mission")
COUNTED_PRESENT = ("present", "late", "half_day", "mission")


class SettingsIn(BaseModel):
    work_days: List[str]
    work_start: str
    work_end: str
    late_grace_minutes: int = 15
    early_leave_grace_minutes: int = 0
    allow_self_checkin: bool = True
    annual_leave_days: int = 30
    holidays: List[dict] = []


class Entry(BaseModel):
    employee_id: str
    status: str
    check_in: Optional[str] = None
    check_out: Optional[str] = None
    note: Optional[str] = ""


class MarkIn(BaseModel):
    date: str
    entries: List[Entry]


class DateIn(BaseModel):
    date: str


def _hm(s: str, what: str = "الوقت") -> int:
    try:
        h, m = s.split(":")[:2]
        return int(h) * 60 + int(m)
    except Exception:
        raise HTTPException(status_code=400, detail=f"{what} غير صحيح (HH:MM)")


def _late(check_in: Optional[str], settings: dict) -> int:
    if not check_in:
        return 0
    return max(0, _hm(check_in) - _hm(settings["work_start"]) - int(settings.get("late_grace_minutes", 0)))


async def _leave_map(db, d: str) -> dict:
    """employee_id -> نوع الإجازة المعتمدة في هذا اليوم"""
    from .hr_leaves import LEAVE_TYPES
    return {l["employee_id"]: LEAVE_TYPES.get(l.get("type"), "إجازة") async for l in db.leave_requests.find({"status": "approved", "start_date": {"$lte": d}, "end_date": {"$gte": d}}, {"employee_id": 1, "type": 1})}


@router.get("/meta")
async def attendance_meta(current_user: dict = Depends(get_current_user)):
    return {"statuses": ATT_STATUS, "manual": list(MANUAL), "days": list(AR_DAYS.values())}


@router.get("/settings")
async def get_settings(current_user: dict = Depends(get_current_user)):
    return await get_hr_settings(get_db())


@router.put("/settings")
async def put_settings(data: SettingsIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ATTEND)
    db = get_db()
    if _hm(data.work_start, "بداية الدوام") >= _hm(data.work_end, "نهاية الدوام"):
        raise HTTPException(status_code=400, detail="نهاية الدوام يجب أن تكون بعد بدايته")
    bad = [d for d in data.work_days if d not in AR_DAYS.values()]
    if bad or not data.work_days:
        raise HTTPException(status_code=400, detail="أيام العمل غير صحيحة")
    hol = []
    for h in data.holidays:
        if not h.get("date"):
            continue
        parse_date(h["date"], "تاريخ العطلة")
        hol.append({"date": h["date"][:10], "name": (h.get("name") or "عطلة رسمية").strip()})
    doc = {**data.dict(), "holidays": sorted(hol, key=lambda x: x["date"]), "updated_by_name": current_user.get("full_name", ""), "updated_at": _now()}
    await db.hr_settings.update_one({"_id": "global"}, {"$set": doc}, upsert=True)
    await log_activity(current_user, "hr_attendance_settings", "hr_settings", "global", "إعدادات الدوام")
    return {"message": "تم حفظ إعدادات الدوام", "settings": await get_hr_settings(db)}


async def _daily_rows(db, d: str, org_unit_id: Optional[str], category: Optional[str], settings: dict) -> List[dict]:
    q: dict = {"status": {"$nin": ["ended", "suspended"]}}
    if org_unit_id:
        q["org_unit_id"] = org_unit_id
    if category:
        q["category"] = category
    emps = await db.employees.find(q, {"full_name": 1, "employee_no": 1, "org_unit_id": 1, "job_title": 1, "category": 1}).sort("full_name", 1).to_list(5000)
    recs = {r["employee_id"]: r for r in await db.hr_attendance.find({"date": d}).to_list(10000)}
    leaves = await _leave_map(db, d)
    day = parse_date(d)
    workday, hol = is_work_day(day, settings), holiday_name(day, settings)
    rows = []
    for e in emps:
        eid = str(e["_id"])
        r = recs.get(eid)
        if r:
            row = {**_ser(r), "recorded": True}
        elif eid in leaves:
            row = {"employee_id": eid, "status": "leave", "note": leaves[eid], "recorded": False}
        elif not workday:
            row = {"employee_id": eid, "status": "holiday", "note": hol or "خارج أيام العمل", "recorded": False}
        else:
            row = {"employee_id": eid, "status": None, "recorded": False}
        row["status_label"] = ATT_STATUS.get(row.get("status") or "", "لم يُسجَّل")
        row["category"] = e.get("category")
        rows.append(row)
    await enrich_employee_refs(db, rows)
    return rows


@router.get("/daily")
async def daily_sheet(date: Optional[str] = None, org_unit_id: Optional[str] = None, category: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    d = date or _today()
    day = parse_date(d)
    settings = await get_hr_settings(db)
    rows = await _daily_rows(db, d, org_unit_id, category, settings)
    summary = {k: 0 for k in [*ATT_STATUS, "unmarked"]}
    for r in rows:
        summary[r["status"] or "unmarked"] += 1
    return {"date": d, "day_name": AR_DAYS[day.weekday()], "is_work_day": is_work_day(day, settings), "holiday": holiday_name(day, settings), "settings": settings, "rows": rows, "summary": summary, "total": len(rows)}


@router.post("/mark")
async def mark_attendance(data: MarkIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ATTEND)
    db = get_db()
    parse_date(data.date)
    if data.date > _today():
        raise HTTPException(status_code=400, detail="لا يمكن تسجيل حضور لتاريخ مستقبلي")
    settings = await get_hr_settings(db)
    saved = 0
    for en in data.entries:
        if en.status not in MANUAL:
            raise HTTPException(status_code=400, detail=f"حالة غير مسموحة: {en.status}")
        _oid(en.employee_id, "الموظف")
        ci = en.check_in or (settings["work_start"] if en.status in ("present", "late", "half_day", "mission") else None)
        co = en.check_out or None
        if ci: _hm(ci, "وقت الحضور")
        if co: _hm(co, "وقت الانصراف")
        late = _late(ci, settings) if en.status in ("present", "late") else 0
        status = "late" if (en.status == "present" and late > 0) else en.status
        await db.hr_attendance.update_one({"employee_id": en.employee_id, "date": data.date}, {"$set": {
            "status": status, "check_in": ci, "check_out": co, "late_minutes": late, "note": (en.note or "").strip(), "source": "manual",
            "by_name": current_user.get("full_name", ""), "updated_at": _now()}, "$setOnInsert": {"created_at": _now()}}, upsert=True)
        saved += 1
    await log_activity(current_user, "hr_attendance_mark", "hr_attendance", data.date, "تسجيل حضور إداري", {"count": saved})
    return {"saved": saved, "message": f"تم حفظ {saved} سجلاً"}


@router.post("/mark-all-present")
async def mark_all_present(data: DateIn, org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """✅ تحديد كل من لم يُسجَّل (وليس في إجازة) حاضراً في وقت الدوام الرسمي"""
    _guard(current_user, P_ATTEND)
    db = get_db()
    if data.date > _today():
        raise HTTPException(status_code=400, detail="لا يمكن تسجيل حضور لتاريخ مستقبلي")
    settings = await get_hr_settings(db)
    if not is_work_day(parse_date(data.date), settings):
        raise HTTPException(status_code=400, detail="هذا اليوم ليس يوم عمل")
    rows = await _daily_rows(db, data.date, org_unit_id, None, settings)
    n = 0
    for r in rows:
        if r["status"] is None:
            await db.hr_attendance.insert_one({"employee_id": r["employee_id"], "date": data.date, "status": "present", "check_in": settings["work_start"], "check_out": settings["work_end"], "late_minutes": 0, "note": "", "source": "bulk", "by_name": current_user.get("full_name", ""), "created_at": _now(), "updated_at": _now()})
            n += 1
    await log_activity(current_user, "hr_attendance_bulk", "hr_attendance", data.date, "تحديد الكل حاضر", {"count": n})
    return {"marked": n, "message": f"تم تحديد {n} موظفاً حاضراً"}


@router.delete("/record/{employee_id}/{d}")
async def delete_record(employee_id: str, d: str, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ATTEND)
    r = await get_db().hr_attendance.delete_one({"employee_id": employee_id, "date": d})
    return {"deleted": r.deleted_count, "message": "تم حذف السجل"}


async def _monthly(db, month: str, org_unit_id: Optional[str], category: Optional[str]) -> dict:
    try:
        y, m = int(month[:4]), int(month[5:7])
        first = date(y, m, 1)
    except Exception:
        raise HTTPException(status_code=400, detail="الشهر غير صحيح (YYYY-MM)")
    last = (date(y + (m == 12), (m % 12) + 1, 1) - timedelta(days=1))
    today = datetime.now(YEMEN_TZ).date()
    end = min(last, today)
    settings = await get_hr_settings(db)
    work_dates = [(first + timedelta(days=i)).strftime("%Y-%m-%d") for i in range((end - first).days + 1) if is_work_day(first + timedelta(days=i), settings)] if end >= first else []
    q: dict = {"status": {"$nin": ["ended"]}}
    if org_unit_id:
        q["org_unit_id"] = org_unit_id
    if category:
        q["category"] = category
    emps = await db.employees.find(q, {"full_name": 1, "employee_no": 1, "org_unit_id": 1, "job_title": 1, "hire_date": 1}).sort("full_name", 1).to_list(5000)
    recs: dict = {}
    async for r in db.hr_attendance.find({"date": {"$gte": first.strftime("%Y-%m-%d"), "$lte": end.strftime("%Y-%m-%d")}}):
        recs.setdefault(r["employee_id"], {})[r["date"]] = r
    leaves: dict = {}
    async for l in db.leave_requests.find({"status": "approved", "start_date": {"$lte": end.strftime("%Y-%m-%d")}, "end_date": {"$gte": first.strftime("%Y-%m-%d")}}, {"employee_id": 1, "start_date": 1, "end_date": 1}):
        leaves.setdefault(l["employee_id"], []).append((l["start_date"], l["end_date"]))
    items = []
    for e in emps:
        eid = str(e["_id"])
        c = {k: 0 for k in ATT_STATUS}
        c["unmarked"] = 0
        late_min = 0
        for d in work_dates:
            if e.get("hire_date") and d < e["hire_date"]:
                continue
            r = recs.get(eid, {}).get(d)
            if r:
                c[r["status"]] = c.get(r["status"], 0) + 1
                late_min += int(r.get("late_minutes") or 0)
            elif any(s <= d <= en for s, en in leaves.get(eid, [])):
                c["leave"] += 1
            else:
                c["unmarked"] += 1
        # سجلات في أيام غير عمل (مهمات مثلاً) تُحسب حضوراً إضافياً
        extra = sum(1 for d, r in recs.get(eid, {}).items() if d not in work_dates and r["status"] in COUNTED_PRESENT)
        present_total = sum(c[k] for k in COUNTED_PRESENT)
        due = max(0, len(work_dates) - c["leave"] - c["excused"])
        items.append({"employee_id": eid, "counts": c, "present_total": present_total, "extra_days": extra, "late_minutes": late_min, "due_days": due,
                      "rate": round(present_total / due * 100, 1) if due else None})
    await enrich_employee_refs(db, items)
    return {"month": month, "work_days": len(work_dates), "from": first.strftime("%Y-%m-%d"), "to": end.strftime("%Y-%m-%d"), "items": items,
            "totals": {"present": sum(i["present_total"] for i in items), "absent": sum(i["counts"]["absent"] for i in items), "late": sum(i["counts"]["late"] for i in items), "leave": sum(i["counts"]["leave"] for i in items)}}


@router.get("/monthly")
async def monthly_report(month: Optional[str] = None, org_unit_id: Optional[str] = None, category: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    return await _monthly(get_db(), month or _today()[:7], org_unit_id, category)


@router.get("/monthly/export")
async def monthly_export(month: Optional[str] = None, org_unit_id: Optional[str] = None, category: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    rep = await _monthly(get_db(), month or _today()[:7], org_unit_id, category)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "الحضور الإداري"; ws.sheet_view.rightToLeft = True
    ws.append([f"تقرير الحضور الإداري — {rep['month']} (أيام العمل: {rep['work_days']})"]); ws.merge_cells("A1:L1")
    ws["A1"].font = Font(bold=True, size=13); ws["A1"].alignment = Alignment(horizontal="center")
    heads = ["الرقم الوظيفي", "الاسم", "الوحدة", "المسمى", "حاضر", "متأخر", "نصف يوم", "مهمة", "غائب", "بعذر", "إجازة", "لم يُسجَّل", "دقائق التأخير", "نسبة الحضور %"]
    ws.append(heads)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440"); c.alignment = Alignment(horizontal="center")
    for i in rep["items"]:
        c = i["counts"]
        ws.append([i["employee_no"], i["employee_name"], i["org_unit_name"], i["job_title"], c["present"], c["late"], c["half_day"], c["mission"], c["absent"], c["excused"], c["leave"], c["unmarked"], i["late_minutes"], i["rate"] if i["rate"] is not None else "—"])
    from openpyxl.utils import get_column_letter
    for ci in range(1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 16
    ws.column_dimensions["B"].width = 30
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=export_headers(export_filename("الحضور الإداري", rep["month"], ext="xlsx")))


@router.get("/employee/{employee_id}")
async def employee_records(employee_id: str, month: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    m = month or _today()[:7]
    recs = [_ser(r) for r in await db.hr_attendance.find({"employee_id": employee_id, "date": {"$regex": f"^{m}"}}).sort("date", -1).to_list(100)]
    for r in recs:
        r["status_label"] = ATT_STATUS.get(r["status"], r["status"])
    return {"month": m, "records": recs}


# ══════════════ الخدمة الذاتية ══════════════

@router.get("/my")
async def my_attendance(month: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    db = get_db()
    emp = await find_my_employee(db, current_user)
    if not emp:
        return {"profile": None}
    settings = await get_hr_settings(db)
    t = _today()
    m = month or t[:7]
    eid = str(emp["_id"])
    today_rec = await db.hr_attendance.find_one({"employee_id": eid, "date": t})
    leaves = await _leave_map(db, t)
    day = parse_date(t)
    recs = [_ser(r) for r in await db.hr_attendance.find({"employee_id": eid, "date": {"$regex": f"^{m}"}}).sort("date", -1).to_list(60)]
    for r in recs:
        r["status_label"] = ATT_STATUS.get(r["status"], r["status"])
    now_hm = datetime.now(YEMEN_TZ).strftime("%H:%M")
    on_leave = eid in leaves
    can_in = settings.get("allow_self_checkin", True) and is_work_day(day, settings) and not on_leave and not today_rec
    can_out = settings.get("allow_self_checkin", True) and bool(today_rec) and today_rec.get("source") in ("self", "manual", "bulk") and not today_rec.get("check_out")
    counts = {k: 0 for k in ATT_STATUS}
    for r in recs:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    return {"profile": {"id": eid, "full_name": emp.get("full_name", "")}, "date": t, "day_name": AR_DAYS[day.weekday()], "now": now_hm, "is_work_day": is_work_day(day, settings),
            "holiday": holiday_name(day, settings), "on_leave": leaves.get(eid), "today": ({**_ser(today_rec), "status_label": ATT_STATUS.get(today_rec["status"], "")} if today_rec else None),
            "can_check_in": can_in, "can_check_out": can_out, "settings": {k: settings[k] for k in ("work_start", "work_end", "late_grace_minutes", "allow_self_checkin")},
            "month": m, "records": recs, "counts": counts, "late_minutes": sum(int(r.get("late_minutes") or 0) for r in recs)}


@router.post("/check-in")
async def check_in(current_user: dict = Depends(get_current_user)):
    db = get_db()
    emp = await find_my_employee(db, current_user)
    if not emp:
        raise HTTPException(status_code=404, detail="لا يوجد ملف إداري مرتبط بحسابك")
    settings = await get_hr_settings(db)
    if not settings.get("allow_self_checkin", True):
        raise HTTPException(status_code=400, detail="التسجيل الذاتي غير مفعّل — يُسجَّل الحضور من شؤون الموظفين")
    t = _today()
    if not is_work_day(parse_date(t), settings):
        raise HTTPException(status_code=400, detail="اليوم ليس يوم عمل")
    eid = str(emp["_id"])
    if eid in await _leave_map(db, t):
        raise HTTPException(status_code=400, detail="أنت في إجازة معتمدة اليوم")
    if await db.hr_attendance.find_one({"employee_id": eid, "date": t}):
        raise HTTPException(status_code=400, detail="تم تسجيل حضورك اليوم بالفعل")
    now_hm = datetime.now(YEMEN_TZ).strftime("%H:%M")
    late = _late(now_hm, settings)
    doc = {"employee_id": eid, "date": t, "status": "late" if late else "present", "check_in": now_hm, "check_out": None, "late_minutes": late, "note": "", "source": "self", "by_name": current_user.get("full_name", ""), "created_at": _now(), "updated_at": _now()}
    await db.hr_attendance.insert_one(doc)
    return {"check_in": now_hm, "status": doc["status"], "late_minutes": late, "message": f"تم تسجيل الحضور {now_hm}" + (f" — متأخر {late} دقيقة" if late else "")}


@router.post("/check-out")
async def check_out(current_user: dict = Depends(get_current_user)):
    db = get_db()
    emp = await find_my_employee(db, current_user)
    if not emp:
        raise HTTPException(status_code=404, detail="لا يوجد ملف إداري مرتبط بحسابك")
    t = _today()
    rec = await db.hr_attendance.find_one({"employee_id": str(emp["_id"]), "date": t})
    if not rec:
        raise HTTPException(status_code=400, detail="لم يُسجَّل حضورك اليوم")
    if rec.get("check_out"):
        raise HTTPException(status_code=400, detail="تم تسجيل الانصراف بالفعل")
    now_hm = datetime.now(YEMEN_TZ).strftime("%H:%M")
    await db.hr_attendance.update_one({"_id": rec["_id"]}, {"$set": {"check_out": now_hm, "updated_at": _now()}})
    return {"check_out": now_hm, "message": f"تم تسجيل الانصراف {now_hm}"}
