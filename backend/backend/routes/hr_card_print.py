"""🖨️ طباعة بطاقات الموظفين دفعة واحدة — بنفس أسلوب بطاقات الطلاب (دفعات، خلفيات، سجل، تقرير)"""
import io
import asyncio
from datetime import datetime, timezone
from typing import Optional, List

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, export_filename, export_headers, has_permission, log_activity
from .hr_common import P_MANAGE, _guard_any, _can_view, _guard, _descendants, hr_scope_units, scope_filter, _now
from .hr import CATEGORIES, STATUSES
from .hr_cards import card_payload, _hr_card_settings, render_employee_png, render_payload
from .student_cards import (
    DEFAULT_PRINT_SETTINGS, ORIENTATIONS, DEFAULT_CARD_FONT, DEFAULT_BACK_SETTINGS,
    _render_card_png, _render_card_back_png, BackSettingsIn,
)

router = APIRouter(prefix="/hr", tags=["شؤون الموظفين - طباعة البطاقات"])

PRINT_SETTINGS_ID = "hr_global"
BACK_SETTINGS_ID = "hr_back_global"
HR_BACK_LINES = [
    "هذه البطاقة ملك للجامعة وتُسلَّم لإدارة شؤون الموظفين عند الطلب أو عند انتهاء الخدمة.",
    "البطاقة شخصية ولا يجوز إعارتها أو استخدامها من قبل الغير.",
    "يجب حمل البطاقة أثناء الدوام الرسمي وإبرازها عند الطلب.",
    "في حال فقدانها يُبلَّغ شؤون الموظفين فوراً، ويُصدَر بدل فاقد وفق اللوائح.",
    "البطاقة صالحة حتى التاريخ المدوَّن عليها ما دام حاملها على رأس العمل.",
    "من يعثر عليها يرجى تسليمها لأقرب مكتب في الجامعة أو الاتصال بالرقم أدناه.",
]
HR_DEFAULT_BACK = {**DEFAULT_BACK_SETTINGS, "lines": HR_BACK_LINES, "title": "تعليمات استخدام البطاقة الوظيفية"}


class HrBatchRequest(BaseModel):
    org_unit_id: Optional[str] = None
    category: Optional[str] = None
    status: Optional[str] = None
    employee_ids: Optional[List[str]] = None
    q: Optional[str] = None
    base_url: Optional[str] = ""
    orientation: Optional[str] = "auto"
    settings: Optional[dict] = None
    exclude_printed: Optional[bool] = False
    only_with_photo: Optional[bool] = False
    exclude_ids: Optional[List[str]] = None
    reprint_batch_no: Optional[int] = None


class HrBackBatchRequest(BaseModel):
    batch_no: Optional[int] = None
    count: Optional[int] = None
    orientation: Optional[str] = "auto"
    settings: Optional[dict] = None


async def _print_settings(db) -> dict:
    doc = await db.card_print_settings.find_one({"_id": PRINT_SETTINGS_ID}) or await db.card_print_settings.find_one({"_id": "global"}) or {}
    st = {**DEFAULT_PRINT_SETTINGS, **{k: v for k, v in doc.items() if k in DEFAULT_PRINT_SETTINGS}}
    st["orientation"] = doc.get("orientation", "auto")
    return st


async def _back_settings(db) -> dict:
    doc = await db.card_settings.find_one({"_id": BACK_SETTINGS_ID}) or {}
    return {**HR_DEFAULT_BACK, **{k: v for k, v in doc.items() if k != "_id"}}


async def _unit_names(db, ids: set) -> dict:
    oids = [ObjectId(x) for x in ids if x and ObjectId.is_valid(x)]
    return {str(u["_id"]): u.get("name", "") for u in await db.org_units.find({"_id": {"$in": oids}}, {"name": 1}).to_list(5000)}


async def _candidates(db, current_user: dict, data: HrBatchRequest) -> List[dict]:
    """الموظفون المرشحون للدفعة (بعد النطاق والفلاتر، قبل فلاتر التكرار/الصورة)."""
    q: dict = await scope_filter(db, current_user)
    if data.employee_ids:
        q["_id"] = {"$in": [ObjectId(x) for x in data.employee_ids if ObjectId.is_valid(x)]}
    else:
        if data.org_unit_id:
            units = await _descendants(db, {data.org_unit_id})
            scope = await hr_scope_units(db, current_user)
            if scope is not None:
                units &= scope
            q["org_unit_id"] = {"$in": list(units)}
        if data.category:
            q["category"] = data.category
        if data.status:
            q["status"] = data.status
        else:
            q["status"] = {"$ne": "ended"}
        if data.q and data.q.strip():
            import re
            rx = {"$regex": re.escape(data.q.strip()), "$options": "i"}
            q["$or"] = [{"full_name": rx}, {"employee_no": rx}, {"job_title": rx}]
    emps = await db.employees.find(q).sort([("org_unit_id", 1), ("full_name", 1)]).to_list(5000)
    names = await _unit_names(db, {e.get("org_unit_id") for e in emps})
    for e in emps:
        e["_unit_name"] = names.get(e.get("org_unit_id") or "", "")
    return emps


def _row(e: dict, included: bool, reason: str) -> dict:
    return {
        "id": str(e["_id"]), "full_name": e.get("full_name", ""), "number": e.get("employee_no", ""),
        "job_title": e.get("job_title", ""), "unit_name": e.get("_unit_name", ""),
        "category": CATEGORIES.get(e.get("category"), e.get("category") or ""), "status": STATUSES.get(e.get("status"), e.get("status") or ""),
        "has_photo": bool(e.get("photo_path")), "printed_at": e.get("card_printed_at"), "print_batch_no": e.get("card_print_batch_no"),
        "print_count": int(e.get("card_print_count") or 0), "included": included, "reason": reason,
    }


def _apply_filters(emps: List[dict], data: HrBatchRequest):
    excl = set(data.exclude_ids or [])
    rows, selected = [], []
    for e in emps:
        sid = str(e["_id"])
        if data.exclude_printed and e.get("card_printed_at"):
            rows.append(_row(e, False, f"طُبعت في دفعة #{e.get('card_print_batch_no') or '؟'}"))
        elif data.only_with_photo and not e.get("photo_path"):
            rows.append(_row(e, False, "لا توجد صورة معتمدة"))
        elif sid in excl:
            rows.append(_row(e, False, "مستبعد يدوياً"))
        else:
            rows.append(_row(e, True, ""))
            selected.append(e)
    summary = {
        "total": len(emps), "included": len(selected), "pages": (len(selected) + 1) // 2,
        "excluded_printed": sum(1 for r in rows if r["reason"].startswith("طُبعت")),
        "excluded_no_photo": sum(1 for r in rows if r["reason"] == "لا توجد صورة معتمدة"),
        "excluded_manual": sum(1 for r in rows if r["reason"] == "مستبعد يدوياً"),
    }
    return rows, selected, summary


@router.get("/cards/print-settings")
async def get_print_settings(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    return await _print_settings(get_db())


@router.post("/cards/batch-preview")
async def batch_preview(data: HrBatchRequest, current_user: dict = Depends(get_current_user)):
    """🧾 معاينة الدفعة: من سيُطبع ومن سيُستبعد ولماذا"""
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    emps = await _candidates(db, current_user, data)
    rows, _sel, summary = _apply_filters(emps, data)
    return {"rows": rows, "summary": summary}


def _compose_pdf(pngs: List[bytes], st: dict, out_portrait: bool) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image as PILImage
    W, H = A4
    cw, ch = (st["card_h"] * mm, st["card_w"] * mm) if out_portrait else (st["card_w"] * mm, st["card_h"] * mm)
    positions = [(st["card1_x"] * mm, st["card1_y"] * mm), (st["card2_x"] * mm, st["card2_y"] * mm)]
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    for i in range(0, len(pngs), 2):
        for j, png in enumerate(pngs[i:i + 2]):
            im = PILImage.open(io.BytesIO(png)).convert("RGB")
            b = io.BytesIO(); im.save(b, format="JPEG", quality=88)
            x, y_top = positions[j]
            c.drawImage(ImageReader(io.BytesIO(b.getvalue())), x, H - y_top - ch, cw, ch)
        c.showPage()
    c.save()
    return buf.getvalue()


def _rotate(png: bytes) -> bytes:
    from PIL import Image as PILImage
    im = PILImage.open(io.BytesIO(png)).rotate(90, expand=True)
    b = io.BytesIO(); im.save(b, format="PNG"); return b.getvalue()


async def _render_employee_png(db, emp: dict, base: str, s: dict) -> bytes:
    return await render_employee_png(db, emp, base, s)


async def _render_employee_jpeg(db, emp: dict, base: str, s: dict, rotate: bool) -> bytes:
    """بطاقة موظف → JPEG جاهز للـPDF (الرسم في thread حتى لا يُحجب الخادم)"""
    from services.storage_service import get_object
    from PIL import Image as PILImage
    p = await card_payload(db, emp, base)
    loop = asyncio.get_event_loop()
    photo_bytes = None
    if emp.get("photo_path"):
        try:
            photo_bytes, _ = await loop.run_in_executor(None, get_object, emp["photo_path"])
        except Exception:
            photo_bytes = None
    payload = render_payload(p, s)

    def _draw():
        png = _render_card_png(payload, photo_bytes, p["verify_url"])
        im = PILImage.open(io.BytesIO(png)).convert("RGB")
        if rotate:
            im = im.rotate(90, expand=True)
        b = io.BytesIO(); im.save(b, format="JPEG", quality=88)
        return b.getvalue()
    return await loop.run_in_executor(None, _draw)


def _compose_pdf_jpegs(jpegs: List[bytes], st: dict, out_portrait: bool) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.utils import ImageReader
    W, H = A4
    cw, ch = (st["card_h"] * mm, st["card_w"] * mm) if out_portrait else (st["card_w"] * mm, st["card_h"] * mm)
    positions = [(st["card1_x"] * mm, st["card1_y"] * mm), (st["card2_x"] * mm, st["card2_y"] * mm)]
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    for i in range(0, len(jpegs), 2):
        for j, jpg in enumerate(jpegs[i:i + 2]):
            x, y_top = positions[j]
            c.drawImage(ImageReader(io.BytesIO(jpg)), x, H - y_top - ch, cw, ch)
        c.showPage()
    c.save()
    return buf.getvalue()


async def _resolve_hr_batch(db, current_user: dict, data: HrBatchRequest):
    """الموظفون النهائيون للدفعة + إعدادات الطباعة (مشترك بين PDF المباشر والمهمة المرحلية)"""
    st = await _print_settings(db)
    orientation = data.orientation if data.orientation in ORIENTATIONS else "auto"
    if data.settings:
        for k in DEFAULT_PRINT_SETTINGS:
            if k in data.settings:
                try:
                    st[k] = float(data.settings[k])
                except (TypeError, ValueError):
                    pass
        await db.card_print_settings.update_one({"_id": PRINT_SETTINGS_ID}, {"$set": {**{k: st[k] for k in DEFAULT_PRINT_SETTINGS}, "orientation": orientation}}, upsert=True)
    reprint = None
    if data.reprint_batch_no:
        reprint = await db.hr_card_print_batches.find_one({"batch_no": int(data.reprint_batch_no)})
        if not reprint:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
        ids = [ObjectId(x) for x in reprint.get("employee_ids", []) if ObjectId.is_valid(x)]
        emps = await db.employees.find({"_id": {"$in": ids}}).to_list(5000)
        order = {x: i for i, x in enumerate(reprint.get("employee_ids", []))}
        emps.sort(key=lambda e: order.get(str(e["_id"]), 0))
        names = await _unit_names(db, {e.get("org_unit_id") for e in emps})
        for e in emps:
            e["_unit_name"] = names.get(e.get("org_unit_id") or "", "")
    else:
        emps_all = await _candidates(db, current_user, data)
        _rows, emps, _summary = _apply_filters(emps_all, data)
    if not emps:
        raise HTTPException(status_code=400, detail="لا يوجد موظفون للطباعة بعد تطبيق الفلاتر")
    s = await _hr_card_settings(db)
    template_portrait = s["template"] != "horizontal"
    out_portrait = template_portrait if orientation == "auto" else orientation == "portrait"
    return emps, st, s, orientation, out_portrait, template_portrait != out_portrait, reprint


async def _register_hr_batch(db, current_user: dict, data: HrBatchRequest, emps: List[dict], st: dict, orientation: str) -> int:
    last = await db.hr_card_print_batches.find_one({}, sort=[("batch_no", -1)])
    batch_no = int((last or {}).get("batch_no", 0)) + 1
    now = datetime.now(timezone.utc).isoformat()
    unit_name = ""
    if data.org_unit_id and ObjectId.is_valid(data.org_unit_id):
        unit_name = ((await db.org_units.find_one({"_id": ObjectId(data.org_unit_id)})) or {}).get("name", "")
    await db.hr_card_print_batches.insert_one({
        "batch_no": batch_no, "created_at": now,
        "by_user_id": str(current_user.get("_id") or current_user.get("id", "")), "by_name": current_user.get("full_name", ""),
        "mode": "ids" if data.employee_ids else "filter", "org_unit_id": data.org_unit_id, "unit_name": unit_name or ("موظفون محددون" if data.employee_ids else "كل الوحدات"),
        "category": data.category, "category_label": CATEGORIES.get(data.category or "", "الكل"), "status": data.status,
        "count": len(emps), "pages": (len(emps) + 1) // 2, "orientation": orientation, "settings": {k: st[k] for k in DEFAULT_PRINT_SETTINGS},
        "employee_ids": [str(e["_id"]) for e in emps],
        "employees": [{"id": str(e["_id"]), "full_name": e.get("full_name", ""), "number": e.get("employee_no", ""), "unit_name": e.get("_unit_name", ""), "job_title": e.get("job_title", "")} for e in emps],
    })
    await db.employees.update_many({"_id": {"$in": [e["_id"] for e in emps]}},
                                   {"$set": {"card_printed_at": now, "card_print_batch_no": batch_no}, "$inc": {"card_print_count": 1}})
    return batch_no


# ==================== 🧩 مهمة طباعة مرحلية (للدفعات الكبيرة) ====================
JOB_STEP = 8


@router.post("/cards/print-jobs")
async def create_hr_print_job(data: HrBatchRequest, current_user: dict = Depends(get_current_user)):
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    emps, st, _s, orientation, out_portrait, rotate, reprint = await _resolve_hr_batch(db, current_user, data)
    now = datetime.now(timezone.utc).isoformat()
    doc = {"status": "running", "created_at": now, "updated_at": now, "by_user_id": str(current_user.get("_id") or current_user.get("id", "")),
           "total": len(emps), "done": 0, "employee_ids": [str(e["_id"]) for e in emps], "request": data.model_dump(),
           "st": {k: st[k] for k in DEFAULT_PRINT_SETTINGS}, "orientation": orientation, "out_portrait": out_portrait, "rotate": rotate,
           "reprint_batch_no": (reprint or {}).get("batch_no")}
    r = await db.hr_card_print_jobs.insert_one(doc)
    return {"job_id": str(r.inserted_id), "total": len(emps), "step": JOB_STEP}


@router.post("/cards/print-jobs/{job_id}/step")
async def hr_print_job_step(job_id: str, current_user: dict = Depends(get_current_user)):
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    job = await db.hr_card_print_jobs.find_one({"_id": ObjectId(job_id)}) if ObjectId.is_valid(job_id) else None
    if not job or job.get("by_user_id") != str(current_user.get("_id") or current_user.get("id", "")):
        raise HTTPException(status_code=404, detail="مهمة الطباعة غير موجودة")
    if job["status"] == "ready":
        return {"status": "ready", "done": job["done"], "total": job["total"], "batch_no": job.get("batch_no")}
    ids = job["employee_ids"][job["done"]:job["done"] + JOB_STEP]
    emps = await db.employees.find({"_id": {"$in": [ObjectId(x) for x in ids]}}).to_list(JOB_STEP)
    emps.sort(key=lambda e: ids.index(str(e["_id"])))
    s = await _hr_card_settings(db)
    base = (job["request"].get("base_url") or "").rstrip("/")
    for idx, e in enumerate(emps, start=job["done"]):
        jpg = await _render_employee_jpeg(db, e, base, s, job["rotate"])
        await db.hr_card_print_job_pages.update_one({"job_id": job_id, "idx": idx}, {"$set": {"data": jpg}}, upsert=True)
    done = job["done"] + len(ids)
    await db.hr_card_print_jobs.update_one({"_id": job["_id"]}, {"$set": {"done": done, "updated_at": datetime.now(timezone.utc).isoformat()}})
    if done < job["total"]:
        return {"status": "running", "done": done, "total": job["total"]}
    pages = await db.hr_card_print_job_pages.find({"job_id": job_id}).sort("idx", 1).to_list(5000)
    pdf = await asyncio.get_event_loop().run_in_executor(None, _compose_pdf_jpegs, [bytes(p["data"]) for p in pages], job["st"], job["out_portrait"])
    from motor.motor_asyncio import AsyncIOMotorGridFSBucket
    fs = AsyncIOMotorGridFSBucket(db, bucket_name="hr_card_print_files")
    file_id = await fs.upload_from_stream(f"hr_cards_{job_id}.pdf", pdf, metadata={"job_id": job_id})
    batch_no = job.get("reprint_batch_no")
    if not batch_no:
        all_emps = await db.employees.find({"_id": {"$in": [ObjectId(x) for x in job["employee_ids"]]}}).to_list(5000)
        names = await _unit_names(db, {e.get("org_unit_id") for e in all_emps})
        for e in all_emps:
            e["_unit_name"] = names.get(e.get("org_unit_id") or "", "")
        all_emps.sort(key=lambda e: job["employee_ids"].index(str(e["_id"])))
        batch_no = await _register_hr_batch(db, current_user, HrBatchRequest(**job["request"]), all_emps, job["st"], job["orientation"])
        await log_activity(current_user, "hr_cards_batch_print", "hr_card_print_batches", str(batch_no), f"دفعة طباعة بطاقات موظفين #{batch_no}", {"count": job["total"]})
    await db.hr_card_print_job_pages.delete_many({"job_id": job_id})
    await db.hr_card_print_jobs.update_one({"_id": job["_id"]}, {"$set": {"status": "ready", "file_id": file_id, "batch_no": batch_no, "size": len(pdf), "is_reprint": bool(job.get("reprint_batch_no"))}})
    return {"status": "ready", "done": done, "total": job["total"], "batch_no": batch_no, "size": len(pdf)}


@router.get("/cards/print-jobs/{job_id}/download")
async def hr_print_job_download(job_id: str, current_user: dict = Depends(get_current_user)):
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    job = await db.hr_card_print_jobs.find_one({"_id": ObjectId(job_id)}) if ObjectId.is_valid(job_id) else None
    if not job or job.get("by_user_id") != str(current_user.get("_id") or current_user.get("id", "")):
        raise HTTPException(status_code=404, detail="مهمة الطباعة غير موجودة")
    if job["status"] != "ready":
        raise HTTPException(status_code=400, detail="الملف لم يكتمل بعد")
    from motor.motor_asyncio import AsyncIOMotorGridFSBucket
    fs = AsyncIOMotorGridFSBucket(db, bucket_name="hr_card_print_files")
    pdf = await (await fs.open_download_stream(job["file_id"])).read()
    label = f"إعادة تنزيل دفعة {job['batch_no']}" if job.get("is_reprint") else f"دفعة {job['batch_no']}"
    headers = export_headers(export_filename("بطاقات الموظفين", label, f"{job['total']} بطاقة", ext="pdf"))
    headers["X-Batch-No"] = str(job["batch_no"]); headers["X-Batch-Count"] = str(job["total"])
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf", headers=headers)


@router.post("/cards/batch-pdf")
async def batch_pdf(data: HrBatchRequest, current_user: dict = Depends(get_current_user)):
    """🖨️ PDF بطاقات الموظفين: بطاقتان في كل ورقة A4 — يسجّل دفعة ويوسم الموظفين كمطبوعين (أو إعادة تنزيل دفعة سابقة)"""
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    st = await _print_settings(db)
    orientation = data.orientation if data.orientation in ORIENTATIONS else "auto"
    if data.settings:
        for k in DEFAULT_PRINT_SETTINGS:
            if k in data.settings:
                try:
                    st[k] = float(data.settings[k])
                except (TypeError, ValueError):
                    pass
        await db.card_print_settings.update_one({"_id": PRINT_SETTINGS_ID}, {"$set": {**{k: st[k] for k in DEFAULT_PRINT_SETTINGS}, "orientation": orientation}}, upsert=True)

    reprint = None
    if data.reprint_batch_no:
        reprint = await db.hr_card_print_batches.find_one({"batch_no": int(data.reprint_batch_no)})
        if not reprint:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
        ids = [ObjectId(x) for x in reprint.get("employee_ids", []) if ObjectId.is_valid(x)]
        emps = await db.employees.find({"_id": {"$in": ids}}).to_list(5000)
        order = {x: i for i, x in enumerate(reprint.get("employee_ids", []))}
        emps.sort(key=lambda e: order.get(str(e["_id"]), 0))
        rows, summary = [], {}
    else:
        emps_all = await _candidates(db, current_user, data)
        rows, emps, summary = _apply_filters(emps_all, data)
    if not emps:
        raise HTTPException(status_code=400, detail="لا يوجد موظفون للطباعة بعد تطبيق الفلاتر")

    s = await _hr_card_settings(db)
    template_portrait = s["template"] != "horizontal"
    out_portrait = template_portrait if orientation == "auto" else orientation == "portrait"
    base = (data.base_url or "").rstrip("/")
    pngs = []
    for e in emps:
        png = await _render_employee_png(db, e, base, s)
        pngs.append(_rotate(png) if template_portrait != out_portrait else png)
    pdf = _compose_pdf(pngs, st, out_portrait)

    if reprint:
        return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                                 headers=export_headers(export_filename("بطاقات الموظفين", f"إعادة تنزيل دفعة {reprint['batch_no']}", f"{len(pngs)} بطاقة", ext="pdf")))

    last = await db.hr_card_print_batches.find_one({}, sort=[("batch_no", -1)])
    batch_no = int((last or {}).get("batch_no", 0)) + 1
    now = datetime.now(timezone.utc).isoformat()
    unit_name = ""
    if data.org_unit_id and ObjectId.is_valid(data.org_unit_id):
        unit_name = ((await db.org_units.find_one({"_id": ObjectId(data.org_unit_id)})) or {}).get("name", "")
    await db.hr_card_print_batches.insert_one({
        "batch_no": batch_no, "created_at": now,
        "by_user_id": str(current_user.get("_id") or current_user.get("id", "")), "by_name": current_user.get("full_name", ""),
        "mode": "ids" if data.employee_ids else "filter", "org_unit_id": data.org_unit_id, "unit_name": unit_name or ("موظفون محددون" if data.employee_ids else "كل الوحدات"),
        "category": data.category, "category_label": CATEGORIES.get(data.category or "", "الكل"), "status": data.status,
        "count": len(emps), "pages": (len(emps) + 1) // 2, "orientation": orientation, "settings": {k: st[k] for k in DEFAULT_PRINT_SETTINGS},
        "employee_ids": [str(e["_id"]) for e in emps],
        "employees": [{"id": str(e["_id"]), "full_name": e.get("full_name", ""), "number": e.get("employee_no", ""), "unit_name": e.get("_unit_name", ""), "job_title": e.get("job_title", "")} for e in emps],
    })
    await db.employees.update_many({"_id": {"$in": [e["_id"] for e in emps]}},
                                   {"$set": {"card_printed_at": now, "card_print_batch_no": batch_no}, "$inc": {"card_print_count": 1}})
    headers = export_headers(export_filename("بطاقات الموظفين", f"دفعة {batch_no}", f"{len(pngs)} بطاقة", ext="pdf"))
    headers["X-Batch-No"] = str(batch_no)
    headers["X-Batch-Count"] = str(len(emps))
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf", headers=headers)


@router.get("/cards/batches")
async def list_batches(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    docs = await get_db().hr_card_print_batches.find({}, {"employees": 0, "employee_ids": 0}).sort("batch_no", -1).to_list(500)
    for d in docs:
        d["id"] = str(d.pop("_id"))
    return {"items": docs}


@router.get("/cards/batches/{batch_no}")
async def get_batch(batch_no: int, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    d = await get_db().hr_card_print_batches.find_one({"batch_no": batch_no})
    if not d:
        raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
    d["id"] = str(d.pop("_id"))
    return d


@router.post("/cards/batches/{batch_no}/reset")
async def reset_batch_marks(batch_no: int, current_user: dict = Depends(get_current_user)):
    """↩️ إلغاء وسم «مطبوع» لموظفي دفعة (مثلاً تلفت الورقة) ليعودوا للظهور في الدفعات القادمة"""
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    d = await db.hr_card_print_batches.find_one({"batch_no": batch_no})
    if not d:
        raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
    ids = [ObjectId(x) for x in d.get("employee_ids", []) if ObjectId.is_valid(x)]
    r = await db.employees.update_many({"_id": {"$in": ids}, "card_print_batch_no": batch_no}, {"$unset": {"card_printed_at": "", "card_print_batch_no": ""}})
    await db.hr_card_print_batches.update_one({"_id": d["_id"]}, {"$set": {"reset_at": _now(), "reset_by": current_user.get("full_name", "")}})
    return {"message": f"أُلغي وسم الطباعة عن {r.modified_count} موظفاً", "modified": r.modified_count}


@router.get("/cards/print-report")
async def print_report(org_unit_id: Optional[str] = None, category: Optional[str] = None, status: Optional[str] = None,
                       current_user: dict = Depends(get_current_user)):
    """📥 تقرير Excel: حالة طباعة بطاقات الموظفين"""
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    db = get_db()
    emps = await _candidates(db, current_user, HrBatchRequest(org_unit_id=org_unit_id, category=category, status=status))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "طباعة بطاقات الموظفين"
    ws.sheet_view.rightToLeft = True
    headers = ["#", "الموظف", "الرقم الوظيفي", "المسمى", "الوحدة", "الفئة", "الحالة", "صورة معتمدة", "حالة الطباعة", "رقم الدفعة", "تاريخ الطباعة", "عدد مرات الطباعة"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill(start_color="0F2440", end_color="0F2440", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for idx, e in enumerate(emps, 1):
        r = _row(e, True, "")
        ws.append([idx, r["full_name"], r["number"], r["job_title"], r["unit_name"], r["category"], r["status"],
                   "نعم" if r["has_photo"] else "لا", "مطبوعة" if r["printed_at"] else "لم تُطبع", r["print_batch_no"] or "",
                   (r["printed_at"] or "")[:16].replace("T", " "), r["print_count"]])
    for i, w in enumerate([4, 26, 12, 20, 22, 14, 14, 10, 12, 10, 16, 10], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=export_headers(export_filename("تقرير طباعة بطاقات الموظفين", f"{len(emps)} موظف", ext="xlsx")))


# ==================== 🔄 خلفية البطاقة الوظيفية ====================
@router.get("/cards/back-settings")
async def get_back_settings(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    return {**await _back_settings(get_db()), "defaults": HR_DEFAULT_BACK, "can_edit": has_permission(current_user, P_MANAGE) or has_permission(current_user, "hr_print_cards")}


def _guard_ok(u: dict) -> bool:
    try:
        _guard(u, P_MANAGE); return True
    except HTTPException:
        return False


@router.put("/cards/back-settings")
async def put_back_settings(data: BackSettingsIn, current_user: dict = Depends(get_current_user)):
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    upd = {k: v for k, v in data.model_dump(exclude_none=True).items()}
    if "lines" in upd:
        upd["lines"] = [str(x).strip() for x in upd["lines"] if str(x).strip()][:10]
    upd["updated_at"] = _now()
    await get_db().card_settings.update_one({"_id": BACK_SETTINGS_ID}, {"$set": upd}, upsert=True)
    return {"message": "تم حفظ خلفية البطاقة الوظيفية", **await _back_settings(get_db())}


@router.get("/cards/back-preview")
async def back_preview(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    back, s = await _back_settings(db), await _hr_card_settings(db)
    uni = await db.university.find_one({}) or {}
    png = _render_card_back_png(back, s["template"], s.get("font") or DEFAULT_CARD_FONT, s["template"] != "horizontal", uni)
    return StreamingResponse(io.BytesIO(png), media_type="image/png")


@router.post("/cards/batch-back-pdf")
async def batch_back_pdf(data: HrBackBatchRequest, current_user: dict = Depends(get_current_user)):
    """🔄 PDF خلفيات دفعة الموظفين — بطاقتان في كل ورقة بمواضع مستقلة تُحفظ"""
    _guard_any(current_user, P_MANAGE, "hr_print_cards")
    db = get_db()
    back = await _back_settings(db)
    if not back.get("enabled"):
        raise HTTPException(status_code=400, detail="خلفية البطاقة معطّلة من الإعدادات")
    if data.batch_no:
        b = await db.hr_card_print_batches.find_one({"batch_no": int(data.batch_no)})
        if not b:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
        n, label = int(b.get("count") or 0), f"خلفيات دفعة {b['batch_no']}"
    else:
        n, label = int(data.count or 0), "خلفيات البطاقات"
    if n <= 0:
        raise HTTPException(status_code=400, detail="عدد البطاقات غير محدد")
    s = await _hr_card_settings(db)
    template_portrait = s["template"] != "horizontal"
    orientation = data.orientation if data.orientation in ORIENTATIONS else "auto"
    out_portrait = template_portrait if orientation == "auto" else orientation == "portrait"
    uni = await db.university.find_one({}) or {}
    png = _render_card_back_png(back, s["template"], s.get("font") or DEFAULT_CARD_FONT, template_portrait, uni)
    if template_portrait != out_portrait:
        png = _rotate(png)
    pst = await _print_settings(db)
    bst = {k: float(back.get(k, HR_DEFAULT_BACK[k])) for k in ("back1_x", "back1_y", "back2_x", "back2_y")}
    if data.settings:
        for k in ("card_w", "card_h"):
            if k in data.settings:
                try: pst[k] = float(data.settings[k])
                except (TypeError, ValueError): pass
        changed = {}
        for k in bst:
            if k in data.settings:
                try: bst[k] = float(data.settings[k]); changed[k] = bst[k]
                except (TypeError, ValueError): pass
        if changed:
            await db.card_settings.update_one({"_id": BACK_SETTINGS_ID}, {"$set": changed}, upsert=True)
    st = {**pst, "card1_x": bst["back1_x"], "card1_y": bst["back1_y"], "card2_x": bst["back2_x"], "card2_y": bst["back2_y"]}
    pdf = _compose_pdf([png] * n, st, out_portrait)
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                             headers=export_headers(export_filename("بطاقات الموظفين", label, f"{n} خلفية", ext="pdf")))


@router.get("/cards/sample-preview")
async def sample_preview(employee_id: Optional[str] = None, template: Optional[str] = None, font: Optional[str] = None,
                         current_user: dict = Depends(get_current_user)):
    """🪪 معاينة حيّة لبطاقة موظف بقالب/خط مختار (قبل الحفظ)"""
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    from .student_cards import TEMPLATES, CARD_FONTS
    db = get_db()
    q = await scope_filter(db, current_user)
    if employee_id and ObjectId.is_valid(employee_id):
        q["_id"] = ObjectId(employee_id)
    emp = await db.employees.find_one({**q, "photo_path": {"$exists": True, "$ne": None}}) or await db.employees.find_one(q)
    if not emp:
        raise HTTPException(status_code=404, detail="لا يوجد موظف للمعاينة")
    s = await _hr_card_settings(db)
    if template in TEMPLATES and template != "custom":
        s["template"] = template
    if font in CARD_FONTS:
        s["font"] = font
    return StreamingResponse(io.BytesIO(await _render_employee_png(db, emp, "", s)), media_type="image/png")
