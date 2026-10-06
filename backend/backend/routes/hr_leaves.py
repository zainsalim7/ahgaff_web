"""🏖️ شؤون الموظفين — الإجازات: طلب / اعتماد (مدير مباشر ← شؤون الموظفين) / رصيد سنوي"""
import io
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, has_permission, log_activity, export_headers, export_filename
from .hr_common import (scope_employee_ids, assert_in_scope, P_LEAVES, YEMEN_TZ, _now, _today, _oid, _ser, _can_view, _guard, parse_date, user_id_of, get_hr_settings,
                        work_days_between, find_my_employee, employee_user_ids, notify_users, hr_manager_user_ids, enrich_employee_refs)

router = APIRouter(prefix="/hr/leaves", tags=["شؤون الموظفين - الإجازات"])

LEAVE_TYPES = {"annual": "سنوية", "sick": "مرضية", "emergency": "اضطرارية", "unpaid": "بدون راتب", "maternity": "وضع / أمومة",
               "hajj": "حج", "study": "دراسية", "mission": "مهمة رسمية / انتداب", "other": "أخرى"}
DEDUCT_BALANCE = {"annual"}  # قديم — المصدر الآن إعدادات الإجازات (hr_leave_types)
from .hr_leave_types import leave_types_map, entitlement_for  # noqa: E402
STATUSES = {"pending": "بانتظار المدير المباشر", "hr_pending": "بانتظار شؤون الموظفين", "approved": "معتمدة", "rejected": "مرفوضة", "cancelled": "ملغاة"}
OPEN = ("pending", "hr_pending")


class LeaveIn(BaseModel):
    employee_id: Optional[str] = None
    type: str = "annual"
    start_date: str
    end_date: str
    reason: Optional[str] = ""
    substitute_employee_id: Optional[str] = None
    contact_during_leave: Optional[str] = ""


class DecisionIn(BaseModel):
    action: str  # approve | reject
    note: Optional[str] = ""


class BalanceIn(BaseModel):
    year: int
    type: str = "annual"
    entitlement: int
    carried_over: int = 0
    note: Optional[str] = ""


async def _balance(db, employee_id: str, year: int, settings: dict = None, type_key: str = "annual", emp: dict = None, types: dict = None) -> dict:
    """رصيد نوع إجازة واحد للموظف: الاستحقاق حسب فئته من إعدادات الإجازات (أو تجاوز يدوي في leave_balances) − المستخدم"""
    settings = settings or await get_hr_settings(db)
    types = types or await leave_types_map(db)
    t = types.get(type_key) or {}
    emp = emp or await db.employees.find_one({"_id": ObjectId(employee_id)}, {"category": 1}) or {}
    tq = {"$in": [None, "annual"]} if type_key == "annual" else type_key
    doc = await db.leave_balances.find_one({"employee_id": employee_id, "year": year, "type": tq}) or {}
    default_ent = entitlement_for(t, emp)
    if default_ent is None and type_key == "annual":
        default_ent = settings.get("annual_leave_days", 30)
    ent = doc.get("entitlement", default_ent)
    carried = doc.get("carried_over", 0)
    # 🔁 ترحيل تلقائي: إن كان النوع يسمح بالترحيل ولم يُحدَّد المرحَّل لهذه السنة بعد → نحسب متبقي السنة السابقة (حتى الحد) ونثبّته
    if t.get("carry_over_enabled") and "carried_over" not in doc and ent is not None and year <= datetime.now(YEMEN_TZ).year + 1:
        prev_doc = await db.leave_balances.find_one({"employee_id": employee_id, "year": year - 1, "type": tq}) or {}
        prev_ent = prev_doc.get("entitlement", default_ent)
        if prev_ent is not None:
            prev_used = sum([l.get("days", 0) async for l in db.leave_requests.find({"employee_id": employee_id, "type": type_key, "status": "approved", "start_date": {"$regex": f"^{year - 1}"}}, {"days": 1})])
            had_any = prev_doc or prev_used > 0 or await db.employees.count_documents({"_id": ObjectId(employee_id), "hire_date": {"$lt": f"{year}-01-01"}}) > 0
            if had_any:
                carried = max(0, prev_ent + prev_doc.get("carried_over", 0) - prev_used)
                if t.get("carry_over_max_days") is not None:
                    carried = min(carried, int(t["carry_over_max_days"]))
                await db.leave_balances.update_one({"employee_id": employee_id, "year": year, "type": type_key}, {"$set": {"carried_over": carried, "auto_carried": True, "carried_at": _now()}, "$setOnInsert": {"type": type_key}}, upsert=True)
                doc = {**doc, "carried_over": carried, "auto_carried": True}
    used = pending = 0
    async for l in db.leave_requests.find({"employee_id": employee_id, "type": type_key, "status": {"$in": ["approved", *OPEN]}, "start_date": {"$regex": f"^{year}"}}, {"days": 1, "status": 1}):
        if l["status"] == "approved":
            used += l.get("days", 0)
        else:
            pending += l.get("days", 0)
    return {"year": year, "type": type_key, "type_label": t.get("name") or LEAVE_TYPES.get(type_key, type_key), "color": t.get("color"), "entitlement": ent, "carried_over": carried, "used": used, "pending": pending,
            "remaining": (ent + carried - used) if ent is not None else None, "unlimited": ent is None, "overridden": bool(doc.get("entitlement") is not None or doc.get("note")), "auto_carried": bool(doc.get("auto_carried")), "note": doc.get("note", "")}


async def _balances(db, emp: dict, year: int, settings: dict = None) -> list:
    """أرصدة كل الأنواع التي لها رصيد محدد لفئة الموظف"""
    settings = settings or await get_hr_settings(db)
    types = await leave_types_map(db, active_only=True)
    eid = str(emp["_id"])
    out = []
    for k, t in types.items():
        if k == "annual" or entitlement_for(t, emp) is not None:
            out.append(await _balance(db, eid, year, settings, k, emp, types))
    return out


async def _manager_user(db, emp: dict) -> Optional[str]:
    mid = emp.get("manager_employee_id")
    if not mid or not ObjectId.is_valid(mid):
        return None
    return (await employee_user_ids(db, [mid])).get(mid)


async def _create_request(db, emp: dict, data: LeaveIn, by: dict, auto_approve: bool = False) -> dict:
    types = await leave_types_map(db, active_only=True)
    if data.type not in types:
        raise HTTPException(status_code=400, detail="نوع الإجازة غير صحيح أو معطّل")
    s, e = parse_date(data.start_date, "تاريخ البداية"), parse_date(data.end_date, "تاريخ النهاية")
    if e < s:
        raise HTTPException(status_code=400, detail="تاريخ النهاية قبل تاريخ البداية")
    if (e - s).days > 365:
        raise HTTPException(status_code=400, detail="مدة الإجازة تتجاوز سنة")
    settings = await get_hr_settings(db)
    days = work_days_between(s, e, settings)
    if days <= 0:
        raise HTTPException(status_code=400, detail="الفترة المحددة لا تحوي أيام عمل")
    eid = str(emp["_id"])
    overlap = await db.leave_requests.find_one({"employee_id": eid, "status": {"$in": ["approved", *OPEN]}, "start_date": {"$lte": data.end_date}, "end_date": {"$gte": data.start_date}})
    if overlap:
        raise HTTPException(status_code=400, detail=f"يوجد طلب إجازة متداخل مع هذه الفترة ({overlap['start_date']} → {overlap['end_date']})")
    if types[data.type].get("deducts_balance") and entitlement_for(types[data.type], emp) is not None or data.type == "annual":
        bal = await _balance(db, eid, s.year, settings, data.type, emp, types)
        if bal["remaining"] is not None and days > bal["remaining"] - bal["pending"]:
            raise HTTPException(status_code=400, detail=f"الرصيد غير كافٍ: المتبقي {bal['remaining']} يوماً (منها {bal['pending']} معلّقة) والمطلوب {days}")
    mgr_uid = await _manager_user(db, emp)
    status = "approved" if auto_approve else ("pending" if mgr_uid and mgr_uid != user_id_of(by) else "hr_pending")
    doc = {"employee_id": eid, "type": data.type, "start_date": data.start_date, "end_date": data.end_date, "days": days, "reason": (data.reason or "").strip(),
           "substitute_employee_id": data.substitute_employee_id or None, "contact_during_leave": data.contact_during_leave or "", "status": status,
           "manager_user_id": mgr_uid, "created_by_user_id": user_id_of(by), "created_by_name": by.get("full_name", ""), "created_at": _now(),
           "history": [{"action": "submitted" if not auto_approve else "registered", "by_name": by.get("full_name", ""), "at": _now(), "note": ""}]}
    if auto_approve:
        doc["hr_decision"] = {"by_name": by.get("full_name", ""), "at": _now(), "note": "تسجيل مباشر من شؤون الموظفين"}
    r = await db.leave_requests.insert_one(doc)
    doc["_id"] = r.inserted_id
    label = types[data.type]["name"]
    if status == "pending" and mgr_uid:
        await notify_users(db, [mgr_uid], "طلب إجازة جديد بانتظار موافقتك", f"{emp.get('full_name', '')} — إجازة {label} من {data.start_date} إلى {data.end_date} ({days} يوم عمل)", "hr_leave", {"leave_id": str(r.inserted_id)})
    elif status == "hr_pending":
        await notify_users(db, await hr_manager_user_ids(db, P_LEAVES), "طلب إجازة بانتظار اعتماد شؤون الموظفين", f"{emp.get('full_name', '')} — إجازة {label} من {data.start_date} إلى {data.end_date} ({days} يوم عمل)", "hr_leave", {"leave_id": str(r.inserted_id)})
    if auto_approve:
        await _apply_status(db, emp, doc)
    return doc


async def _apply_status(db, emp: dict, leave: dict):
    """تحديث حالة الموظف إلى «في إجازة» إذا كانت الإجازة المعتمدة تشمل اليوم"""
    t = _today()
    if leave["status"] == "approved" and leave["start_date"] <= t <= leave["end_date"] and emp.get("status") in ("active", "probation"):
        await db.employees.update_one({"_id": emp["_id"]}, {"$set": {"status": "leave", "status_before_leave": emp.get("status")}})


async def sync_leave_statuses(db):
    """مزامنة يومية idempotent: من يبدأ إجازته يتحوّل لـ leave، ومن انتهت يعود لحالته السابقة"""
    t = _today()
    on_leave = {l["employee_id"] async for l in db.leave_requests.find({"status": "approved", "start_date": {"$lte": t}, "end_date": {"$gte": t}}, {"employee_id": 1})}
    async for e in db.employees.find({"status": {"$in": ["active", "probation", "leave"]}}, {"status": 1, "status_before_leave": 1}):
        eid = str(e["_id"])
        if eid in on_leave and e["status"] != "leave":
            await db.employees.update_one({"_id": e["_id"]}, {"$set": {"status": "leave", "status_before_leave": e["status"]}})
        elif eid not in on_leave and e["status"] == "leave" and e.get("status_before_leave"):
            await db.employees.update_one({"_id": e["_id"]}, {"$set": {"status": e["status_before_leave"]}, "$unset": {"status_before_leave": ""}})


def _view(l: dict) -> dict:
    d = _ser(l)
    d["type_label"] = (_TYPE_NAMES.get(d.get("type")) or LEAVE_TYPES.get(d.get("type"), d.get("type")))
    d["status_label"] = STATUSES.get(d.get("status"), d.get("status"))
    d["deducts_balance"] = d.get("type") in _DEDUCT
    d["has_attachment"] = bool(d.get("attachment_path"))
    d["requires_attachment"] = d.get("type") in _REQ_ATTACH
    d.pop("attachment_path", None)
    return d


_TYPE_NAMES: dict = {}
_DEDUCT: set = set()
_REQ_ATTACH: set = set()


async def _refresh_type_cache(db):
    types = await leave_types_map(db)
    _TYPE_NAMES.clear(); _TYPE_NAMES.update({k: t["name"] for k, t in types.items()})
    _DEDUCT.clear(); _DEDUCT.update({k for k, t in types.items() if t.get("deducts_balance")})
    _REQ_ATTACH.clear(); _REQ_ATTACH.update({k for k, t in types.items() if t.get("requires_attachment")})
    return types


@router.get("/meta")
async def leaves_meta(current_user: dict = Depends(get_current_user)):
    types = await _refresh_type_cache(get_db())
    active = {k: t["name"] for k, t in types.items() if t.get("is_active")}
    return {"types": active, "statuses": STATUSES, "deduct_balance": sorted(_DEDUCT), "type_defs": list(types.values()), "categories": __import__('routes.hr_leave_types', fromlist=['CATEGORIES']).CATEGORIES}


@router.get("/my")
async def my_leaves(year: Optional[int] = None, current_user: dict = Depends(get_current_user)):
    """👤 الخدمة الذاتية: رصيدي + طلباتي + طلبات فريقي (إن كنت مديراً مباشراً)"""
    db = get_db()
    emp = await find_my_employee(db, current_user)
    year = year or datetime.now(YEMEN_TZ).year
    uid = user_id_of(current_user)
    team = [_view(l) for l in await db.leave_requests.find({"manager_user_id": uid, "status": "pending"}).sort("created_at", -1).to_list(200)]
    await enrich_employee_refs(db, team)
    if not emp:
        return {"profile": None, "balance": None, "requests": [], "team_pending": team}
    eid = str(emp["_id"])
    reqs = [_view(l) for l in await db.leave_requests.find({"employee_id": eid}).sort("created_at", -1).limit(100).to_list(100)]
    return {"profile": {"id": eid, "full_name": emp.get("full_name", ""), "employee_no": emp.get("employee_no", ""), "has_manager": bool(emp.get("manager_employee_id"))},
            "balance": await _balance(db, eid, year, None, "annual", emp), "balances": await _balances(db, emp, year), "requests": reqs, "team_pending": team}


@router.post("/my")
async def submit_my_leave(data: LeaveIn, current_user: dict = Depends(get_current_user)):
    db = get_db()
    emp = await find_my_employee(db, current_user)
    if not emp:
        raise HTTPException(status_code=404, detail="لا يوجد ملف إداري مرتبط بحسابك — راجع شؤون الموظفين")
    if emp.get("status") in ("suspended", "ended"):
        raise HTTPException(status_code=400, detail="لا يمكن تقديم طلب إجازة في الحالة الوظيفية الحالية")
    doc = await _create_request(db, emp, data, current_user)
    await log_activity(current_user, "hr_leave_submit", "leave", str(doc["_id"]), emp.get("full_name", ""), {"days": doc["days"], "type": data.type})
    return {"id": str(doc["_id"]), "status": doc["status"], "days": doc["days"], "message": f"تم تقديم الطلب ({doc['days']} يوم عمل) — {STATUSES[doc['status']]}"}


@router.get("/today")
async def on_leave_today(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    t = _today()
    items = [_view(l) for l in await db.leave_requests.find({"status": "approved", "start_date": {"$lte": t}, "end_date": {"$gte": t}}).to_list(1000)]
    return {"date": t, "items": await enrich_employee_refs(db, items)}


@router.get("/balances")
async def all_balances(year: Optional[int] = None, org_unit_id: Optional[str] = None, type: str = "annual", current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    year = year or datetime.now(YEMEN_TZ).year
    settings = await get_hr_settings(db)
    q: dict = {"status": {"$ne": "ended"}}
    if org_unit_id:
        q["org_unit_id"] = org_unit_id
    out = []
    types = await _refresh_type_cache(db)
    if type not in types:
        raise HTTPException(status_code=400, detail="نوع الإجازة غير معروف")
    async for e in db.employees.find(q, {"full_name": 1, "employee_no": 1, "org_unit_id": 1, "job_title": 1, "category": 1}).sort("full_name", 1):
        eid = str(e["_id"])
        out.append({"employee_id": eid, "category": e.get("category"), **(await _balance(db, eid, year, settings, type, e, types))})
    await enrich_employee_refs(db, out)
    t = types[type]
    return {"year": year, "type": type, "type_label": t["name"], "entitlements": t.get("entitlements"), "default_entitlement": settings.get("annual_leave_days", 30), "items": out,
            "balance_types": [{"key": k, "name": v["name"], "color": v.get("color")} for k, v in types.items() if v.get("is_active") and (k == "annual" or any(x is not None for x in (v.get("entitlements") or {}).values()))]}


@router.get("/balances/export")
async def balances_export(year: Optional[int] = None, org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """📥 تقرير الأرصدة السنوي Excel: لكل موظف ولكل نوع إجازة — الاستحقاق / المرحَّل / المستخدم / المعلّق / المتبقي"""
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    year = year or datetime.now(YEMEN_TZ).year
    settings = await get_hr_settings(db)
    types = await _refresh_type_cache(db)
    keys = [k for k, v in types.items() if v.get("is_active") and (k == "annual" or any(x is not None for x in (v.get("entitlements") or {}).values()))]
    q: dict = {"status": {"$ne": "ended"}}
    if org_unit_id:
        q["org_unit_id"] = org_unit_id
    units = {str(u["_id"]): u.get("name", "") for u in await db.org_units.find({}, {"name": 1}).to_list(2000)}
    cat_label = {"academic": "أكاديمي", "administrative": "إداري", "technical": "فني", "service": "خدمات"}
    rows = []
    async for e in db.employees.find(q, {"full_name": 1, "employee_no": 1, "org_unit_id": 1, "job_title": 1, "category": 1}).sort("full_name", 1):
        eid = str(e["_id"])
        bals = [await _balance(db, eid, year, settings, k, e, types) for k in keys]
        rows.append((e, bals))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = f"الأرصدة {year}"; ws.sheet_view.rightToLeft = True
    base = ["الرقم الوظيفي", "الاسم", "الوحدة", "المسمى", "الفئة"]
    sub = ["الاستحقاق", "مرحَّل", "مستخدم", "معلّق", "المتبقي"]
    total_cols = len(base) + len(keys) * len(sub)
    ws.append([f"تقرير أرصدة الإجازات السنوي — {year}" + (f" — {units.get(org_unit_id, '')}" if org_unit_id else "") + f" ({len(rows)} موظف)"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws["A1"].font = Font(bold=True, size=13); ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(base + [x for k in keys for x in [types[k]["name"]] + [""] * (len(sub) - 1)])
    ws.append([""] * len(base) + sub * len(keys))
    thin = Side(style="thin", color="CBD5E1"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    palette = ["1565C0", "2E7D32", "C62828", "6A1B9A", "EF6C00", "00838F", "4E342E", "37474F"]
    for ci in range(1, len(base) + 1):
        ws.merge_cells(start_row=2, start_column=ci, end_row=3, end_column=ci)
        c = ws.cell(row=2, column=ci); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440"); c.alignment = Alignment(horizontal="center", vertical="center")
    for ti, k in enumerate(keys):
        start = len(base) + 1 + ti * len(sub)
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=start + len(sub) - 1)
        col = palette[ti % len(palette)]
        h = ws.cell(row=2, column=start); h.font = Font(bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor=col); h.alignment = Alignment(horizontal="center")
        for j in range(len(sub)):
            c = ws.cell(row=3, column=start + j); c.font = Font(bold=True, color=col); c.fill = PatternFill("solid", fgColor="F1F5F9"); c.alignment = Alignment(horizontal="center"); c.border = border
    for e, bals in rows:
        line = [e.get("employee_no", ""), e.get("full_name", ""), units.get(e.get("org_unit_id") or "", ""), e.get("job_title", ""), cat_label.get(e.get("category"), e.get("category") or "")]
        for b in bals:
            line += ["غير محدود" if b["unlimited"] else b["entitlement"], b["carried_over"], b["used"], b["pending"], "—" if b["remaining"] is None else b["remaining"]]
        ws.append(line)
        r = ws.max_row
        for ti, b in enumerate(bals):
            rc = ws.cell(row=r, column=len(base) + ti * len(sub) + len(sub))
            rc.font = Font(bold=True, color="C62828" if (b["remaining"] is not None and b["remaining"] <= 0) else "1B5E20")
    for ci in range(1, total_cols + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 11
    ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 22; ws.column_dimensions["D"].width = 18
    ws.freeze_panes = "F4"
    # ورقة ثانية بصيغة طولية (سطر لكل موظف ونوع) للتحليل
    ws2 = wb.create_sheet("تفصيلي"); ws2.sheet_view.rightToLeft = True
    ws2.append(["الرقم الوظيفي", "الاسم", "الوحدة", "الفئة", "نوع الإجازة", "الاستحقاق", "مرحَّل", "مستخدم", "معلّق", "المتبقي", "رصيد مخصّص", "ملاحظة"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440"); c.alignment = Alignment(horizontal="center")
    for e, bals in rows:
        for b in bals:
            ws2.append([e.get("employee_no", ""), e.get("full_name", ""), units.get(e.get("org_unit_id") or "", ""), cat_label.get(e.get("category"), e.get("category") or ""), b["type_label"],
                        "غير محدود" if b["unlimited"] else b["entitlement"], b["carried_over"], b["used"], b["pending"], "—" if b["remaining"] is None else b["remaining"], "نعم" if b["overridden"] else "", b.get("note", "")])
    for ci in range(1, 13):
        ws2.column_dimensions[get_column_letter(ci)].width = 13
    ws2.column_dimensions["B"].width = 28; ws2.column_dimensions["L"].width = 30
    await log_activity(current_user, "hr_leave_balances_export", "hr_leaves", str(year), "تصدير تقرير الأرصدة السنوي", {"count": len(rows)})
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=export_headers(export_filename("تقرير الأرصدة السنوي", str(year), units.get(org_unit_id, "") if org_unit_id else "كل الوحدات", f"{len(rows)} موظف", ext="xlsx")))


@router.put("/balances/{employee_id}")
async def set_balance(employee_id: str, data: BalanceIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_LEAVES)
    db = get_db()
    if not await db.employees.find_one({"_id": _oid(employee_id, "الموظف")}):
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    if data.entitlement < 0 or data.carried_over < 0 or data.entitlement > 365:
        raise HTTPException(status_code=400, detail="قيم الرصيد غير منطقية")
    tq = {"$in": [None, "annual"]} if data.type == "annual" else data.type
    await db.leave_balances.update_one({"employee_id": employee_id, "year": data.year, "type": tq}, {"$set": {"type": data.type, "entitlement": data.entitlement, "carried_over": data.carried_over, "note": data.note or "", "updated_by_name": current_user.get("full_name", ""), "updated_at": _now()}}, upsert=True)
    await log_activity(current_user, "hr_leave_balance", "employee", employee_id, "تعديل رصيد الإجازات", {"year": data.year, "entitlement": data.entitlement, "carried_over": data.carried_over})
    return {"message": "تم تحديث الرصيد", "balance": await _balance(db, employee_id, data.year, None, data.type)}


@router.get("")
async def list_leaves(status: Optional[str] = None, type: Optional[str] = None, employee_id: Optional[str] = None, year: Optional[int] = None,
                      month: Optional[str] = None, org_unit_id: Optional[str] = None, today: bool = False, page: int = 1, per_page: int = 40, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    await sync_leave_statuses(db)
    q: dict = {}
    if status:
        q["status"] = {"$in": status.split(",")}
    if type:
        q["type"] = type
    if today:
        t = _today()
        q["status"] = "approved"
        q["start_date"] = {"$lte": t}
        q["end_date"] = {"$gte": t}
    if employee_id:
        q["employee_id"] = employee_id
    if org_unit_id:
        ids = [str(e["_id"]) for e in await db.employees.find({"org_unit_id": org_unit_id}, {"_id": 1}).to_list(5000)]
        q["employee_id"] = {"$in": ids}
    sc_ids = await scope_employee_ids(db, current_user)
    if sc_ids is not None:
        cur = q.get("employee_id")
        allowed = set(sc_ids)
        if isinstance(cur, dict):
            allowed &= set(cur.get("$in", []))
        elif isinstance(cur, str):
            allowed = allowed & {cur}
        q["employee_id"] = {"$in": list(allowed)}
    if month:
        q["start_date"] = {"$lte": f"{month}-31"}
        q["end_date"] = {"$gte": f"{month}-01"}
    elif year:
        q["start_date"] = {"$regex": f"^{year}"}
    total = await db.leave_requests.count_documents(q)
    per_page = max(1, min(per_page, 200))
    items = [_view(l) for l in await db.leave_requests.find(q).sort([("status", 1), ("created_at", -1)]).skip((page - 1) * per_page).limit(per_page).to_list(per_page)]
    await enrich_employee_refs(db, items)
    t = _today()
    y = str(datetime.now(YEMEN_TZ).year)
    stats = {"pending": await db.leave_requests.count_documents({"status": {"$in": list(OPEN)}}),
             "on_leave_today": await db.leave_requests.count_documents({"status": "approved", "start_date": {"$lte": t}, "end_date": {"$gte": t}}),
             "approved_year": await db.leave_requests.count_documents({"status": "approved", "start_date": {"$regex": f"^{y}"}}),
             "rejected_year": await db.leave_requests.count_documents({"status": "rejected", "start_date": {"$regex": f"^{y}"}})}
    return {"items": items, "total": total, "page": page, "per_page": per_page, "stats": stats}


@router.post("")
async def register_leave(data: LeaveIn, current_user: dict = Depends(get_current_user)):
    """📝 تسجيل إجازة بالنيابة عن موظف (تُعتمد مباشرة)"""
    _guard(current_user, P_LEAVES)
    db = get_db()
    emp = await db.employees.find_one({"_id": _oid(data.employee_id or "", "الموظف")})
    if not emp:
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    doc = await _create_request(db, emp, data, current_user, auto_approve=True)
    uid = (await employee_user_ids(db, [str(emp["_id"])])).get(str(emp["_id"]))
    if uid:
        await notify_users(db, [uid], "تم تسجيل إجازة لك", f"إجازة {_TYPE_NAMES.get(data.type, data.type)} من {data.start_date} إلى {data.end_date} ({doc['days']} يوم عمل) — معتمدة", "hr_leave", {"leave_id": str(doc["_id"])})
    await log_activity(current_user, "hr_leave_register", "leave", str(doc["_id"]), emp.get("full_name", ""), {"days": doc["days"], "type": data.type})
    return {"id": str(doc["_id"]), "days": doc["days"], "message": f"تم تسجيل الإجازة واعتمادها ({doc['days']} يوم عمل)"}


@router.get("/{leave_id}")
async def get_leave(leave_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    l = await db.leave_requests.find_one({"_id": _oid(leave_id)})
    if not l:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    uid = user_id_of(current_user)
    emp = await find_my_employee(db, current_user)
    mine = emp and str(emp["_id"]) == l["employee_id"]
    if not (_can_view(current_user) or mine or l.get("manager_user_id") == uid):
        raise HTTPException(status_code=403, detail="غير مصرح")
    d = _view(l)
    await enrich_employee_refs(db, [d])
    if d.get("substitute_employee_id"):
        s = await db.employees.find_one({"_id": ObjectId(d["substitute_employee_id"])}, {"full_name": 1}) if ObjectId.is_valid(d["substitute_employee_id"]) else None
        d["substitute_name"] = (s or {}).get("full_name", "")
    d["balance"] = await _balance(db, l["employee_id"], parse_date(l["start_date"]).year)
    return d


@router.post("/{leave_id}/decide")
async def decide_leave(leave_id: str, data: DecisionIn, current_user: dict = Depends(get_current_user)):
    """✅ قرار المدير المباشر (pending → hr_pending) أو شؤون الموظفين (hr_pending → approved) / رفض"""
    db = get_db()
    l = await db.leave_requests.find_one({"_id": _oid(leave_id)})
    if not l:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    if l["status"] not in OPEN:
        raise HTTPException(status_code=400, detail=f"الطلب {STATUSES.get(l['status'])} ولا يمكن تغييره")
    if data.action not in ("approve", "reject"):
        raise HTTPException(status_code=400, detail="الإجراء غير صحيح")
    uid = user_id_of(current_user)
    is_hr = has_permission(current_user, P_LEAVES)
    if is_hr and l.get("status") == "hr_pending":
        await assert_in_scope(db, current_user, employee_id=l["employee_id"])
    is_mgr = l.get("manager_user_id") == uid
    if l["status"] == "pending" and not (is_mgr or is_hr):
        raise HTTPException(status_code=403, detail="هذا الطلب بانتظار المدير المباشر")
    if l["status"] == "hr_pending" and not is_hr:
        raise HTTPException(status_code=403, detail="اعتماد الإجازة من صلاحية شؤون الموظفين")
    emp = await db.employees.find_one({"_id": ObjectId(l["employee_id"])})
    decision = {"by_name": current_user.get("full_name", ""), "by_user_id": uid, "at": _now(), "note": (data.note or "").strip(), "action": data.action}
    step = "manager" if (l["status"] == "pending" and is_mgr) else "hr"
    new_status = "rejected" if data.action == "reject" else ("hr_pending" if step == "manager" else "approved")
    if new_status == "approved":
        _types_now = await _refresh_type_cache(db)
        if (_types_now.get(l["type"]) or {}).get("requires_attachment") and not l.get("attachment_path"):
            raise HTTPException(status_code=400, detail="هذا النوع يتطلب مرفقاً (تقرير/مستند) — لا يمكن الاعتماد قبل رفع المرفق")
        bal = await _balance(db, l["employee_id"], parse_date(l["start_date"]).year, None, l["type"], emp)
        if bal["remaining"] is not None and l["days"] > bal["remaining"]:
            raise HTTPException(status_code=400, detail=f"الرصيد الحالي ({bal['remaining']}) لا يغطي الطلب ({l['days']})")
    upd = {"status": new_status, f"{step}_decision": decision, "updated_at": _now()}
    await db.leave_requests.update_one({"_id": l["_id"]}, {"$set": upd, "$push": {"history": {"action": f"{step}_{data.action}", "by_name": decision["by_name"], "at": decision["at"], "note": decision["note"]}}})
    l.update(upd)
    if emp:
        await _apply_status(db, emp, l)
    emp_uid = (await employee_user_ids(db, [l["employee_id"]])).get(l["employee_id"])
    label, period = (_TYPE_NAMES.get(l["type"]) or LEAVE_TYPES.get(l["type"], l["type"])), f"{l['start_date']} → {l['end_date']}"
    if new_status == "rejected" and emp_uid:
        await notify_users(db, [emp_uid], "تم رفض طلب الإجازة", f"إجازة {label} ({period}) — رُفضت من {decision['by_name']}{(': ' + decision['note']) if decision['note'] else ''}", "hr_leave", {"leave_id": leave_id})
    elif new_status == "approved" and emp_uid:
        await notify_users(db, [emp_uid], "تم اعتماد إجازتك ✅", f"إجازة {label} ({period}) معتمدة من شؤون الموظفين", "hr_leave", {"leave_id": leave_id})
    elif new_status == "hr_pending":
        if emp_uid:
            await notify_users(db, [emp_uid], "وافق مديرك على طلب الإجازة", f"إجازة {label} ({period}) — بانتظار اعتماد شؤون الموظفين", "hr_leave", {"leave_id": leave_id})
        await notify_users(db, await hr_manager_user_ids(db, P_LEAVES), "طلب إجازة بانتظار اعتماد شؤون الموظفين", f"{(emp or {}).get('full_name', '')} — إجازة {label} ({period}) وافق عليها المدير المباشر", "hr_leave", {"leave_id": leave_id})
    await log_activity(current_user, f"hr_leave_{data.action}", "leave", leave_id, (emp or {}).get("full_name", ""), {"to": new_status})
    return {"status": new_status, "status_label": STATUSES[new_status], "message": f"تم — الطلب الآن: {STATUSES[new_status]}"}


@router.post("/{leave_id}/cancel")
async def cancel_leave(leave_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    l = await db.leave_requests.find_one({"_id": _oid(leave_id)})
    if not l:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    emp = await find_my_employee(db, current_user)
    mine = emp and str(emp["_id"]) == l["employee_id"]
    is_hr = has_permission(current_user, P_LEAVES)
    if not (mine or is_hr):
        raise HTTPException(status_code=403, detail="غير مصرح")
    if l["status"] in ("rejected", "cancelled"):
        raise HTTPException(status_code=400, detail="الطلب مغلق بالفعل")
    if l["status"] == "approved" and not is_hr:
        if l["start_date"] <= _today():
            raise HTTPException(status_code=400, detail="لا يمكن إلغاء إجازة معتمدة بدأت — راجع شؤون الموظفين")
    await db.leave_requests.update_one({"_id": l["_id"]}, {"$set": {"status": "cancelled", "updated_at": _now()}, "$push": {"history": {"action": "cancelled", "by_name": current_user.get("full_name", ""), "at": _now(), "note": ""}}})
    e = await db.employees.find_one({"_id": ObjectId(l["employee_id"])})
    if e and e.get("status") == "leave":
        await sync_leave_statuses(db)
    await log_activity(current_user, "hr_leave_cancel", "leave", leave_id, (e or {}).get("full_name", ""))
    return {"message": "تم إلغاء الطلب"}


# ══════════════ 📎 مرفق طلب الإجازة ══════════════
LEAVE_ATTACH_ALLOWED = {"application/pdf": "pdf", "image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}


async def _leave_access(db, l: dict, current_user: dict) -> bool:
    if _can_view(current_user) or has_permission(current_user, P_LEAVES) or l.get("manager_user_id") == user_id_of(current_user):
        return True
    emp = await find_my_employee(db, current_user)
    return bool(emp and str(emp["_id"]) == l.get("employee_id"))


@router.post("/{leave_id}/attachment")
async def upload_leave_attachment(leave_id: str, file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    """📎 رفع مرفق (تقرير طبي…) لطلب إجازة — صاحب الطلب أو شؤون الموظفين"""
    db = get_db()
    l = await db.leave_requests.find_one({"_id": _oid(leave_id, "الطلب")})
    if not l:
        raise HTTPException(status_code=404, detail="الطلب غير موجود")
    if not await _leave_access(db, l, current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    if l.get("status") in ("rejected", "cancelled"):
        raise HTTPException(status_code=400, detail="لا يمكن إرفاق ملف لطلب مرفوض/ملغى")
    if file.content_type not in LEAVE_ATTACH_ALLOWED:
        raise HTTPException(status_code=400, detail="صيغة المرفق غير مدعومة (PDF أو صورة)")
    data = await file.read()
    if not data or len(data) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="المرفق فارغ أو يتجاوز 8MB")
    from services.storage_service import upload_file
    stored = upload_file(data, file.filename or "attachment", file.content_type, f"hr_leaves/{l['employee_id']}")
    await db.leave_requests.update_one({"_id": l["_id"]}, {"$set": {"attachment_path": stored["storage_path"], "attachment_name": stored["original_filename"], "attachment_type": file.content_type, "attachment_at": _now()},
                                                           "$push": {"history": {"action": "attachment", "by_name": current_user.get("full_name", ""), "at": _now(), "note": stored["original_filename"]}}})
    return {"message": "تم رفع المرفق", "attachment_name": stored["original_filename"]}


@router.get("/{leave_id}/attachment")
async def get_leave_attachment(leave_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    l = await db.leave_requests.find_one({"_id": _oid(leave_id, "الطلب")})
    if not l or not l.get("attachment_path"):
        raise HTTPException(status_code=404, detail="لا يوجد مرفق")
    if not await _leave_access(db, l, current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    from services.storage_service import get_object
    from urllib.parse import quote
    import io
    from fastapi.responses import StreamingResponse
    try:
        data, ct = get_object(l["attachment_path"])
    except Exception:
        raise HTTPException(status_code=404, detail="تعذّر جلب المرفق")
    return StreamingResponse(io.BytesIO(data), media_type=l.get("attachment_type") or ct, headers={"Content-Disposition": f"inline; filename*=UTF-8''{quote(l.get('attachment_name') or 'attachment')}"})
