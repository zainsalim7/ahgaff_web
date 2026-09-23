"""📊 شؤون الموظفين — التقرير السنوي: القوى العاملة، اتجاه الحضور الشهري، استخدام الإجازات، توزيع تقديرات التقييم حسب الوحدة، المهام"""
import io
from datetime import datetime, date, timedelta
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from bson import ObjectId

from .deps import get_db, get_current_user, export_headers, export_filename
from .hr_common import YEMEN_TZ, _today, _can_view, get_hr_settings, is_work_day

router = APIRouter(prefix="/hr/reports", tags=["شؤون الموظفين - التقارير"])
AR_MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]


async def build_annual(db, year: int, org_unit_id: Optional[str]) -> dict:
    from .hr import CATEGORIES, STATUSES, CONTRACT_TYPES
    from .hr_leaves import LEAVE_TYPES
    from .hr_appraisals import STATUSES as APPR_STATUSES
    settings = await get_hr_settings(db)
    y0, y1 = f"{year}-01-01", f"{year}-12-31"
    q: dict = {}
    unit_name = "كل الجامعة"
    if org_unit_id:
        ids, frontier = {org_unit_id}, [org_unit_id]
        while frontier:
            kids = [str(u["_id"]) for u in await db.org_units.find({"parent_id": {"$in": frontier}}, {"_id": 1}).to_list(2000)]
            frontier = [k for k in kids if k not in ids]
            ids |= set(kids)
        q["org_unit_id"] = {"$in": list(ids)}
        u = await db.org_units.find_one({"_id": ObjectId(org_unit_id)}, {"name": 1}) if ObjectId.is_valid(org_unit_id) else None
        unit_name = (u or {}).get("name", "")
    emps = await db.employees.find(q, {"full_name": 1, "category": 1, "status": 1, "contract_type": 1, "org_unit_id": 1, "hire_date": 1, "gender": 1}).to_list(10000)
    eids = [str(e["_id"]) for e in emps]
    emp_map = {str(e["_id"]): e for e in emps}
    units = {str(u["_id"]): u.get("name", "") for u in await db.org_units.find({}, {"name": 1}).to_list(5000)}

    def count_by(key, labels):
        c: dict = {}
        for e in emps:
            c[e.get(key) or "—"] = c.get(e.get(key) or "—", 0) + 1
        return [{"key": k, "label": labels.get(k, k), "count": v} for k, v in sorted(c.items(), key=lambda x: -x[1])]

    hires = sum(1 for e in emps if (e.get("hire_date") or "")[:4] == str(year))
    exits = await db.employee_history.count_documents({"employee_id": {"$in": eids}, "action": "updated", "details.status.to": "ended", "at": {"$gte": y0, "$lte": y1 + "T23:59"}})
    by_unit_head: dict = {}
    for e in emps:
        if e.get("status") != "ended":
            by_unit_head[e.get("org_unit_id") or ""] = by_unit_head.get(e.get("org_unit_id") or "", 0) + 1

    # ── اتجاه الحضور الشهري
    today = datetime.now(YEMEN_TZ).date()
    att: dict = {}
    async for r in db.hr_attendance.find({"employee_id": {"$in": eids}, "date": {"$gte": y0, "$lte": y1}}, {"date": 1, "status": 1, "late_minutes": 1}):
        m = att.setdefault(r["date"][:7], {"present": 0, "late": 0, "absent": 0, "excused": 0, "half_day": 0, "mission": 0, "late_minutes": 0})
        m[r["status"]] = m.get(r["status"], 0) + 1
        m["late_minutes"] += int(r.get("late_minutes") or 0)
    leaves_all = await db.leave_requests.find({"employee_id": {"$in": eids}, "status": "approved", "start_date": {"$lte": y1}, "end_date": {"$gte": y0}}, {"employee_id": 1, "type": 1, "days": 1, "start_date": 1, "end_date": 1}).to_list(20000)
    months = []
    active_n = max(1, sum(1 for e in emps if e.get("status") != "ended"))
    for mi in range(1, 13):
        first = date(year, mi, 1)
        if first > today:
            break
        last = min(date(year + (mi == 12), (mi % 12) + 1, 1) - timedelta(days=1), today)
        wd = sum(1 for i in range((last - first).days + 1) if is_work_day(first + timedelta(days=i), settings))
        key = f"{year}-{mi:02d}"
        m = att.get(key, {})
        present = m.get("present", 0) + m.get("late", 0) + m.get("half_day", 0) + m.get("mission", 0)
        leave_days = 0
        for l in leaves_all:
            s, e2 = max(l["start_date"], first.strftime("%Y-%m-%d")), min(l["end_date"], last.strftime("%Y-%m-%d"))
            if s <= e2:
                leave_days += sum(1 for i in range((datetime.strptime(e2, "%Y-%m-%d").date() - datetime.strptime(s, "%Y-%m-%d").date()).days + 1) if is_work_day(datetime.strptime(s, "%Y-%m-%d").date() + timedelta(days=i), settings))
        due = max(0, wd * active_n - leave_days - m.get("excused", 0))
        recorded = present + m.get("absent", 0)
        months.append({"month": key, "label": AR_MONTHS[mi - 1], "work_days": wd, "present": present, "late": m.get("late", 0), "absent": m.get("absent", 0), "excused": m.get("excused", 0),
                       "leave_days": leave_days, "late_minutes": m.get("late_minutes", 0), "rate": round(present / recorded * 100, 1) if recorded else None, "coverage": round(recorded / due * 100, 1) if due else None})

    # ── الإجازات
    by_type: dict = {}
    by_unit_leave: dict = {}
    for l in leaves_all:
        by_type[l["type"]] = by_type.get(l["type"], 0) + l.get("days", 0)
        u = (emp_map.get(l["employee_id"]) or {}).get("org_unit_id") or ""
        by_unit_leave[u] = by_unit_leave.get(u, 0) + l.get("days", 0)
    total_leave_days = sum(by_type.values())
    req_stats = {s["_id"]: s["n"] async for s in db.leave_requests.aggregate([{"$match": {"employee_id": {"$in": eids}, "start_date": {"$regex": f"^{year}"}}}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}])}

    # ── التقييمات
    apps = await db.hr_appraisals.find({"employee_id": {"$in": eids}, "year": year}).to_list(10000)
    grades = ["ممتاز", "جيد جداً", "جيد", "مقبول", "ضعيف"]
    dist = {g: 0 for g in grades}
    per_unit: dict = {}
    scores = []
    for a in apps:
        if a["status"] in ("approved", "acknowledged") and a.get("grade"):
            dist[a["grade"]] = dist.get(a["grade"], 0) + 1
            scores.append(a.get("total_score") or 0)
            u = (emp_map.get(a["employee_id"]) or {}).get("org_unit_id") or ""
            pu = per_unit.setdefault(u, {"unit": units.get(u, "بدون وحدة"), "count": 0, "sum": 0, **{g: 0 for g in grades}})
            pu["count"] += 1; pu["sum"] += a.get("total_score") or 0; pu[a["grade"]] += 1
    appr_status = {k: sum(1 for a in apps if a["status"] == k) for k in APPR_STATUSES}
    units_rows = []
    for u, pu in sorted(per_unit.items(), key=lambda x: -x[1]["count"]):
        units_rows.append({**pu, "avg": round(pu["sum"] / pu["count"], 1) if pu["count"] else None, "headcount": by_unit_head.get(u, 0)})

    # ── المهام
    t_stats = {s["_id"]: s["n"] async for s in db.hr_tasks.aggregate([{"$match": {"assignee_employee_id": {"$in": eids}, "created_at": {"$regex": f"^{year}"}}}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}])}
    on_time = await db.hr_tasks.count_documents({"assignee_employee_id": {"$in": eids}, "created_at": {"$regex": f"^{year}"}, "status": "done", "completed_on_time": True})

    return {"year": year, "unit_name": unit_name, "generated_at": _today(),
            "workforce": {"total": len(emps), "active": sum(1 for e in emps if e.get("status") not in ("ended",)), "hires": hires, "exits": exits,
                          "by_category": count_by("category", CATEGORIES), "by_status": count_by("status", STATUSES), "by_contract": count_by("contract_type", CONTRACT_TYPES),
                          "by_unit": sorted([{"unit": units.get(u, "بدون وحدة"), "count": n, "leave_days": by_unit_leave.get(u, 0)} for u, n in by_unit_head.items()], key=lambda x: -x["count"])[:25]},
            "attendance": {"months": months, "avg_rate": round(sum(m["rate"] for m in months if m["rate"] is not None) / max(1, sum(1 for m in months if m["rate"] is not None)), 1) if any(m["rate"] is not None for m in months) else None,
                           "total_late_minutes": sum(m["late_minutes"] for m in months), "total_absent": sum(m["absent"] for m in months)},
            "leaves": {"total_days": total_leave_days, "avg_per_employee": round(total_leave_days / active_n, 1), "by_type": [{"type": k, "label": LEAVE_TYPES.get(k, k), "days": v} for k, v in sorted(by_type.items(), key=lambda x: -x[1])],
                       "requests": req_stats, "default_entitlement": settings.get("annual_leave_days", 30)},
            "appraisals": {"distribution": [{"grade": g, "count": dist[g]} for g in grades], "avg": round(sum(scores) / len(scores), 1) if scores else None, "completed": len(scores), "status": appr_status, "per_unit": units_rows},
            "tasks": {"total": sum(t_stats.values()), "done": t_stats.get("done", 0), "on_time": on_time, "open": t_stats.get("open", 0) + t_stats.get("in_progress", 0), "cancelled": t_stats.get("cancelled", 0)}}


@router.get("/annual")
async def annual(year: Optional[int] = None, org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    return await build_annual(get_db(), year or datetime.now(YEMEN_TZ).year, org_unit_id)


@router.get("/annual/export")
async def annual_export(fmt: str = "pdf", year: Optional[int] = None, org_unit_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    r = await build_annual(get_db(), year or datetime.now(YEMEN_TZ).year, org_unit_id)
    w, at, lv, ap, tk = r["workforce"], r["attendance"], r["leaves"], r["appraisals"], r["tasks"]
    kpis = [("الموظفون", w["total"]), ("على رأس العمل", w["active"]), ("تعيينات السنة", w["hires"]), ("متوسط الحضور", f"{at['avg_rate']}%" if at["avg_rate"] is not None else "—"), ("أيام الإجازات", lv["total_days"]), ("متوسط التقييم", ap["avg"] if ap["avg"] is not None else "—"), ("مهام مُنجزة", f"{tk['done']}/{tk['total']}")]
    sections = [
        {"title": "القوى العاملة حسب الفئة", "rows": [["الفئة", "العدد"]] + [[c["label"], c["count"]] for c in w["by_category"]]},
        {"title": "حسب الوحدة التنظيمية", "rows": [["الوحدة", "الموظفون", "أيام الإجازة"]] + [[u["unit"], u["count"], u["leave_days"]] for u in w["by_unit"]]},
        {"title": "اتجاه الحضور الشهري", "rows": [["الشهر", "أيام العمل", "حاضر", "متأخر", "غائب", "بعذر", "أيام إجازة", "دقائق التأخير", "نسبة الحضور"]] + [[m["label"], m["work_days"], m["present"], m["late"], m["absent"], m["excused"], m["leave_days"], m["late_minutes"], f"{m['rate']}%" if m["rate"] is not None else "—"] for m in at["months"]]},
        {"title": "استخدام الإجازات حسب النوع", "rows": [["النوع", "الأيام"]] + [[t["label"], t["days"]] for t in lv["by_type"]]},
        {"title": "توزيع تقديرات التقييم", "rows": [["التقدير", "العدد"]] + [[d["grade"], d["count"]] for d in ap["distribution"]]},
        {"title": "التقييم حسب الوحدة", "rows": [["الوحدة", "مقيَّمون", "المتوسط", "ممتاز", "جيد جداً", "جيد", "مقبول", "ضعيف"]] + [[u["unit"], u["count"], u["avg"], u["ممتاز"], u["جيد جداً"], u["جيد"], u["مقبول"], u["ضعيف"]] for u in ap["per_unit"]]},
    ]
    title = f"التقرير السنوي لشؤون الموظفين {r['year']}"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "ملخص"; ws.sheet_view.rightToLeft = True
        ws.append([title, r["unit_name"]]); ws["A1"].font = Font(bold=True, size=13)
        for k, v in kpis:
            ws.append([k, v])
        for s in sections:
            w2 = wb.create_sheet(s["title"][:30]); w2.sheet_view.rightToLeft = True
            for i, row in enumerate(s["rows"]):
                w2.append(row)
                if i == 0:
                    for c in w2[1]:
                        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440")
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=export_headers(export_filename("تقرير HR السنوي", str(r["year"]), ext="xlsx")))
    from .report_pdf import build_report_pdf
    buf = build_report_pdf(title, r["unit_name"], kpis, sections, generated_by=current_user.get("full_name", ""))
    return StreamingResponse(buf, media_type="application/pdf", headers=export_headers(export_filename("تقرير HR السنوي", str(r["year"]), ext="pdf")))
