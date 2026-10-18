"""🔑 بيانات الدخول الافتراضية (اسم المستخدم + كلمة المرور الأولية) للطلاب/المعلمين/الموظفين — أدمن فقط، معاينة + Excel"""
import io
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from .deps import get_db, get_current_user, log_activity, export_filename, export_headers

router = APIRouter(prefix="/admin/default-credentials", tags=["بيانات الدخول الافتراضية"])

KINDS = {"students": "الطلاب", "teachers": "المعلمون", "employees": "الموظفون"}
STATUS_LABEL = {
    "default": "حساب مفعّل — كلمة المرور الافتراضية سارية",
    "changed": "غُيّرت كلمة المرور من قبل المستخدم",
    "no_account": "لا حساب مفعّل — القيم المتوقعة عند التفعيل",
    "disabled": "الحساب موقوف",
}
STATUS_COLOR = {"default": "E8F5E9", "changed": "FFF3E0", "no_account": "F1F5F9", "disabled": "FFEBEE"}


def _guard(user: dict):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="هذه الأداة متاحة لمدير النظام فقط")


def _oid(s: str):
    try:
        return ObjectId(s)
    except Exception:
        raise HTTPException(status_code=400, detail="معرّف غير صالح")


async def _users_by_id(db, ids, numbers=None):
    """خريطة user_id → مستخدم، مع احتياط بالبحث عن username = الرقم لمن لا user_id له (مفتاح `u:{username}`)"""
    proj = {"username": 1, "must_change_password": 1, "is_active": 1}
    oids = [ObjectId(i) for i in ids if i and ObjectId.is_valid(i)]
    out = {str(u["_id"]): u for u in await db.users.find({"_id": {"$in": oids}}, proj).to_list(None)} if oids else {}
    nums = [n for n in (numbers or []) if n]
    if nums:
        for u in await db.users.find({"username": {"$in": nums}}, proj).to_list(None):
            out.setdefault(f"u:{u['username']}", u)
    return out


def _find_user(users: dict, user_id, number):
    return users.get(str(user_id or "")) or users.get(f"u:{number}")


def _account(user: Optional[dict], default_pw: str, default_username: str) -> dict:
    if not user:
        return {"username": default_username, "password": default_pw, "status": "no_account"}
    if user.get("is_active") is False:
        st = "disabled"
    else:
        st = "default" if user.get("must_change_password", False) else "changed"
    return {"username": user.get("username") or default_username, "password": default_pw if st != "changed" else "", "status": st}


async def _names(db):
    facs = {str(f["_id"]): f.get("name", "") for f in await db.faculties.find({}, {"name": 1}).to_list(None)}
    deps = {str(d["_id"]): d for d in await db.departments.find({}, {"name": 1, "faculty_id": 1}).to_list(None)}
    return facs, deps


async def _students(db, faculty_id, department_id, level):
    facs, deps = await _names(db)
    q: dict = {}
    if department_id:
        q["department_id"] = department_id
    elif faculty_id:
        q["department_id"] = {"$in": [k for k, d in deps.items() if str(d.get("faculty_id") or "") == faculty_id]}
    if level is not None:
        q["level"] = level
    studs = await db.students.find(q, {"student_id": 1, "reference_number": 1, "full_name": 1, "department_id": 1, "level": 1, "section": 1, "user_id": 1, "is_active": 1}).sort([("level", 1), ("full_name", 1)]).to_list(None)
    users = await _users_by_id(db, [s.get("user_id") for s in studs], [s.get("student_id") for s in studs])
    rows = []
    for s in studs:
        d = deps.get(str(s.get("department_id") or ""), {})
        acc = _account(_find_user(users, s.get("user_id"), s.get("student_id")), s.get("student_id", ""), s.get("student_id", ""))
        rows.append({"id": str(s["_id"]), "number": s.get("student_id", ""), "reference_number": s.get("reference_number", ""), "name": s.get("full_name", ""),
                     "faculty": facs.get(str(d.get("faculty_id") or ""), ""), "department": d.get("name", ""), "level": s.get("level"), "section": s.get("section", ""),
                     "record_active": s.get("is_active", True) is not False, **acc})
    return rows


async def _teachers(db, faculty_id, department_id):
    facs, deps = await _names(db)
    q: dict = {}
    if department_id:
        q["$or"] = [{"department_id": department_id}, {"department_ids": department_id}]
    elif faculty_id:
        dep_ids = [k for k, d in deps.items() if str(d.get("faculty_id") or "") == faculty_id]
        q["$or"] = [{"faculty_id": faculty_id}, {"department_id": {"$in": dep_ids}}, {"department_ids": {"$in": dep_ids}}]
    teachers = await db.teachers.find(q, {"teacher_id": 1, "full_name": 1, "department_id": 1, "department_ids": 1, "faculty_id": 1, "user_id": 1, "is_active": 1}).sort("full_name", 1).to_list(None)
    users = await _users_by_id(db, [t.get("user_id") for t in teachers], [t.get("teacher_id") for t in teachers])
    rows = []
    for t in teachers:
        dep_ids = t.get("department_ids") or ([t["department_id"]] if t.get("department_id") else [])
        fac = t.get("faculty_id") or (deps.get(dep_ids[0], {}).get("faculty_id") if dep_ids else "")
        acc = _account(_find_user(users, t.get("user_id"), t.get("teacher_id")), t.get("teacher_id", ""), t.get("teacher_id", ""))
        rows.append({"id": str(t["_id"]), "number": t.get("teacher_id", ""), "name": t.get("full_name", ""), "faculty": facs.get(str(fac or ""), ""),
                     "department": "، ".join(deps.get(x, {}).get("name", "") for x in dep_ids if x in deps), "record_active": t.get("is_active", True) is not False, **acc})
    return rows


async def _employees(db, org_unit_id, include_ended: bool):
    units = {str(u["_id"]): u.get("name", "") for u in await db.org_units.find({}, {"name": 1}).to_list(None)}
    q: dict = {}
    if org_unit_id:
        q["org_unit_id"] = org_unit_id
    if not include_ended:
        q["status"] = {"$ne": "ended"}
    emps = await db.employees.find(q, {"employee_no": 1, "full_name": 1, "org_unit_id": 1, "job_title": 1, "user_id": 1, "status": 1, "teacher_id": 1}).sort("full_name", 1).to_list(None)
    users = await _users_by_id(db, [e.get("user_id") for e in emps], [e.get("employee_no") for e in emps])
    rows = []
    for e in emps:
        acc = _account(_find_user(users, e.get("user_id"), e.get("employee_no")), e.get("employee_no", ""), e.get("employee_no", ""))
        rows.append({"id": str(e["_id"]), "number": e.get("employee_no", ""), "name": e.get("full_name", ""), "org_unit": units.get(str(e.get("org_unit_id") or ""), ""),
                     "job_title": e.get("job_title", ""), "is_teacher": bool(e.get("teacher_id")), "record_active": e.get("status") != "ended", **acc})
    return rows


async def _rows(db, kind, faculty_id, department_id, level, org_unit_id, include_ended):
    if kind not in KINDS:
        raise HTTPException(status_code=400, detail="النوع غير معروف")
    if kind == "students":
        return await _students(db, faculty_id, department_id, level)
    if kind == "teachers":
        return await _teachers(db, faculty_id, department_id)
    return await _employees(db, org_unit_id, include_ended)


def _summary(rows):
    s = {k: 0 for k in STATUS_LABEL}
    for r in rows:
        s[r["status"]] += 1
    return {"total": len(rows), **s}


async def _scope_label(db, kind, faculty_id, department_id, level, org_unit_id) -> str:
    parts = []
    if kind == "employees":
        if org_unit_id:
            u = await db.org_units.find_one({"_id": _oid(org_unit_id)}, {"name": 1})
            parts.append(u.get("name", "") if u else "")
    else:
        if department_id:
            d = await db.departments.find_one({"_id": _oid(department_id)}, {"name": 1})
            parts.append(d.get("name", "") if d else "")
        elif faculty_id:
            f = await db.faculties.find_one({"_id": _oid(faculty_id)}, {"name": 1})
            parts.append(f.get("name", "") if f else "")
        if level is not None:
            parts.append(f"المستوى {level}")
    return " - ".join(p for p in parts if p) or "الكل"


@router.get("/preview")
async def preview(kind: str, faculty_id: Optional[str] = None, department_id: Optional[str] = None, level: Optional[int] = None,
                  org_unit_id: Optional[str] = None, include_ended: bool = False, current_user: dict = Depends(get_current_user)):
    _guard(current_user)
    db = get_db()
    rows = await _rows(db, kind, faculty_id, department_id, level, org_unit_id, include_ended)
    return {"kind": kind, "rows": rows, "summary": _summary(rows), "status_labels": STATUS_LABEL}


@router.get("/export")
async def export_xlsx(kind: str, faculty_id: Optional[str] = None, department_id: Optional[str] = None, level: Optional[int] = None,
                      org_unit_id: Optional[str] = None, include_ended: bool = False, status: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    _guard(current_user)
    db = get_db()
    rows = await _rows(db, kind, faculty_id, department_id, level, org_unit_id, include_ended)
    if status:
        rows = [r for r in rows if r["status"] == status]
    scope = await _scope_label(db, kind, faculty_id, department_id, level, org_unit_id)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if kind == "students":
        cols = ["#", "رقم القيد", "الرقم المرجعي", "الاسم", "الكلية", "القسم", "المستوى", "الشعبة", "اسم المستخدم", "كلمة المرور الافتراضية", "حالة الحساب"]
        line = lambda i, r: [i, r["number"], r.get("reference_number", ""), r["name"], r["faculty"], r["department"], r.get("level"), r.get("section", ""), r["username"], r["password"], STATUS_LABEL[r["status"]]]
    elif kind == "teachers":
        cols = ["#", "الرقم الوظيفي", "الاسم", "الكلية", "القسم", "اسم المستخدم", "كلمة المرور الافتراضية", "حالة الحساب"]
        line = lambda i, r: [i, r["number"], r["name"], r["faculty"], r["department"], r["username"], r["password"], STATUS_LABEL[r["status"]]]
    else:
        cols = ["#", "الرقم الوظيفي", "الاسم", "الوحدة الإدارية", "المسمى الوظيفي", "اسم المستخدم", "كلمة المرور الافتراضية", "حالة الحساب"]
        line = lambda i, r: [i, r["number"], r["name"], r["org_unit"], r["job_title"], r["username"], r["password"], STATUS_LABEL[r["status"]]]

    wb = Workbook(); ws = wb.active; ws.title = KINDS[kind]; ws.sheet_view.rightToLeft = True
    ws.append([f"بيانات الدخول الافتراضية — {KINDS[kind]} — {scope} ({len(rows)})"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws["A1"].font = Font(bold=True, size=13); ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["⚠️ كلمات المرور مشفّرة في النظام: تظهر هنا كلمة المرور الافتراضية فقط لمن لم يغيّرها بعد؛ من غيّرها تبقى خليته فارغة. سرّي — لا يُشارك إلا مع صاحب الحساب."])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws["A2"].font = Font(size=10, color="92400E"); ws["A2"].alignment = Alignment(horizontal="right")
    ws.append(cols)
    thin = Side(style="thin", color="CBD5E1"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440"); c.alignment = Alignment(horizontal="center"); c.border = border
    pw_col = cols.index("كلمة المرور الافتراضية") + 1
    for i, r in enumerate(rows, 1):
        ws.append(line(i, r))
        fill = PatternFill("solid", fgColor=STATUS_COLOR[r["status"]])
        for c in ws[ws.max_row]:
            c.fill = fill; c.border = border; c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=ws.max_row, column=pw_col).font = Font(bold=True, color="1B5E20" if r["password"] else "94A3B8")
        if not r["password"]:
            ws.cell(row=ws.max_row, column=pw_col).value = "—"
    for ci in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 16
    ws.column_dimensions[get_column_letter(cols.index("الاسم") + 1)].width = 30
    ws.column_dimensions[get_column_letter(len(cols))].width = 44
    ws.freeze_panes = "A4"
    await log_activity(current_user, "export_default_credentials", kind, scope, f"تصدير بيانات الدخول الافتراضية — {KINDS[kind]}", {"count": len(rows), "scope": scope})
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=export_headers(export_filename("بيانات الدخول الافتراضية", KINDS[kind], scope, f"{len(rows)}", ext="xlsx")))
