"""البطاقة الرقمية للطالب: صورة + QR تحقق عام + قوالب تصميم لكل كلية + اعتماد صور الطلاب.
صلاحية البطاقة: العام الجامعي النشط — تتجدد تلقائياً مع كل عام.
"""
import io
import uuid
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends, Query, UploadFile, File, Request
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, log_activity, export_filename, export_headers
from .statements import _can_issue as _can_manage, get_verify_base

router = APIRouter()

TEMPLATES = ("green", "dark", "horizontal", "official", "custom")
CARD_SCALE = 2  # دقة مضاعفة للطباعة (1280×2020)

# 🔤 خطوط البطاقة: (ملف الخط، محاور الخط المتغير [slnt, wght] أو [wght]) — كلها عريضة للطباعة
CARD_FONTS = {
    "amiri": ("Amiri-Bold.ttf", None),
    "cairo": ("Cairo-Variable.ttf", [700, 0]),
    "tajawal": ("Tajawal-Bold.ttf", None),
    "kufi": ("NotoKufiArabic-Variable.ttf", [700]),
    "almarai": ("Almarai-Bold.ttf", None),
    "changa": ("Changa-Variable.ttf", [700]),
}
DEFAULT_CARD_FONT = "kufi"
CARD_FONT_LABELS = {"amiri": "أميري عريض (نسخي كلاسيكي)", "cairo": "القاهرة عريض (عصري)", "tajawal": "تجوّل عريض (بسيط واضح)", "kufi": "نوتو كوفي عريض (كوفي رسمي)", "almarai": "المرعي عريض (هندسي)", "changa": "تشانغا عريض (مميز)"}

# 📐 المواضع الافتراضية لعناصر القالب المخصص (نسب مئوية من أبعاد البطاقة)
DEFAULT_CUSTOM_LAYOUT = {
    "photo": {"x": 66, "y": 16, "w": 26, "h": 24},
    "qr": {"x": 8, "y": 66, "w": 18},
    "name": {"x": 50, "y": 46, "size": 34, "color": "#1a2540"},
    "enrollment": {"x": 50, "y": 55, "size": 26, "color": "#1a2540"},
    "dept": {"x": 50, "y": 62, "size": 24, "color": "#1a2540"},
    "year": {"x": 50, "y": 69, "size": 22, "color": "#455a64"},
}


def _payload_portrait(p: dict) -> bool:
    t = p.get("template", "green")
    if t == "horizontal":
        return False
    if t == "custom":
        return p.get("custom_orientation") != "landscape"
    return True
LEVEL_AR = {1: "الأول", 2: "الثاني", 3: "الثالث", 4: "الرابع", 5: "الخامس", 6: "السادس", 7: "السابع", 8: "الثامن"}


async def _active_academic_year(db) -> str:
    sem = await db.semesters.find_one({"status": "active"})
    ay = (sem or {}).get("academic_year") or ""
    if not ay:
        y = datetime.now(timezone.utc).year
        ay = f"{y}-{y + 1}"
    return ay


async def _ensure_card(db, student: dict) -> dict:
    """يضمن وجود توكن بطاقة صالح للعام الجامعي النشط (تجديد تلقائي)."""
    year = await _active_academic_year(db)
    if student.get("card_token") and student.get("card_academic_year") == year:
        return {"token": student["card_token"], "academic_year": year}
    token = uuid.uuid4().hex
    await db.students.update_one({"_id": student["_id"]}, {"$set": {
        "card_token": token, "card_academic_year": year,
        "card_issued_at": datetime.now(timezone.utc).isoformat(),
    }})
    return {"token": token, "academic_year": year}


async def _student_or_404(db, student_id: str) -> dict:
    try:
        student = await db.students.find_one({"_id": ObjectId(student_id)})
    except Exception:
        student = None
    if not student:
        raise HTTPException(status_code=404, detail="الطالب غير موجود")
    return student


async def _resolve_faculty(db, student: dict):
    dept = None
    if student.get("department_id"):
        try:
            dept = await db.departments.find_one({"_id": ObjectId(student["department_id"])})
        except Exception:
            dept = None
    fid = student.get("faculty_id") or (dept or {}).get("faculty_id", "")
    faculty = None
    if fid:
        try:
            faculty = await db.faculties.find_one({"_id": ObjectId(fid)})
        except Exception:
            faculty = None
    return fid, (faculty or {}).get("name", ""), (dept or {}).get("name", "")


# ==================== إعدادات التصميم لكل كلية ====================
class CardSettings(BaseModel):
    template: str = "green"
    font: Optional[str] = None
    custom_bg_base64: Optional[str] = None
    custom_layout: Optional[dict] = None


@router.get("/cards/settings/{faculty_id}")
async def get_card_settings(faculty_id: str, current_user: dict = Depends(get_current_user)):
    if not _can_manage(current_user, faculty_id):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    db = get_db()
    doc = await db.card_settings.find_one({"_id": f"faculty_{faculty_id}"}) or {}
    return {
        "template": doc.get("template", "green"),
        "font": doc.get("font") or DEFAULT_CARD_FONT,
        "fonts": [{"key": k, "label": v} for k, v in CARD_FONT_LABELS.items()],
        "custom_bg_base64": doc.get("custom_bg_base64", ""),
        "custom_layout": doc.get("custom_layout") or DEFAULT_CUSTOM_LAYOUT,
        "custom_orientation": doc.get("custom_orientation", "portrait"),
    }


@router.put("/cards/settings/{faculty_id}")
async def update_card_settings(faculty_id: str, data: CardSettings, current_user: dict = Depends(get_current_user)):
    if not _can_manage(current_user, faculty_id):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    if data.template not in TEMPLATES:
        raise HTTPException(status_code=400, detail="قالب غير معروف")
    db = get_db()
    if data.font and data.font not in CARD_FONTS:
        raise HTTPException(status_code=400, detail="خط غير معروف")
    update = {"template": data.template, "font": data.font or DEFAULT_CARD_FONT}
    if data.custom_layout is not None:
        update["custom_layout"] = data.custom_layout
    if data.custom_bg_base64:
        raw = data.custom_bg_base64.split(",")[-1]
        try:
            import base64 as _b
            from PIL import Image as _I
            im = _I.open(io.BytesIO(_b.b64decode(raw)))
            update["custom_orientation"] = "landscape" if im.width > im.height else "portrait"
            update["custom_bg_base64"] = raw
        except Exception:
            raise HTTPException(status_code=400, detail="تعذر قراءة صورة التصميم — ارفع PNG أو JPG صالحة")
    if data.template == "custom":
        existing = await db.card_settings.find_one({"_id": f"faculty_{faculty_id}"}) or {}
        if not (data.custom_bg_base64 or existing.get("custom_bg_base64")):
            raise HTTPException(status_code=400, detail="ارفع صورة تصميم البطاقة أولاً")
    await db.card_settings.update_one({"_id": f"faculty_{faculty_id}"}, {"$set": update}, upsert=True)
    return {"message": "تم حفظ تصميم البطاقة"}


# ==================== بيانات البطاقة ====================
async def _card_payload(db, student: dict, base_url: str) -> dict:
    fid, faculty_name, dept_name = await _resolve_faculty(db, student)
    card = await _ensure_card(db, student)
    settings = await db.card_settings.find_one({"_id": f"faculty_{fid}"}) or {}
    verify_base = (await get_verify_base(db)) or (base_url or "").rstrip("/")
    verify_url = f"{verify_base}/verify-card?token={card['token']}" if verify_base else card["token"]
    return {
        "student_db_id": str(student["_id"]),
        "student_name": student.get("full_name", ""),
        "enrollment_no": student.get("student_id", ""),
        "reference_number": student.get("reference_number", ""),
        "nationality": student.get("nationality") or "يمني",
        "level": student.get("level") or 1,
        "section": student.get("section", ""),
        "faculty_id": fid,
        "faculty_name": faculty_name,
        "department_name": dept_name,
        "academic_year": card["academic_year"],
        "template": settings.get("template", "green"),
        "font": settings.get("font") or DEFAULT_CARD_FONT,
        "photo_upload_used": bool(student.get("photo_upload_used")),
        "photo_upload_allowed": bool(student.get("photo_upload_allowed")),
        "can_upload_photo": (not student.get("photo_upload_used")) or bool(student.get("photo_upload_allowed")),
        "photo_guidelines": {
            "title": "📸 صورتك على بطاقتك",
            "lines": [
                "صورتك تعبّر عن شخصيتك وستُطبع على بطاقتك الجامعية الرسمية",
                "استخدم صورة شخصية واضحة بمقاس 4×6 (عمودية)",
                "الخلفية بيضاء أو فاتحة موحّدة",
                "وجه واضح بمواجهة الكاميرا — بدون فلاتر أو نظارة شمسية",
            ],
            "confirm_text": "فهمت، اختيار الصورة",
        },
        "custom_bg_base64": settings.get("custom_bg_base64", ""),
        "custom_layout": settings.get("custom_layout") or {},
        "custom_orientation": settings.get("custom_orientation", "portrait"),
        "photo_path": student.get("photo_path", ""),
        "pending_photo_path": student.get("pending_photo_path", ""),
        "card_token": card["token"],
        "verify_url": verify_url,
        "status": student.get("status", "active"),
    }


@router.get("/students/me/card")
async def get_my_card(base_url: Optional[str] = Query(None), current_user: dict = Depends(get_current_user)):
    db = get_db()
    student = await db.students.find_one({"user_id": current_user.get("id")})
    if not student:
        raise HTTPException(status_code=404, detail="لا يوجد ملف طالب مرتبط بحسابك")
    return await _card_payload(db, student, base_url or "")


@router.get("/students/{student_id}/card")
async def get_student_card(student_id: str, base_url: Optional[str] = Query(None), current_user: dict = Depends(get_current_user)):
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    return await _card_payload(db, student, base_url or "")


# ==================== صور الطلاب ====================
ALLOWED_PHOTO_TYPES = ("image/jpeg", "image/png", "image/webp")


async def _store_photo(file: UploadFile) -> str:
    if file.content_type not in ALLOWED_PHOTO_TYPES:
        raise HTTPException(status_code=400, detail="نوع الصورة غير مدعوم (JPEG/PNG/WebP)")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="حجم الصورة يتجاوز 5MB")
    from services.storage_service import upload_file
    result = upload_file(data, file.filename or "photo.jpg", file.content_type, "student_photos")
    return result["storage_path"]


@router.post("/students/me/photo")
async def upload_my_photo(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    """الطالب يرفع صورته من تطبيقه — تبقى معلقة حتى يعتمدها المسجل."""
    db = get_db()
    student = await db.students.find_one({"user_id": current_user.get("id")})
    if not student:
        raise HTTPException(status_code=404, detail="لا يوجد ملف طالب مرتبط بحسابك")
    # 🔒 حد الرفع: مرة واحدة فقط — الرفع الثاني بإذن الإدارة
    if student.get("photo_upload_used") and not student.get("photo_upload_allowed"):
        raise HTTPException(status_code=403, detail="لقد استخدمت فرصة رفع الصورة المتاحة — لتغيير صورتك راجع إدارة الكلية للسماح برفع جديد")
    path = await _store_photo(file)
    await db.students.update_one({"_id": student["_id"]}, {"$set": {
        "pending_photo_path": path,
        "pending_photo_at": datetime.now(timezone.utc).isoformat(),
        "photo_upload_used": True,
    }, "$unset": {"photo_upload_allowed": ""}})
    return {"message": "تم رفع صورتك — بانتظار اعتماد المسجل", "pending_photo_path": path}


@router.post("/students/{student_id}/photo")
async def upload_student_photo(student_id: str, file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    """الإدارة ترفع صورة الطالب — تُعتمد مباشرة."""
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    path = await _store_photo(file)
    await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_path": path}, "$unset": {"pending_photo_path": "", "pending_photo_at": ""}})
    await log_activity(current_user, "set_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {})
    return {"message": "تم حفظ صورة الطالب", "photo_path": path}


@router.get("/pending-photos")
async def list_pending_photos(current_user: dict = Depends(get_current_user)):
    db = get_db()
    if current_user.get("role") in ("teacher", "student"):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    q = {"pending_photo_path": {"$exists": True, "$ne": ""}}
    if current_user.get("role") != "admin":
        fids = set(current_user.get("faculty_ids") or [])
        if current_user.get("faculty_id"):
            fids.add(current_user["faculty_id"])
        if fids:
            q["faculty_id"] = {"$in": list(fids)}
    items = []
    async for s in db.students.find(q).limit(300):
        items.append({
            "id": str(s["_id"]),
            "full_name": s.get("full_name", ""),
            "student_id": s.get("student_id", ""),
            "level": s.get("level"),
            "pending_photo_path": s.get("pending_photo_path", ""),
            "pending_photo_at": s.get("pending_photo_at", ""),
            "current_photo_path": s.get("photo_path", ""),
        })
    return items


@router.get("/approved-photos")
async def list_approved_photos(search: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """🖼️ الصور المعتمدة (الظاهرة على البطاقات)"""
    db = get_db()
    if current_user.get("role") in ("teacher", "student"):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    q: dict = {"photo_path": {"$exists": True, "$ne": ""}}
    if current_user.get("role") != "admin":
        fids = set(current_user.get("faculty_ids") or [])
        if current_user.get("faculty_id"):
            fids.add(current_user["faculty_id"])
        if fids:
            q["faculty_id"] = {"$in": list(fids)}
    if search:
        q["$or"] = [
            {"full_name": {"$regex": search, "$options": "i"}},
            {"student_id": {"$regex": search, "$options": "i"}},
        ]
    items = []
    async for s in db.students.find(q).sort("full_name", 1).limit(500):
        items.append({
            "id": str(s["_id"]),
            "full_name": s.get("full_name", ""),
            "student_id": s.get("student_id", ""),
            "level": s.get("level"),
            "photo_path": s.get("photo_path", ""),
        })
    return items


@router.post("/students/{student_id}/photo/revoke")
async def revoke_student_photo(student_id: str, current_user: dict = Depends(get_current_user)):
    """↩️ إلغاء اعتماد الصورة: سحبها من البطاقة + فتح فرصة رفع جديدة للطالب"""
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    if not student.get("photo_path"):
        raise HTTPException(status_code=400, detail="لا توجد صورة معتمدة لهذا الطالب")
    await db.students.update_one(
        {"_id": student["_id"]},
        {"$set": {"photo_upload_allowed": True}, "$unset": {"photo_path": ""}},
    )
    await log_activity(current_user, "revoke_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {})
    await _notify_photo_decision(db, student, approved=False)
    return {"message": "تم إلغاء اعتماد الصورة وفتح فرصة رفع جديدة للطالب"}


# ==================== إشعار قرار الصورة ====================
async def _notify_photo_decision(db, student: dict, approved: bool):
    """إشعار داخل التطبيق + Push للطالب عند اعتماد/رفض صورته."""
    user_id = student.get("user_id")
    if not user_id:
        return
    if approved:
        title = "✅ تم اعتماد صورتك الشخصية"
        message = "تم اعتماد صورتك الشخصية من المسجل وأصبحت ظاهرة على بطاقتك الرقمية."
    else:
        title = "❌ لم يتم اعتماد صورتك الشخصية"
        message = "نعتذر، لم يتم اعتماد الصورة التي رفعتها. يرجى رفع صورة شخصية واضحة وبخلفية مناسبة من شاشة البطاقة."
    from datetime import timedelta
    await db.notifications.insert_one({
        "student_id": str(student["_id"]),
        "user_id": user_id,
        "title": title,
        "message": message,
        "type": "photo_approved" if approved else "photo_rejected",
        "is_read": False,
        "created_at": datetime.now(timezone(timedelta(hours=3))).isoformat(),
    })
    try:
        from services.firebase_service import send_notification_to_many
        tokens = [d["token"] async for d in db.fcm_tokens.find({"user_id": user_id})]
        if tokens:
            await send_notification_to_many(tokens, title, message)
    except Exception:
        pass


@router.post("/students/{student_id}/photo/approve")
async def approve_student_photo(student_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    pending = student.get("pending_photo_path")
    if not pending:
        raise HTTPException(status_code=400, detail="لا توجد صورة معلقة لهذا الطالب")
    await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_path": pending}, "$unset": {"pending_photo_path": "", "pending_photo_at": ""}})
    await log_activity(current_user, "approve_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {})
    await _notify_photo_decision(db, student, approved=True)
    return {"message": "تم اعتماد صورة الطالب"}


@router.post("/students/{student_id}/photo/reject")
async def reject_student_photo(student_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    # الرفض يعيد فتح فرصة الرفع للطالب تلقائياً ليرفع صورة أفضل
    await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_upload_allowed": True}, "$unset": {"pending_photo_path": "", "pending_photo_at": ""}})
    await log_activity(current_user, "reject_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {})
    if student.get("pending_photo_path"):
        await _notify_photo_decision(db, student, approved=False)
    return {"message": "تم رفض الصورة المعلقة"}


@router.post("/photos/bulk-decision")
async def bulk_photo_decision(request: Request, current_user: dict = Depends(get_current_user)):
    """✅❌ اعتماد/رفض جماعي لصور الطلاب المعلقة"""
    data = await request.json()
    student_ids = data.get("student_ids") or []
    action = data.get("action")
    if action not in ("approve", "reject") or not student_ids:
        raise HTTPException(status_code=400, detail="حدد الطلاب والإجراء")
    db = get_db()
    done, skipped = 0, 0
    for sid in student_ids[:300]:
        try:
            student = await _student_or_404(db, sid)
        except HTTPException:
            skipped += 1
            continue
        fid, _, _ = await _resolve_faculty(db, student)
        if not _can_manage(current_user, fid):
            skipped += 1
            continue
        pending = student.get("pending_photo_path")
        if action == "approve":
            if not pending:
                skipped += 1
                continue
            await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_path": pending}, "$unset": {"pending_photo_path": "", "pending_photo_at": ""}})
            await log_activity(current_user, "approve_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {"bulk": True})
            await _notify_photo_decision(db, student, approved=True)
        else:
            await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_upload_allowed": True}, "$unset": {"pending_photo_path": "", "pending_photo_at": ""}})
            await log_activity(current_user, "reject_student_photo", "student", str(student["_id"]), student.get("full_name", ""), {"bulk": True})
            if pending:
                await _notify_photo_decision(db, student, approved=False)
        done += 1
    verb = "اعتماد" if action == "approve" else "رفض"
    return {"message": f"تم {verb} {done} صورة" + (f" (تخطي {skipped})" if skipped else ""), "done": done, "skipped": skipped}


@router.post("/students/{student_id}/photo/allow-upload")
async def allow_photo_upload(student_id: str, current_user: dict = Depends(get_current_user)):
    """🔓 الإدارة تسمح للطالب برفع صورة جديدة (فرصة واحدة تُستهلك عند الرفع)."""
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    await db.students.update_one({"_id": student["_id"]}, {"$set": {"photo_upload_allowed": True}})
    await log_activity(current_user, "allow_photo_upload", "student", str(student["_id"]), student.get("full_name", ""), {})
    if student.get("user_id"):
        from datetime import timedelta as _td
        await db.notifications.insert_one({
            "student_id": str(student["_id"]),
            "user_id": student["user_id"],
            "title": "🔓 سُمح لك برفع صورة جديدة",
            "message": "سمحت لك إدارة الكلية برفع صورة شخصية جديدة لبطاقتك — ارفعها من شاشة البطاقة (فرصة واحدة).",
            "type": "photo_upload_allowed",
            "is_read": False,
            "created_at": datetime.now(timezone(_td(hours=3))).isoformat(),
        })
    return {"message": "تم السماح للطالب برفع صورة جديدة"}


# ==================== التحقق العام ====================
@router.get("/verify/card/{token}")
async def verify_card(token: str):
    """تحقق عام من البطاقة — بدون تسجيل دخول."""
    db = get_db()
    s = await db.students.find_one({"card_token": token})
    if not s:
        return {"valid": False, "message": "لا توجد بطاقة بهذا الرمز — قد تكون البطاقة غير صحيحة"}
    year = await _active_academic_year(db)
    _, faculty_name, dept_name = await _resolve_faculty(db, s)
    base = {
        "student_name": s.get("full_name", ""),
        "enrollment_no": s.get("student_id", ""),
        "faculty_name": faculty_name,
        "department_name": dept_name,
        "level": s.get("level"),
        "nationality": s.get("nationality") or "يمني",
        "academic_year": s.get("card_academic_year", ""),
        "has_photo": bool(s.get("photo_path")),
    }
    if s.get("is_alumni") or s.get("status") in ("expelled",):
        return {"valid": False, "message": "هذه البطاقة لم تعد سارية — الطالب غير مقيد حالياً", **base}
    if s.get("card_academic_year") != year:
        return {"valid": False, "message": f"انتهت صلاحية هذه البطاقة (كانت للعام الجامعي {s.get('card_academic_year', '')})", **base}
    return {"valid": True, "message": "بطاقة طالب سارية صادرة رسمياً من جامعة الأحقاف", **base}


@router.get("/public/card-photo/{token}")
async def public_card_photo(token: str):
    """صورة الطالب المعتمدة لصفحة التحقق العامة (التوكن غير قابل للتخمين)."""
    db = get_db()
    s = await db.students.find_one({"card_token": token})
    if not s or not s.get("photo_path"):
        raise HTTPException(status_code=404, detail="لا توجد صورة")
    from services.storage_service import get_object
    try:
        data, content_type = get_object(s["photo_path"])
    except Exception:
        raise HTTPException(status_code=404, detail="لا توجد صورة")
    return Response(content=data, media_type=content_type or "image/jpeg")


# ==================== الطباعة الدفعية ====================
DEFAULT_PRINT_SETTINGS = {
    "card_w": 85.6, "card_h": 54.0,
    "card1_x": 62.0, "card1_y": 40.0,
    "card2_x": 62.0, "card2_y": 180.0,
}
ORIENTATIONS = ("auto", "portrait", "landscape")


@router.get("/cards/print-settings")
async def get_print_settings(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") in ("teacher", "student"):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    db = get_db()
    doc = await db.card_print_settings.find_one({"_id": "global"}) or {}
    return {**DEFAULT_PRINT_SETTINGS, "orientation": doc.get("orientation", "auto"),
            **{k: v for k, v in doc.items() if k not in ("_id", "orientation")}}


class BatchPrintRequest(BaseModel):
    department_id: Optional[str] = None
    student_ids: Optional[list] = None
    level: Optional[int] = None
    section: Optional[str] = None
    base_url: Optional[str] = None
    orientation: Optional[str] = "auto"
    settings: Optional[dict] = None
    # 🖨️ تتبّع الطباعة (دفعات): استبعاد المطبوعين سابقاً / فقط من له صورة معتمدة / استبعاد أفراد
    exclude_printed: bool = True
    only_with_photo: bool = True
    exclude_ids: Optional[list] = None
    reprint_batch_no: Optional[int] = None  # إعادة تنزيل دفعة سابقة كما هي (لا تُسجَّل كطباعة جديدة)


def _print_flags(s: dict, exclude_printed: bool, only_with_photo: bool, excluded: set) -> dict:
    """تصنيف الطالب ضمن دفعة الطباعة: مشمول أو مستبعد (مع السبب)"""
    printed_at = s.get("card_printed_at")
    has_photo = bool(s.get("photo_path"))
    reason = ""
    if str(s["_id"]) in excluded:
        reason = "مستبعد يدوياً"
    elif exclude_printed and printed_at:
        reason = f"طُبعت سابقاً (دفعة #{s.get('card_print_batch_no', '')} — {str(printed_at)[:10]})"
    elif only_with_photo and not has_photo:
        reason = "بانتظار اعتماد الصورة"
    return {"id": str(s["_id"]), "full_name": s.get("full_name", ""), "student_id": s.get("student_id", ""), "level": s.get("level"), "section": s.get("section", ""),
            "has_photo": has_photo, "printed_at": printed_at, "batch_no": s.get("card_print_batch_no"), "print_count": s.get("card_print_count", 0), "included": not reason, "reason": reason}


async def _batch_candidates(db, data: "BatchPrintRequest", current_user: dict):
    """الطلاب المرشحون للدفعة (قسم/مستوى/شعبة أو IDs) مع بيانات الكلية/القسم"""
    ids_mode = bool(data.student_ids)
    dept = None
    if ids_mode:
        oids = []
        for i in (data.student_ids or [])[:400]:
            try:
                oids.append(ObjectId(str(i)))
            except Exception:
                pass
        students = [s async for s in db.students.find({"_id": {"$in": oids}}).sort("full_name", 1)]
    else:
        if not data.department_id:
            raise HTTPException(status_code=400, detail="اختر القسم أو حدد طلاباً")
        try:
            dept = await db.departments.find_one({"_id": ObjectId(data.department_id)})
        except Exception:
            dept = None
        if not dept:
            raise HTTPException(status_code=404, detail="القسم غير موجود")
        if not _can_manage(current_user, dept.get("faculty_id", "")):
            raise HTTPException(status_code=403, detail="غير مصرح لك")
        q = {"department_id": data.department_id, "is_alumni": {"$ne": True}}
        if data.level:
            q["level"] = data.level
        if data.section:
            q["section"] = data.section
        students = [s async for s in db.students.find(q).sort("full_name", 1)]
    return ids_mode, dept, students


@router.post("/cards/batch-preview")
async def batch_preview(data: BatchPrintRequest, current_user: dict = Depends(get_current_user)):
    """🔎 معاينة الدفعة قبل التوليد: من سيُطبع ومن سيُستبعد (طُبع سابقاً / بلا صورة) — لا يغيّر شيئاً"""
    db = get_db()
    _, _, students = await _batch_candidates(db, data, current_user)
    excluded = {str(x) for x in (data.exclude_ids or [])}
    rows = [_print_flags(s, data.exclude_printed, data.only_with_photo, excluded) for s in students]
    inc = [r for r in rows if r["included"]]
    return {"rows": rows, "summary": {"total": len(rows), "included": len(inc), "pages": (len(inc) + 1) // 2,
                                      "printed_before": sum(1 for r in rows if r["printed_at"]), "no_photo": sum(1 for r in rows if not r["has_photo"]),
                                      "excluded_printed": sum(1 for r in rows if not r["included"] and r["reason"].startswith("طُبعت")),
                                      "excluded_no_photo": sum(1 for r in rows if not r["included"] and r["reason"].startswith("بانتظار")),
                                      "excluded_manual": sum(1 for r in rows if r["reason"] == "مستبعد يدوياً")}}


@router.post("/cards/batch-pdf")
async def batch_print_cards(data: BatchPrintRequest, current_user: dict = Depends(get_current_user)):
    """PDF واحد: بطاقتان في كل ورقة A4 بمواضع قابلة للضبط (ملم) — حسب القسم أو طلاب محددين."""
    db = get_db()
    if data.reprint_batch_no:
        rb = await db.card_print_batches.find_one({"batch_no": int(data.reprint_batch_no)})
        if not rb:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
        data.student_ids = rb.get("student_ids") or []
    ids_mode = bool(data.student_ids)
    dept = None
    if ids_mode:
        oids = []
        for i in (data.student_ids or [])[:400]:
            try:
                oids.append(ObjectId(str(i)))
            except Exception:
                pass
        students = [s async for s in db.students.find({"_id": {"$in": oids}}).sort("full_name", 1)]
        if not students:
            raise HTTPException(status_code=404, detail="لا يوجد طلاب مطابقون")
        tpl_map = {}
        for s in students:
            fid_s, fac_name_s, dept_name_s = await _resolve_faculty(db, s)
            if not _can_manage(current_user, fid_s):
                raise HTTPException(status_code=403, detail=f"غير مصرح لك بطباعة بطاقات كلية الطالب {s.get('full_name', '')}")
            if fid_s not in tpl_map:
                tpl_map[fid_s] = await db.card_settings.find_one({"_id": f"faculty_{fid_s}"}) or {}
            s["_fid"], s["_fac_name"], s["_dept_name"] = fid_s, fac_name_s, dept_name_s
        fid = students[0].get("_fid", "")
        tpl_settings = tpl_map.get(fid, {})
        tpl = tpl_settings.get("template", "green")
    else:
        if not data.department_id:
            raise HTTPException(status_code=400, detail="اختر القسم أو حدد طلاباً")
        try:
            dept = await db.departments.find_one({"_id": ObjectId(data.department_id)})
        except Exception:
            dept = None
        if not dept:
            raise HTTPException(status_code=404, detail="القسم غير موجود")
        fid = dept.get("faculty_id", "")
        if not _can_manage(current_user, fid):
            raise HTTPException(status_code=403, detail="غير مصرح لك")

    # حفظ إعدادات المواضع للاستخدام القادم
    orientation = data.orientation if data.orientation in ORIENTATIONS else "auto"
    st = {**DEFAULT_PRINT_SETTINGS}
    if data.settings:
        for k in DEFAULT_PRINT_SETTINGS:
            try:
                st[k] = float(data.settings.get(k, st[k]))
            except (TypeError, ValueError):
                pass
        await db.card_print_settings.update_one({"_id": "global"}, {"$set": {**st, "orientation": orientation}}, upsert=True)

    if not ids_mode:
        q = {"department_id": data.department_id, "is_alumni": {"$ne": True}}
        if data.level:
            q["level"] = data.level
        if data.section:
            q["section"] = data.section
        students = [s async for s in db.students.find(q).sort("full_name", 1)]
        if not students:
            raise HTTPException(status_code=404, detail="لا يوجد طلاب مطابقون")
        if len(students) > 400:
            raise HTTPException(status_code=400, detail="العدد يتجاوز 400 طالب — قسّم الطلبات حسب المستوى")

    # 🖨️ تتبّع الدفعات: استبعاد المطبوعين سابقاً / بلا صورة / المستبعدين يدوياً (إلا في إعادة تنزيل دفعة سابقة)
    reprint = None
    if data.reprint_batch_no:
        reprint = await db.card_print_batches.find_one({"batch_no": int(data.reprint_batch_no)})
        if not reprint:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
    else:
        excluded = {str(x) for x in (data.exclude_ids or [])}
        flagged = [(s, _print_flags(s, data.exclude_printed, data.only_with_photo, excluded)) for s in students]
        skipped = [f for _, f in flagged if not f["included"]]
        students = [s for s, f in flagged if f["included"]]
        if not students:
            n_p = sum(1 for f in skipped if f["reason"].startswith("طُبعت")); n_ph = sum(1 for f in skipped if f["reason"].startswith("بانتظار"))
            raise HTTPException(status_code=400, detail=f"لا توجد بطاقات جديدة للطباعة — {n_p} طُبعت سابقاً، {n_ph} بانتظار الصورة" + (f"، {len(skipped) - n_p - n_ph} مستبعد يدوياً" if len(skipped) - n_p - n_ph else ""))

    faculty = await db.faculties.find_one({"_id": ObjectId(fid)}) if fid and not ids_mode else None
    faculty_name = (faculty or {}).get("name", "")
    if not ids_mode:
        tpl_settings = await db.card_settings.find_one({"_id": f"faculty_{fid}"}) or {}
        tpl = tpl_settings.get("template", "green")
    year = await _active_academic_year(db)
    base = (await get_verify_base(db)) or (data.base_url or "").rstrip("/")

    from services.storage_service import get_object

    def make_payload(s, card):
        s_settings = (tpl_map.get(s.get("_fid"), {}) if ids_mode else tpl_settings) or {}
        s_tpl = s_settings.get("template", "green")
        return {
            "student_name": s.get("full_name", ""),
            "enrollment_no": s.get("student_id", ""),
            "reference_number": s.get("reference_number", ""),
            "nationality": s.get("nationality") or "يمني",
            "level": s.get("level") or 1,
            "section": s.get("section", ""),
            "department_name": s.get("_dept_name", "") if ids_mode else dept.get("name", ""),
            "faculty_name": s.get("_fac_name", "") if ids_mode else faculty_name,
            "academic_year": card["academic_year"],
            "template": s_tpl,
            "font": s_settings.get("font") or DEFAULT_CARD_FONT,
            "custom_bg_base64": s_settings.get("custom_bg_base64", ""),
            "custom_layout": s_settings.get("custom_layout") or {},
            "custom_orientation": s_settings.get("custom_orientation", "portrait"),
        }

    pngs = []
    for s in students:
        card = await _ensure_card(db, s)
        verify_url = f"{base}/verify-card?token={card['token']}" if base else card["token"]
        photo_bytes = None
        if s.get("photo_path"):
            try:
                photo_bytes, _ct = get_object(s["photo_path"])
            except Exception:
                photo_bytes = None
        payload = make_payload(s, card)
        pngs.append((_render_card_png(payload, photo_bytes, verify_url), _payload_portrait(payload)))

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.utils import ImageReader

    W, H = A4
    # اتجاه الإخراج: auto = حسب القالب، أو فرض عمودي/أفقي مع تدوير البطاقة عند الحاجة
    template_portrait = _payload_portrait({"template": tpl, "custom_orientation": (tpl_settings or {}).get("custom_orientation", "portrait")})
    if orientation == "auto":
        out_portrait = template_portrait
    else:
        out_portrait = orientation == "portrait"
    from PIL import Image as PILImage
    final_pngs = []
    for png, tp in pngs:
        if tp != out_portrait:
            im = PILImage.open(io.BytesIO(png)).rotate(90, expand=True)
            b = io.BytesIO()
            im.save(b, format="PNG")
            png = b.getvalue()
        final_pngs.append(png)
    pngs = final_pngs
    if out_portrait:
        cw, ch = st["card_h"] * mm, st["card_w"] * mm
    else:
        cw, ch = st["card_w"] * mm, st["card_h"] * mm
    positions = [(st["card1_x"] * mm, st["card1_y"] * mm), (st["card2_x"] * mm, st["card2_y"] * mm)]
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    for i in range(0, len(pngs), 2):
        for j, png in enumerate(pngs[i:i + 2]):
            x, y_top = positions[j]
            c.drawImage(ImageReader(io.BytesIO(png)), x, H - y_top - ch, cw, ch)
        c.showPage()
    c.save()

    if reprint:
        return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/pdf",
                                 headers=export_headers(export_filename("بطاقات الطلاب", f"إعادة تنزيل دفعة {reprint['batch_no']}", f"{len(pngs)} بطاقة", ext="pdf")))

    # 🖨️ تسجيل الدفعة ووسم الطلاب كمطبوعين (تنزيل PDF = طباعة)
    last = await db.card_print_batches.find_one({}, sort=[("batch_no", -1)])
    batch_no = int((last or {}).get("batch_no", 0)) + 1
    now = datetime.now(timezone.utc).isoformat()
    await db.card_print_batches.insert_one({
        "batch_no": batch_no, "created_at": now, "by_user_id": str(current_user.get("_id") or current_user.get("id", "")), "by_name": current_user.get("full_name", ""),
        "mode": "ids" if ids_mode else "department", "department_id": data.department_id if not ids_mode else None, "department_name": (dept or {}).get("name", "") if not ids_mode else "طلاب محددون",
        "faculty_name": faculty_name if not ids_mode else (students[0].get("_fac_name", "") if students else ""), "level": data.level, "section": data.section,
        "count": len(students), "pages": (len(students) + 1) // 2, "orientation": orientation, "settings": st,
        "student_ids": [str(s["_id"]) for s in students], "students": [{"id": str(s["_id"]), "full_name": s.get("full_name", ""), "student_id": s.get("student_id", ""), "level": s.get("level"), "section": s.get("section", "")} for s in students],
    })
    await db.students.update_many({"_id": {"$in": [s["_id"] for s in students]}}, {"$set": {"card_printed_at": now, "card_print_batch_no": batch_no}, "$inc": {"card_print_count": 1}})
    await log_activity(current_user, "cards_batch_print", "card_print_batches", str(batch_no), f"دفعة طباعة بطاقات #{batch_no}", {"count": len(students), "department": (dept or {}).get("name", "")})
    return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/pdf",
                             headers={**export_headers(export_filename("بطاقات الطلاب", f"دفعة {batch_no}", f"{len(pngs)} بطاقة", ext="pdf")), "X-Batch-No": str(batch_no), "X-Batch-Count": str(len(pngs))})


@router.get("/cards/batches")
async def list_print_batches(limit: int = 50, current_user: dict = Depends(get_current_user)):
    """📚 الدفعات السابقة (لإعادة التنزيل والمراجعة)"""
    db = get_db()
    items = []
    async for b in db.card_print_batches.find({}, {"students": 0, "settings": 0}).sort("batch_no", -1).limit(limit):
        b["id"] = str(b.pop("_id"))
        items.append(b)
    return {"items": items}


@router.get("/cards/batches/{batch_no}")
async def get_print_batch(batch_no: int, current_user: dict = Depends(get_current_user)):
    db = get_db()
    b = await db.card_print_batches.find_one({"batch_no": batch_no})
    if not b:
        raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
    b["id"] = str(b.pop("_id"))
    return b


@router.get("/cards/print-report")
async def cards_print_report(department_id: Optional[str] = None, faculty_id: Optional[str] = None, level: Optional[int] = None, current_user: dict = Depends(get_current_user)):
    """📥 تقرير Excel: حالة طباعة البطاقات لكل طالب (مطبوع / بانتظار الصورة / لم يُطبع)"""
    db = get_db()
    q: dict = {"is_alumni": {"$ne": True}}
    if department_id:
        q["department_id"] = department_id
    elif faculty_id:
        dept_ids = [str(d["_id"]) async for d in db.departments.find({"faculty_id": faculty_id}, {"_id": 1})]
        q["department_id"] = {"$in": dept_ids}
    else:
        raise HTTPException(status_code=400, detail="اختر الكلية أو القسم")
    if level:
        q["level"] = level
    depts = {str(d["_id"]): d.get("name", "") for d in await db.departments.find({}, {"name": 1}).to_list(2000)}
    rows = [s async for s in db.students.find(q, {"full_name": 1, "student_id": 1, "department_id": 1, "level": 1, "section": 1, "photo_path": 1, "card_printed_at": 1, "card_print_batch_no": 1, "card_print_count": 1}).sort([("department_id", 1), ("level", 1), ("full_name", 1)])]
    def status(s):
        if s.get("card_printed_at"):
            return "مطبوعة"
        return "بانتظار الصورة" if not s.get("photo_path") else "لم تُطبع بعد"
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "حالة الطباعة"; ws.sheet_view.rightToLeft = True
    n_p = sum(1 for s in rows if s.get("card_printed_at")); n_np = sum(1 for s in rows if not s.get("photo_path") and not s.get("card_printed_at"))
    heads = ["#", "رقم القيد", "الاسم", "القسم", "المستوى", "الشعبة", "الصورة", "حالة البطاقة", "تاريخ الطباعة", "رقم الدفعة", "عدد مرات الطباعة"]
    ws.append([f"تقرير طباعة البطاقات — {len(rows)} طالب: مطبوعة {n_p} · بانتظار الصورة {n_np} · لم تُطبع بعد {len(rows) - n_p - n_np}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads)); ws["A1"].font = Font(bold=True, size=12)
    ws.append(heads)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="00838F"); c.alignment = Alignment(horizontal="center")
    colors = {"مطبوعة": "E8F5E9", "بانتظار الصورة": "FFF8E1", "لم تُطبع بعد": "F1F5F9"}
    for i, s in enumerate(rows, 1):
        stt = status(s)
        ws.append([i, s.get("student_id", ""), s.get("full_name", ""), depts.get(str(s.get("department_id") or ""), ""), s.get("level"), s.get("section", ""), "معتمدة" if s.get("photo_path") else "لا توجد", stt, str(s.get("card_printed_at") or "")[:16].replace("T", " "), s.get("card_print_batch_no") or "", s.get("card_print_count") or 0])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=colors[stt])
    for i, w in enumerate([5, 14, 28, 24, 8, 8, 10, 16, 18, 10, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=export_headers(export_filename("تقرير طباعة البطاقات", f"{len(rows)} طالب", ext="xlsx")))


# ==================== 🔄 خلفية البطاقة (موحّدة على مستوى الجامعة) ====================
DEFAULT_BACK_LINES = [
    "هذه البطاقة ملك للجامعة وتُسلَّم عند الطلب أو عند انتهاء العلاقة بالجامعة.",
    "البطاقة شخصية ولا يجوز إعارتها أو استخدامها من قبل الغير.",
    "يجب إبرازها عند الدخول إلى الحرم الجامعي وفي الامتحانات والمكتبة.",
    "في حال فقدانها يُبلَّغ شؤون الطلاب فوراً، ويُصدَر بدل فاقد وفق اللوائح.",
    "البطاقة صالحة للعام الدراسي المدوَّن عليها فقط.",
    "من يعثر عليها يرجى تسليمها لأقرب مكتب في الجامعة أو الاتصال بالرقم أدناه.",
]
DEFAULT_BACK_SETTINGS = {"enabled": True, "title": "تعليمات استخدام البطاقة", "lines": DEFAULT_BACK_LINES, "footer_note": "", "show_contact": True,
                         "back1_x": 62.0, "back1_y": 40.0, "back2_x": 62.0, "back2_y": 180.0}


class BackSettingsIn(BaseModel):
    enabled: Optional[bool] = True
    title: Optional[str] = None
    lines: Optional[list] = None
    footer_note: Optional[str] = ""
    show_contact: Optional[bool] = True
    back1_x: Optional[float] = None
    back1_y: Optional[float] = None
    back2_x: Optional[float] = None
    back2_y: Optional[float] = None


async def _back_settings(db) -> dict:
    doc = await db.card_settings.find_one({"_id": "back_global"}) or {}
    return {**DEFAULT_BACK_SETTINGS, **{k: v for k, v in doc.items() if k != "_id"}}


@router.get("/cards/back-settings")
async def get_back_settings(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") in ("teacher", "student"):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    return {**await _back_settings(get_db()), "defaults": DEFAULT_BACK_SETTINGS, "can_edit": current_user.get("role") == "admin"}


@router.put("/cards/back-settings")
async def put_back_settings(data: BackSettingsIn, current_user: dict = Depends(get_current_user)):
    """تحرير التعليمات الموحدة لخلفية البطاقة (مدير النظام)"""
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="تحرير خلفية البطاقة متاح لمدير النظام فقط")
    db = get_db()
    lines = [str(x).strip() for x in (data.lines or []) if str(x).strip()][:10]
    if not lines:
        raise HTTPException(status_code=400, detail="أضف بنداً واحداً على الأقل")
    upd = {"enabled": bool(data.enabled), "title": (data.title or "").strip() or DEFAULT_BACK_SETTINGS["title"], "lines": lines, "footer_note": (data.footer_note or "").strip()[:300], "show_contact": bool(data.show_contact)}
    for k in ("back1_x", "back1_y", "back2_x", "back2_y"):
        v = getattr(data, k)
        if v is not None:
            upd[k] = float(v)
    await db.card_settings.update_one({"_id": "back_global"}, {"$set": upd}, upsert=True)
    await log_activity(current_user, "card_back_settings_update", "card_settings", "back_global", "تحديث خلفية البطاقة", {"lines": len(lines)})
    return {"message": "تم حفظ إعدادات خلفية البطاقة", **await _back_settings(db)}


def _render_card_back_png(back: dict, template: str, font_key: str, portrait: bool, uni: dict) -> bytes:
    """🔄 خلفية البطاقة: شريط بلون القالب + شعار + عنوان + تعليمات مرقّمة + بيانات التواصل — بنفس مقاس الوجه الأمامي"""
    from pathlib import Path
    from PIL import Image, ImageDraw, ImageFont, features
    HAS_RAQM = features.check("raqm")
    if not HAS_RAQM:
        import arabic_reshaper
        from bidi.algorithm import get_display

    def ar(t):
        t = str(t or "")
        return t if HAS_RAQM else get_display(arabic_reshaper.reshape(t))
    _dir = {"direction": "rtl"} if HAS_RAQM else {}
    S = CARD_SCALE
    theme = THEMES.get(template if template in THEMES else "green", THEMES["green"])
    if template == "custom":
        theme = THEMES["green"]
    font_file, font_axes = CARD_FONTS.get(font_key or DEFAULT_CARD_FONT, CARD_FONTS[DEFAULT_CARD_FONT])
    font_path = str(Path(__file__).parent.parent / "fonts" / font_file)
    cache: dict = {}

    def F(size):
        if size not in cache:
            f = ImageFont.truetype(font_path, int(size * S))
            if font_axes:
                try:
                    f.set_variation_by_axes(font_axes)
                except Exception:
                    pass
            cache[size] = f
        return cache[size]

    W, H = (640, 1010) if portrait else (1010, 640)
    img = Image.new("RGB", (W * S, H * S), theme["bg"])
    d = ImageDraw.Draw(img)
    band_h = 120 if portrait else 96
    d.rectangle([0, 0, W * S, band_h * S], fill=theme["band"])
    logo_path = Path(__file__).parent.parent / "assets" / "university_logo.jpeg"
    logo = Image.open(logo_path).convert("RGBA") if logo_path.exists() else None
    uni_name = (uni or {}).get("name") or "جامعة الأحقاف"
    if logo:
        ls = band_h - 28
        img.paste(logo.resize((ls * S, ls * S), Image.LANCZOS), (14 * S, 14 * S), logo.resize((ls * S, ls * S), Image.LANCZOS))
    d.text(((W - 24) * S, (band_h // 2) * S), ar(uni_name), font=F(34 if portrait else 32), fill=theme["band_text"], anchor="rm", **_dir)
    # علامة مائية
    if logo:
        wm_s = 300 if portrait else 260
        wm = logo.resize((wm_s * S, wm_s * S), Image.LANCZOS).convert("RGBA")
        a = wm.split()[3].point(lambda v: int(v * 0.07))
        wm.putalpha(a)
        img.paste(wm, ((W // 2 - wm_s // 2) * S, (H // 2 - wm_s // 2) * S), wm)
    y = band_h + (34 if portrait else 22)
    d.text(((W // 2) * S, y * S), ar(back.get("title") or DEFAULT_BACK_SETTINGS["title"]), font=F(30 if portrait else 28), fill=theme["accent"], anchor="mm", **_dir)
    y += 26
    d.line([(60 * S, y * S), ((W - 60) * S, y * S)], fill=theme["accent"], width=2 * S)
    y += 22
    # التعليمات مع التفاف النص + ملاءمة تلقائية لحجم الخط حسب المساحة المتاحة
    max_w = (W - 90)
    right = W - 44
    text_col = theme["text"]
    foot_h = 118 if portrait else 96
    lines = back.get("lines") or DEFAULT_BACK_LINES
    avail = H - foot_h - 10 - y

    def wrap(txt, f):
        words, out, cur = txt.split(), [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if d.textlength(ar(t), font=f) / S <= max_w:
                cur = t
            else:
                out.append(cur); cur = w
        if cur:
            out.append(cur)
        return out
    body_size = 18
    for size in range(28 if portrait else 24, 17, -1):
        f = F(size)
        total = sum(len(wrap(ln, f)) * (size + 11) + 8 for ln in lines)
        if total <= avail:
            body_size = size
            break
    f_body = F(body_size)
    lh = body_size + 11
    for i, ln in enumerate(lines, 1):
        for j, part in enumerate(wrap(ln, f_body)):
            if y + lh > H - foot_h - 10:
                break
            prefix = f"{i}. " if j == 0 else "    "
            d.text((right * S, y * S), ar(prefix + part), font=f_body, fill=text_col, anchor="ra", **_dir)
            y += lh
        y += 8
    # التذييل: بيانات التواصل
    d.rectangle([0, (H - foot_h) * S, W * S, H * S], fill=theme["strip"])
    fy = H - foot_h + 16
    f_small = F(17 if portrait else 16)
    if back.get("show_contact", True):
        contact = [str(x) for x in [(uni or {}).get("address")] if x] + [f"\u202a{x}\u202c" for x in [(uni or {}).get("phone"), (uni or {}).get("website")] if x]
        # بيانات التواصل في سطر أفقي واحد مفصولة بنقطة (وتلتفّ لسطر ثانٍ إن طالت)
        line, rows_ = "", []
        for c in contact[:3]:
            t = f"{line}  ·  {c}" if line else c
            if d.textlength(ar(t), font=f_small) / S <= W - 60 or not line:
                line = t
            else:
                rows_.append(line); line = c
        if line:
            rows_.append(line)
        for r_ in rows_[:2]:
            d.text(((W // 2) * S, fy * S), ar(r_), font=f_small, fill=theme["strip_text"], anchor="ma", **_dir)
            fy += 26
    if back.get("footer_note"):
        d.text(((W // 2) * S, fy * S), ar(back["footer_note"]), font=f_small, fill=theme["strip_text"], anchor="ma", **_dir)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@router.get("/cards/back-preview/{faculty_id}")
async def back_preview(faculty_id: str, current_user: dict = Depends(get_current_user)):
    """معاينة PNG لخلفية البطاقة بقالب الكلية المختار"""
    if not _can_manage(current_user, faculty_id):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    db = get_db()
    tpl_settings = await db.card_settings.find_one({"_id": f"faculty_{faculty_id}"}) or {}
    back = await _back_settings(db)
    uni = await db.university.find_one({}) or {}
    portrait = _payload_portrait({"template": tpl_settings.get("template", "green"), "custom_orientation": tpl_settings.get("custom_orientation", "portrait")})
    png = _render_card_back_png(back, tpl_settings.get("template", "green"), tpl_settings.get("font") or DEFAULT_CARD_FONT, portrait, uni)
    return Response(content=png, media_type="image/png", headers={"Cache-Control": "no-store"})


class BackBatchRequest(BaseModel):
    batch_no: Optional[int] = None
    department_id: Optional[str] = None
    count: Optional[int] = None
    orientation: Optional[str] = "auto"
    settings: Optional[dict] = None  # card_w/card_h + back1_x/back1_y/back2_x/back2_y


@router.post("/cards/batch-back-pdf")
async def batch_back_pdf(data: BackBatchRequest, current_user: dict = Depends(get_current_user)):
    """🔄 PDF خلفيات الدفعة: بطاقتان في كل ورقة A4 بمواضع مستقلة (تُحفظ) — نفس عدد وترتيب الوجه الأمامي للدفعة"""
    db = get_db()
    if current_user.get("role") in ("teacher", "student"):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    back = await _back_settings(db)
    if not back.get("enabled"):
        raise HTTPException(status_code=400, detail="خلفية البطاقة معطّلة من الإعدادات")
    if data.batch_no:
        b = await db.card_print_batches.find_one({"batch_no": int(data.batch_no)})
        if not b:
            raise HTTPException(status_code=404, detail="الدفعة غير موجودة")
        n = int(b.get("count") or 0)
        dept_id = b.get("department_id")
        if not dept_id and b.get("student_ids"):
            first = await db.students.find_one({"_id": ObjectId(b["student_ids"][0])})
            dept_id = (first or {}).get("department_id")
        label = f"خلفيات دفعة {b['batch_no']}"
    else:
        n = int(data.count or 0)
        dept_id = data.department_id
        label = "خلفيات البطاقات"
    if n <= 0:
        raise HTTPException(status_code=400, detail="عدد البطاقات غير محدد")
    fid = ""
    if dept_id:
        dept = await db.departments.find_one({"_id": ObjectId(dept_id)})
        fid = (dept or {}).get("faculty_id", "")
    tpl_settings = (await db.card_settings.find_one({"_id": f"faculty_{fid}"}) if fid else None) or {}
    tpl = tpl_settings.get("template", "green")
    template_portrait = _payload_portrait({"template": tpl, "custom_orientation": tpl_settings.get("custom_orientation", "portrait")})
    orientation = data.orientation if data.orientation in ORIENTATIONS else "auto"
    out_portrait = template_portrait if orientation == "auto" else orientation == "portrait"
    uni = await db.university.find_one({}) or {}
    png = _render_card_back_png(back, tpl, tpl_settings.get("font") or DEFAULT_CARD_FONT, template_portrait, uni)
    if template_portrait != out_portrait:
        from PIL import Image as PILImage
        im = PILImage.open(io.BytesIO(png)).rotate(90, expand=True)
        bb = io.BytesIO(); im.save(bb, format="PNG"); png = bb.getvalue()
    pst = {**DEFAULT_PRINT_SETTINGS, **{k: v for k, v in ((await db.card_print_settings.find_one({"_id": "global"})) or {}).items() if k in DEFAULT_PRINT_SETTINGS}}
    bst = {k: float(back.get(k, DEFAULT_BACK_SETTINGS[k])) for k in ("back1_x", "back1_y", "back2_x", "back2_y")}
    if data.settings:
        for k in ("card_w", "card_h"):
            if k in data.settings:
                try: pst[k] = float(data.settings[k])
                except (TypeError, ValueError): pass
        changed = {}
        for k in bst:
            if k in data.settings:
                try: bst[k] = float(data.settings[k]); changed[k] = bst[k]
                except (TypeError, ValueError): pass
        if changed:
            await db.card_settings.update_one({"_id": "back_global"}, {"$set": changed}, upsert=True)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.utils import ImageReader
    W, H = A4
    cw, ch = (pst["card_h"] * mm, pst["card_w"] * mm) if out_portrait else (pst["card_w"] * mm, pst["card_h"] * mm)
    positions = [(bst["back1_x"] * mm, bst["back1_y"] * mm), (bst["back2_x"] * mm, bst["back2_y"] * mm)]
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    reader = ImageReader(io.BytesIO(png))
    for i in range(0, n, 2):
        for j in range(min(2, n - i)):
            x, y_top = positions[j]
            c.drawImage(reader, x, H - y_top - ch, cw, ch)
        c.showPage()
    c.save()
    return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/pdf",
                             headers=export_headers(export_filename("بطاقات الطلاب", label, f"{n} خلفية", ext="pdf")))


@router.get("/cards/batch-count")
async def batch_count(department_id: str, level: Optional[int] = None, section: Optional[str] = None,
                      current_user: dict = Depends(get_current_user)):
    db = get_db()
    q = {"department_id": department_id, "is_alumni": {"$ne": True}}
    if level:
        q["level"] = level
    if section:
        q["section"] = section
    n = await db.students.count_documents(q)
    return {"count": n, "pages": (n + 1) // 2}


# ==================== توليد PNG / PDF ====================
THEMES = {
    "green": {"band": (27, 94, 32), "bg": (255, 255, 255), "text": (26, 37, 64), "band_text": (255, 255, 255), "accent": (27, 94, 32), "muted": (91, 102, 120), "strip": (232, 245, 233), "strip_text": (27, 94, 32)},
    "dark": {"band": (7, 20, 23), "bg": (15, 32, 39), "text": (255, 255, 255), "band_text": (255, 255, 255), "accent": (77, 182, 172), "muted": (176, 190, 197), "strip": (7, 20, 23), "strip_text": (77, 182, 172)},
    "official": {"band": (27, 94, 32), "bg": (255, 255, 255), "text": (26, 37, 64), "band_text": (255, 255, 255), "accent": (27, 94, 32), "muted": (91, 102, 120), "strip": (232, 245, 233), "strip_text": (27, 94, 32)},
}


def _render_card_png(p: dict, photo_bytes: Optional[bytes], verify_url: str) -> bytes:
    import qrcode
    from pathlib import Path
    from PIL import Image, ImageDraw, ImageFont, features

    HAS_RAQM = features.check("raqm")
    if not HAS_RAQM:
        import arabic_reshaper
        from bidi.algorithm import get_display

    def ar(t):
        t = str(t or "")
        if HAS_RAQM:
            return t
        return get_display(arabic_reshaper.reshape(t))

    _dir = {"direction": "rtl"} if HAS_RAQM else {}

    _fonts_dir = Path(__file__).parent.parent / "fonts"
    font_key = p.get("font") or DEFAULT_CARD_FONT
    font_file, font_axes = CARD_FONTS.get(font_key, CARD_FONTS[DEFAULT_CARD_FONT])
    font_path = str(_fonts_dir / font_file)
    _font_cache: dict = {}

    def F(size):
        size = int(size)
        if size not in _font_cache:
            f = ImageFont.truetype(font_path, size * CARD_SCALE)
            if font_axes:
                try:
                    f.set_variation_by_axes(font_axes)
                except Exception:
                    pass
            _font_cache[size] = f
        return _font_cache[size]

    S = CARD_SCALE

    class _SDraw:
        """رسم بمقياس مضاعف: الإحداثيات تُضرب في S، والخطوط مُهيّأة مسبقاً بالمقياس"""
        def __init__(self, d): self.d = d
        def _sc(self, xy): return [tuple(v * S for v in pt) if isinstance(pt, (tuple, list)) else pt * S for pt in xy]
        def rectangle(self, xy, **kw):
            if kw.get("width"): kw["width"] = int(kw["width"] * S)
            self.d.rectangle(self._sc(xy), **kw)
        def rounded_rectangle(self, xy, radius=0, **kw):
            if kw.get("width"): kw["width"] = int(kw["width"] * S)
            self.d.rounded_rectangle(self._sc(xy), radius=radius * S, **kw)
        def ellipse(self, xy, **kw):
            if kw.get("width"): kw["width"] = int(kw["width"] * S)
            self.d.ellipse(self._sc(xy), **kw)
        def line(self, xy, **kw):
            if kw.get("width"): kw["width"] = int(kw["width"] * S)
            self.d.line(self._sc(xy), **kw)
        def text(self, xy, t, **kw): self.d.text((xy[0] * S, xy[1] * S), t, **kw)
        def textlength(self, t, **kw): return self.d.textlength(t, **kw) / S

    class _SImage:
        def __init__(self, im): self.im = im
        def paste(self, src, box, mask=None):
            same = mask is src
            if not getattr(src, "_scaled", False):
                src = src.resize((src.width * S, src.height * S), Image.LANCZOS)
            if same:
                mask = src
            elif mask is not None and mask.size != src.size:
                mask = mask.resize(src.size, Image.LANCZOS)
            self.im.paste(src, (int(box[0] * S), int(box[1] * S)), mask)
        def save(self, *a, **kw): return self.im.save(*a, **kw)

    def _canvas(w, h, color):
        return _SImage(Image.new("RGB", (w * S, h * S), color))

    def _draw(img):
        return _SDraw(ImageDraw.Draw(img.im if isinstance(img, _SImage) else img))

    def _mark(im):
        im._scaled = True
        return im

    def _logo(size):
        return _mark(logo.resize((size * S, size * S), Image.LANCZOS)) if logo else None

    def _qr(size):
        return _mark(qr_img.resize((size * S, size * S), Image.NEAREST))

    def rtl(draw, right_x, y, text, font, fill):
        t = ar(text)
        w = draw.textlength(t, font=font, **_dir)
        draw.text((right_x - w, y), t, font=font, fill=fill, **_dir)

    def center(draw, cx, y, text, font, fill):
        t = ar(text)
        w = draw.textlength(t, font=font, **_dir)
        draw.text((cx - w / 2, y), t, font=font, fill=fill, **_dir)

    template = p.get("template", "green")
    horizontal = template == "horizontal"
    theme = THEMES.get(template if template in THEMES else "green", THEMES["green"])

    logo = None
    logo_path = Path(__file__).parent.parent / "assets" / "university_logo.jpeg"
    if logo_path.exists():
        logo = Image.open(logo_path).convert("RGBA")

    def paste_watermark(im, cx, cy, size, alpha=26):
        if not logo:
            return
        wm = logo.resize((size * S, size * S), Image.LANCZOS).convert("RGBA")
        mask = wm.convert("L").point(lambda v: 0 if v > 235 else alpha)
        wm.putalpha(mask)
        im.im.paste(wm, ((cx - size // 2) * S, (cy - size // 2) * S), wm)

    # صورة الطالب
    photo = None
    if photo_bytes:
        try:
            photo = Image.open(io.BytesIO(photo_bytes)).convert("RGB")
        except Exception:
            photo = None

    def fit_photo(img, w, h):
        w, h = w * S, h * S
        ratio = max(w / img.width, h / img.height)
        img = img.resize((int(img.width * ratio) + 1, int(img.height * ratio) + 1), Image.LANCZOS)
        left = (img.width - w) // 2
        top = (img.height - h) // 2
        return _mark(img.crop((left, top, left + w, top + h)))

    qr_img = qrcode.make(verify_url, box_size=6, border=1).convert("RGB")

    level_ar = LEVEL_AR.get(p.get("level") or 1, str(p.get("level")))
    is_emp = p.get("kind") in ("employee", "academic")
    title_ar = p.get("title_ar") or ("بطاقة موظف" if p.get("kind") == "employee" else "بطاقة أكاديمية" if p.get("kind") == "academic" else "بطاقة طالب")
    title_en = "STAFF ID CARD" if is_emp else "STUDENT ID CARD"
    if is_emp:
        rows = [(("رقم أكاديمي" if p.get("kind") == "academic" and p.get("academic_no") else "رقم وظيفي"), p.get("academic_no") or p.get("enrollment_no", ""))]
        if (p.get("job_title") or "").strip():
            rows.append(("المسمى", p["job_title"]))
        if (p.get("org_unit_name") or p.get("department_name") or "").strip():
            rows.append(("الوحدة" if p.get("kind") == "employee" else "القسم", p.get("org_unit_name") or p.get("department_name")))
        if (p.get("category_label") or "").strip():
            rows.append(("الفئة", p["category_label"]))
        rows.append(("الجنسية", p.get("nationality", "")))
    else:
        rows = [
            ("رقم القيد", p.get("enrollment_no", "")),
            ("التخصص", p.get("department_name", "")),
            ("المستوى", f"المستوى {level_ar}"),
        ]
        if (p.get("section") or "").strip():
            rows.append(("الشعبة", str(p["section"]).strip()))
        rows.append(("الجنسية", p.get("nationality", "")))
    rows = [(k, v) for k, v in rows if str(v or "").strip()]
    valid_text = p.get("validity_text") or f"صالحة للعام الجامعي {p.get('academic_year', '')}"

    if template == "custom" and p.get("custom_bg_base64"):
        import base64 as _b64
        try:
            bg = Image.open(io.BytesIO(_b64.b64decode(str(p["custom_bg_base64"]).split(",")[-1]))).convert("RGB")
        except Exception:
            bg = None
        if bg is not None:
            W, H = (1010, 640) if bg.width > bg.height else (640, 1010)
            img = _SImage(fit_photo(bg, W, H))
            d = _draw(img)
            L = {**DEFAULT_CUSTOM_LAYOUT, **{k: v for k, v in (p.get("custom_layout") or {}).items() if isinstance(v, dict)}}
            if photo:
                ph = L["photo"]
                pw, phh = int(ph.get("w", 26) / 100 * W), int(ph.get("h", 24) / 100 * H)
                pxp, pyp = int(ph.get("x", 66) / 100 * W), int(ph.get("y", 16) / 100 * H)
                img.paste(fit_photo(photo, pw, phh), (pxp, pyp))
                d.rectangle([pxp, pyp, pxp + pw, pyp + phh], outline=(255, 255, 255), width=3)
            q = L["qr"]
            qs = max(40, int(q.get("w", 18) / 100 * W))
            img.paste(_qr(qs), (int(q.get("x", 8) / 100 * W), int(q.get("y", 66) / 100 * H)))

            def _txt(key, value):
                el = L.get(key) or DEFAULT_CUSTOM_LAYOUT[key]
                center(d, int(el.get("x", 50) / 100 * W), int(el.get("y", 50) / 100 * H), value, F(int(el.get("size", 24))), el.get("color", "#1a2540"))

            _txt("name", p.get("student_name", ""))
            _txt("enrollment", f"{rows[0][0]}: {rows[0][1]}" if rows else "")
            _txt("dept", f"{p.get('department_name', '')} — المستوى {level_ar}" if not is_emp else " — ".join(v for k, v in rows[1:3]))
            _txt("year", f"العام الجامعي: {p.get('academic_year', '')}")
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            return buf.getvalue()

    if template == "official":
        DG = (27, 94, 32)
        MG = (46, 125, 50)
        LG = (232, 245, 233)
        MUT = (96, 125, 102)
        NAVY = (26, 37, 64)
        W, H = 640, 1010
        img = _canvas(W, H, (255, 255, 255))
        d = _draw(img)
        # زخارف دوائر خفيفة
        d.ellipse([-150, -150, 170, 170], outline=LG, width=3)
        d.ellipse([-95, -95, 115, 115], outline=LG, width=3)
        d.ellipse([W - 170, H - 240, W + 150, H + 80], outline=LG, width=3)
        # الأشرطة الجانبية
        d.rectangle([W - 18, 0, W, H], fill=DG)
        d.rectangle([W - 26, 0, W - 22, H], fill=MG)
        d.rectangle([0, 0, 6, H], fill=LG)
        paste_watermark(img, W // 2, 660, 480)
        # الشعار داخل حلقة خضراء
        cy = 96
        d.ellipse([W // 2 - 76, cy - 76, W // 2 + 76, cy + 76], outline=LG, width=10)
        d.ellipse([W // 2 - 66, cy - 66, W // 2 + 66, cy + 66], outline=DG, width=3)
        if logo:
            lg = _logo(108)
            img.paste(lg, (W // 2 - 54, cy - 54), lg)
        center(d, W // 2, 176, "جامعة الأحقاف", F(40), DG)
        center(d, W // 2, 234, "AL-AHGAFF UNIVERSITY", F(18), MUT)
        # اسم الكلية في كبسولة فاتحة
        fac_w = d.textlength(ar(p.get("faculty_name", "")), font=F(22), **_dir)
        d.rounded_rectangle([W // 2 - fac_w / 2 - 26, 274, W // 2 + fac_w / 2 + 26, 318], radius=22, fill=LG)
        center(d, W // 2, 281, p.get("faculty_name", ""), F(22), DG)
        # شريط بطاقة طالب
        d.rounded_rectangle([W // 2 - 108, 334, W // 2 + 108, 378], radius=10, fill=DG)
        center(d, W // 2, 340, title_ar, F(24), (255, 255, 255))
        # الصورة بإطار أخضر مزدوج
        py = 404
        d.rounded_rectangle([W // 2 - 114, py - 14, W // 2 + 114, py + 264], radius=14, fill=LG)
        d.rectangle([W // 2 - 104, py - 4, W // 2 + 104, py + 254], fill=DG)
        if photo:
            ph = fit_photo(photo, 200, 250)
            img.paste(ph, (W // 2 - 100, py))
        else:
            d.rectangle([W // 2 - 100, py, W // 2 + 100, py + 250], fill=(244, 248, 245))
            center(d, W // 2, py + 110, "لا توجد صورة", F(20), MUT)
        # الاسم
        center(d, W // 2, py + 272, p.get("student_name", ""), F(32), DG)
        d.rectangle([130, py + 326, W - 130, py + 328], fill=LG)
        # البيانات يميناً + QR يساراً
        y = py + 344
        spacing = min(34, 190 // max(len(rows), 1))
        for label, value in rows:
            rtl(d, W - 62, y, f"{label}:", F(19), MG)
            rtl(d, W - 216, y, value, F(21), NAVY)
            y += spacing
        d.rounded_rectangle([36, 748, 198, 948], radius=10, outline=LG, width=4)
        q = _qr(140)
        img.paste(q, (47, 760))
        center(d, 117, 908, "امسح للتحقق", F(16), MUT)
        # الشريط السفلي
        d.rectangle([0, H - 64, W, H - 60], fill=DG)
        d.rectangle([0, H - 60, W, H], fill=LG)
        center(d, W // 2, H - 50, valid_text, F(23), DG)
    elif not horizontal:
        W, H = 640, 1010
        img = _canvas(W, H, theme["bg"])
        d = _draw(img)
        # الشريط العلوي
        d.rectangle([0, 0, W, 180], fill=theme["band"])
        paste_watermark(img, W // 2, 620, 460)
        if logo:
            lg = _logo(104)
            white = Image.new("RGB", (116, 116), (255, 255, 255))
            img.paste(white, (W // 2 - 58, 14))
            img.paste(lg, (W // 2 - 52, 20), lg)
        center(d, W // 2 - 170, 40, "AL-AHGAFF", F(26), theme["band_text"])
        center(d, W // 2 - 170, 76, "UNIVERSITY", F(26), theme["band_text"])
        center(d, W // 2 + 170, 44, "جامعة الأحقاف", F(34), theme["band_text"])
        center(d, W // 2 + 170, 96, p.get("faculty_name", ""), F(22), theme["band_text"])
        center(d, W // 2, 138, title_ar, F(26), theme["band_text"])
        # صورة الطالب
        py = 210
        if photo:
            ph = fit_photo(photo, 240, 300)
            d.rectangle([W // 2 - 124, py - 4, W // 2 + 124, py + 304], fill=theme["accent"])
            img.paste(ph, (W // 2 - 120, py))
        else:
            d.rectangle([W // 2 - 120, py, W // 2 + 120, py + 300], fill=(230, 234, 242))
            center(d, W // 2, py + 135, "لا توجد صورة", F(22), (120, 130, 145))
        # الاسم
        center(d, W // 2, py + 320, p.get("student_name", ""), F(32), theme["text"])
        # البيانات
        y = py + 380
        for label, value in rows:
            rtl(d, W - 50, y, f"{label}:", F(22), theme["muted"])
            rtl(d, W - 210, y, value, F(24), theme["text"])
            y += 44
        # QR
        q = _qr(150)
        img.paste(q, (40, H - 226))
        rtl(d, W - 50, H - 190, "امسح الرمز للتحقق", F(20), theme["muted"])
        rtl(d, W - 50, H - 158, "من صحة البطاقة", F(20), theme["muted"])
        # شريط الصلاحية
        d.rectangle([0, H - 56, W, H], fill=theme["strip"])
        center(d, W // 2, H - 46, valid_text, F(24), theme["strip_text"])
    else:
        W, H = 1010, 640
        theme = THEMES["green"]
        img = _canvas(W, H, theme["bg"])
        d = _draw(img)
        d.rectangle([0, 0, W, 120], fill=theme["band"])
        paste_watermark(img, 420, 380, 380)
        if logo:
            lg = _logo(92)
            white = Image.new("RGB", (100, 100), (255, 255, 255))
            img.paste(white, (W // 2 - 50, 14))
            img.paste(lg, (W // 2 - 46, 18), lg)
        rtl(d, W - 30, 22, "جامعة الأحقاف", F(32), theme["band_text"])
        rtl(d, W - 30, 70, p.get("faculty_name", ""), F(22), theme["band_text"])
        d.text((30, 28), "AL-AHGAFF UNIVERSITY", font=F(24), fill=theme["band_text"])
        d.text((30, 66), title_en, font=F(20), fill=theme["band_text"])
        # الصورة يميناً
        px, py = W - 290, 160
        if photo:
            ph = fit_photo(photo, 230, 290)
            d.rectangle([px - 4, py - 4, px + 234, py + 294], fill=theme["accent"])
            img.paste(ph, (px, py))
        else:
            d.rectangle([px, py, px + 230, py + 290], fill=(230, 234, 242))
            center(d, px + 115, py + 130, "لا توجد صورة", F(20), (120, 130, 145))
        # البيانات يساراً (من اليمين للصورة)
        rtl(d, px - 40, 160, p.get("student_name", ""), F(34), theme["text"])
        y = 230
        for label, value in rows:
            rtl(d, px - 40, y, f"{label}:", F(22), theme["muted"])
            rtl(d, px - 200, y, value, F(24), theme["text"])
            y += 48
        # QR أسفل يسار مع مساحة فاصلة
        q = _qr(130)
        img.paste(q, (36, H - 196))
        center(d, 101, H - 226, "امسح للتحقق", F(18), theme["muted"])
        d.rectangle([0, H - 52, W, H], fill=theme["strip"])
        center(d, W // 2, H - 44, valid_text, F(22), theme["strip_text"])

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@router.get("/students/{student_id}/card/download")
async def download_student_card(
    student_id: str,
    fmt: str = Query("png", pattern="^(png|pdf)$"),
    base_url: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user),
):
    db = get_db()
    student = await _student_or_404(db, student_id)
    fid, _, _ = await _resolve_faculty(db, student)
    if not _can_manage(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    p = await _card_payload(db, student, base_url or "")
    photo_bytes = None
    if p.get("photo_path"):
        from services.storage_service import get_object
        try:
            photo_bytes, _ct = get_object(p["photo_path"])
        except Exception:
            photo_bytes = None
    png = _render_card_png(p, photo_bytes, p["verify_url"])
    _card_label = ("بطاقة الطالب", p.get("full_name") or student.get("full_name") or "", p.get("enrollment_no", ""))
    if fmt == "png":
        return StreamingResponse(io.BytesIO(png), media_type="image/png",
                                 headers=export_headers(export_filename(*_card_label, ext="png")))
    # PDF بمقاس البطاقة القياسي CR80
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.utils import ImageReader
    horizontal = p.get("template") == "horizontal"
    page = (85.6 * mm, 54 * mm) if horizontal else (54 * mm, 85.6 * mm)
    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=page)
    c.drawImage(ImageReader(io.BytesIO(png)), 0, 0, page[0], page[1])
    c.showPage()
    c.save()
    return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/pdf",
                             headers=export_headers(export_filename(*_card_label, ext="pdf")))


@router.get("/cards/preview/{faculty_id}")
async def card_preview(faculty_id: str, template: str = "green", font: str = DEFAULT_CARD_FONT, current_user: dict = Depends(get_current_user)):
    """🖼️ معاينة فورية للقالب والخط ببيانات تجريبية (بلا حفظ)"""
    if not _can_manage(current_user, faculty_id):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    if template not in TEMPLATES or font not in CARD_FONTS:
        raise HTTPException(status_code=400, detail="قالب أو خط غير معروف")
    db = get_db()
    fac = None
    try:
        fac = await db.faculties.find_one({"_id": ObjectId(faculty_id)}, {"name": 1})
    except Exception:
        pass
    saved = await db.card_settings.find_one({"_id": f"faculty_{faculty_id}"}) or {}
    p = {"template": template, "font": font, "student_name": "محمد عبدالله سالم باعباد", "enrollment_no": "20231045", "department_name": "التخصص التجريبي",
         "faculty_name": (fac or {}).get("name") or "الكلية", "level": 2, "section": "أ", "nationality": "يمني", "academic_year": "2025-2026",
         "custom_bg_base64": saved.get("custom_bg_base64", ""), "custom_layout": saved.get("custom_layout") or {}}
    png = _render_card_png(p, None, "https://ahgaff.net/verify-card?token=preview")
    return StreamingResponse(io.BytesIO(png), media_type="image/png", headers={"Cache-Control": "no-store"})
