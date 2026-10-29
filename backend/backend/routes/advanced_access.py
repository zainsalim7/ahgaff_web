"""
مركز الصلاحيات المتقدم — Advanced Access Control (SHADOW MODE)
مسار مستقل تمامًا: /api/advanced-access/* — لا يُستخدم في أي قرار وصول حقيقي بالنظام القديم.
الوصول محصور بالأدمن أو مستخدمين مصرَّح لهم صراحة عبر advanced_access_testers (انظر has_sub_permission).
"""
from datetime import datetime, timezone
import re
from typing import Optional, List

from fastapi import APIRouter, Depends, HTTPException
from bson import ObjectId

from .deps import get_db, get_current_user
from models.advanced_access import (
    AdvancedRoleCreate, AdvancedRoleUpdate, AssignmentCreate, TesterGrant, CheckPermissionRequest, DirectPermissionGrant,
)
from services.advanced_access_service import (
    seed_permission_registry, create_advanced_access_indexes, import_legacy_roles_reference,
    has_sub_permission, has_advanced_access, compute_effective_permissions, check_permission,
    list_user_work_contexts, write_audit_log, SUB_PERMISSIONS, get_tester_record,
    compare_decision, compare_summary_for_assignment,
)
from services.pilot_service import run_pilot, reset_pilot

router = APIRouter(prefix="/advanced-access", tags=["مركز الصلاحيات المتقدم (تجريبي)"])


@router.get("/me/access")
async def my_access(current_user: dict = Depends(get_current_user)):
    """بدون بوابة — يخبر الواجهة فقط إن كان المستخدم الحالي يملك حق الدخول لعرض الرابط في القائمة"""
    db = get_db()
    is_admin = current_user.get("role") == "admin"
    tester = None if is_admin else await get_tester_record(db, current_user["id"])
    return {
        "is_admin": is_admin,
        "has_access": is_admin or bool(tester),
        "sub_permissions": SUB_PERMISSIONS if is_admin else ((tester or {}).get("sub_permissions") or []),
    }


@router.get("/me/contexts")
async def my_contexts(current_user: dict = Depends(get_current_user)):
    db = get_db()
    if not await has_advanced_access(db, current_user):
        raise HTTPException(status_code=403, detail="لا تملك صلاحية الوصول إلى مركز الصلاحيات المتقدم (تجريبي)")
    return await list_user_work_contexts(db, current_user["id"], active_only=True)


def require_adv_access(sub_permission: Optional[str] = None):
    async def _dep(current_user: dict = Depends(get_current_user)):
        db = get_db()
        if not await has_sub_permission(db, current_user, sub_permission):
            raise HTTPException(status_code=403, detail="لا تملك صلاحية الوصول إلى مركز الصلاحيات المتقدم (تجريبي)")
        return current_user
    return _dep


def _oid(v: str, what: str = "المعرّف"):
    if not v or not ObjectId.is_valid(v):
        raise HTTPException(status_code=400, detail=f"{what} غير صحيح")
    return ObjectId(v)


def _ser(d: dict) -> dict:
    d = dict(d)
    d["id"] = str(d.pop("_id"))
    return d


# ══════════════ سجل الصلاحيات والمجموعات ══════════════

@router.get("/permissions")
async def list_permissions(group_key: Optional[str] = None, current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    q = {"is_active": True}
    if group_key:
        q["group_key"] = group_key
    return await db.advanced_permissions.find(q, {"_id": 0}).to_list(2000)


@router.get("/permission-groups")
async def list_permission_groups(current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    return await db.advanced_permission_groups.find({}, {"_id": 0}).to_list(200)


@router.post("/seed-registry")
async def seed_registry(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="للأدمن فقط")
    db = get_db()
    await create_advanced_access_indexes(db)
    result = await seed_permission_registry(db)
    await write_audit_log(db, current_user, "permission_registry_seeded", after=result)
    return result


@router.post("/import-legacy-roles")
async def import_legacy_roles(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="للأدمن فقط")
    db = get_db()
    result = await import_legacy_roles_reference(db)
    await write_audit_log(db, current_user, "legacy_roles_imported", after=result)
    return result


# ══════════════ الأدوار المتقدمة (تجريبية) ══════════════

@router.get("/roles")
async def list_advanced_roles(current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    return [_ser(r) async for r in db.advanced_roles.find({})]


@router.get("/roles/{role_id}")
async def get_advanced_role(role_id: str, current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    role = await db.advanced_roles.find_one({"_id": _oid(role_id, "الدور")})
    if not role:
        raise HTTPException(status_code=404, detail="الدور غير موجود")
    return _ser(role)


@router.post("/roles")
async def create_advanced_role(data: AdvancedRoleCreate, current_user: dict = Depends(require_adv_access("manage_roles"))):
    db = get_db()
    valid_keys = {p["key"] for p in await db.advanced_permissions.find({"is_active": True}, {"key": 1}).to_list(2000)}
    invalid = [k for k in data.permission_keys if k not in valid_keys]
    if invalid:
        raise HTTPException(status_code=400, detail=f"صلاحيات غير معروفة: {invalid[:5]}")
    doc = {
        "key": None, "name_ar": data.name_ar, "description": data.description or "",
        "permission_keys": data.permission_keys, "status": "draft", "is_system": False,
        "imported_from_legacy_role_id": None,
        "created_at": datetime.now(timezone.utc), "created_by": current_user["id"],
        "updated_at": datetime.now(timezone.utc),
    }
    res = await db.advanced_roles.insert_one(doc)
    await db.advanced_roles.update_one({"_id": res.inserted_id}, {"$set": {"key": f"custom_{res.inserted_id}"}})
    await write_audit_log(db, current_user, "advanced_role_created", after={"role_id": str(res.inserted_id), **data.dict()})
    return {"id": str(res.inserted_id)}


@router.put("/roles/{role_id}")
async def update_advanced_role(role_id: str, data: AdvancedRoleUpdate, current_user: dict = Depends(require_adv_access("manage_roles"))):
    db = get_db()
    role = await db.advanced_roles.find_one({"_id": _oid(role_id, "الدور")})
    if not role:
        raise HTTPException(status_code=404, detail="الدور غير موجود")
    if current_user.get("role") != "admin" and role.get("imported_from_legacy_role_id"):
        raise HTTPException(status_code=403, detail="تعديل الأدوار المستوردة من النظام القديم للأدمن فقط")
    update = {k: v for k, v in data.dict(exclude_unset=True).items()}
    if "status" in update and update["status"] not in ("draft", "shadow", "ready_for_review"):
        raise HTTPException(status_code=400, detail="حالة غير صالحة (draft/shadow/ready_for_review فقط — لا حالة إنتاج ضمن هذا المشروع)")
    update["updated_at"] = datetime.now(timezone.utc)
    await db.advanced_roles.update_one({"_id": role["_id"]}, {"$set": update})
    await write_audit_log(db, current_user, "advanced_role_updated", before=role.get("permission_keys"), after=update)
    return {"message": "تم التحديث"}


# ══════════════ نطاقات متاحة (تُقرأ مباشرة — تدعم وحدات مستقبلية تلقائيًا) ══════════════

@router.get("/scope-options/{scope_type}")
async def scope_options(scope_type: str, current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    if scope_type in ("faculty", "multi_faculty"):
        items = await db.faculties.find({}, {"name": 1}).to_list(500)
    elif scope_type in ("department", "multi_department"):
        items = await db.departments.find({}, {"name": 1, "faculty_id": 1}).to_list(1000)
    elif scope_type in ("unit", "multi_unit"):
        items = await db.org_units.find({"is_active": {"$ne": False}}, {"name": 1, "type": 1}).to_list(2000)
    elif scope_type == "course":
        items = await db.courses.find({}, {"name": 1}).to_list(5000)
    else:
        return []
    return [{"id": str(i["_id"]), "name": i.get("name", "")} for i in items]


# ══════════════ قائمة مستخدمين خفيفة (لاختيار المستهدف بالتكليف أو داخل المحاكي) ══════════════

async def _users_list_access(current_user: dict = Depends(get_current_user)):
    db = get_db()
    if not (await has_sub_permission(db, current_user, "manage_assignments") or await has_sub_permission(db, current_user, "run_tests")):
        raise HTTPException(status_code=403, detail="لا تملك صلاحية الوصول إلى مركز الصلاحيات المتقدم (تجريبي)")
    return current_user


@router.get("/users")
async def list_light_users(current_user: dict = Depends(_users_list_access)):
    db = get_db()
    users = await db.users.find({}, {"username": 1, "full_name": 1, "role": 1}).to_list(5000)
    return [{"id": str(u["_id"]), "username": u.get("username"), "full_name": u.get("full_name"), "role": u.get("role")} for u in users]


# ══════════════ التكليفات المتعددة (الجوهر الفعلي لتعدد المناصب) ══════════════

@router.get("/assignments")
async def list_all_assignments(status: Optional[str] = None, current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    q = {}
    if status:
        q["status"] = status
    out = []
    async for a in db.advanced_user_assignments.find(q).sort("created_at", -1):
        user_doc = await db.users.find_one({"_id": ObjectId(a["user_id"])}) if ObjectId.is_valid(a["user_id"]) else None
        role_doc = await db.advanced_roles.find_one({"_id": ObjectId(a["advanced_role_id"])}) if ObjectId.is_valid(a["advanced_role_id"]) else None
        out.append({**_ser(a), "user_name": (user_doc or {}).get("full_name") or (user_doc or {}).get("username"),
                    "role_name_ar": (role_doc or {}).get("name_ar")})
    return out


@router.get("/users/{user_id}/assignments")
async def get_user_assignments(user_id: str, current_user: dict = Depends(require_adv_access("view"))):
    db = get_db()
    return await list_user_work_contexts(db, user_id, active_only=False)


@router.get("/users/{user_id}/contexts")
async def get_user_active_contexts(user_id: str, current_user: dict = Depends(require_adv_access("context_switch"))):
    """سياقات العمل المتاحة فعليًا للتبديل بينها (التكليفات النشطة فقط)"""
    db = get_db()
    return await list_user_work_contexts(db, user_id, active_only=True)


@router.post("/assignments")
async def create_assignment(data: AssignmentCreate, current_user: dict = Depends(require_adv_access("manage_assignments"))):
    db = get_db()
    if current_user.get("role") != "admin":
        if data.user_id == current_user["id"]:
            raise HTTPException(status_code=403, detail="لا يمكنك إسناد تكليف لنفسك")
        target_role = await db.advanced_roles.find_one({"_id": _oid(data.advanced_role_id, "الدور")})
        if target_role and target_role.get("imported_from_legacy_role_id"):
            raise HTTPException(status_code=403, detail="لا يمكن إسناد الأدوار المستوردة من النظام القديم إلا من قِبل الأدمن")
    if not await db.users.find_one({"_id": _oid(data.user_id, "المستخدم")}):
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if not await db.advanced_roles.find_one({"_id": _oid(data.advanced_role_id, "الدور")}):
        raise HTTPException(status_code=404, detail="الدور المتقدم غير موجود")
    doc = {
        "user_id": data.user_id,
        "assignment_type": data.assignment_type,
        "advanced_role_id": data.advanced_role_id,
        "context_label_ar": data.context_label_ar,
        "scope": data.scope.dict(),
        "direct_permissions": [dp.dict() for dp in data.direct_permissions],
        "status": "active",
        "started_at": data.started_at or datetime.now(timezone.utc),
        "ended_at": data.ended_at,
        "created_by": current_user["id"],
        "created_at": datetime.now(timezone.utc),
    }
    res = await db.advanced_user_assignments.insert_one(doc)
    audit_doc = {k: v for k, v in doc.items() if k != "_id"}
    await write_audit_log(db, current_user, "assignment_created", target_user_id=data.user_id, assignment_id=str(res.inserted_id), after=audit_doc)
    return {"id": str(res.inserted_id)}


@router.post("/assignments/{assignment_id}/end")
async def end_assignment(assignment_id: str, current_user: dict = Depends(require_adv_access("manage_assignments"))):
    db = get_db()
    assignment = await db.advanced_user_assignments.find_one({"_id": _oid(assignment_id, "التكليف")})
    if not assignment:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    now = datetime.now(timezone.utc)
    await db.advanced_user_assignments.update_one({"_id": assignment["_id"]}, {"$set": {"status": "ended", "ended_at": now}})
    await write_audit_log(db, current_user, "assignment_ended", target_user_id=assignment["user_id"], assignment_id=assignment_id, before=assignment.get("status"), after="ended")
    return {"message": "تم إنهاء التكليف — بقية تكليفات المستخدم لم تتأثر"}


@router.get("/assignments/{assignment_id}/context")
async def get_assignment_context(assignment_id: str, current_user: dict = Depends(require_adv_access("context_switch"))):
    """تفاصيل سياق العمل عند اختيار هذا التكليف — للعرض فقط، لا يغيّر أي توكن حقيقي"""
    db = get_db()
    assignment = await db.advanced_user_assignments.find_one({"_id": _oid(assignment_id, "التكليف")})
    if not assignment:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    is_admin = current_user.get("role") == "admin"
    is_owner = assignment.get("user_id") == current_user["id"]
    if not is_admin and not is_owner and not await has_sub_permission(db, current_user, "run_tests"):
        raise HTTPException(status_code=403, detail="لا يمكنك الاطلاع على سياق عمل مستخدم آخر")
    role = await db.advanced_roles.find_one({"_id": ObjectId(assignment["advanced_role_id"])}) if ObjectId.is_valid(assignment["advanced_role_id"]) else None
    effective = await compute_effective_permissions(db, assignment)
    await write_audit_log(db, current_user, "context_switched", target_user_id=assignment["user_id"], assignment_id=assignment_id)
    return {
        "assignment_id": assignment_id,
        "context_label_ar": assignment.get("context_label_ar"),
        "assignment_type": assignment.get("assignment_type"),
        "role_name_ar": (role or {}).get("name_ar"),
        "scope": assignment.get("scope"),
        "effective_permissions": sorted(effective),
        "status": assignment.get("status"),
    }


# ══════════════ الصلاحيات المباشرة لتكليف (استثناءات دون تعديل الدور) ══════════════

@router.post("/assignments/{assignment_id}/direct-permissions")
async def grant_direct_permission(assignment_id: str, data: DirectPermissionGrant, current_user: dict = Depends(require_adv_access("manage_permissions"))):
    db = get_db()
    assignment = await db.advanced_user_assignments.find_one({"_id": _oid(assignment_id, "التكليف")})
    if not assignment:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    if not await db.advanced_permissions.find_one({"key": data.permission_key, "is_active": True}):
        raise HTTPException(status_code=400, detail="صلاحية غير معروفة")
    grant = data.dict()
    grant["granted_by"] = current_user["id"]
    grant["granted_at"] = datetime.now(timezone.utc)
    await db.advanced_user_assignments.update_one({"_id": assignment["_id"]}, {"$push": {"direct_permissions": grant}})
    await write_audit_log(db, current_user, "direct_permission_granted", target_user_id=assignment["user_id"], assignment_id=assignment_id, after=grant["permission_key"])
    return {"message": "تم منح الصلاحية المباشرة"}


@router.delete("/assignments/{assignment_id}/direct-permissions/{permission_key}")
async def revoke_direct_permission(assignment_id: str, permission_key: str, current_user: dict = Depends(require_adv_access("manage_permissions"))):
    db = get_db()
    assignment = await db.advanced_user_assignments.find_one({"_id": _oid(assignment_id, "التكليف")})
    if not assignment:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    await db.advanced_user_assignments.update_one({"_id": assignment["_id"]}, {"$pull": {"direct_permissions": {"permission_key": permission_key}}})
    await write_audit_log(db, current_user, "direct_permission_revoked", target_user_id=assignment["user_id"], assignment_id=assignment_id, before=permission_key)
    return {"message": "تم سحب الصلاحية المباشرة"}


# ══════════════ محاكي التحقق ══════════════

@router.post("/check")
async def run_check(data: CheckPermissionRequest, current_user: dict = Depends(require_adv_access("run_tests"))):
    db = get_db()
    return await check_permission(db, data.user_id, data.assignment_id, data.permission_key, data.target_scope_type, data.target_ref_id)


# ══════════════ مقارنة النظام القديم بالجديد (قراءة فقط) ══════════════

@router.post("/compare")
async def compare_one(data: CheckPermissionRequest, current_user: dict = Depends(require_adv_access("run_tests"))):
    db = get_db()
    return await compare_decision(db, data.user_id, data.assignment_id, data.permission_key)


@router.get("/compare/summary/{assignment_id}")
async def compare_summary(assignment_id: str, current_user: dict = Depends(require_adv_access("run_tests"))):
    db = get_db()
    assignment = await db.advanced_user_assignments.find_one({"_id": _oid(assignment_id, "التكليف")})
    if not assignment:
        raise HTTPException(status_code=404, detail="التكليف غير موجود")
    return await compare_summary_for_assignment(db, assignment["user_id"], assignment_id)


# ══════════════ المختبرون (الوصول للمركز نفسه) ══════════════

@router.get("/testers")
async def list_testers(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="للأدمن فقط")
    db = get_db()
    return [_ser(t) async for t in db.advanced_access_testers.find({})]


@router.post("/testers")
async def grant_tester(data: TesterGrant, current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="للأدمن فقط")
    db = get_db()
    if not await db.users.find_one({"_id": _oid(data.user_id, "المستخدم")}):
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    invalid = [s for s in data.sub_permissions if s not in SUB_PERMISSIONS]
    if invalid:
        raise HTTPException(status_code=400, detail=f"صلاحيات فرعية غير معروفة: {invalid}")
    now = datetime.now(timezone.utc)
    await db.advanced_access_testers.update_one(
        {"user_id": data.user_id},
        {"$set": {"sub_permissions": data.sub_permissions, "is_active": True, "granted_by": current_user["id"], "granted_at": now, "revoked_at": None}},
        upsert=True,
    )
    await write_audit_log(db, current_user, "tester_access_granted", target_user_id=data.user_id, after=data.sub_permissions)
    return {"message": "تم منح الوصول التجريبي"}


@router.delete("/testers/{user_id}")
async def revoke_tester(user_id: str, current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="للأدمن فقط")
    db = get_db()
    await db.advanced_access_testers.update_one({"user_id": user_id}, {"$set": {"is_active": False, "revoked_at": datetime.now(timezone.utc)}})
    await write_audit_log(db, current_user, "tester_access_revoked", target_user_id=user_id)
    return {"message": "تم سحب الوصول التجريبي"}


# ══════════════ سجل التدقيق ══════════════

@router.get("/audit-logs")
async def list_audit_logs(limit: int = 100, current_user: dict = Depends(require_adv_access("view_audit"))):
    db = get_db()
    logs = await db.advanced_permission_audit_logs.find({}).sort("timestamp", -1).limit(min(limit, 500)).to_list(500)
    return [_ser(logv) for logv in logs]


# ══════════════ المرحلة 5: التجربة التشغيلية المحدودة (Pilot) — SHADOW MODE، مستخدمون حقيقيون محدودون ══════════════

def _require_admin():
    async def _dep(current_user: dict = Depends(get_current_user)):
        if current_user.get("role") != "admin":
            raise HTTPException(status_code=403, detail="هذا الإجراء للأدمن فقط")
        return current_user
    return _dep


@router.post("/pilot/run")
async def pilot_run(current_user: dict = Depends(_require_admin())):
    """يُنشئ أدوارًا وتكليفات تجريبية (PILOT_) لعدد محدود من المستخدمين الحقيقيين المعتمدين، ويُسجِّل نتائج فحوصات فعلية"""
    db = get_db()
    return await run_pilot(db, current_user)


@router.post("/pilot/reset")
async def pilot_reset(current_user: dict = Depends(_require_admin())):
    """إيقاف فوري للتجربة: إنهاء كل التكليفات التجريبية وحذف الأدوار والنتائج — لا يلمس النظام القديم"""
    db = get_db()
    result = await reset_pilot(db)
    await write_audit_log(db, current_user, "pilot_reset_executed")
    return result


@router.get("/pilot/summary")
async def pilot_summary(current_user: dict = Depends(require_adv_access("view_pilot"))):
    db = get_db()
    run = await db.advanced_pilot_runs.find_one({})
    if not run:
        return {"has_run": False}
    run["id"] = str(run.pop("_id"))
    run["has_run"] = True
    return run


@router.get("/pilot/results")
async def pilot_results(
    user_key: Optional[str] = None, test_type: Optional[str] = None, category: Optional[str] = None,
    severity: Optional[str] = None, permission_key: Optional[str] = None, scope_type: Optional[str] = None,
    status: Optional[str] = None,
    current_user: dict = Depends(require_adv_access("view_pilot")),
):
    db = get_db()
    q = {}
    if user_key:
        q["user_key"] = user_key
    if test_type:
        q["test_type"] = test_type
    if category:
        q["category"] = category
    if severity:
        q["severity"] = severity
    if status:
        q["status"] = status
    if permission_key:
        q["permission_key"] = {"$regex": re.escape(permission_key)}
    if scope_type:
        q["target_scope_type"] = scope_type
    docs = await db.advanced_pilot_results.find(q).sort("created_at", -1).to_list(2000)
    return [_ser(d) for d in docs]


@router.get("/pilot/export.xlsx")
async def pilot_export_xlsx(current_user: dict = Depends(require_adv_access("view_pilot"))):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from fastapi.responses import StreamingResponse

    db = get_db()
    docs = await db.advanced_pilot_results.find({}).sort("created_at", -1).to_list(5000)
    wb = Workbook()
    ws = wb.active
    ws.title = "نتائج التجربة التشغيلية"
    ws.sheet_view.rightToLeft = True
    headers = ["المستخدم", "نوع الاختبار", "التكليف", "الصلاحية", "النطاق المستهدف", "المتوقَّع", "الفعلي",
               "القرار القديم", "القرار الجديد", "التصنيف", "الخطورة", "الحالة", "السبب"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
    for d in docs:
        ws.append([
            d.get("username"), d.get("test_type"), d.get("assignment_id") or "-", d.get("permission_key"),
            d.get("target_scope_type") or "-", str(d.get("expected", "-")), str(d.get("actual", "-")),
            str(d.get("old_decision")), str(d.get("new_decision")), d.get("category") or "-",
            d.get("severity"), d.get("status"), d.get("reason"),
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", wrap_text=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=pilot_results.xlsx"},
    )
