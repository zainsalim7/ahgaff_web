"""🏢 شؤون الموظفين — المرحلة 1 (الأساس): الهيكل التنظيمي + سجل الموظفين الموحّد (معلمون وإداريون) + استيراد Excel + حسابات الخدمة الذاتية"""
import io
from datetime import datetime, timedelta, timezone
from typing import Optional, List

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, has_permission, log_activity, get_password_hash
from .hr_common import _today

router = APIRouter(prefix="/hr", tags=["شؤون الموظفين"])
YEMEN_TZ = timezone(timedelta(hours=3))

P_VIEW, P_MANAGE, P_ORG = "hr_view_employees", "hr_manage_employees", "hr_manage_org"
UNIT_TYPES = {"presidency": "رئاسة الجامعة", "faculty": "كلية", "department": "قسم أكاديمي", "administration": "إدارة", "office": "مكتب / وحدة"}
CONTRACT_TYPES = {"permanent": "دائم", "contract": "متعاقد", "part_time": "جزئي", "hourly": "بالساعة", "visiting": "متعاون / زائر"}
STATUSES = {"probation": "تحت التجربة", "active": "على رأس العمل", "leave": "في إجازة", "suspended": "موقوف", "ended": "منتهية خدمته"}
CATEGORIES = {"academic": "أكاديمي (معلم)", "administrative": "إداري", "technical": "فني", "service": "خدمات مساندة"}


def _now():
    return datetime.now(YEMEN_TZ).isoformat()


def _oid(v: str, what: str = "المعرّف"):
    if not v or not ObjectId.is_valid(v):
        raise HTTPException(status_code=400, detail=f"{what} غير صحيح")
    return ObjectId(v)


def _ser(d: dict) -> dict:
    d = dict(d)
    d["id"] = str(d.pop("_id"))
    return d


def _can_view(u: dict) -> bool:
    return has_permission(u, P_VIEW) or has_permission(u, P_MANAGE)


def _guard(u: dict, perm: str):
    if not has_permission(u, perm):
        raise HTTPException(status_code=403, detail="ليست لديك صلاحية على شؤون الموظفين")


# ══════════════ الهيكل التنظيمي ══════════════

class OrgUnitIn(BaseModel):
    name: str
    type: str = "administration"
    parent_id: Optional[str] = None
    code: Optional[str] = ""
    head_employee_id: Optional[str] = None
    description: Optional[str] = ""


@router.get("/org-units")
async def list_org_units(current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user) and not has_permission(current_user, P_ORG):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    units = [_ser(u) for u in await db.org_units.find({"is_active": {"$ne": False}}).sort([("order", 1), ("name", 1)]).to_list(2000)]
    counts = {c["_id"]: c["n"] async for c in db.employees.aggregate([{"$match": {"status": {"$ne": "ended"}}}, {"$group": {"_id": "$org_unit_id", "n": {"$sum": 1}}}])}
    for u in units:
        u["employees_count"] = counts.get(u["id"], 0)
        u["type_label"] = UNIT_TYPES.get(u.get("type"), u.get("type"))
    return {"units": units, "types": UNIT_TYPES}


@router.post("/org-units/sync-academic")
async def sync_academic_units(current_user: dict = Depends(get_current_user)):
    """🔄 توليد وحدات الهيكل من الكليات والأقسام الموجودة (idempotent)"""
    _guard(current_user, P_ORG)
    db = get_db()
    root = await db.org_units.find_one({"type": "presidency"})
    if not root:
        r = await db.org_units.insert_one({"name": "رئاسة الجامعة", "type": "presidency", "parent_id": None, "code": "PRES", "order": 0, "is_active": True, "created_at": _now()})
        root_id = str(r.inserted_id)
    else:
        root_id = str(root["_id"])
    created = 0
    fac_map = {}
    async for f in db.faculties.find({}):
        ex = await db.org_units.find_one({"type": "faculty", "faculty_id": str(f["_id"])})
        if not ex:
            r = await db.org_units.insert_one({"name": f.get("name", ""), "type": "faculty", "parent_id": root_id, "faculty_id": str(f["_id"]), "code": f.get("code", ""), "order": 10, "is_active": True, "created_at": _now()})
            fac_map[str(f["_id"])] = str(r.inserted_id); created += 1
        else:
            fac_map[str(f["_id"])] = str(ex["_id"])
            if ex.get("name") != f.get("name"):
                await db.org_units.update_one({"_id": ex["_id"]}, {"$set": {"name": f.get("name", "")}})
    async for d in db.departments.find({}):
        ex = await db.org_units.find_one({"type": "department", "department_id": str(d["_id"])})
        parent = fac_map.get(str(d.get("faculty_id") or ""), root_id)
        if not ex:
            await db.org_units.insert_one({"name": d.get("name", ""), "type": "department", "parent_id": parent, "department_id": str(d["_id"]), "faculty_id": str(d.get("faculty_id") or ""), "code": d.get("code", ""), "order": 20, "is_active": True, "created_at": _now()})
            created += 1
        elif ex.get("name") != d.get("name") or ex.get("parent_id") != parent:
            await db.org_units.update_one({"_id": ex["_id"]}, {"$set": {"name": d.get("name", ""), "parent_id": parent}})
    rl = await relink_teacher_units(db)
    await log_activity(current_user, "hr_sync_org", "org_unit", root_id, "الهيكل التنظيمي", {"created": created, "linked": rl["linked"]})
    return {"created": created, "linked": rl["linked"], "message": f"تمت مزامنة الهيكل — أُضيفت {created} وحدة جديدة" + (f" وربط {rl['linked']} معلماً بوحداتهم" if rl["linked"] else "")}


@router.post("/org-units")
async def create_org_unit(data: OrgUnitIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ORG)
    db = get_db()
    if data.type not in UNIT_TYPES:
        raise HTTPException(status_code=400, detail="نوع الوحدة غير صحيح")
    if data.parent_id and not await db.org_units.find_one({"_id": _oid(data.parent_id, "الوحدة الأم")}):
        raise HTTPException(status_code=404, detail="الوحدة الأم غير موجودة")
    if not data.name.strip():
        raise HTTPException(status_code=400, detail="اسم الوحدة مطلوب")
    doc = {**data.dict(), "name": data.name.strip(), "order": 50, "is_active": True, "created_at": _now(), "created_by_name": current_user.get("full_name", "")}
    r = await db.org_units.insert_one(doc)
    await log_activity(current_user, "hr_create_org_unit", "org_unit", str(r.inserted_id), data.name)
    return {"id": str(r.inserted_id), "message": "تمت إضافة الوحدة"}


@router.put("/org-units/{unit_id}")
async def update_org_unit(unit_id: str, data: OrgUnitIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ORG)
    db = get_db()
    oid = _oid(unit_id)
    if data.parent_id == unit_id:
        raise HTTPException(status_code=400, detail="لا يمكن أن تكون الوحدة أماً لنفسها")
    # منع الحلقات: الأم الجديدة لا تكون من أبناء الوحدة
    cur = data.parent_id
    for _ in range(50):
        if not cur:
            break
        if cur == unit_id:
            raise HTTPException(status_code=400, detail="لا يمكن نقل الوحدة تحت أحد أبنائها")
        p = await db.org_units.find_one({"_id": ObjectId(cur)}, {"parent_id": 1}) if ObjectId.is_valid(cur) else None
        cur = (p or {}).get("parent_id")
    upd = {k: v for k, v in data.dict().items() if k != "type" or not (await db.org_units.find_one({"_id": oid})).get("faculty_id")}
    await db.org_units.update_one({"_id": oid}, {"$set": {**upd, "updated_at": _now()}})
    return {"message": "تم تحديث الوحدة"}


@router.delete("/org-units/{unit_id}")
async def delete_org_unit(unit_id: str, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_ORG)
    db = get_db()
    oid = _oid(unit_id)
    u = await db.org_units.find_one({"_id": oid})
    if not u:
        raise HTTPException(status_code=404, detail="الوحدة غير موجودة")
    if u.get("faculty_id") or u.get("department_id") or u.get("type") == "presidency":
        raise HTTPException(status_code=400, detail="الوحدات الأكاديمية تُدار من شاشات الكليات والأقسام")
    if await db.org_units.count_documents({"parent_id": unit_id, "is_active": {"$ne": False}}):
        raise HTTPException(status_code=400, detail="انقل أو احذف الوحدات الفرعية أولاً")
    if await db.employees.count_documents({"org_unit_id": unit_id}):
        raise HTTPException(status_code=400, detail="انقل موظفي الوحدة إلى وحدة أخرى أولاً")
    await db.org_units.update_one({"_id": oid}, {"$set": {"is_active": False, "deleted_at": _now()}})
    await log_activity(current_user, "hr_delete_org_unit", "org_unit", unit_id, u.get("name", ""))
    return {"message": "تم حذف الوحدة"}


# ══════════════ الموظفون ══════════════

class EmployeeIn(BaseModel):
    employee_no: str
    full_name: str
    category: str = "administrative"
    job_title: Optional[str] = ""
    grade: Optional[str] = ""
    org_unit_id: Optional[str] = None
    manager_employee_id: Optional[str] = None
    contract_type: str = "permanent"
    hire_date: Optional[str] = None
    contract_end_date: Optional[str] = None
    status: str = "active"
    gender: Optional[str] = ""
    nationality: Optional[str] = ""
    national_id: Optional[str] = ""
    id_expiry_date: Optional[str] = None
    birth_date: Optional[str] = None
    phone: Optional[str] = ""
    email: Optional[str] = ""
    address: Optional[str] = ""
    emergency_contact: Optional[str] = ""
    qualification: Optional[str] = ""
    specialization: Optional[str] = ""
    notes: Optional[str] = ""
    teacher_id: Optional[str] = None


def _validate(data: EmployeeIn):
    if not data.employee_no.strip() or not data.full_name.strip():
        raise HTTPException(status_code=400, detail="الرقم الوظيفي والاسم مطلوبان")
    if data.category not in CATEGORIES:
        raise HTTPException(status_code=400, detail="فئة الموظف غير صحيحة")
    if data.contract_type not in CONTRACT_TYPES:
        raise HTTPException(status_code=400, detail="نوع التعاقد غير صحيح")
    if data.status not in STATUSES:
        raise HTTPException(status_code=400, detail="الحالة غير صحيحة")
    for f in ("hire_date", "contract_end_date", "id_expiry_date", "birth_date"):
        v = getattr(data, f)
        if v:
            try:
                datetime.strptime(v, "%Y-%m-%d")
            except ValueError:
                raise HTTPException(status_code=400, detail=f"صيغة التاريخ غير صحيحة في {f} (YYYY-MM-DD)")


async def _enrich(db, emps: List[dict]) -> List[dict]:
    unit_ids = {e.get("org_unit_id") for e in emps if e.get("org_unit_id") and ObjectId.is_valid(e["org_unit_id"])}
    units = {str(u["_id"]): u for u in await db.org_units.find({"_id": {"$in": [ObjectId(i) for i in unit_ids]}}, {"name": 1, "type": 1}).to_list(2000)} if unit_ids else {}
    mgr_ids = {e.get("manager_employee_id") for e in emps if e.get("manager_employee_id") and ObjectId.is_valid(e["manager_employee_id"])}
    mgrs = {str(m["_id"]): m.get("full_name", "") for m in await db.employees.find({"_id": {"$in": [ObjectId(i) for i in mgr_ids]}}, {"full_name": 1}).to_list(2000)} if mgr_ids else {}
    today = datetime.now(YEMEN_TZ).date()
    out = []
    for e in emps:
        d = _ser(e)
        u = units.get(d.get("org_unit_id") or "", {})
        d["org_unit_name"] = u.get("name", "")
        d["manager_name"] = mgrs.get(d.get("manager_employee_id") or "", "")
        d["category_label"] = CATEGORIES.get(d.get("category"), d.get("category"))
        d["contract_type_label"] = CONTRACT_TYPES.get(d.get("contract_type"), d.get("contract_type"))
        d["status_label"] = STATUSES.get(d.get("status"), d.get("status"))
        d["has_account"] = bool(d.get("user_id"))
        alerts = []
        for f, label in (("id_expiry_date", "الهوية/الإقامة"), ("contract_end_date", "العقد")):
            if d.get(f):
                try:
                    days = (datetime.strptime(d[f], "%Y-%m-%d").date() - today).days
                    if days < 0:
                        alerts.append(f"{label} منتهية")
                    elif days <= 60:
                        alerts.append(f"{label} تنتهي خلال {days} يوماً")
                except ValueError:
                    pass
        d["alerts"] = alerts
        out.append(d)
    return out


@router.get("/employees/meta")
async def employees_meta(current_user: dict = Depends(get_current_user)):
    return {"categories": CATEGORIES, "contract_types": CONTRACT_TYPES, "statuses": STATUSES, "unit_types": UNIT_TYPES}


async def _build_query(db, search, org_unit_id, category, status, contract_type) -> dict:
    q: dict = {}
    if search:
        import re
        rx = {"$regex": re.escape(search.strip()), "$options": "i"}
        q["$or"] = [{"full_name": rx}, {"employee_no": rx}, {"job_title": rx}, {"phone": rx}, {"national_id": rx}]
    if org_unit_id:
        ids, frontier = {org_unit_id}, [org_unit_id]
        while frontier:
            kids = [str(u["_id"]) for u in await db.org_units.find({"parent_id": {"$in": frontier}}, {"_id": 1}).to_list(2000)]
            frontier = [k for k in kids if k not in ids]
            ids |= set(kids)
        q["org_unit_id"] = {"$in": list(ids)}
    for k, v in (("category", category), ("status", status), ("contract_type", contract_type)):
        if v:
            q[k] = v
    return q


@router.get("/employees/export")
async def export_employees(search: Optional[str] = None, org_unit_id: Optional[str] = None, category: Optional[str] = None, status: Optional[str] = None,
                           contract_type: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """📊 تصدير سجل الموظفين (بنفس الفلاتر) إلى Excel"""
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    q = await _build_query(db, search, org_unit_id, category, status, contract_type)
    emps = await _enrich(db, await db.employees.find(q).sort("full_name", 1).to_list(10000))
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from .deps import export_headers, export_filename
    wb = Workbook(); ws = wb.active; ws.title = "الموظفون"; ws.sheet_view.rightToLeft = True
    heads = ["الرقم الوظيفي", "الاسم", "الفئة", "المسمى الوظيفي", "الدرجة", "الوحدة التنظيمية", "المدير المباشر", "نوع التعاقد", "تاريخ التعيين", "انتهاء العقد", "الحالة", "الجنس", "الهاتف", "البريد", "الرقم الوطني", "انتهاء الهوية", "المؤهل", "التخصص", "تاريخ الميلاد", "حساب دخول", "ملاحظات"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2440"); c.alignment = Alignment(horizontal="center")
    for e in emps:
        ws.append([e.get("employee_no"), e.get("full_name"), CATEGORIES.get(e.get("category"), e.get("category")), e.get("job_title"), e.get("grade"), e.get("org_unit_name"), e.get("manager_name"),
                   CONTRACT_TYPES.get(e.get("contract_type"), e.get("contract_type")), e.get("hire_date"), e.get("contract_end_date"), STATUSES.get(e.get("status"), e.get("status")), e.get("gender"), e.get("phone"), e.get("email"),
                   e.get("national_id"), e.get("id_expiry_date"), e.get("qualification"), e.get("specialization"), e.get("birth_date"), "نعم" if e.get("has_account") else "لا", e.get("notes")])
    for i in range(1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions["B"].width = 30
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=export_headers(export_filename("سجل الموظفين", _today(), ext="xlsx")))


@router.get("/employees")
async def list_employees(search: Optional[str] = None, org_unit_id: Optional[str] = None, category: Optional[str] = None, status: Optional[str] = None,
                         contract_type: Optional[str] = None, page: int = 1, per_page: int = 30, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    q = await _build_query(db, search, org_unit_id, category, status, contract_type)
    total = await db.employees.count_documents(q)
    per_page = max(1, min(per_page, 200))
    emps = await db.employees.find(q).sort("full_name", 1).skip((page - 1) * per_page).limit(per_page).to_list(per_page)
    stats_raw = {s["_id"]: s["n"] async for s in db.employees.aggregate([{"$group": {"_id": "$status", "n": {"$sum": 1}}}])}
    cat_raw = {s["_id"]: s["n"] async for s in db.employees.aggregate([{"$group": {"_id": "$category", "n": {"$sum": 1}}}])}
    return {"employees": await _enrich(db, emps), "total": total, "page": page, "per_page": per_page,
            "stats": {"total": sum(stats_raw.values()), "by_status": stats_raw, "by_category": cat_raw}}


@router.get("/employees/me")
async def my_employee_profile(current_user: dict = Depends(get_current_user)):
    """👤 الخدمة الذاتية: ملفي الإداري"""
    db = get_db()
    from .hr_common import find_my_employee
    e = await find_my_employee(db, current_user)
    if not e:
        return {"profile": None}
    return {"profile": (await _enrich(db, [e]))[0]}


class BulkIn(BaseModel):
    ids: List[str]
    action: str  # move_unit | set_status | set_manager | set_category | set_contract_type | delete | register_leave
    value: Optional[str] = None
    leave: Optional[dict] = None


BULK_LABELS = {"move_unit": "نقل إلى وحدة", "set_status": "تغيير الحالة", "set_manager": "تعيين مدير مباشر", "set_category": "تغيير الفئة", "set_contract_type": "تغيير نوع التعاقد", "delete": "حذف", "register_leave": "تسجيل إجازة"}


@router.post("/employees/bulk")
async def bulk_action(data: BulkIn, current_user: dict = Depends(get_current_user)):
    """⚡ إجراء جماعي على عدة موظفين"""
    _guard(current_user, P_MANAGE)
    db = get_db()
    if data.action not in BULK_LABELS:
        raise HTTPException(status_code=400, detail="إجراء غير معروف")
    ids = [ObjectId(i) for i in data.ids if ObjectId.is_valid(i)]
    if not ids:
        raise HTTPException(status_code=400, detail="لم يُحدَّد موظفون")
    emps = await db.employees.find({"_id": {"$in": ids}}).to_list(len(ids))
    label_to = None
    if data.action == "move_unit":
        u = await db.org_units.find_one({"_id": _oid(data.value or "", "الوحدة")})
        if not u:
            raise HTTPException(status_code=404, detail="الوحدة غير موجودة")
        label_to = u.get("name")
    elif data.action == "set_status" and data.value not in STATUSES:
        raise HTTPException(status_code=400, detail="حالة غير صحيحة")
    elif data.action == "set_category" and data.value not in CATEGORIES:
        raise HTTPException(status_code=400, detail="فئة غير صحيحة")
    elif data.action == "set_contract_type" and data.value not in CONTRACT_TYPES:
        raise HTTPException(status_code=400, detail="نوع تعاقد غير صحيح")
    elif data.action == "set_manager":
        if data.value:
            m = await db.employees.find_one({"_id": _oid(data.value, "المدير")}, {"full_name": 1})
            if not m:
                raise HTTPException(status_code=404, detail="المدير غير موجود")
            label_to = m.get("full_name")
    elif data.action == "register_leave":
        from .hr_leaves import LeaveIn, _create_request
        lv = data.leave or {}
        if not (lv.get("start_date") and lv.get("end_date")):
            raise HTTPException(status_code=400, detail="حدد فترة الإجازة")
    ok, errors = 0, []
    for e in emps:
        eid = str(e["_id"])
        try:
            if data.action == "delete":
                await _delete_employee(db, e, current_user, force=True)
            elif data.action == "register_leave":
                await _create_request(db, e, LeaveIn(employee_id=eid, type=lv.get("type", "annual"), start_date=lv["start_date"], end_date=lv["end_date"], reason=lv.get("reason", "")), current_user, auto_approve=True)
            else:
                field = {"move_unit": "org_unit_id", "set_status": "status", "set_manager": "manager_employee_id", "set_category": "category", "set_contract_type": "contract_type"}[data.action]
                if data.action == "set_manager" and data.value == eid:
                    raise HTTPException(status_code=400, detail="لا يمكن أن يكون الموظف مديراً لنفسه")
                old = e.get(field)
                if old != (data.value or None):
                    await db.employees.update_one({"_id": e["_id"]}, {"$set": {field: data.value or None, "updated_at": _now()}})
                    await _history(db, eid, "updated", current_user, {field: {"from": old, "to": data.value or None}, "bulk": BULK_LABELS[data.action]})
            ok += 1
        except HTTPException as ex:
            errors.append({"id": eid, "name": e.get("full_name", ""), "error": ex.detail})
        except Exception as ex:
            errors.append({"id": eid, "name": e.get("full_name", ""), "error": str(ex)})
    await log_activity(current_user, "hr_bulk", "employee", "", BULK_LABELS[data.action], {"count": ok, "value": label_to or data.value})
    return {"ok": ok, "errors": errors, "message": f"{BULK_LABELS[data.action]}: نجح {ok}" + (f" · تعذّر {len(errors)}" if errors else "")}


@router.get("/employees/{emp_id}")
async def get_employee(emp_id: str, current_user: dict = Depends(get_current_user)):
    if not _can_view(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    e = await db.employees.find_one({"_id": _oid(emp_id)})
    if not e:
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    d = (await _enrich(db, [e]))[0]
    d["subordinates"] = [{"id": str(s["_id"]), "full_name": s.get("full_name", ""), "job_title": s.get("job_title", "")} for s in await db.employees.find({"manager_employee_id": emp_id}, {"full_name": 1, "job_title": 1}).to_list(200)]
    d["history"] = [_ser(h) for h in await db.employee_history.find({"employee_id": emp_id}).sort("at", -1).limit(50).to_list(50)]
    d["role_id"], d["role_name"], d["account_username"] = None, None, None
    if d.get("user_id") and ObjectId.is_valid(d["user_id"]):
        u = await db.users.find_one({"_id": ObjectId(d["user_id"])}, {"role_id": 1, "role": 1, "username": 1})
        if u:
            d["account_username"] = u.get("username")
            if u.get("role_id") and ObjectId.is_valid(u["role_id"]):
                rd = await db.roles.find_one({"_id": ObjectId(u["role_id"])}, {"name": 1})
                d["role_id"], d["role_name"] = u["role_id"], (rd or {}).get("name")
            elif u.get("role") and u["role"] not in ("employee",):
                d["role_name"] = {"admin": "مدير النظام", "dean": "عميد", "department_head": "رئيس قسم", "registrar": "مسجّل", "registration_manager": "مدير التسجيل", "university_president": "رئيس الجامعة", "teacher": "معلم"}.get(u["role"], u["role"])
    if d.get("teacher_id") and ObjectId.is_valid(d["teacher_id"]):
        t = await db.teachers.find_one({"_id": ObjectId(d["teacher_id"])}, {"teacher_id": 1, "department_id": 1})
        if t:
            d["teacher"] = {"id": d["teacher_id"], "teacher_id": t.get("teacher_id", ""), "courses_count": await db.courses.count_documents({"teacher_id": d["teacher_id"], "is_active": True})}
    return d


async def _history(db, emp_id: str, action: str, by: dict, details: dict = None):
    await db.employee_history.insert_one({"employee_id": emp_id, "action": action, "details": details or {}, "by_name": by.get("full_name", ""), "at": _now()})


@router.post("/employees")
async def create_employee(data: EmployeeIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    db = get_db()
    _validate(data)
    if await db.employees.find_one({"employee_no": data.employee_no.strip()}):
        raise HTTPException(status_code=400, detail="الرقم الوظيفي مستخدم لموظف آخر")
    if data.teacher_id and await db.employees.find_one({"teacher_id": data.teacher_id}):
        raise HTTPException(status_code=400, detail="هذا المعلم مرتبط بملف إداري آخر")
    doc = {**data.dict(), "employee_no": data.employee_no.strip(), "full_name": data.full_name.strip(), "user_id": None, "created_at": _now(), "created_by_name": current_user.get("full_name", "")}
    r = await db.employees.insert_one(doc)
    await _history(db, str(r.inserted_id), "created", current_user)
    await log_activity(current_user, "hr_create_employee", "employee", str(r.inserted_id), data.full_name)
    return {"id": str(r.inserted_id), "message": "تمت إضافة الموظف"}


TRACKED = ("job_title", "grade", "org_unit_id", "manager_employee_id", "contract_type", "status", "category", "contract_end_date")


@router.put("/employees/{emp_id}")
async def update_employee(emp_id: str, data: EmployeeIn, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    db = get_db()
    oid = _oid(emp_id)
    old = await db.employees.find_one({"_id": oid})
    if not old:
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    _validate(data)
    dup = await db.employees.find_one({"employee_no": data.employee_no.strip(), "_id": {"$ne": oid}})
    if dup:
        raise HTTPException(status_code=400, detail="الرقم الوظيفي مستخدم لموظف آخر")
    if data.manager_employee_id == emp_id:
        raise HTTPException(status_code=400, detail="لا يمكن أن يكون الموظف مديراً لنفسه")
    upd = {**data.dict(), "employee_no": data.employee_no.strip(), "full_name": data.full_name.strip(), "updated_at": _now()}
    await db.employees.update_one({"_id": oid}, {"$set": upd})
    changes = {k: {"from": old.get(k), "to": upd.get(k)} for k in TRACKED if (old.get(k) or None) != (upd.get(k) or None)}
    if changes:
        await _history(db, emp_id, "updated", current_user, changes)
    if old.get("user_id") and (old.get("full_name") != upd["full_name"] or old.get("phone") != upd.get("phone")):
        await db.users.update_one({"_id": ObjectId(old["user_id"])}, {"$set": {"full_name": upd["full_name"], "phone": upd.get("phone", "")}})
    await log_activity(current_user, "hr_update_employee", "employee", emp_id, data.full_name, changes)
    return {"message": "تم تحديث بيانات الموظف", "changes": changes}


@router.delete("/employees/{emp_id}")
async def delete_employee(emp_id: str, force: bool = False, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    db = get_db()
    e = await db.employees.find_one({"_id": _oid(emp_id)})
    if not e:
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    return await _delete_employee(db, e, current_user, force)


async def _delete_employee(db, e: dict, current_user: dict, force: bool = False) -> dict:
    """حذف ناعم إلى سلة المحذوفات + تعطيل الحساب. ملف المعلم يُحذف فقط بـ force (سجل المعلم الأكاديمي يبقى)"""
    emp_id, oid = str(e["_id"]), e["_id"]
    if e.get("teacher_id") and not force:
        raise HTTPException(status_code=400, detail="هذا ملف معلم — احذفه بتأكيد «حذف الملف الإداري فقط» (يبقى سجله الأكاديمي في شاشة المعلمين)")
    if await db.employees.count_documents({"manager_employee_id": emp_id}):
        raise HTTPException(status_code=400, detail="هذا الموظف مدير لموظفين آخرين — غيّر مديرهم أولاً")
    backup = {"backup_type": "employee_backup", "deleted_at": _now(), "deleted_by": current_user.get("full_name", ""), "employee": {**e, "_id": str(e["_id"])}}
    await db.trash.insert_one({"item_type": "employee", "item_name": e.get("full_name", ""), "backup_data": backup, "deleted_by": current_user.get("full_name", ""),
                               "deleted_at": datetime.now(YEMEN_TZ), "expires_at": datetime.now(YEMEN_TZ) + timedelta(days=30)})
    await db.employees.delete_one({"_id": oid})
    if e.get("user_id") and ObjectId.is_valid(e["user_id"]) and not e.get("teacher_id"):
        await db.users.update_one({"_id": ObjectId(e["user_id"])}, {"$set": {"is_active": False}})
    await log_activity(current_user, "hr_delete_employee", "employee", emp_id, e.get("full_name", ""))
    return {"message": "تم حذف الموظف (يمكن استعادته من سلة المحذوفات) وتعطيل حسابه إن وُجد", "backup": backup}


async def relink_teacher_units(db, force: bool = False) -> dict:
    """🔗 ربط ملفات المعلمين بوحدة القسم (أو الكلية) وفق بيانات المعلم — idempotent"""
    linked, skipped, unresolved = 0, 0, []
    dept_units = {u["department_id"]: str(u["_id"]) async for u in db.org_units.find({"type": "department", "department_id": {"$ne": None}}, {"department_id": 1})}
    fac_units = {u["faculty_id"]: str(u["_id"]) async for u in db.org_units.find({"type": "faculty", "faculty_id": {"$ne": None}}, {"faculty_id": 1})}
    dept_fac = {str(d["_id"]): str(d.get("faculty_id") or "") async for d in db.departments.find({}, {"faculty_id": 1})}
    async for e in db.employees.find({"teacher_id": {"$ne": None}}, {"full_name": 1, "org_unit_id": 1, "teacher_id": 1}):
        if e.get("org_unit_id") and not force:
            skipped += 1
            continue
        t = await db.teachers.find_one({"_id": ObjectId(e["teacher_id"])}, {"department_id": 1, "department_ids": 1, "faculty_id": 1}) if ObjectId.is_valid(e["teacher_id"]) else None
        if not t:
            unresolved.append({"id": str(e["_id"]), "name": e.get("full_name", ""), "reason": "سجل المعلم غير موجود"})
            continue
        dept = str(t.get("department_id") or (t.get("department_ids") or [None])[0] or "")
        fac = str(t.get("faculty_id") or dept_fac.get(dept) or "")
        unit = dept_units.get(dept) or fac_units.get(fac)
        if not unit:
            unresolved.append({"id": str(e["_id"]), "name": e.get("full_name", ""), "reason": "لا قسم/كلية للمعلم أو لم تُزامَن وحدته بعد"})
            continue
        if unit != e.get("org_unit_id"):
            await db.employees.update_one({"_id": e["_id"]}, {"$set": {"org_unit_id": unit, "updated_at": _now()}})
            linked += 1
        else:
            skipped += 1
    return {"linked": linked, "skipped": skipped, "unresolved": unresolved}


@router.post("/employees/link-units")
async def link_units(force: bool = False, current_user: dict = Depends(get_current_user)):
    """🔗 ربط المعلمين بوحداتهم التنظيمية (القسم أو الكلية). force=true يعيد الربط حتى للمرتبطين"""
    _guard(current_user, P_MANAGE)
    res = await relink_teacher_units(get_db(), force)
    await log_activity(current_user, "hr_link_units", "employee", "", "ربط المعلمين بالوحدات", {"linked": res["linked"]})
    msg = f"تم ربط {res['linked']} موظفاً بوحداتهم" + (f" · {res['skipped']} مرتبط مسبقاً" if res["skipped"] else "") + (f" · {len(res['unresolved'])} بلا وحدة (اربطهم يدوياً من «تعديل»)" if res["unresolved"] else "")
    return {**res, "message": msg}


@router.post("/employees/sync-teachers")
async def sync_teachers(current_user: dict = Depends(get_current_user)):
    """🔄 إنشاء ملف إداري لكل معلم موجود ليس له ملف (idempotent)"""
    _guard(current_user, P_MANAGE)
    db = get_db()
    linked = {e["teacher_id"] for e in await db.employees.find({"teacher_id": {"$ne": None}}, {"teacher_id": 1}).to_list(10000)}
    created = 0
    async for t in db.teachers.find({}):
        tid = str(t["_id"])
        if tid in linked:
            continue
        unit = await db.org_units.find_one({"type": "department", "department_id": str(t.get("department_id") or "")}, {"_id": 1})
        emp_no = (t.get("teacher_id") or f"T-{tid[-6:]}").strip()
        if await db.employees.find_one({"employee_no": emp_no}):
            emp_no = f"{emp_no}-T"
        await db.employees.insert_one({
            "employee_no": emp_no, "full_name": t.get("full_name", ""), "category": "academic", "job_title": t.get("academic_rank") or t.get("title") or "عضو هيئة تدريس",
            "grade": "", "org_unit_id": str(unit["_id"]) if unit else None, "manager_employee_id": None, "contract_type": "permanent",
            "hire_date": None, "contract_end_date": None, "status": "active" if t.get("is_active", True) else "ended",
            "gender": t.get("gender", ""), "nationality": "", "national_id": "", "id_expiry_date": None, "birth_date": None,
            "phone": t.get("phone", ""), "email": t.get("email", ""), "address": "", "emergency_contact": "", "qualification": t.get("qualification", ""),
            "specialization": t.get("specialization", ""), "notes": "", "teacher_id": tid, "user_id": t.get("user_id"),
            "created_at": _now(), "created_by_name": "مزامنة المعلمين",
        })
        created += 1
    rl = await relink_teacher_units(db)
    await log_activity(current_user, "hr_sync_teachers", "employee", None, "مزامنة المعلمين", {"created": created, "linked": rl["linked"]})
    return {"created": created, "linked": rl["linked"], "unresolved": rl["unresolved"], "message": f"تم إنشاء {created} ملفاً إدارياً للمعلمين" + (f" وربط {rl['linked']} بوحداتهم" if rl["linked"] else "") + (f" · {len(rl['unresolved'])} بلا وحدة" if rl["unresolved"] else "")}


class AccountIn(BaseModel):
    role_id: Optional[str] = None


async def _resolve_role(db, role_id: Optional[str]) -> Optional[dict]:
    if not role_id:
        return None
    r = await db.roles.find_one({"_id": _oid(role_id, "الدور")})
    if not r or r.get("system_key") == "admin":
        raise HTTPException(status_code=400, detail="الدور غير صالح")
    return r


@router.get("/roles")
async def hr_roles(current_user: dict = Depends(get_current_user)):
    """الأدوار التي يمكن إسنادها لحساب موظف (كل الأدوار ما عدا مدير النظام/الطالب/المعلم)"""
    _guard(current_user, P_MANAGE)
    db = get_db()
    out = []
    async for r in db.roles.find({"system_key": {"$nin": ["admin", "student", "teacher"]}}).sort("name", 1):
        out.append({"id": str(r["_id"]), "name": r["name"], "description": r.get("description", ""), "permissions_count": len(r.get("permissions", [])), "preset_key": r.get("preset_key"), "is_system": r.get("is_system", False), "hr": any(p.startswith("hr_") for p in r.get("permissions", []))})
    return {"roles": sorted(out, key=lambda x: (not x["hr"], x["name"]))}


@router.post("/employees/{emp_id}/account")
async def create_employee_account(emp_id: str, data: AccountIn = None, current_user: dict = Depends(get_current_user)):
    """🔐 إنشاء حساب دخول (خدمة ذاتية) — اسم المستخدم = الرقم الوظيفي، كلمة مرور أولية = الرقم الوظيفي مع إلزام التغيير، مع دور اختياري"""
    _guard(current_user, P_MANAGE)
    db = get_db()
    e = await db.employees.find_one({"_id": _oid(emp_id)})
    if not e:
        raise HTTPException(status_code=404, detail="الموظف غير موجود")
    if e.get("user_id"):
        raise HTTPException(status_code=400, detail="لدى الموظف حساب بالفعل")
    if e.get("teacher_id"):
        raise HTTPException(status_code=400, detail="حساب المعلم يُفعَّل من شاشة المعلمين")
    role = await _resolve_role(db, (data.role_id if data else None))
    username = e["employee_no"]
    if await db.users.find_one({"username": username}):
        raise HTTPException(status_code=400, detail="يوجد مستخدم بهذا الرقم الوظيفي")
    user = {"username": username, "password": get_password_hash(username), "hashed_password": get_password_hash(username), "full_name": e.get("full_name", ""), "role": "employee",
            "role_id": str(role["_id"]) if role else None, "email": e.get("email") or None, "phone": e.get("phone") or "", "employee_record_id": emp_id, "permissions": [], "custom_permissions": [],
            "must_change_password": True, "is_active": True, "created_at": datetime.now(YEMEN_TZ)}
    r = await db.users.insert_one(user)
    await db.employees.update_one({"_id": e["_id"]}, {"$set": {"user_id": str(r.inserted_id)}})
    await _history(db, emp_id, "account_created", current_user, {"username": username, **({"role": role["name"]} if role else {})})
    await log_activity(current_user, "hr_create_account", "employee", emp_id, e.get("full_name", ""), {"role": role["name"] if role else "موظف"})
    return {"username": username, "temp_password": username, "must_change_password": True, "role_name": role["name"] if role else None,
            "message": f"تم إنشاء الحساب: اسم المستخدم وكلمة المرور الأولية = {username}" + (f" — بدور «{role['name']}»" if role else "") + " (يُطلب تغييرها عند أول دخول)"}


@router.put("/employees/{emp_id}/account-role")
async def set_account_role(emp_id: str, data: AccountIn, current_user: dict = Depends(get_current_user)):
    """تغيير دور حساب الموظف (أو إزالته ليصبح موظفاً عادياً)"""
    _guard(current_user, P_MANAGE)
    db = get_db()
    e = await db.employees.find_one({"_id": _oid(emp_id)})
    if not e or not e.get("user_id"):
        raise HTTPException(status_code=404, detail="لا يوجد حساب لهذا الموظف")
    u = await db.users.find_one({"_id": ObjectId(e["user_id"])}) if ObjectId.is_valid(e["user_id"]) else None
    if not u:
        raise HTTPException(status_code=404, detail="حساب المستخدم غير موجود")
    if u.get("role") == "admin":
        raise HTTPException(status_code=400, detail="لا يمكن تغيير دور مدير النظام من هنا")
    role = await _resolve_role(db, data.role_id)
    old = await db.roles.find_one({"_id": ObjectId(u["role_id"])}, {"name": 1}) if u.get("role_id") and ObjectId.is_valid(u["role_id"]) else None
    await db.users.update_one({"_id": u["_id"]}, {"$set": {"role_id": str(role["_id"]) if role else None}})
    await _history(db, emp_id, "role_changed", current_user, {"الدور": {"from": (old or {}).get("name") or "موظف", "to": role["name"] if role else "موظف (بدون دور إداري)"}})
    await log_activity(current_user, "hr_account_role", "employee", emp_id, e.get("full_name", ""), {"role": role["name"] if role else None})
    return {"role_name": role["name"] if role else None, "message": f"تم تعيين الدور: {role['name'] if role else 'موظف بدون دور إداري'} — يسري عند تسجيل الدخول التالي"}


# ══════════════ استيراد Excel ══════════════

IMPORT_COLUMNS = [("employee_no", "الرقم الوظيفي*"), ("full_name", "الاسم الكامل*"), ("category", "الفئة (إداري/فني/خدمات)"), ("job_title", "المسمى الوظيفي"),
                  ("grade", "الدرجة"), ("org_unit", "الوحدة التنظيمية (اسم)"), ("contract_type", "نوع التعاقد (دائم/متعاقد/جزئي/بالساعة)"), ("hire_date", "تاريخ التعيين (YYYY-MM-DD)"),
                  ("contract_end_date", "نهاية العقد"), ("status", "الحالة (تحت التجربة/على رأس العمل/في إجازة/موقوف)"), ("gender", "الجنس"), ("nationality", "الجنسية"),
                  ("national_id", "رقم الهوية/الجواز"), ("id_expiry_date", "انتهاء الهوية"), ("birth_date", "تاريخ الميلاد"), ("phone", "الهاتف"), ("email", "البريد"),
                  ("qualification", "المؤهل"), ("specialization", "التخصص"), ("manager_no", "الرقم الوظيفي للمدير المباشر"), ("notes", "ملاحظات")]
_REV = lambda m: {v: k for k, v in m.items()}


@router.get("/employees/import/template")
async def import_template(current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "الموظفون"; ws.sheet_view.rightToLeft = True
    for i, (_, label) in enumerate(IMPORT_COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=label); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1565C0"); c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(16, len(label) + 4)
    ws.append(["EMP-001", "أحمد محمد سالم", "إداري", "أخصائي موارد بشرية", "الرابعة", "إدارة الموارد البشرية", "دائم", "2022-09-01", "", "على رأس العمل", "ذكر", "يمني", "01234567", "2028-01-01", "1990-05-10", "777000000", "a@ahgaff.edu", "بكالوريوس", "إدارة أعمال", "", ""])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    from .deps import export_headers, export_filename
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=export_headers(export_filename("نموذج استيراد الموظفين", ext="xlsx")))


async def _parse_import(db, file: UploadFile) -> List[dict]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="تعذّر قراءة الملف — يجب أن يكون Excel (.xlsx)")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="الملف فارغ")
    header = [str(h or "").strip() for h in rows[0]]
    label_to_key = {label.replace("*", "").strip(): key for key, label in IMPORT_COLUMNS}
    label_to_key.update({label: key for key, label in IMPORT_COLUMNS})
    col_keys = [label_to_key.get(h.replace("*", "").strip()) or (h if h in dict(IMPORT_COLUMNS) else None) for h in header]
    if "employee_no" not in col_keys or "full_name" not in col_keys:
        raise HTTPException(status_code=400, detail="الملف يجب أن يحوي عمودَي «الرقم الوظيفي» و«الاسم الكامل» — استخدم النموذج")
    units = {u["name"].strip(): str(u["_id"]) for u in await db.org_units.find({"is_active": {"$ne": False}}, {"name": 1}).to_list(2000)}
    cat_rev, ct_rev, st_rev = _REV(CATEGORIES), _REV(CONTRACT_TYPES), _REV(STATUSES)
    cat_rev.update({"إداري": "administrative", "أكاديمي": "academic", "معلم": "academic", "فني": "technical", "خدمات": "service"})
    existing_nos = {e["employee_no"]: str(e["_id"]) for e in await db.employees.find({}, {"employee_no": 1}).to_list(50000)}
    out, seen = [], set()

    def _d(v):
        if v in (None, ""):
            return None
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d")
        s = str(v).strip()[:10]
        try:
            datetime.strptime(s, "%Y-%m-%d"); return s
        except ValueError:
            return "INVALID"

    for idx, row in enumerate(rows[1:], start=2):
        rec = {k: (str(v).strip() if v is not None else "") for k, v in zip(col_keys, row) if k}
        if not any(rec.values()):
            continue
        errs = []
        no, name = rec.get("employee_no", ""), rec.get("full_name", "")
        if not no or not name:
            errs.append("الرقم الوظيفي والاسم مطلوبان")
        if no in seen:
            errs.append("رقم وظيفي مكرر داخل الملف")
        seen.add(no)
        cat = cat_rev.get(rec.get("category", ""), "administrative" if not rec.get("category") else None)
        if cat is None: errs.append(f"فئة غير معروفة: {rec.get('category')}")
        ct = ct_rev.get(rec.get("contract_type", ""), "permanent" if not rec.get("contract_type") else None)
        if ct is None: errs.append(f"نوع تعاقد غير معروف: {rec.get('contract_type')}")
        st = st_rev.get(rec.get("status", ""), "active" if not rec.get("status") else None)
        if st is None: errs.append(f"حالة غير معروفة: {rec.get('status')}")
        unit_id = None
        if rec.get("org_unit"):
            unit_id = units.get(rec["org_unit"])
            if not unit_id: errs.append(f"وحدة تنظيمية غير موجودة: {rec['org_unit']}")
        dates = {}
        for f in ("hire_date", "contract_end_date", "id_expiry_date", "birth_date"):
            dates[f] = _d(rec.get(f))
            if dates[f] == "INVALID": errs.append(f"تاريخ غير صحيح في {dict(IMPORT_COLUMNS)[f]}")
        out.append({"row": idx, "errors": errs, "exists": no in existing_nos, "existing_id": existing_nos.get(no), "data": {
            "employee_no": no, "full_name": name, "category": cat or "administrative", "job_title": rec.get("job_title", ""), "grade": rec.get("grade", ""),
            "org_unit_id": unit_id, "org_unit_name": rec.get("org_unit", ""), "contract_type": ct or "permanent", "status": st or "active", **dates,
            "gender": rec.get("gender", ""), "nationality": rec.get("nationality", ""), "national_id": rec.get("national_id", ""), "phone": rec.get("phone", ""),
            "email": rec.get("email", ""), "qualification": rec.get("qualification", ""), "specialization": rec.get("specialization", ""),
            "manager_no": rec.get("manager_no", ""), "notes": rec.get("notes", ""),
        }})
    return out


@router.post("/employees/import/preview")
async def import_preview(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    rows = await _parse_import(get_db(), file)
    return {"rows": rows, "total": len(rows), "valid": sum(1 for r in rows if not r["errors"]), "invalid": sum(1 for r in rows if r["errors"]),
            "new": sum(1 for r in rows if not r["errors"] and not r["exists"]), "updates": sum(1 for r in rows if not r["errors"] and r["exists"])}


@router.post("/employees/import")
async def import_employees(file: UploadFile = File(...), update_existing: bool = True, current_user: dict = Depends(get_current_user)):
    _guard(current_user, P_MANAGE)
    db = get_db()
    rows = await _parse_import(db, file)
    created, updated, skipped = 0, 0, 0
    for r in rows:
        if r["errors"]:
            skipped += 1; continue
        d = {k: v for k, v in r["data"].items() if k not in ("org_unit_name", "manager_no")}
        d.update({"address": "", "emergency_contact": "", "teacher_id": None, "manager_employee_id": None, "_manager_no": r["data"].get("manager_no", "")})
        if r["exists"]:
            if not update_existing:
                skipped += 1; continue
            await db.employees.update_one({"_id": ObjectId(r["existing_id"])}, {"$set": {**{k: v for k, v in d.items() if k not in ("_manager_no", "user_id", "teacher_id", "manager_employee_id")}, "updated_at": _now()}})
            updated += 1
        else:
            await db.employees.insert_one({**{k: v for k, v in d.items() if k != "_manager_no"}, "user_id": None, "created_at": _now(), "created_by_name": f"استيراد ({current_user.get('full_name', '')})"})
            created += 1
    # ربط المدير المباشر بعد إدراج الجميع
    nos = {e["employee_no"]: str(e["_id"]) for e in await db.employees.find({}, {"employee_no": 1}).to_list(50000)}
    for r in rows:
        m = r["data"].get("manager_no")
        if not r["errors"] and m and nos.get(m) and nos.get(r["data"]["employee_no"]) != nos.get(m):
            await db.employees.update_one({"employee_no": r["data"]["employee_no"]}, {"$set": {"manager_employee_id": nos[m]}})
    await log_activity(current_user, "hr_import_employees", "employee", None, "استيراد Excel", {"created": created, "updated": updated, "skipped": skipped})
    return {"created": created, "updated": updated, "skipped": skipped, "message": f"تم الاستيراد: {created} جديد، {updated} مُحدَّث، {skipped} مُتجاوَز"}
